"""Render one slide carrying every primitive, 24 glyphs, vendor marks, a monogram and an SVG picture,
once per theme, so the eye can check them: glyphs crisp, unfilled, upright; cards flat, no shadow.

    $PY scripts/render_prims.py            # writes tests/out/prims-<theme>.pptx, renders to .cache/prims/

The pptx files are git-ignored test output. The PNGs land in .cache/prims/<theme>-1.png.
"""
from __future__ import annotations

import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

SKILL = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SKILL / "kit"))

from pptx import Presentation  # noqa: E402

from deckkit import glyphs, icons, package, prims, text  # noqa: E402
from deckkit.canvas import Box  # noqa: E402
from deckkit.text import Para, Run  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

OUT = SKILL / "tests" / "out"
CACHE = SKILL / ".cache" / "prims"
BASE = SKILL / "kit" / "assets" / "base-wide.pptx"

GLYPH_SET = ["bot", "brain", "database", "cloud", "server", "shield-check", "lock", "users", "building", "factory",
             "truck", "chart-column", "trending-up", "target", "gauge", "clock", "calendar", "circle-check",
             "triangle-alert", "search", "workflow", "git-branch", "handshake", "stethoscope"]

SAMPLE_SVG = (b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 48 48" width="48" height="48">'
              b'<rect x="4" y="4" width="40" height="40" rx="8" fill="#4B5960"/>'
              b'<path d="M14 30 L22 18 L28 26 L34 16" stroke="#FFFFFF" stroke-width="3" fill="none" '
              b'stroke-linecap="round" stroke-linejoin="round"/></svg>')


def build(theme: Theme, path: Path):
    prs = Presentation(str(BASE))
    s = prs.slides.add_slide(prs.slide_layouts[0])
    with prims.use(theme):
        text.textbox(s, Box(0.5, 0.18, 12.3, 0.42),
                     [Para([Run("Primitive sheet: every creator, one theme", role="ink-deep", bold=True)])],
                     "title", theme, size=21.0)
        # row 1: cards, rect, zone, tab
        y = 0.85
        prims.card(s, Box(0.5, y, 1.9, 1.3), fill="paper", line="rule-strong", name="row1.card.outline")
        prims.card(s, Box(2.6, y, 1.9, 1.3), fill="card", line="hairline", name="row1.card.filled")
        prims.card(s, Box(4.7, y, 1.9, 1.3), fill="accent-tint", line=None, name="row1.card.tint")
        prims.rect(s, Box(6.8, y, 1.4, 1.3), fill="card", line="hairline", name="row1.rect")
        prims.zone(s, Box(8.4, y, 2.3, 1.3), label="Grouping zone", name="row1.zone")
        prims.pill(s, Box(8.6, y + 0.55, 1.1, 0.28), "Tag", fill="paper", line="hairline", name="row1.pill")
        prims.tab(s, Box(10.9, y, 1.95, 0.34), "Header tab", name="row1.tab")
        prims.card(s, Box(10.9, y + 0.34, 1.95, 0.96), fill="card", line="hairline", radius_in=0,
                   name="row1.tab.body")
        # row 2: step header, discs, person, quote
        y = 2.4
        prims.step_header(s, Box(0.5, y, 6.2, 0.5), ["Scope", "Build", "Prove", "Scale"], active=1,
                          name="row2.steps")
        for i in range(3):
            prims.disc(s, Box(7.0 + i * 0.55, y + 0.05, 0.4, 0.4), i + 1,
                       fill="accent" if i == 0 else "ink-strong", name="row2.disc.%d" % (i + 1))
        prims.person(s, Box(8.75, y - 0.05, 0.45, 0.6), name="row2.person")
        prims.quote(s, Box(9.4, y - 0.05, 3.45, 1.05), "The review took one afternoon, not three weeks.",
                    "Programme lead, fictional sample", name="row2.quote")
        # row 3: stats, caps label, hairline, connectors, elbow
        y = 3.65
        prims.stat(s, Box(0.5, y, 1.9, 1.25), "42%", "Fewer manual checks per wave", kind="measured",
                   name="row3.stat.1")
        prims.stat(s, Box(2.6, y, 1.9, 1.25), "3.1x", "Faster first draft", kind="illustrative", name="row3.stat.2")
        prims.stat(s, Box(4.7, y, 1.9, 1.25), "6 wk", "To a working pilot", kind="target", name="row3.stat.3")
        prims.caps_label(s, Box(6.9, y, 2.0, 0.22), "Caps label", name="row3.caps")
        prims.hairline(s, (6.9, y + 0.32), (8.9, y + 0.32), name="row3.rule")
        a = Box(9.2, y, 1.2, 0.5)
        b = Box(11.6, y, 1.2, 0.5)
        c = Box(11.6, y + 0.8, 1.2, 0.5)
        for bx, nm in ((a, "a"), (b, "b"), (c, "c")):
            prims.card(s, bx, fill="card", line="hairline", name="row3.node.%s" % nm)
        prims.connect(s, (a.r, a.cy), (b.x, b.cy), name="row3.arrow.ab")
        prims.elbow(s, [(a.cx, a.b), (a.cx, c.cy), (c.x, c.cy)], dash="dash", name="row3.elbow.ac")
        prims.connect(s, (b.cx, b.b), (c.cx, c.y), head="both", role="accent", weight_pt=1.0, name="row3.arrow.bc")
        # row 4: 24 glyphs
        y = 5.35
        for i, g in enumerate(GLYPH_SET):
            glyphs.glyph(s, g, Box(0.5 + i * 0.515, y, 0.4, 0.4), role="accent" if i % 6 == 0 else "ink-strong")
        # row 5: vendor marks, monograms, svg picture, large glyphs
        y = 6.15
        for i, k in enumerate(["aws/compute/ec2", "aws/storage/simple-storage-service-s3", "azure/database/sql-databases",
                               "gcp/analytics/bigquery"]):
            icons.vendor(s, k, Box(0.5 + i * 0.7, y, 0.5, 0.5))
        icons.monogram(s, Box(3.4, y, 0.5, 0.5), "Dx", "accent", name="row5.mono.1")
        icons.monogram(s, Box(4.1, y, 0.5, 0.5), "QA", "ink-strong", name="row5.mono.2")
        icons.svg_picture(s, SAMPLE_SVG, Box(4.8, y, 0.5, 0.5), "auto", name="row5.svg", alt="Sample trend mark")
        for i, g in enumerate(["database", "shield-check", "chart-pie", "map-pin"]):
            glyphs.glyph(s, g, Box(5.6 + i * 0.9, y - 0.1, 0.7, 0.7), role="navy")
        glyphs.glyph(s, "sparkles", Box(9.3, y + 0.15, 0.2, 0.2))
        glyphs.glyph(s, "key-round", Box(9.7, y, 0.5, 0.5), role="accent")
    path.parent.mkdir(parents=True, exist_ok=True)
    package.save(prs, path, title="Primitive sheet", theme=theme, overwrite=True)
    return path


def render(pptx: Path, stem: str) -> Path:
    CACHE.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        prof = "file://" + os.path.join(td, "lo-profile")
        subprocess.run(["soffice", "-env:UserInstallation=" + prof, "--headless", "--convert-to", "pdf",
                        "--outdir", td, str(pptx)], check=True, capture_output=True, timeout=180)
        pdf = Path(td) / (pptx.stem + ".pdf")
        subprocess.run(["pdftoppm", "-png", "-r", "150", str(pdf), str(CACHE / stem)], check=True, timeout=120)
    pngs = sorted(CACHE.glob(stem + "*.png"))
    return pngs[0]


def main() -> int:
    for name in ("sigmoid", "neutral"):
        th = Theme.load(SKILL / "themes" / ("%s.json" % name))
        p = build(th, OUT / ("prims-%s.pptx" % name))
        ok, msg = package.validate_xsd(p)
        png = render(p, name)
        print("%s: %s  xsd %s  render %s" % (name, p.relative_to(SKILL), "PASS" if ok else "FAIL " + msg[:200], png))
    return 0


if __name__ == "__main__":
    sys.exit(main())
