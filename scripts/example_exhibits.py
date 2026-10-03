"""Six slides, one per exhibit component, each a real archetype use (ruling 22, track A2):

    1 lifecycle-loop   exhibits.cycle_ring (wedges) between two side rails
    2 reference-table  tables.table with glyph-plus-word status cells
    3 credentials      compose.stat_band of delivered facts over exhibits.mark_wall
    4 coverage-matrix  exhibits.cell_grid with a key, over a shared-control band
    5 team-roster      exhibits.portrait pairs across a counterpart line, exhibits.quote_portrait
    6 product-screen   exhibits.screen_frame (browser) on a capture made by scripts/capture.py

    $PY scripts/example_exhibits.py [--out .cache/exhibits/exhibits.pptx] [--theme themes/sigmoid.json] [--recapture]

The capture is made once (Playwright) into <out folder>/assets/ and reused, so a rebuild (gate P21) is byte-stable.
Fictional programme: a weekly competitor brief for a field force; no client, brand or person. The firm facts on the
credentials slide are the approved claims (CLAUDE.md section 5).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
import zipfile
from pathlib import Path

SKILL = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SKILL / "kit"))
sys.path.insert(0, str(SKILL / "scripts"))

from pptx import Presentation  # noqa: E402

from deckkit import compose, exhibits, package, prims, tables, tag  # noqa: E402
from deckkit.canvas import WIDE, Box  # noqa: E402
from deckkit.glyphs import glyph  # noqa: E402
from deckkit.text import Para, Run, need_height, textbox, title  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

C = WIDE.content


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[0])


def rail(s, box, header, items, th, name):
    """A side rail: a solid header tab over a card of glyph rows, rows spread to the card's foot."""
    tab_h = 0.42
    prims.tab(s, Box(box.x, box.y, box.w, tab_h), header, fill="ink-strong", name="%s.tab" % name, theme=th)
    body = Box(box.x, box.y + tab_h, box.w, box.h - tab_h)
    prims.card(s, body, fill="card", line=None, radius_in=0, name="%s.body" % name, theme=th)
    gs, pad = 0.40, 0.18
    tw = body.w - 2 * pad - gs - 0.14
    stacks = [[Para((Run(h, "ink-deep", bold=True, size=11.0),)), Para((Run(b, "ink", size=10.5),), space_before=2.0)]
              for _g, h, b in items]
    hs = [max(gs, need_height(p, tw, th)) for p in stacks]
    free = body.h - 2 * pad - sum(hs)
    step = free / (len(hs) - 1)
    y = body.y + pad
    for i, ((g, _h, _b), p, h) in enumerate(zip(items, stacks, hs), 1):
        glyph(s, g, Box(body.x + pad, y, gs, gs), role="accent", shape_name="%s.%d.icon" % (name, i), theme=th)
        textbox(s, Box(body.x + pad + gs + 0.14, y, tw, need_height(p, tw, th)), p, "%s.%d.text" % (name, i), th,
                background="card")
        y += h + step


# ------------------------------------------------------------------ 1 lifecycle-loop

def s_loop(prs, th):
    s = blank(prs)
    tag.tag(s, "lifecycle-loop", "wedge-ring-side-rails")
    title(s, "Every weekly cycle feeds the next, so the brief sharpens with each release",
          keywords=("sharpens",), canvas=WIDE, theme=th)
    rw = 2.30
    rail(s, Box(C.x, C.y, rw, C.h), "What feeds it", [
        ("database", "Approved source list", "Registries, notices, press, congress"),
        ("map", "Territory plans", "Brand priority by region"),
        ("message-square", "Rep ratings", "Useful or not, on every brief"),
        ("calendar-days", "Release calendar", "Congress dates, launch windows")], th, "feeds")
    rail(s, Box(C.r - rw, C.y, rw, C.h), "What holds it", [
        ("shield-check", "Medical sign-off", "No brief ships unsigned"),
        ("file-check", "A source per claim", "Checked before review opens"),
        ("clipboard-list", "Release log", "Who saw which brief, and when"),
        ("gauge", "Weekly scorecard", "Speed, sourcing, rep ratings")], th, "holds")
    ring = Box(C.x + rw + 0.20, C.y, C.w - 2 * (rw + 0.20), C.h)
    exhibits.cycle_ring(s, ring, [
        ("scan-search", "Scan", "Approved feeds, every weekday"),
        ("list-ordered", "Rank", "Scored by territory plan"),
        ("file-text", "Draft", "One page, a source per claim"),
        ("user-check", "Review", "Medical signs or rejects"),
        ("send", "Release", "Monday, by territory"),
        ("refresh-cw", "Learn", "Ratings retune the ranking")],
        hub=("Weekly brief", "One cycle, every week"), name="loop", theme=th)


# ------------------------------------------------------------------ 2 reference-table

def s_register(prs, th):
    s = blank(prs)
    tag.tag(s, "reference-table", "source-register")
    title(s, "Every source behind a brief has an owner, a cadence and a check", keywords=("owner", "check"),
          canvas=WIDE, theme=th)
    tables.table(s, C, ["Source", "What it carries", "Cadence", "Owner", "Check before use", "Status"], [
        ["Trial registries", "New and amended competitor studies", "Daily", "Data steward",
         "Study id resolves in the registry", "Live"],
        ["Regulatory notices", "Approvals, label changes, safety letters", "Daily", "Regulatory lead",
         "Agency page cited and dated", "Live"],
        ["Company filings", "Guidance and pipeline changes", "Quarterly", "Finance partner",
         "Filing section cited", "Live"],
        ["Press releases", "Launches, deals and pricing moves", "Hourly", "Signal analyst",
         "Issuer domain verified", "Live"],
        ["Congress abstracts", "Data readouts by indication", "Per congress", "Medical lead",
         "Abstract number cited", "Pilot"],
        ["Payer updates", "Formulary and coverage changes", "Weekly", "Market access lead",
         "Plan document attached", "Pilot"],
        ["Field notes", "What reps hear in calls", "Daily", "Field excellence",
         "Two reports before it counts", "Planned"],
        ["Analyst notes", "Forecast and share commentary", "Weekly", "Insights lead",
         "Licensed source only", "Planned"]],
        cell_glyphs={"Live": "circle-check", "Pilot": ("clock", "ink-strong"), "Planned": ("calendar", "ink-muted")},
        widths=[1.7, 3.0, 1.15, 1.6, 2.75, 1.2], name="register", theme=th)


# ------------------------------------------------------------------ 3 credentials

def s_credentials(prs, th, numbers):
    s = blank(prs)
    tag.tag(s, "credentials", "stat-band-mark-wall")
    title(s, "The team already ships on the major cloud, data and agent platforms", keywords=("already ships",), canvas=WIDE,
          theme=th)
    band_h = 1.30
    names = compose.stat_band(s, Box(C.x, C.y, C.w, band_h), [
        ("100+", "AI products delivered", "delivered"),
        ("20+", "production-ready agents", "delivered"),
        ("~1,700", "people", "measured")], name="facts", theme=th)
    numbers.update({n: {"value": v, "kind": kd} for n, v, kd in
                    zip(names, ("100+", "20+", "~1,700"), ("delivered", "delivered", "measured"))})
    rec_h = 0.86
    wall = Box(C.x, C.y + band_h + 0.22, C.w, C.h - band_h - 0.22 - rec_h - 0.22)
    recognition(s, Box(C.x, C.b - rec_h, C.w, rec_h), th)
    exhibits.mark_wall(s, wall, [
        ("Cloud and models", [("vendor:aws/ml/bedrock", "Amazon Bedrock"),
                              ("vendor:azure/ml/azure-open-ai", "Azure OpenAI"),
                              ("vendor:gcp/ml/vertex-ai", "Vertex AI"),
                              ("mark:claude-spark", "Claude")]),
        ("Data platforms", [("vendor:onprem/analytics/databricks", "Databricks"),
                            ("vendor:saas/analytics/snowflake", "Snowflake"),
                            ("vendor:onprem/analytics/dbt", "dbt"),
                            ("vendor:onprem/database/postgresql", "PostgreSQL"),
                            ("mark:dataiku-logo", "Dataiku")]),
        ("Agents and evaluation", [("mark:langgraph-logo", "LangGraph"),
                                   ("mark:langchain-logo", "LangChain"),
                                   ("mark:langfuse-icon", "Langfuse"),
                                   ("mark:anthropic-symbol", "Anthropic API")]),
        ("Delivery", [("vendor:programming/language/python", "Python"),
                      ("vendor:programming/framework/fastapi", "FastAPI"),
                      ("vendor:programming/framework/react", "React"),
                      ("vendor:onprem/iac/terraform", "Terraform"),
                      ("vendor:onprem/ci/github-actions", "GitHub Actions"),
                      ("vendor:aws/compute/elastic-kubernetes-service", "Amazon EKS")])],
        name="wall", theme=th)


def recognition(s, box, th):
    """Recognition and plain facts (approved claims, CLAUDE.md section 5): a glyph, a bold head, its source. Every
    card's content starts at one offset (P07), centred on the fullest card."""
    items = [("handshake", "Claude Partner Network", "Member firm"),
             ("chart-scatter", "Gartner, Emerging Tech: AI Vendor Race", "Named in the 2026 report"),
             ("landmark", "Series B, founded 2013", "Head office in San Francisco")]
    cols = box.split_cols(3, 0.20)
    gs = 0.40
    tw = cols[0].w - 0.22 - gs - 0.16 - 0.18
    ps = [[Para((Run(h, "ink-deep", bold=True, size=12.0),)), Para((Run(b, "ink", size=10.5),), space_before=1.0)]
          for _g, h, b in items]
    ph = max(need_height(p, tw, th) for p in ps)
    top = box.y + (box.h - ph) / 2
    for i, ((g, _h, _b), c, p) in enumerate(zip(items, cols, ps), 1):
        prims.card(s, c, fill="card", line=None, name="recog.%d.card" % i, theme=th)
        glyph(s, g, Box(c.x + 0.22, top, gs, gs), role="accent", shape_name="recog.%d.icon" % i, theme=th)
        textbox(s, Box(c.x + 0.22 + gs + 0.16, top, tw, need_height(p, tw, th)), p, "recog.%d.text" % i, th,
                background="card")


# ------------------------------------------------------------------ 4 coverage-matrix

def s_coverage(prs, th):
    s = blank(prs)
    tag.tag(s, "coverage-matrix", "state-cells-control-band")
    title(s, "Each market goes live one capability at a time, every cell naming its state",
          keywords=("one capability at a time",), canvas=WIDE, theme=th)
    band_h = 0.80
    L, P, N, X = "live", "pilot", "next", "out"
    exhibits.cell_grid(s, Box(C.x, C.y, C.w, C.h - band_h - 0.20), [
        ("Signal scan", "Approved feeds, daily"), ("Ranking", "By territory plan"),
        ("Brief drafting", "A source per claim"), ("Medical review", "Sign, edit or reject"),
        ("Field release", "Monday, by territory"), ("Rep feedback", "Rated per brief")],
        ["United States", "Germany", "Japan", "Brazil", "France"], [
            [L, L, L, P, (N, "Wave 2")],
            [L, L, P, P, (N, "Wave 2")],
            [L, P, P, (N, "Wave 2"), (N, "Wave 3")],
            [L, L, P, (N, "Wave 2"), (N, "Wave 3")],
            [L, P, (N, "Wave 2"), (N, "Wave 3"), X],
            [P, (N, "Wave 2"), (N, "Wave 3"), X, X]],
        {L: ("circle-check", "Live", "accent"), P: ("clock", "Pilot", "ink-strong"),
         N: ("calendar", "Next wave", "card"), X: ("circle-x", "Not planned", "paper")},
        name="cover", theme=th)
    compose.pressure_band(s, Box(C.x, C.b - band_h, C.w, band_h), "Same in every market", [
        ("Source list", "One approved list, kept by the data steward"),
        ("Sign-off", "A medical reviewer signs each brief"),
        ("Release log", "Every brief traced to its readers")], fill="ink-deep", label_w=2.2, name="shared",
        theme=th)


# ------------------------------------------------------------------ 5 team-roster

PAIRS = [
    (("EL", "Engagement lead", "Scope, plan and the weekly readout"),
     "Weekly steering",
     ("FE", "Field excellence director", "Sets priorities and signs off scope")),
    (("SA", "Solution architect", "Pipeline design and its controls"),
     "Design reviews",
     ("IT", "Platform owner", "Hosting, access and security review")),
    (("DS", "Data steward", "Keeps the source list and its checks"),
     "Source approvals",
     ("DG", "Data governance lead", "Approves each new source")),
    (("AE", "Agent engineer", "Builds the scan, rank and brief agents"),
     "Daily review loop",
     ("MR", "Medical reviewer", "Signs every brief before release")),
    (("FA", "Field analyst", "Measures use and rep ratings"),
     "Weekly ratings",
     ("FM", "Regional field managers", "Bring the reps' view each week")),
    (("QA", "Evaluation lead", "Tests every agent against a gold set"),
     "Monthly audit",
     ("CO", "Compliance officer", "Reviews the release log each month")),
]


def s_team(prs, th):
    s = blank(prs)
    tag.tag(s, "team-roster", "paired-counterparts-quote")
    title(s, "Each delivery lead works beside one named counterpart from week one",
          keywords=("one named counterpart",), canvas=WIDE, theme=th)
    quote = exhibits.quote_portrait(s, Box(C.x, C.y, C.w, 0.95),
                                    "Every brief carries its sources, and nothing reaches a rep before medical has "
                                    "signed it.", "Engagement lead, delivery team", initials="EL",
                                    portrait_fill="ink-deep", portrait_in=0.80, name="team.quote", theme=th)
    col_w, mid_w = 4.55, C.w - 2 * 4.55
    lx, mx, rx = C.x, C.x + col_w, C.x + col_w + mid_w
    tab_h = 0.40
    ty = quote.b + 0.22
    prims.tab(s, Box(lx, ty, col_w, tab_h), "Delivery team", fill="ink-deep", name="team.left.tab", theme=th)
    prims.tab(s, Box(rx, ty, col_w, tab_h), "Client counterparts", fill="accent", name="team.right.tab", theme=th)
    top = ty + tab_h + 0.08
    row_h = (C.b - top) / len(PAIRS)
    d = min(0.62, row_h - 0.08)
    tw = col_w - d - 0.20
    for i, ((li, lr, ll), link, (ri, rr, rl)) in enumerate(PAIRS, 1):
        cy = top + (i - 0.5) * row_h
        # left: words, then the disc on the inner edge; right: the disc on the inner edge, then words
        exhibits.portrait(s, Box(mx - d, cy - d / 2, d, d), initials=li, fill="ink-deep", name="team.%d.left" % i,
                          theme=th)
        exhibits.portrait(s, Box(rx, cy - d / 2, d, d), initials=ri, fill="accent", name="team.%d.right" % i,
                          theme=th)
        lp = [Para((Run(lr, "ink-deep", bold=True, size=12.0),), align="right"),
              Para((Run(ll, "ink", size=10.5),), align="right", space_before=1.0)]
        rp = [Para((Run(rr, "ink-deep", bold=True, size=12.0),)),
              Para((Run(rl, "ink", size=10.5),), space_before=1.0)]
        h = max(need_height(lp, tw, th), need_height(rp, tw, th))
        textbox(s, Box(lx, cy - h / 2, tw, h), lp, "team.%d.left.text" % i, th)
        textbox(s, Box(rx + d + 0.20, cy - h / 2, tw, h), rp, "team.%d.right.text" % i, th)
        prims.hairline(s, (mx, cy), (rx, cy), role="rule-strong", name="team.%d.link" % i, theme=th)
        lab = [Para((Run(link, "ink-muted", bold=True, size=10.0),), align="center")]
        lh = need_height(lab, mid_w - 0.6, th)
        textbox(s, Box(mx + 0.3, cy - lh - 0.04, mid_w - 0.6, lh), lab, "team.%d.link.text" % i, th)


# ------------------------------------------------------------------ 6 product-screen

QUEUE_HTML = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Brief review queue</title>
<style>
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body { width: 1280px; height: 570px; font-family: Arial, Helvetica, sans-serif; font-size: 15px; color: #1f2933;
         background: #f4f5f7; overflow: hidden; }
  header { height: 52px; background: #1f2933; color: #fff; display: flex; align-items: center; padding: 0 22px;
           gap: 28px; }
  header b { font-size: 18px; }
  header span { font-size: 15px; color: #c9d1d9; }
  header .me { margin-left: auto; color: #fff; }
  main { display: grid; grid-template-columns: 330px 1fr 290px; height: 518px; }
  .list { background: #fff; border-right: 1px solid #d9dde3; overflow: hidden; }
  .list h2, .detail h2, .side h2 { font-size: 15px; text-transform: uppercase; letter-spacing: .04em; color: #52606d;
           padding: 16px 18px 10px; }
  .item { padding: 12px 18px; border-top: 1px solid #eceef1; }
  .item.on { background: #e8edf5; }
  .item b { display: block; font-size: 16px; }
  .item span { font-size: 15px; color: #52606d; }
  .detail { padding: 0 24px; overflow: hidden; }
  .detail h1 { font-size: 22px; padding: 6px 0 4px; }
  .meta { font-size: 15px; color: #52606d; padding-bottom: 12px; }
  .claim { background: #fff; border: 1px solid #d9dde3; border-radius: 6px; padding: 12px 14px; margin-bottom: 10px;
           display: flex; gap: 14px; align-items: flex-start; }
  .claim p { font-size: 16px; line-height: 1.35; flex: 1; }
  .src { font-size: 15px; color: #1e3c7d; border: 1px solid #b8c4dc; border-radius: 4px; padding: 3px 8px;
         white-space: nowrap; }
  .hist { width: 100%; border-collapse: collapse; font-size: 15px; background: #fff; border: 1px solid #d9dde3; }
  .hist td { padding: 8px 12px; border-top: 1px solid #eceef1; }
  .note { margin-top: 18px; font-size: 15px; color: #52606d; border: 1px solid #d9dde3; border-radius: 6px;
          padding: 10px 12px; height: 92px; }
  .side { background: #fff; border-left: 1px solid #d9dde3; padding: 0 18px; }
  .check { font-size: 15px; padding: 8px 0; border-top: 1px solid #eceef1; display: flex; justify-content: space-between; }
  .ok { color: #2f6b3a; font-weight: bold; }
  .btn { display: block; width: 100%; text-align: center; padding: 11px; border-radius: 6px; font-size: 16px;
         font-weight: bold; margin-top: 12px; }
  .approve { background: #1e3c7d; color: #fff; }
  .edit { border: 1px solid #1e3c7d; color: #1e3c7d; }
  .reject { border: 1px solid #9aa5b1; color: #3e4c59; }
</style></head>
<body>
<header><b>Brief desk</b><span>Review queue</span><span>Week 41</span><span class="me">Medical reviewer</span></header>
<main>
  <section class="list"><h2>Waiting for review (6)</h2>
    <div class="item on"><b>Northeast, oncology</b><span>3 claims, drafted 06:40</span></div>
    <div class="item"><b>Midwest, oncology</b><span>2 claims, drafted 06:42</span></div>
    <div class="item"><b>Southeast, immunology</b><span>4 claims, drafted 06:45</span></div>
    <div class="item"><b>West, immunology</b><span>2 claims, drafted 06:47</span></div>
    <div class="item"><b>Southwest, cardiology</b><span>3 claims, drafted 06:51</span></div>
    <div class="item"><b>Northwest, cardiology</b><span>1 claim, drafted 06:55</span></div>
  </section>
  <section class="detail"><h2>Brief</h2>
    <h1>Northeast, oncology: two competitor moves this week</h1>
    <div class="meta">Drafted by the brief agent at 06:40, ranked 1 of 14 signals for this territory</div>
    <div class="claim"><p>A competitor filed an amended phase 3 protocol that widens enrolment to second-line patients.</p>
      <span class="src">Trial registry</span></div>
    <div class="claim"><p>The same company announced a co-promotion deal covering three northeast academic centres.</p>
      <span class="src">Press release</span></div>
    <div class="claim"><p>A regional payer moved the competing product to preferred status from the first of next month.</p>
      <span class="src">Payer update</span></div>
    <h2 style="padding-left:0">Last three briefs for this territory</h2>
    <table class="hist">
      <tr><td>Week 40</td><td>2 claims</td><td>Approved by medical, 08:12</td><td class="ok">Released</td></tr>
      <tr><td>Week 39</td><td>3 claims</td><td>Edited, one claim removed</td><td class="ok">Released</td></tr>
      <tr><td>Week 38</td><td>1 claim</td><td>Rejected, source not on the list</td><td>Withdrawn</td></tr>
    </table>
  </section>
  <aside class="side"><h2>Checks</h2>
    <div class="check"><span>Every claim has a source</span><span class="ok">Passed</span></div>
    <div class="check"><span>Sources on approved list</span><span class="ok">Passed</span></div>
    <div class="check"><span>No off-label wording</span><span class="ok">Passed</span></div>
    <div class="check"><span>Territory matches plan</span><span class="ok">Passed</span></div>
    <span class="btn approve">Approve and release</span>
    <span class="btn edit">Edit before release</span>
    <span class="btn reject">Reject with reason</span>
    <div class="note">Reviewer note: confirm the payer date against the plan document before release.</div>
  </aside>
</main>
</body></html>
"""


def screen_asset(folder: Path, recapture: bool) -> Path:
    """The review-queue capture: made once by scripts/capture.py, then reused (byte-stable rebuilds)."""
    folder.mkdir(parents=True, exist_ok=True)
    html = folder / "review-queue.html"
    png = folder / "review-queue.png"
    if recapture or not png.exists() or not (folder / "review-queue.capture.json").exists() \
            or not html.exists() or html.read_text(encoding="utf-8") != QUEUE_HTML:
        html.write_text(QUEUE_HTML, encoding="utf-8")
        import capture
        capture.capture(html, png, width=1280, height=570, scale=2.0, placed_width_in=12.0)
    return png


def s_screen(prs, th, png):
    s = blank(prs)
    tag.tag(s, "product-screen", "browser-framed-full-width")
    title(s, "Reviewers clear every brief in one queue, claims beside sources",
          keywords=("one queue",), canvas=WIDE, theme=th)
    return exhibits.screen_frame(s, C, png, chrome="browser", alt="Brief review queue with three sourced claims",
                                 name="screen", theme=th)


# ------------------------------------------------------------------ build

def mark_media(deck: Path) -> dict:
    """{sha256: reason} for every embedded picture that is a registered vendor or brand mark, byte for byte. Gate P18
    matches public marks against the reference decks, which embed the same artwork; P22 verifies each one."""
    from deckkit import icons
    table = json.loads((SKILL / "kit" / "assets" / "vendor_hashes.json").read_text(encoding="utf-8"))["keys"]
    keys = {h: "official vendor mark %s from the diagrams package" % k for k, sizes in table.items()
            for h in sizes.values()}
    for k, v in icons.mark_registry().items():
        keys[v["sha256"]] = "brand mark %s from %s (ruling 21)" % (k, v["source"])
    out = {}
    with zipfile.ZipFile(deck) as z:
        for n in sorted(z.namelist()):
            if n.startswith("ppt/media/"):
                h = hashlib.sha256(z.read(n)).hexdigest()
                if h in keys:
                    out[h] = keys[h] + ", byte-verified by P22; the reference decks embed the same public mark"
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(SKILL / ".cache" / "exhibits" / "exhibits.pptx"))
    ap.add_argument("--theme", default=str(SKILL / "themes" / "sigmoid.json"))
    ap.add_argument("--assets", default=str(SKILL / ".cache" / "exhibits" / "assets"))
    ap.add_argument("--recapture", action="store_true")
    a = ap.parse_args()
    th = Theme.load(a.theme)
    png = screen_asset(Path(a.assets), a.recapture)
    prs = Presentation(str(SKILL / "kit" / "assets" / "base-wide.pptx"))
    prims.bind(th)
    numbers: dict = {}
    s_loop(prs, th)
    s_register(prs, th)
    s_credentials(prs, th, numbers)
    s_coverage(prs, th)
    s_team(prs, th)
    shot = s_screen(prs, th, png)
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    sha = package.save(prs, out, title="Exhibits", theme=th, manifest={"numbers": numbers,
                       "captures": {"screen": {"type_pt": shot["type_pt"]}}}, overwrite=True, date="2026-10-04")
    dj = {"type": "example", "canvas": "wide", "theme": a.theme, "client": None, "spelling": "US",
          "media_allow": mark_media(out), "overrides": {}, "series": [], "deliverable": out.name,
          "build": ["{python}", str(Path(__file__).resolve()), "--theme", a.theme, "--assets", a.assets,
                    "--out", "{out}"]}
    (out.parent / "deck.json").write_text(json.dumps(dj, indent=2) + "\n", encoding="utf-8")
    print("screen type %.2fpt as placed" % shot["type_pt"])
    print("OK %s sha256 %s" % (out, sha[:12]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
