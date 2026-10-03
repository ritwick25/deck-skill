"""Helpers for a device agent's build.py (ruling 14): open the lifted draft, inspect it, rename, retext, move,
delete, add glyphs, save deterministically.

    import sys; sys.path.insert(0, "<skill>/scripts"); from devwork import *
    prs, s, th = open_lifted("DL07")
    dump(s)                                   # name, kind, box, fill role, text, size for every shape
    rename(s, "dev.12", "slot.card.1.head")
    text(s, "slot.card.1.head", "Confirm the pilot region", th)          # keeps the shape's size and bold
    move(s, "dev.40", 0.5, 6.2, 12.3, 0.5); delete(s, "dev.41")
    icon(s, "slot.card.1.icon", "map-pin", 1.0, 1.0, 0.42, th)
    save(prs, "DL07", th)                     # devices/work/DL07/device.pptx, deterministic (P21)
"""
from __future__ import annotations

import sys
from pathlib import Path

SKILL = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SKILL / "kit"))

from pptx import Presentation  # noqa: E402

from deckkit import ooxml, package, text as txt  # noqa: E402
from deckkit.canvas import Box  # noqa: E402
from deckkit.geom import find  # noqa: E402
from deckkit.glyphs import glyph  # noqa: E402
from deckkit.read.ns import qn  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

WORK = SKILL / "devices" / "work"
EMU = 914400


def open_lifted(dv: str, theme: str = "sigmoid"):
    th = Theme.load(SKILL / "themes" / ("%s.json" % theme))
    prs = Presentation(str(WORK / dv / "lifted.pptx"))
    return prs, prs.slides[0], th


def _walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:                       # group
            yield from _walk(sh.shapes)


def dump(s) -> None:
    for sh in _walk(s.shapes):
        el = sh._element
        fill = el.find(".//" + qn("p:spPr") + "/" + qn("a:solidFill") + "/" + qn("a:schemeClr"))
        r = el.find(".//" + qn("a:r") + "/" + qn("a:rPr"))
        size = int(r.get("sz")) / 100 if r is not None and r.get("sz") else ""
        t = sh.text_frame.text.replace("\n", " / ")[:60] if sh.has_text_frame and sh.text_frame.text.strip() else ""
        geo = el.find(".//" + qn("a:prstGeom"))
        kind = (geo.get("prst") if geo is not None else ("custGeom" if el.find(".//" + qn("a:custGeom")) is not None
                                                          else el.tag.split("}")[1]))
        print("%-22s %-14s x%5.2f y%5.2f w%5.2f h%5.2f fill=%-8s %4s %s" % (
            sh.name, kind, sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU,
            fill.get("val") if fill is not None else "-", size, t))


def rename(s, old: str, new: str) -> None:
    find(s, old).name = new


def delete(s, name: str) -> None:
    el = find(s, name)._element
    el.getparent().remove(el)


def move(s, name: str, x: float, y: float, w: float | None = None, h: float | None = None) -> None:
    sh = find(s, name)
    sh.left, sh.top = int(x * EMU), int(y * EMU)
    if w is not None:
        sh.width = int(w * EMU)
    if h is not None:
        sh.height = int(h * EMU)


def insets_of(shape):
    """The text frame's own insets (bodyPr lIns/tIns/rIns/bIns, PowerPoint defaults when absent), in inches."""
    from deckkit import measure as _m
    bp = shape._element.find(".//" + qn("a:bodyPr"))
    d = {"lIns": 91440, "tIns": 45720, "rIns": 91440, "bIns": 45720}
    v = {k: int(bp.get(k)) if bp is not None and bp.get(k) is not None else d[k] for k in d}
    return _m.Insets(v["lIns"] / 914400, v["tIns"] / 914400, v["rIns"] / 914400, v["bIns"] / 914400)


def text(s, name: str, value, th, *, role: str | None = None, size: float | None = None, bold: bool | None = None,
         background: str = "paper", anchor: str = "top") -> None:
    """Rewrite a text shape in its own box, keeping its first run's size and bold unless given; fit-checked."""
    sh = find(s, name)
    r = sh._element.find(".//" + qn("a:r") + "/" + qn("a:rPr"))
    size = size or (int(r.get("sz")) / 100 if r is not None and r.get("sz") else 11.0)
    bold = bold if bold is not None else (r is not None and r.get("b") in ("1", "true"))
    role = role or "ink"
    vals = [value] if isinstance(value, str) else list(value)
    paras = [v if isinstance(v, txt.Para) else txt.Para((txt.Run(v, role, bold=bold, size=size),),
                                                        space_before=0 if i == 0 else 3) for i, v in enumerate(vals)]
    box = Box(sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU)
    ins = insets_of(sh)
    txt.check_fit("slot %s" % name, paras, box, th, ins, size=size)
    txt.write(sh.text_frame, paras, th, ins, size=size, anchor=anchor, background=background)


def icon(s, name: str, glyph_name: str, x: float, y: float, side: float, th, role: str = "accent") -> None:
    glyph(s, glyph_name, Box(x, y, side, side), role=role, shape_name=name, theme=th)


def save(prs, dv: str, th, archetype: str | None = None, variant: str | None = None) -> Path:
    from deckkit import tag
    if archetype:
        tag.tag(prs.slides[0], archetype, variant or "device")
    for sh in prs.slides[0].shapes:
        ooxml.finish(sh)
    out = WORK / dv / "device.pptx"
    package.save(prs, out, title="Device %s" % dv, theme=th, overwrite=True, date="2026-10-03")
    return out
