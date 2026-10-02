"""Fit the measurer's width and line multipliers against LibreOffice, and record them.

Procedure (what this script does, in order):
  1. Build a throwaway pptx in a temp directory: one slide per (face, point size). Each slide holds
     the known strings below, one per text box, no wrap, zero insets, explicit size and face; plus a
     block of six one-line paragraphs at 100% line spacing with no space before or after.
  2. Convert it to PDF with LibreOffice headless under an isolated profile
     (-env:UserInstallation), so a parallel build cannot collide with this one.
  3. Read word boxes with `pdftotext -bbox-layout`. A PDF point is 1/72 inch, the same unit the
     measurer reports in, so no page scaling is involved.
  4. Width ratio per string = rendered line width (xMax - xMin) / measured advance-width sum
     (measure._units, before any calibration). Line ratio = median baseline pitch of the six-line
     block / (pt x the font's own line factor).
  5. Record min, median and max per point size and the median over every sample as the fitted
     factor, per face, in kit/assets/metrics.json. measure.calibration_factor() applies a fitted
     factor only when it is above 1.0: the font's design metric is the floor of the estimate.

SLACK (1.15) is separate and is not fitted here: it covers PowerPoint setting text wider than
LibreOffice renders it, which a LibreOffice run cannot observe.

Run:   $PY scripts/calibrate_metrics.py                       (Arial and Arial Bold at 8, 10, 14, 21)
       $PY scripts/calibrate_metrics.py --faces "Arial" --pts 8 10 14 21 --dry
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import statistics
import subprocess
import sys
import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET

SKILL = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SKILL / "kit"))

from deckkit import measure as M  # noqa: E402

STRINGS = [
    "The quick brown fox jumps over the lazy dog",
    "Regional demand review for the third quarter",
    "Revenue grew 18% across 4 markets in 2025",
    "WAVE MAVEN AVOW TYPE",
    "illicit little lilies fill it",
    "0123456789 1,284 37.5% 2,040",
    "Hamburgefonstiv Hamburgefonstiv",
]
PITCH_ROWS = 6
PITCH_TEXT = "Pitch row {} Hxg"


def build_pptx(path: Path, cases: list[tuple[str, float]]) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from lxml import etree

    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for key, pt in cases:
        face = M.FACES[key]
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        y = 0.3
        step = max(0.45, pt * 2.2 / 72)

        def box(ty, paras):
            tb = slide.shapes.add_textbox(Inches(0.3), Inches(ty), Inches(12.5), Inches(0.4))
            body = tb.text_frame._txBody
            bp = body.find(A + "bodyPr")
            for k in ("lIns", "tIns", "rIns", "bIns"):
                bp.set(k, "0")
            bp.set("wrap", "none")
            for child in list(bp):
                bp.remove(child)
            tf = tb.text_frame
            for i, text in enumerate(paras):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                pPr = p._p.get_or_add_pPr()
                ln = etree.SubElement(pPr, A + "lnSpc")
                etree.SubElement(ln, A + "spcPct").set("val", "100000")
                for tag in ("spcBef", "spcAft"):
                    etree.SubElement(etree.SubElement(pPr, A + tag), A + "spcPts").set("val", "0")
                r = p.add_run()
                r.text = text
                r.font.size = Pt(pt)
                r.font.name = face.family
                r.font.bold = face.bold
                r.font.italic = face.italic

        for s in STRINGS:
            box(y, [s])
            y += step
        box(y + 0.1, [PITCH_TEXT.format(i + 1) for i in range(PITCH_ROWS)])
    prs.save(str(path))


def render(pptx: Path, out: Path) -> Path:
    profile = out / "lo-profile"
    cmd = ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
           "--outdir", str(out), str(pptx)]
    subprocess.run(cmd, check=True, capture_output=True, timeout=240)
    pdf = out / (pptx.stem + ".pdf")
    if not pdf.exists():
        raise SystemExit(f"LibreOffice produced no PDF at {pdf}")
    return pdf


def lines_by_page(pdf: Path) -> list[list[dict]]:
    xml = subprocess.run(["pdftotext", "-bbox-layout", str(pdf), "-"], check=True,
                         capture_output=True, text=True).stdout
    xml = re.sub(r"<!DOCTYPE[^>]*>", "", xml)
    xml = re.sub(r'xmlns="[^"]*"', "", xml)
    root = ET.fromstring(xml)
    pages = []
    for page in root.iter("page"):
        rows = []
        for line in page.iter("line"):
            ws = list(line.iter("word"))
            if not ws:
                continue
            rows.append({"text": " ".join(w.text or "" for w in ws),
                         "x0": float(line.get("xMin")), "x1": float(line.get("xMax")),
                         "y0": float(line.get("yMin")), "y1": float(line.get("yMax"))})
        pages.append(rows)
    return pages


def fit(cases, pages):
    per_face: dict[str, dict] = {}
    for (key, pt), rows in zip(cases, pages):
        face = M.FACES[key]
        by_text = {r["text"]: r for r in rows}
        w_ratios = []
        for s in STRINGS:
            r = by_text.get(s)
            if r is None:
                raise SystemExit(f"{key} {pt}pt: string {s!r} not found in the PDF text layer")
            measured = M._units(s, key) * pt          # points, uncalibrated advance sum
            w_ratios.append((r["x1"] - r["x0"]) / measured)
        ys = []
        for i in range(PITCH_ROWS):
            r = by_text.get(PITCH_TEXT.format(i + 1))
            if r is None:
                raise SystemExit(f"{key} {pt}pt: pitch row {i + 1} not found in the PDF text layer")
            ys.append(r["y0"])
        pitches = [b - a for a, b in zip(ys, ys[1:])]
        factor = M._font_metrics(face).line_factor
        line_ratio = statistics.median(pitches) / (pt * factor)
        rec = per_face.setdefault(key, {"samples": {}})
        rec["samples"][f"{pt:g}"] = {
            "width_min": round(min(w_ratios), 4), "width_median": round(statistics.median(w_ratios), 4),
            "width_max": round(max(w_ratios), 4), "line_pitch_pt": round(statistics.median(pitches), 3),
            "line_ratio": round(line_ratio, 4)}
    for key, rec in per_face.items():
        ws = [v["width_median"] for v in rec["samples"].values()]
        ls = [v["line_ratio"] for v in rec["samples"].values()]
        rec["width"] = round(statistics.median(ws), 4)
        rec["line"] = round(statistics.median(ls), 4)
        rec["font_line_factor"] = round(M._font_metrics(M.FACES[key]).line_factor, 4)
        rec["applied"] = {"width": max(1.0, rec["width"]), "line": max(1.0, rec["line"])}
    return per_face


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--faces", nargs="+", default=["Arial", "Arial Bold"])
    ap.add_argument("--pts", nargs="+", type=float, default=[8, 10, 14, 21])
    ap.add_argument("--dry", action="store_true", help="print the fit, do not write metrics.json")
    a = ap.parse_args(argv)
    for k in a.faces:
        if k not in M.FACES:
            raise SystemExit(f"unknown face {k!r}; known: {sorted(M.FACES)}")
    cases = [(k, pt) for k in a.faces for pt in a.pts]
    lo = subprocess.run(["soffice", "--version"], capture_output=True, text=True).stdout.strip()
    with tempfile.TemporaryDirectory(prefix="deck-calib-") as td:
        td = Path(td)
        pptx = td / "calibration.pptx"
        build_pptx(pptx, cases)
        pages = lines_by_page(render(pptx, td))
    if len(pages) != len(cases):
        raise SystemExit(f"expected {len(cases)} PDF pages, got {len(pages)}")
    faces = fit(cases, pages)
    for key, rec in faces.items():
        print(f"{key}: width x{rec['width']} line x{rec['line']} (font line factor {rec['font_line_factor']})")
        for pt, s in rec["samples"].items():
            print(f"  {pt:>4}pt width min/median/max {s['width_min']}/{s['width_median']}/{s['width_max']}"
                  f"  pitch {s['line_pitch_pt']}pt  line x{s['line_ratio']}")
    if a.dry:
        return 0
    path = M.METRICS_JSON
    data = json.loads(path.read_text()) if path.exists() else {}
    data.setdefault("faces", {}).update(faces)
    data["status"] = "calibrated"
    data["renderer"] = lo
    data["date"] = dt.date.today().isoformat()
    data["procedure"] = ("scripts/calibrate_metrics.py: known strings rendered through LibreOffice "
                         "headless to PDF, widths and line pitch read with pdftotext -bbox-layout, "
                         "compared with fontTools advance sums and the font's own line factor. "
                         "Applied factor = max(1.0, fitted). SLACK 1.15 is separate and not fitted.")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, sort_keys=True) + "\n")
    print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
