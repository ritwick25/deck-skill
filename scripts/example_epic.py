"""Three slides in the epic layer (ruling 10): a near-black stat band with one basis label, numbered panels under
solid header bands, a pressure band over a stage chain, and an engine-drawn L0 master (deckkit.arch, ruling 18)
beside a coded autonomy ladder. The pattern for dense, high-contrast slides under the house rules.

    $PY scripts/example_epic.py [--out .cache/epic/epic.pptx] [--theme themes/sigmoid.json]

  1 problem-framing      two numbered panels under solid header bands, the figures in a near-black band beneath
  2 process-flow         pressure band at the top, a five-step chain whose facts reach the foot (no foot band)
  3 layered-architecture an L0 master from the engine (product marks, no arrows) and the autonomy ladder as a rail

Fictional client and figures; every stat is illustrative. Builds on any theme (a dark client primary included:
stat_band switches the numerals to accent-bright where the brand does not clear 3:1 on the band).
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

SKILL = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SKILL / "kit"))

from pptx import Presentation  # noqa: E402

from deckkit import arch, compose, package, prims, tag  # noqa: E402
from deckkit.canvas import WIDE, Box  # noqa: E402
from deckkit.text import title  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

C = WIDE.content


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[0])


def problem(prs, th, numbers):
    s = blank(prs)
    tag.tag(s, "problem-framing", "banded-panels-stat-band")
    title(s, "Field briefs arrive late because no one owns the signal", keywords=("arrive late",), canvas=WIDE,
          theme=th)
    half = (C.w - 0.2) / 2
    ph = 3.95
    compose.numbered_rows(s, Box(C.x, C.y, half, ph), [
        ("Signals are scattered", "Registries, press, congress and payer notices, read by no one person"),
        ("Reps learn last", "Customers mention a competitor move before the field hears of it"),
        ("Notes lack sources", "Ad hoc summaries reach reps with no source and no medical check"),
        ("Every region rebuilds it", "Six teams write six versions of the same week")],
        header="Why briefs arrive late", header_glyph="clock", name="why", theme=th)
    compose.numbered_rows(s, Box(C.x + half + 0.2, C.y, half, ph), [
        ("Manual triage", "Analysts sift signals by hand, so the backlog grows every week"),
        ("No single source list", "Each team trusts different feeds and none is approved"),
        ("Review comes too late", "Medical sees the brief after reps have used it"),
        ("No record of what changed", "Edits and rejections are lost in email threads")],
        header="Why it is hard to fix", header_glyph="triangle-alert", name="hard", theme=th)
    figures = [("40+", "public competitor signals a week", "illustrative"),
               ("5 days", "from a competitor move to the field", "illustrative"),
               ("6", "regional versions of one brief", "illustrative")]
    top = C.y + ph + 0.12
    names = compose.stat_band(s, Box(C.x, top, C.w, C.b - top), figures, name="stat", theme=th)
    numbers.update({n: {"value": v, "kind": k} for n, (v, _l, k) in zip(names, figures)})


def process(prs, th):
    s = blank(prs)
    tag.tag(s, "process-flow", "pressure-band-stage-chain")
    title(s, "Five steps turn public signals into an approved brief each week", keywords=("approved brief",),
          canvas=WIDE, theme=th)
    compose.pressure_band(s, Box(C.x, C.y, C.w, 1.0), "Field pressures", [
        ("Volume", "40+ signals a week, most of them noise"),
        ("Speed", "Reps need the move before the Monday calls"),
        ("Trust", "Each claim links to its public source"),
        ("Control", "Medical signs before anything ships")], name="press", theme=th)
    compose.stage_chain(s, Box(C.x, C.y + 1.2, C.w, C.h - 1.2), [
        ("Scan", "scan-search", "Read the source list", "Approved public feeds only, every weekday"),
        ("Rank", "list-ordered", "Score each signal", "Relevance by territory and brand plan"),
        ("Draft", "file-text", "Write one page", "Every claim linked to its source"),
        ("Review", "user-check", "Medical decides", "Approve, edit or reject with a reason"),
        ("Release", "send", "Ship by territory", "Each rep gets the moves that matter")],
        facts=[[("Output", "New signals, no repeats"), ("Control", "Approved source list"), ("Owner", "Scan agent"),
                ("Measure", "Feeds read on time")],
               [("Output", "A ranked shortlist"), ("Control", "Relevance threshold"), ("Owner", "Scan agent"),
                ("Measure", "Signals kept per week")],
               [("Output", "Draft with sources"), ("Control", "No claim without a source"), ("Owner", "Brief agent"),
                ("Measure", "Claims with a live source")],
               [("Output", "Signed brief"), ("Control", "Reviewer signature"), ("Owner", "Medical reviewer"),
                ("Measure", "Unsigned releases: none")],
               [("Output", "Field brief and record"), ("Control", "Release log"), ("Owner", "Field excellence"),
                ("Measure", "Reps who rate it useful")]],
        name="chain", theme=th)


# The L0 master: one name and one mark per component (P26.v), product marks only, no arrows and no numbers.
SOURCES = (("web", "vendor:aws/general/internet-alt1", "Public web"), ("reg", "glyph:file-search", "Trial feeds"),
           ("pay", "glyph:landmark", "Payers"), ("media", "glyph:megaphone", "Congress media"),
           ("partner", "glyph:folder-open", "Partner feeds"))
IDENTITY = (("cognito", "vendor:aws/security/cognito", "Cognito"),
            ("iam", "vendor:aws/security/identity-and-access-management-iam", "AWS IAM"),
            ("kms", "vendor:aws/security/key-management-service", "AWS KMS"),
            ("secrets", "vendor:aws/security/secrets-manager", "Secrets"))
INGEST = (("sched", "vendor:aws/integration/eventbridge-scheduler", "EventBridge"),
          ("collect", "vendor:aws/compute/lambda", "Lambda"),
          ("ocr", "vendor:aws/ml/textract", "Textract"),
          ("queue", "vendor:aws/integration/simple-queue-service-sqs", "Amazon SQS"))
STORE = (("s3", "vendor:aws/storage/simple-storage-service-s3", "Amazon S3"),
         ("search", "vendor:aws/analytics/amazon-opensearch-service", "OpenSearch"),
         ("rds", "vendor:aws/database/rds", "RDS"),
         ("glue", "vendor:aws/analytics/glue", "AWS Glue"))
AGENTS = (("bedrock", "vendor:aws/ml/bedrock", "Bedrock", 0.62), ("claude", "mark:claude-spark", "Claude", None),
          ("graph", "mark:langgraph-logo", "LangGraph", None),
          ("scan", "vendor:aws/compute/elastic-container-service-task", "Scan agent", None),
          ("writer", "vendor:aws/compute/elastic-container-service-task", "Writer agent", None))
SERVE = (("api", "vendor:aws/network/api-gateway", "API Gateway"),
         ("ui", "vendor:programming/framework/react", "Review app"),
         ("ddb", "vendor:aws/database/dynamodb", "DynamoDB"),
         ("ses", "vendor:aws/engagement/simple-email-service-ses", "Amazon SES"))
DELIVERY = (("actions", "vendor:onprem/ci/github-actions", "GitHub Actions"),
            ("ecr", "vendor:aws/compute/ec2-container-registry", "Amazon ECR"),
            ("tf", "vendor:onprem/iac/terraform", "Terraform"), ("langfuse", "mark:langfuse-icon", "Langfuse"),
            ("watch", "vendor:aws/management/cloudwatch", "CloudWatch"))


def architecture(prs, th):
    s = blank(prs)
    tag.tag(s, "layered-architecture", "l0-ladder-rail")
    title(s, "One governed platform feeds every brief, and autonomy is earned", keywords=("autonomy is earned",),
          canvas=WIDE, theme=th)
    lw = 8.85
    d = arch.Diagram("house", l0=True, grid=(1, 2), icon=0.38, head_air=0.12)
    d.group("src", "region", "Sources · public and partner", at=(0, 0, 1, 1), grid=(5, 1))
    for k, (nid, mark, nm) in enumerate(SOURCES):
        d.node(nid, mark, nm, group="src", at=(k, 0), caption="right", size=0.34)
    d.group("plat", "boundary", "Brief platform · one AWS account, eu-west-1", at=(0, 1, 1, 1), grid=(4, 3),
            weights=[0.95, 2.45, 2.45, 1.6], row_gap=[0.03, 0.03])
    d.group("idrail", "rail", "Identity", parent="plat", at=(0, 0, 1, 2), grid=(1, 4))
    for r, (nid, mark, nm) in enumerate(IDENTITY):
        d.node(nid, mark, nm, group="idrail", at=(0, r), size=0.34)
    d.group("ingest", "region", "Ingest", parent="plat", at=(1, 0, 1, 1), grid=(2, 2))
    for k, (nid, mark, nm) in enumerate(INGEST):
        d.node(nid, mark, nm, group="ingest", at=(k % 2, k // 2))
    d.group("store", "region", "Store", parent="plat", at=(2, 0, 1, 1), grid=(2, 2))
    for k, (nid, mark, nm) in enumerate(STORE):
        d.node(nid, mark, nm, group="store", at=(k % 2, k // 2))
    d.group("agents", "region", "Agents", parent="plat", at=(1, 1, 2, 1), grid=(5, 1))
    for k, (nid, mark, nm, size) in enumerate(AGENTS):
        d.node(nid, mark, nm, group="agents", at=(k, 0), size=size)
    d.group("serve", "region", "Review and serve", parent="plat", at=(3, 0, 1, 2), grid=(1, 4))
    for r, (nid, mark, nm) in enumerate(SERVE):
        d.node(nid, mark, nm, group="serve", at=(0, r))
    d.group("band", "band", "Delivery and operations", parent="plat", at=(0, 2, 4, 1), grid=(5, 1))
    for k, (nid, mark, nm) in enumerate(DELIVERY):
        d.node(nid, mark, nm, group="band", at=(k, 0), caption="right", size=0.34)
    d.render(s, Box(C.x, C.y, lw, C.h), th, name="l0")
    rx = C.x + lw + 0.25
    rw = C.r - rx
    prims.tab(s, Box(rx, C.y, rw, 0.42), "Autonomy is earned", fill="accent", name="ladder.header", theme=th)
    compose.ladder(s, Box(rx, C.y + 0.55, rw, C.h - 0.55), [
        ("L0", "Inform", "Reps read the brief"),
        ("L1", "Recommend", "Agent ranks the signals"),
        ("L2", "Draft", "Agent writes, medical signs"),
        ("L3", "Release on rule", "Low-risk updates ship alone"),
        ("L4", "Autonomous", "Policy-bound, audited")], name="ladder", theme=th)


def vendor_media(deck: Path) -> dict:
    """{sha256: reason} for every embedded picture that is a registered vendor mark, byte for byte (P18 allowance;
    P22 has verified each against kit/assets/vendor_hashes.json)."""
    import hashlib
    import zipfile
    table = json.loads((SKILL / "kit" / "assets" / "vendor_hashes.json").read_text(encoding="utf-8"))["keys"]
    keys = {h: k for k, sizes in table.items() for h in sizes.values()}
    out = {}
    with zipfile.ZipFile(deck) as z:
        for n in sorted(z.namelist()):
            if n.startswith("ppt/media/"):
                h = hashlib.sha256(z.read(n)).hexdigest()
                if h in keys:
                    out[h] = "official vendor mark %s, byte-verified by P22 against vendor_hashes.json" % keys[h]
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(SKILL / ".cache" / "epic" / "epic.pptx"))
    ap.add_argument("--theme", default=str(SKILL / "themes" / "sigmoid.json"))
    a = ap.parse_args()
    th = Theme.load(a.theme)
    prs = Presentation(str(SKILL / "kit" / "assets" / "base-wide.pptx"))
    prims.bind(th)
    numbers: dict = {}
    problem(prs, th, numbers)
    process(prs, th)
    architecture(prs, th)
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    sha = package.save(prs, out, title="Epic layer", theme=th, manifest={"numbers": numbers}, overwrite=True,
                       date="2026-10-03")
    (out.parent / "deck.json").write_text(json.dumps({
        "type": "example", "canvas": "wide", "theme": a.theme, "client": None, "spelling": "US",
        "media_allow": vendor_media(out), "overrides": {}, "series": [], "deliverable": out.name,
        "build": ["{python}", str(Path(__file__).resolve()), "--theme", a.theme, "--out", "{out}"]}, indent=2) + "\n")
    print("OK %s sha256 %s" % (out, sha[:12]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
