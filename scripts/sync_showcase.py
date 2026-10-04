"""Copy the accepted showcase deck into the skill (ruling 24): showcase/showcase.pptx, its slide index and 4-up sheets.

    $PY scripts/sync_showcase.py            # copy the deck, write showcase/index.json, render and tile the sheets
    $PY scripts/sync_showcase.py --check    # exit 1 when the copy is stale against the library's build (selftest)

The source is the deck library's acceptance build, <library>/../acceptance/showcase, the deck its deck.json names as
"deliverable". The index maps every slide to its section, title, archetype, variant, family, the library device it
places (by variant) and the code that drew it: a skill example function where the slide is wired from one, else the
part function in the acceptance build. Reuse a slide's code, or clone the slide, before drawing one from scratch.
The copy is a reference, never a deliverable: it is rebuilt in the library and synced here, never edited here.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from paths import ARCH_JSON, LIB, PY, SKILL  # noqa: E402

sys.path.insert(0, str(SKILL / "kit"))
from deckkit import tag  # noqa: E402

SRC = LIB.parent / "acceptance" / "showcase"
DST = SKILL / "showcase"
DECK = DST / "showcase.pptx"
INDEX = DST / "index.json"
SHEETS = DST / "sheets"
PER_SHEET = 4
EXAMPLE = re.compile(r"\((example_\w+)\.(\w+)\)")


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def source_deck() -> Path | None:
    dj = SRC / "deck.json"
    if not dj.is_file():
        return None
    p = SRC / json.loads(dj.read_text())["deliverable"]
    return p if p.is_file() else None


def titles(deck: Path) -> list[str]:
    from pptx import Presentation

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                yield from walk(sh.shapes)
            else:
                yield sh
    out = []
    for s in Presentation(str(deck)).slides:
        texts = [(sh.name, " ".join(sh.text_frame.text.split())) for sh in walk(s.shapes) if sh.has_text_frame]
        texts = [(n, t) for n, t in texts if t]
        named = [t for n, t in texts if n in ("title", "slot.title", "slot.headline")]
        out.append((named or sorted((t for _, t in texts), key=lambda t: -len(t.split())) or [""])[0])
    return out


def index(src: Path) -> dict:
    plan = json.loads((SRC / "plan.json").read_text())["slides"]
    tags = tag.read(str(DECK))
    if len(plan) != len(tags):
        sys.exit("plan.json has %d rows, the deck %d slides: the acceptance build is out of step" % (len(plan), len(tags)))
    fam = {s: a["family"] for s, a in json.loads(ARCH_JSON.read_text())["archetypes"].items()}
    dev = {d["variant"]: d["id"] for d in json.loads((SKILL / "devices" / "devices.json").read_text())["devices"].values()}
    heads = titles(DECK)
    slides = []
    for row, t, head in zip(plan, tags, heads):
        if t["archetype"] != row["slug"]:
            sys.exit("slide %d is tagged %s, plan.json says %s" % (row["n"], t["archetype"], row["slug"]))
        m = EXAMPLE.search(row.get("note") or "")
        code = ("scripts/%s.py:%s" % m.groups()) if m else ("library acceptance build, parts/%s.py:%s" % (row["part"], row["fn"]))
        k = row["n"] - 1
        slides.append({"n": row["n"], "section": row["section"], "title": head, "archetype": t["archetype"],
                       "variant": t["variant"], "family": fam.get(t["archetype"]), "device": dev.get(t["variant"]),
                       "code": code, "sheet": "showcase/sheets/sheet-%02d.png" % (k // PER_SHEET + 1),
                       "sheet_pos": ("top-left", "top-right", "bottom-left", "bottom-right")[k % PER_SHEET]})
    return {"_about": "The accepted showcase deck (ruling 24): a fictional proposal on the house theme, every archetype, "
                      "every engine diagram type and the native charts, 27 gates passing. A reference to look at and "
                      "reuse code from, never a deliverable. Synced from the library by scripts/sync_showcase.py.",
            "source": "acceptance/showcase/%s in the deck library" % src.name,
            "sha256": sha(DECK), "slides": slides}


def tile(renders: list[Path]) -> None:
    from PIL import Image, ImageDraw
    if SHEETS.exists():
        shutil.rmtree(SHEETS)
    SHEETS.mkdir(parents=True)
    w, h = 1100, 619
    for k in range(0, len(renders), PER_SHEET):
        sheet = Image.new("RGB", (w * 2 + 30, (h + 28) * 2 + 10), "white")
        d = ImageDraw.Draw(sheet)
        for j, f in enumerate(renders[k:k + PER_SHEET]):
            im = Image.open(f).convert("RGB").resize((w, h), Image.LANCZOS)
            x, y = 10 + (j % 2) * (w + 10), 5 + (j // 2) * (h + 28)
            sheet.paste(im, (x, y + 22))
            d.text((x, y + 6), "slide %d" % (k + j + 1), fill="black")
            d.rectangle([x, y + 22, x + w, y + 22 + h], outline="#999999")
        sheet.save(SHEETS / ("sheet-%02d.png" % (k // PER_SHEET + 1)), optimize=True)


def check() -> list[str]:
    """Why the copy is stale or broken; [] when it is current (or the library is absent and the copy is whole)."""
    fails = []
    if not DECK.is_file() or not INDEX.is_file():
        return ["showcase/ is incomplete: run scripts/sync_showcase.py"]
    doc = json.loads(INDEX.read_text())
    if doc.get("sha256") != sha(DECK):
        fails.append("showcase/showcase.pptx differs from the sha256 its index records")
    n = len(doc.get("slides") or [])
    sheets = sorted(SHEETS.glob("sheet-*.png"))
    if len(sheets) != (n + PER_SHEET - 1) // PER_SHEET:
        fails.append("showcase/sheets holds %d sheets for %d slides" % (len(sheets), n))
    src = source_deck()
    if src is not None and sha(src) != sha(DECK):
        fails.append("the showcase was rebuilt in the library: run scripts/sync_showcase.py")
    return fails


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true")
    a = ap.parse_args()
    if a.check:
        fails = check()
        for f in fails:
            print("FAIL", f)
        print("%s showcase copy" % ("OK" if not fails else "STALE"))
        return 1 if fails else 0
    src = source_deck()
    if src is None:
        sys.exit("no acceptance build at %s (set DECK_LIBRARY)" % SRC)
    DST.mkdir(exist_ok=True)
    shutil.copy2(src, DECK)
    doc = index(src)
    with tempfile.TemporaryDirectory(prefix="showcase-") as tmp:
        r = subprocess.run([PY, str(SKILL / "scripts" / "render.py"), str(DECK), "--out", tmp], capture_output=True,
                           text=True)
        if r.returncode:
            sys.exit("render failed: " + (r.stdout + r.stderr)[-400:])
        renders = sorted(Path(tmp).glob("slide-*.png"), key=lambda p: int(p.stem.split("-")[1]))
        if len(renders) != len(doc["slides"]):
            sys.exit("rendered %d slides for %d in the index" % (len(renders), len(doc["slides"])))
        tile(renders)
    INDEX.write_text(json.dumps(doc, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print("OK %d slides, %d sheets, sha256 %s" % (len(doc["slides"]), len(list(SHEETS.glob("*.png"))),
                                                  doc["sha256"][:12]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
