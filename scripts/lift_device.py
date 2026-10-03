"""Lift one reference slide into a sanitised device draft (ruling 3, steps P0, P1, P2, P4, P5 of the plan).

    $PY scripts/lift_device.py dNN-sNNN DEVICE_ID [--theme themes/sigmoid.json] [--out devices/work]

Mechanical, deterministic, no judgement:
  P0  a fresh package from kit/assets/base-wide.pptx; the slide's own shapes only (no layout or master chrome),
      copied by deckkit.clone.copy_shapes under DEVICE_STRIP (no alt text, links, tags, glue, artistic layers)
  P1  every picture and picture fill removed (no corpus raster is ever reused) and listed with its box, so the
      device agent places a curated glyph or vendor mark there; hidden and off-canvas shapes removed; fields
      (slide numbers, dates) removed
  P2  every word replaced by filler of the same word count (no reference text survives the lift); every run
      gets the explicit size it rendered at in the source (autofit baked in), floored at 8pt, and the theme
      face; autofit removed
  P4  every colour resolved against the source slide's own theme and colour map, then written as the nearest
      house role (ruling 12: brand fills, near-black and greys; text in ink, accent-ink or on-accent);
      gradients flattened to their middle stop; pattern and picture fills made solid; effects and 3D removed
  P5  a 10-inch slide is scaled to the wide canvas (geometry and type x 1.333); groups baked to identity
Writes devices/work/DEVICE_ID/lifted.pptx and lift.json (counts, the removed pictures, the colour map).
The device agent then retexts, places glyphs, names slots, fits the floor and gates it (steps P3, P6 to P9).
"""
from __future__ import annotations

import argparse
import colorsys
import json
import re
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from paths import KIT, LIB, SKILL  # noqa: E402

sys.path.insert(0, str(KIT))
from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402

from deckkit import ooxml, package  # noqa: E402
from deckkit.measure import FLOOR  # noqa: E402
from deckkit.clone import DEVICE_STRIP, copy_shapes  # noqa: E402
from deckkit.geom import bake_groups  # noqa: E402
from deckkit.read.colour import parse_clrmap, resolve  # noqa: E402
from deckkit.read.ns import local  # noqa: E402
from deckkit.read.opc import DeckRead  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
EMU = 914400
WIDE_W = 13.333
FILLER = ("lorem ipsum dolor sit amet consectetur adipiscing elit sed do eiusmod tempor incididunt ut labore et "
          "dolore magna aliqua enim ad minim veniam quis nostrud exercitation ullamco laboris nisi").split()
COLOUR_TAGS = ("srgbClr", "schemeClr", "prstClr", "sysClr", "scrgbClr", "hslClr")
TEXT_CTX = ("rPr", "defRPr", "endParaRPr", "buClr")


def hls(hexv: str):
    r, g, b = (int(hexv[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    return h, l, s


def role_for(hexv: str, ctx: str) -> str:
    """Nearest house role by lightness and saturation (ruling 12: one brand hue, near-black, greys)."""
    _h, l, s = hls(hexv)
    vivid = s >= 0.35 and 0.12 < l < 0.92
    if ctx == "text":
        if vivid:
            return "accent-ink"
        if l >= 0.80:
            return "on-accent"
        if l < 0.30:
            return "ink-deep"
        if l < 0.45:
            return "ink"
        return "ink-muted"
    if vivid:
        if l < 0.35:
            return "accent-deep"
        if l < 0.70:
            return "accent"
        if l < 0.86:
            return "accent-tint"
        return "accent-wash"
    if ctx == "line":
        if l < 0.45:
            return "ink-muted"
        if l < 0.70:
            return "rule-strong"
        return "hairline"
    for lim, role in ((0.16, "ink-deep"), (0.30, "ink-strong"), (0.62, "rule-strong"), (0.80, "hairline"),
                      (0.90, "hairline-soft"), (0.975, "card")):
        if l < lim:
            return role
    return "paper"


def ctx_of(el) -> str:
    for anc in el.iterancestors():
        n = local(anc)
        if n in TEXT_CTX:
            return "text"
        if n == "ln":
            return "line"
        if n in ("spPr", "tcPr", "bg", "bgPr", "grpSpPr"):
            return "fill"
    return "fill"


def slide_cctx(dr: DeckRead, index1: int):
    sp = dr.slide_part(index1)
    root = dr.xml(sp)
    lp = dr.layout_part_of(sp)
    li = dr.layout(lp) if lp else None
    mi = dr.master(li.master_part) if li and li.master_part else None
    ovr = parse_clrmap(root.find(P + "clrMapOvr"))
    cm = ovr or (li.clrmap_ovr if li else None) or (mi.clrmap if mi else None)
    return mi.colour.with_clrmap(cm)


def recolour(tree, cctx, th: Theme, hist: Counter) -> dict:
    """Flatten gradients, pattern and picture fills; drop effects; then every colour element becomes a role."""
    n = Counter()
    for g in list(tree.iter(A + "gradFill")):
        stops = g.findall(A + "gsLst/" + A + "gs")
        mid = min(stops, key=lambda s: abs(int(s.get("pos", "0")) - 50000)) if stops else None
        solid = etree.Element(A + "solidFill")
        if mid is not None:
            c = next((x for x in mid if local(x) in COLOUR_TAGS), None)
            if c is not None:
                solid.append(c)
        g.addprevious(solid)
        g.getparent().remove(g)
        n["gradients"] += 1
    for pf in list(tree.iter(A + "pattFill")):
        solid = etree.Element(A + "solidFill")
        fg = pf.find(A + "fgClr")
        c = next((x for x in fg if local(x) in COLOUR_TAGS), None) if fg is not None else None
        if c is not None:
            solid.append(c)
        pf.addprevious(solid)
        pf.getparent().remove(pf)
        n["patterns"] += 1
    for parent_tag in ("spPr", "tcPr"):
        for sppr in tree.iter(A + parent_tag if parent_tag == "tcPr" else P + parent_tag):
            for bf in list(sppr.findall(A + "blipFill")):
                solid = etree.Element(A + "solidFill")
                solid.append(etree.SubElement(etree.Element("x"), A + "srgbClr", val="F2F2F2"))
                bf.addprevious(solid)
                sppr.remove(bf)
                n["picture_fills"] += 1
    for tag in ("effectLst", "effectDag", "scene3d", "sp3d"):
        for e in list(tree.iter(A + tag)):
            e.getparent().remove(e)
            n["effects"] += 1
    for c in list(tree.iter()):
        if not isinstance(c.tag, str) or local(c) not in COLOUR_TAGS:
            continue
        par = c.getparent()
        if par is None or local(par) in ("gs",):            # a stop left over inside a removed gradient
            continue
        try:
            r = resolve(c, cctx)
        except Exception:
            continue
        hexv = (r.get("hex") or "000000").upper()
        role = role_for(hexv, ctx_of(c))
        hist["%s->%s" % (hexv, role)] += 1
        new = th.spec(role)
        par.replace(c, new)
        n["colours"] += 1
    return dict(n)


def leaves_by_id(sid: str) -> dict:
    d, s = sid.split("-")
    sh = json.loads((LIB / "shapes" / d / ("%s.json" % s)).read_text())
    return {L["id"]: L for L in sh["shapes"] if isinstance(L.get("id"), int)}, sh


def retext(tree, leaves: dict, k: float, n: Counter) -> None:
    """Filler for every word, explicit floor-checked sizes from the source's rendered sizes, theme face."""
    wi = 0
    for sp in tree.iter(P + "sp", P + "graphicFrame"):
        cnv = sp.find(".//" + P + "cNvPr")
        L = leaves.get(int(cnv.get("id"))) if cnv is not None and (cnv.get("id") or "").isdigit() else None
        src_runs = []
        if L and isinstance(L.get("text"), dict):
            for p in L["text"].get("paragraphs") or []:
                for r in p.get("runs") or []:
                    src_runs.append(r.get("size_rendered_pt") or r.get("size_pt"))
        ri = 0
        for body in sp.iter(A + "bodyPr"):
            for tag in ("normAutofit", "spAutoFit"):
                for e in body.findall(A + tag):
                    body.remove(e)
                    n["autofit"] += 1
            if body.find(A + "noAutofit") is None:
                etree.SubElement(body, A + "noAutofit")
        for fld in list(sp.iter(A + "fld")):
            fld.getparent().remove(fld)
            n["fields"] += 1
        for r in sp.iter(A + "r"):
            t = r.find(A + "t")
            if t is None:
                continue
            words = re.findall(r"\S+", t.text or "")
            out = []
            for w in words:
                out.append("00" if re.search(r"\d", w) else FILLER[wi % len(FILLER)])
                wi += 1
            lead = " " if (t.text or "").startswith(" ") else ""
            trail = " " if (t.text or "").endswith(" ") and words else ""
            t.text = lead + " ".join(out) + trail
            rpr = r.find(A + "rPr")
            if rpr is None:
                rpr = etree.Element(A + "rPr")
                r.insert(0, rpr)
            pt = src_runs[ri] if ri < len(src_runs) and src_runs[ri] else None
            ri += 1
            if pt is None and rpr.get("sz"):
                pt = int(rpr.get("sz")) / 100.0
            pt = (pt or 11.0) * k
            if pt < FLOOR:
                n["raised_to_floor"] += 1
                pt = FLOOR
            rpr.set("sz", str(int(round(pt * 2)) * 50))
            for tag in ("latin", "ea", "cs", "sym"):
                for e in rpr.findall(A + tag):
                    rpr.remove(e)
            n["runs"] += 1


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("slide")
    ap.add_argument("device")
    ap.add_argument("--theme", default=str(SKILL / "themes" / "sigmoid.json"))
    ap.add_argument("--out", default=str(SKILL / "devices" / "work"))
    a = ap.parse_args()
    deck_id, sno = a.slide.split("-")
    index1 = int(sno[1:])
    man = json.loads((LIB / "manifest.json").read_text())["decks"][deck_id]
    src_path = LIB / "_work" / "src" / ("%s.pptx" % deck_id)
    src_path = src_path if src_path.is_file() else Path(man["path"])
    th = Theme.load(a.theme)
    src = Presentation(str(src_path))
    s_slide = src.slides[index1 - 1]
    dr = DeckRead(str(src_path))
    cctx = slide_cctx(dr, index1)
    leaves, _sh = leaves_by_id(a.slide)
    k = WIDE_W / (src.slide_width / EMU)
    tree = s_slide.shapes._spTree                       # in memory only: the source file is never written
    hist: Counter = Counter()
    removed = []
    for pic in list(tree.iter(P + "pic")):
        cnv = pic.find(".//" + P + "cNvPr")
        L = leaves.get(int(cnv.get("id"))) if cnv is not None else None
        b = (L or {}).get("bbox") or {}
        removed.append({"box_in": [round(b.get(x, 0) * k, 3) for x in ("x_in", "y_in", "w_in", "h_in")],
                        "kind": ((L or {}).get("picture") or {}).get("kind_guess")})
        pic.getparent().remove(pic)
    for el in list(tree):
        if not isinstance(el.tag, str) or local(el) not in ("sp", "grpSp", "cxnSp", "graphicFrame"):
            continue
        cnv = el.find(".//" + P + "cNvPr")
        if cnv is not None and cnv.get("hidden") in ("1", "true"):
            tree.remove(el)
            hist["hidden_removed"] += 1
    counts = recolour(tree, cctx, th, hist)
    n = Counter(counts)
    retext(tree, leaves, k, n)
    prs = Presentation(str(SKILL / "kit" / "assets" / "base-wide.pptx"))
    d_slide = prs.slides.add_slide(prs.slide_layouts[0])
    copy_shapes(s_slide, d_slide, strip=DEVICE_STRIP)
    dtree = d_slide.shapes._spTree
    if abs(k - 1) > 1e-3:
        for el in dtree:
            if not isinstance(el.tag, str):
                continue
            xf = el.find("./" + P + "spPr/" + A + "xfrm") or el.find("./" + P + "grpSpPr/" + A + "xfrm") \
                or el.find("./" + P + "xfrm")
            if xf is None:
                continue
            for tag in ("off", "ext"):
                e = xf.find(A + tag)
                if e is not None:
                    for at in ("x", "y", "cx", "cy"):
                        if e.get(at) is not None:
                            e.set(at, str(int(round(int(e.get(at)) * k))))
            if local(el) == "graphicFrame":
                for g in el.iter(A + "gridCol"):
                    g.set("w", str(int(round(int(g.get("w")) * k))))
                for tr in el.iter(A + "tr"):
                    tr.set("h", str(int(round(int(tr.get("h")) * k))))
        n["scaled"] = round(k, 4)
    try:
        bake_groups(d_slide)
    except Exception as e:                       # a rotated or flipped group: the gates compose transforms anyway
        n["bake_skipped"] = str(e)[:120]
    W, H = prs.slide_width, prs.slide_height
    for el in list(dtree):
        xf = el.find(".//" + A + "xfrm") if isinstance(el.tag, str) else None
        off = xf.find(A + "off") if xf is not None else None
        ext = xf.find(A + "ext") if xf is not None else None
        if off is None or ext is None:
            continue
        x, y, cx, cy = (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))
        if x >= W or y >= H or x + cx <= 0 or y + cy <= 0:
            dtree.remove(el)
            n["offcanvas_removed"] += 1
    for i, cnv in enumerate(dtree.iter(P + "cNvPr"), 1):
        if cnv.getparent() is not None and local(cnv.getparent()) == "nvGrpSpPr" and cnv.getparent().getparent() is dtree:
            continue
        cnv.set("name", "dev.%d" % i)
    for shp in d_slide.shapes:
        ooxml.finish(shp)
    out = Path(a.out) / a.device
    out.mkdir(parents=True, exist_ok=True)
    package.save(prs, out / "lifted.pptx", title="Device %s" % a.device, theme=th, overwrite=True, date="2026-10-03")
    report = {"slide": a.slide, "device": a.device, "scale": round(k, 4), "counts": dict(n),
              "pictures_removed": removed, "colour_map": dict(hist.most_common())}
    (out / "lift.json").write_text(json.dumps(report, indent=1) + "\n")
    print("OK %s -> %s  %s  pictures removed %d" % (a.slide, out / "lifted.pptx", dict(n), len(removed)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
