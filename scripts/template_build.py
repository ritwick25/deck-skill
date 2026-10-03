"""__NAME__: the build script, and the deck's source of truth. Never hand-edit the pptx; change this file and rebuild.

    $PY build.py [--out "__OUT__"] [--theme __THEME__]
    $PY ~/.claude/skills/deck/gates/run_all.py --deck "__OUT__"
    $PY ~/.claude/skills/deck/scripts/render.py "__OUT__"         # then read every slide

Scaffolded by the deck skill's new_deck.py in the epic layer (ruling 10): a composed slide with a near-black stat
band, an engine-drawn L0 architecture (deckkit.arch, ruling 18) and a library device. Every visible word lives in
COPY below. A value still in {{double braces}} is a placeholder: gate P19 refuses the deck until each one is replaced.
Replace the three slides with the storyline's own (SKILL.md steps 3 to 7), keeping the order of work: tag, title,
device, words, figures in the manifest, then gate and read the render.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

SKILL = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SKILL / "kit"))

from pptx import Presentation  # noqa: E402

from deckkit import arch, compose, devices, package, prims, tag  # noqa: E402
from deckkit.canvas import WIDE, Box  # noqa: E402
from deckkit.text import title  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

C = WIDE.content            # x 0.50, y 0.85, w 12.40, h 5.99: the band every content slide fills
DECK_TITLE = "__NAME__"
DATE = "2026-10-04"          # the package date: fixed, so a rebuild is byte-identical (gate P21)

# --- COPY ---
COPY = {
    "summary": {
        "title": ("{{One claim: what the proposal changes}}", ()),
        "left": ("{{What changes}}", [("{{Change one}}", "{{One line on what it means}}"),
                                      ("{{Change two}}", "{{One line on what it means}}"),
                                      ("{{Change three}}", "{{One line on what it means}}"),
                                      ("{{Change four}}", "{{One line on what it means}}")]),
        "right": ("{{What it takes}}", [("{{Need one}}", "{{One line on who and when}}"),
                                        ("{{Need two}}", "{{One line on who and when}}"),
                                        ("{{Need three}}", "{{One line on who and when}}"),
                                        ("{{Need four}}", "{{One line on who and when}}")]),
        "stats": [("{{n1}}", "{{what the first figure counts}}", "illustrative"),
                  ("{{n2}}", "{{what the second figure counts}}", "illustrative"),
                  ("{{n3}}", "{{what the third figure counts}}", "illustrative")],
    },
    "platform": {
        "title": ("{{One claim: how the platform is layered}}", ()),
        "sources": ("{{Sources}} · {{one fact}}", [
            ("vendor:aws/general/internet-alt1", "{{Web}}"), ("glyph:file-search", "{{Feeds}}"),
            ("glyph:landmark", "{{Records}}"), ("glyph:megaphone", "{{Media}}"), ("glyph:folder-open", "{{Files}}")]),
        "boundary": "{{Platform}} · {{one fact}}",
        "rail": ("{{Identity}}", [("vendor:aws/security/cognito", "{{Cognito}}"),
                                  ("vendor:aws/security/identity-and-access-management-iam", "{{IAM}}"),
                                  ("vendor:aws/security/key-management-service", "{{KMS}}"),
                                  ("vendor:aws/security/secrets-manager", "{{Secrets}}")]),
        "ingest": ("{{Ingest}}", [("vendor:aws/integration/eventbridge-scheduler", "{{Scheduler}}"),
                                  ("vendor:aws/compute/lambda", "{{Lambda}}"),
                                  ("vendor:aws/ml/textract", "{{Textract}}"),
                                  ("vendor:aws/integration/simple-queue-service-sqs", "{{SQS}}")]),
        "store": ("{{Store}}", [("vendor:aws/storage/simple-storage-service-s3", "{{S3}}"),
                                ("vendor:aws/analytics/amazon-opensearch-service", "{{OpenSearch}}"),
                                ("vendor:aws/database/rds", "{{RDS}}"),
                                ("vendor:aws/analytics/glue", "{{Glue}}")]),
        "agents": ("{{Agents}}", [("vendor:aws/ml/bedrock", "{{Bedrock}}"), ("mark:claude-spark", "{{Claude}}"),
                                  ("mark:langgraph-logo", "{{LangGraph}}"),
                                  ("vendor:aws/compute/elastic-container-service-task", "{{Agent one}}"),
                                  ("vendor:aws/compute/elastic-container-service-task", "{{Agent two}}")]),
        "serve": ("{{Serve}}", [("vendor:aws/network/api-gateway", "{{API}}"),
                                ("vendor:programming/framework/react", "{{App}}"),
                                ("vendor:aws/database/dynamodb", "{{DynamoDB}}"),
                                ("vendor:aws/engagement/simple-email-service-ses", "{{SES}}")]),
        "band": ("{{Delivery and operations}}", [("vendor:onprem/ci/github-actions", "{{CI}}"),
                                                 ("vendor:aws/compute/ec2-container-registry", "{{Registry}}"),
                                                 ("vendor:onprem/iac/terraform", "{{Terraform}}"),
                                                 ("mark:langfuse-icon", "{{Langfuse}}"),
                                                 ("vendor:aws/management/cloudwatch", "{{CloudWatch}}")]),
    },
    "ask": {
        "title": ("{{One claim: the decisions that start the work}}", ()),
        "cards": [("check", "{{Decision one}}", "{{Why it matters, in one line}}"),
                  ("user-check", "{{Decision two}}", "{{Why it matters, in one line}}"),
                  ("database", "{{Decision three}}", "{{Why it matters, in one line}}"),
                  ("shield-check", "{{Decision four}}", "{{Why it matters, in one line}}")],
        "rail": "{{Outcomes by date}}",
        "milestones": [("{{Date one}}", "{{What is true by then}}"), ("{{Date two}}", "{{What is true by then}}"),
                       ("{{Date three}}", "{{What is true by then}}"), ("{{Date four}}", "{{What is true by then}}")],
    },
}
# --- END COPY ---


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[0])          # the base template's one blank layout


def summary(prs, th, numbers):
    """executive-summary: two numbered panels under solid header bands, the figures in a near-black band beneath
    that states its basis once. Composed (deckkit.compose)."""
    c = COPY["summary"]
    s = blank(prs)
    tag.tag(s, "executive-summary", "banded-panels-stat-band")
    title(s, c["title"][0], keywords=c["title"][1], canvas=WIDE, theme=th)
    half = (C.w - 0.2) / 2
    ph = 3.95
    for i, (key, glyph_name) in enumerate((("left", "target"), ("right", "list-checks"))):
        head, rows = c[key]
        compose.numbered_rows(s, Box(C.x + i * (half + 0.2), C.y, half, ph), rows, header=head,
                              header_glyph=glyph_name, name=key, theme=th)
    top = C.y + ph + 0.12
    names = compose.stat_band(s, Box(C.x, top, C.w, C.b - top), c["stats"], name="stat", theme=th)
    numbers.update({n: {"value": v, "kind": k} for n, (v, _l, k) in zip(names, c["stats"])})


def platform(prs, th):
    """layered-architecture: the L0 master drawn by the engine (house grammar, l0=True: product marks, no arrows and
    no numbers). Swap the marks for the client's real stack: `icons.vendor_keys()` lists the vendor keys and
    kit/assets/marks/registry.json the brand marks. Gate P26 holds the grammar."""
    c = COPY["platform"]
    s = blank(prs)
    tag.tag(s, "layered-architecture", "l0-house")
    title(s, c["title"][0], keywords=c["title"][1], canvas=WIDE, theme=th)
    d = arch.Diagram("house", l0=True, grid=(1, 2), icon=0.38, head_air=0.12)
    label, nodes = c["sources"]
    d.group("src", "region", label, at=(0, 0, 1, 1), grid=(len(nodes), 1))
    for k, (mark, nm) in enumerate(nodes):
        d.node("src%d" % k, mark, nm, group="src", at=(k, 0), caption="right", size=0.34)
    d.group("plat", "boundary", c["boundary"], at=(0, 1, 1, 1), grid=(4, 3), weights=[1.0, 2.6, 2.6, 1.6],
            row_gap=[0.03, 0.03])
    label, nodes = c["rail"]
    d.group("rail", "rail", label, parent="plat", at=(0, 0, 1, 2), grid=(1, len(nodes)))
    for r, (mark, nm) in enumerate(nodes):
        d.node("rail%d" % r, mark, nm, group="rail", at=(0, r), size=0.34)
    for gid, at in (("ingest", (1, 0, 1, 1)), ("store", (2, 0, 1, 1))):
        label, nodes = c[gid]
        d.group(gid, "region", label, parent="plat", at=at, grid=(2, 2))
        for k, (mark, nm) in enumerate(nodes):
            d.node("%s%d" % (gid, k), mark, nm, group=gid, at=(k % 2, k // 2))
    label, nodes = c["agents"]
    d.group("agents", "region", label, parent="plat", at=(1, 1, 2, 1), grid=(len(nodes), 1))
    for k, (mark, nm) in enumerate(nodes):
        d.node("agents%d" % k, mark, nm, group="agents", at=(k, 0), size=0.62 if k == 0 else None)
    label, nodes = c["serve"]
    d.group("serve", "region", label, parent="plat", at=(3, 0, 1, 2), grid=(1, len(nodes)))
    for r, (mark, nm) in enumerate(nodes):
        d.node("serve%d" % r, mark, nm, group="serve", at=(0, r))
    label, nodes = c["band"]
    d.group("band", "band", label, parent="plat", at=(0, 2, 4, 1), grid=(len(nodes), 1))
    for k, (mark, nm) in enumerate(nodes):
        d.node("band%d" % k, mark, nm, group="band", at=(k, 0), caption="right", size=0.34)
    d.render(s, C, th, name="l0")


def ask(prs, th):
    """the-ask: library device DL01 (decision cards over a dated rail), every slot filled. A device that serves more
    than one archetype is placed with archetype=; `scripts/find.py the-ask` lists the devices with renders."""
    c = COPY["ask"]
    fill = {"slot.title": c["title"], "slot.rail.label": c["rail"]}
    for i, (g, head, body) in enumerate(c["cards"], 1):
        fill.update({"slot.card.%d.num" % i: "%02d" % i, "slot.card.%d.icon" % i: g,
                     "slot.card.%d.head" % i: head, "slot.card.%d.body" % i: body})
    for i, (when, what) in enumerate(c["milestones"], 1):
        fill["slot.milestone.%d" % i] = [when, what]
    devices.place(prs, "DL01", fill, theme=th, archetype="the-ask")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(SKILL / ".cache" / "template" / "template.pptx"))
    ap.add_argument("--theme", default=str(SKILL / "themes" / "sigmoid.json"))
    a = ap.parse_args()
    th = Theme.load(a.theme)
    prs = Presentation(str(SKILL / "kit" / "assets" / "base-wide.pptx"))
    prims.bind(th)
    numbers: dict = {}
    summary(prs, th, numbers)
    platform(prs, th)
    ask(prs, th)
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    sha = package.save(prs, out, title=DECK_TITLE, theme=th, manifest={"numbers": numbers}, overwrite=True,
                       date=DATE)
    dj = out.parent / "deck.json"                 # new_deck.py writes it; a run from the skill writes its own
    if not dj.is_file():
        dj.write_text(json.dumps({
            "type": "deck", "canvas": "wide", "theme": a.theme, "client": None, "spelling": "US", "overrides": {},
            "budgets": {}, "series": [], "deliverable": out.name,
            "build": ["{python}", str(Path(__file__).resolve()), "--theme", a.theme, "--out", "{out}"]},
            indent=2) + "\n", encoding="utf-8")
    print("OK %s sha256 %s" % (out, sha[:12]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
