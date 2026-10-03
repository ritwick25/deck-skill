"""Five flow diagrams drawn by deckkit.arch_flow on the fictional Quillmont brief service (spec 3.6, 3.7, 3.9, 3.11;
ruling 21's agent state machine).

    $PY scripts/example_arch_flow.py [--out .cache/arch_flow/arch_flow.pptx] [--theme themes/sigmoid.json] [--draft]

  1 process-flow:arch-cicd                 one signed image through four gates, a digest rail, a blocked console path
  2 solution-architecture:arch-mlops       a prompt change gated by the golden set, the registry hinged on the zones
  3 layered-architecture:arch-medallion    five sources by stream and batch, bronze to gold, a lineage inset
  4 convergence-hub:arch-events            four event types on one EventBridge bus, rules to targets, a dead letter
  5 decision-logic:arch-state              the brief graph, conditional edges, a human-in-the-loop interrupt

Stack canon: Decks-For-Visuals/acceptance/showcase/canon.json "stack" (environments dev, test, prod; S3 zones raw,
refined, certified; agents on ECS Fargate; Bedrock with Claude; OpenSearch; RDS; Langfuse; GitHub Actions with
CodePipeline). Names agree with scripts/example_arch.py where a component appears in both (house rule 13). Fictional
client and figures; every number on glass is illustrative.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

SKILL = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SKILL / "kit"))

from pptx import Presentation  # noqa: E402

from deckkit import arch_flow as af, package, prims, tag  # noqa: E402
from deckkit.canvas import WIDE  # noqa: E402
from deckkit.text import title  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

C = WIDE.content
DRAFT = False

# one name, role and mark per component (example_arch.py HN agrees where the component recurs)
GITHUB = ("repo", "vendor:onprem/vcs/github", "GitHub", "brief repo")
ACTIONS = ("actions", "vendor:onprem/ci/github-actions", "GitHub Actions", "OIDC to AWS")
ECR = ("ecr", "vendor:aws/compute/ec2-container-registry", "Amazon ECR", "images")
PIPE = ("cp", "vendor:aws/devtools/codepipeline", "CodePipeline", "promotion")
KMS = ("kms", "vendor:aws/security/key-management-service", "AWS KMS", "signing key")
FARGATE = "vendor:aws/compute/fargate"
LANGFUSE = "mark:langfuse-icon"
CW = "vendor:aws/management/cloudwatch"


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[0])


def cicd(prs, th):
    s = blank(prs)
    tag.tag(s, "process-flow", "arch-cicd")
    title(s, "One signed image reaches prod through four gates and is never rebuilt", keywords=("four gates",),
          canvas=WIDE, theme=th)
    spec = {
        "source": {"label": "GitHub · SaaS", "repo": GITHUB, "ci": ACTIONS, "push": "git push"},
        "ci_gate": {"label": "SAST clean", "fail": "build fails"},
        "tooling": {"label": "Tooling account", "fact": "shared", "registry": ECR, "pipeline": PIPE, "signer": KMS,
                    "sign": "sign", "handoff": "digest"},
        "envs": [{"id": "dev", "label": "Dev account", "service": ("dev-ecs", FARGATE, "ECS Fargate", "brief service"),
                  "evidence": ("evals", LANGFUSE, "Langfuse", "golden set"), "evidence_label": "traces",
                  "deploy": "deploy"},
                 {"id": "test", "label": "Test account",
                  "service": ("test-ecs", FARGATE, "ECS Fargate", "brief service")}],
        "gates": [{"kind": "quality", "label": "eval ≥ baseline", "fail": "release held"},
                  {"kind": "approval", "label": "release lead", "fail": "reject"}],
        "prod": {"label": "Prod account", "fact": "2 waves",
                 "waves": [("w1", FARGATE, "ECS Fargate", "wave 1 · 10%"), ("w2", FARGATE, "ECS Fargate",
                                                                            "wave 2 · 100%")],
                 "bake": {"label": "bake 60 min", "fail": "roll back"},
                 "monitor": ("cw", CW, "CloudWatch", "alarms"), "monitor_label": "metrics"},
        "blocked": ("console", "vendor:aws/management/management-console", "AWS Console", "manual deploy"),
        "digest": {"label": "image sha256", "text": "4f9a3c21"},
        "steps": ["Push a commit", "Push the signed image", "Deploy to dev", "Promote to test", "Release wave 1",
                  "Release wave 2"],
        "keys": {"data": "Build or deploy", "feedback": "Fail exit or rollback"},
    }
    return af.cicd(C, th, spec, name="cicd", draft=DRAFT).render(s)


def mlops(prs, th):
    s = blank(prs)
    tag.tag(s, "solution-architecture", "arch-mlops")
    title(s, "A prompt change ships only when the golden set scores at or above baseline",
          keywords=("golden set",), canvas=WIDE, theme=th)
    spec = {
        "build": {"label": "Build and evaluate · dev account", "activity": "edit prompt",
                  "repo": ("gh", "vendor:onprem/vcs/github", "GitHub", "prompt repo"),
                  "runner": ("runner", "vendor:aws/devtools/codebuild", "CodeBuild", "eval runner")},
        "eval": {"label": "Offline eval", "fact": "judge and score",
                 "dataset": ("golden", "vendor:aws/storage/simple-storage-service-s3", "Amazon S3", "golden set"),
                 "dataset_label": "240 cases",
                 "judge": ("judge", "mark:claude-spark", "Claude", "judge model"), "judge_label": "answers",
                 "candidate": ("cand", "vendor:aws/ml/bedrock", "Bedrock", "candidate run"),
                 "candidate_label": "prompt v14",
                 "scores": ("scores", "vendor:aws/database/dynamodb", "DynamoDB", "eval scores")},
        "gate": {"label": "≥ baseline", "revise": "revise"},
        "registry": ("reg", "vendor:aws/ml/bedrock", "Bedrock", "prompt registry"),
        "release": {"label": "Release and monitor · prod account",
                    "pipeline": ("cp", "vendor:aws/devtools/codepipeline", "CodePipeline", "release"),
                    "promote": "promote",
                    "notify": ("sns", "vendor:aws/integration/simple-notification-service-sns", "Amazon SNS",
                               "approval topic"), "notify_label": "approval", "relay_label": "notify",
                    "approver": ("teams", "vendor:saas/chat/teams", "Teams", "approval ask"),
                    "approve_label": "approve", "approval": "release lead",
                    "flag": ("flag", "vendor:aws/management/systems-manager-app-config", "AppConfig", "prompt flag"),
                    "monitor": ("cw", CW, "CloudWatch", "alarms"), "monitor_label": "alarms",
                    "canary": "canary 10%",
                    "service": ("ecs", FARGATE, "ECS Fargate", "writer agent"),
                    "online": ("lf", LANGFUSE, "Langfuse", "online eval"), "online_label": "5% of traces",
                    "loop_label": "failed traces"},
        "steps": ["Commit the prompt", "Score the golden set", "Register the version", "Promote", "Approve and flag",
                  "Serve the canary", "Sample live traces"],
        "keys": {"data": "Build or release", "feedback": "Revise or retrain"},
    }
    return af.mlops(C, th, spec, name="mlops", draft=DRAFT).render(s)


def medallion(prs, th):
    s = blank(prs)
    tag.tag(s, "layered-architecture", "arch-medallion")
    title(s, "Every brief cites certified tables built from five sources, weekly and by stream",
          keywords=("certified tables",), canvas=WIDE, theme=th)
    S3 = "vendor:aws/storage/simple-storage-service-s3"
    spec = {
        "platform": {"label": "Signal platform", "fact": "one AWS account"},
        "sources": [("press", "glyph:megaphone", "Congress media", "abstracts"),
                    ("web", "vendor:aws/general/internet-alt1", "Public web", "press"),
                    ("trials", "glyph:file-search", "Trial feeds", "registries"),
                    ("payers", "glyph:landmark", "Payers", "coverage notices"),
                    ("partners", "glyph:folder-open", "Partner feeds", "file drops")],
        "ingest": {"stream": ("kinesis", "vendor:aws/analytics/kinesis-data-streams", "Kinesis", "press stream"),
                   "batch": ("lambda", "vendor:aws/compute/lambda", "Lambda", "collectors")},
        "stream_hop": ("firehose", "vendor:aws/analytics/kinesis-data-firehose", "Firehose", "60 s batches"),
        "hops": {"stream": "stream", "archive": "archive", "index": "index · 60 s", "append": "append",
                 "refine": "dedupe on key", "certify": "certify"},
        "cells": ["Bronze · raw", "Silver · refined", "Gold · certified"],
        "stores": [("raw", S3, "Amazon S3", "raw zone"), ("refined", S3, "Amazon S3", "refined zone"),
                   ("certified", S3, "Amazon S3", "certified zone")],
        "excluded": {"label": "unlicensed items dropped", "edge": "filter"},
        "serve": {"index": ("search", "vendor:aws/analytics/amazon-opensearch-service", "OpenSearch", "claim index"),
                  "ledger": ("rds", "vendor:aws/database/rds", "RDS", "brief ledger"),
                  "query": ("athena", "vendor:aws/analytics/athena", "Athena", "ad hoc")},
        "serve_labels": {"index": "index claims", "ledger": "nightly", "query": "query"},
        "consume": [("writer", "vendor:aws/compute/elastic-container-service-task", "Writer agent", None),
                    ("review", "vendor:programming/framework/react", "Review app", "React"),
                    ("qs", "vendor:aws/analytics/quicksight", "QuickSight", "signal trends")],
        "consume_labels": ["claims", "review", "trends"],
        "band": {"label": "Processing and governance",
                 "marks": [("glue", "vendor:aws/analytics/glue", "AWS Glue", "jobs"),
                           ("sched", "vendor:aws/integration/eventbridge-scheduler", "EventBridge", "weekly run"),
                           ("textract", "vendor:aws/ml/textract", "Textract", "PDF text"),
                           ("comprehend", "vendor:aws/ml/comprehend", "Comprehend", "entities"),
                           ("kms", "vendor:aws/security/key-management-service", "AWS KMS", "keys"),
                           ("macie", "vendor:aws/security/macie", "Macie", "PII scan"),
                           ("lf", "vendor:aws/analytics/lake-formation", "Lake Formation", "grants"),
                           ("trail", "vendor:aws/management/cloudtrail", "CloudTrail", "audit")]},
        "lineage": {"fact": "certified.signal_claims", "more": "+14 more",
                    "datasets": [("d1", "raw.", "press_items", 0, 1), ("d2", "raw.", "trial_records", 1, 1),
                                 ("d3", "refined.", "claims", 0, 3), ("d4", "refined.", "trials", 1, 3),
                                 ("d5", "certified.", "signal_claims", 0, 4)],
                    "edges": [("d1", "d3", "extract_claims"), ("d2", "d4", "dedupe_trials"), ("d3", "d5", "certify"),
                              ("d4", "d5", "join")]},
        "keys": {"data": "Batch", "async": "Stream"},
    }
    return af.medallion(C, th, spec, name="medal", draft=DRAFT).render(s)


def events(prs, th):
    s = blank(prs)
    tag.tag(s, "convergence-hub", "arch-events")
    title(s, "Four consumers take brief events and failed sends wait in a dead-letter queue",
          keywords=("dead-letter queue",), canvas=WIDE, theme=th)
    TASK = "vendor:aws/compute/elastic-container-service-task"
    SQS = "vendor:aws/integration/simple-queue-service-sqs"
    spec = {
        "bus": {"label": "EventBridge bus", "fact": "brief-events", "mark": "vendor:aws/integration/eventbridge"},
        "schema": ("schema", "vendor:aws/integration/eventbridge-schema", "Schema registry", "brief events v3"),
        "schema_labels": ("validate", "bindings"),
        "producers": [{"node": ("scan", TASK, "Scan agent", "LangGraph"), "events": [("signals.scanned", "scan done")]},
                      {"node": ("writer", TASK, "Writer agent", "LangGraph"),
                       "events": [("brief.drafted", "draft saved"), ("brief.held", "claim unsupported")]},
                      {"node": ("app", "vendor:programming/framework/react", "Review app", "React"),
                       "events": [("brief.approved", "reviewer approves")]}],
        "rules": {"signals.scanned": {"label": "start draft", "target": ("sfn", "vendor:aws/integration/step-functions",
                                                                          "Step Functions", "draft run")},
                  "brief.drafted": {"label": "stream copy",
                                    "target": ("hose", "vendor:aws/analytics/kinesis-data-firehose", "Firehose",
                                               "analytics"),
                                    "then": [(("s3", "vendor:aws/storage/simple-storage-service-s3", "Amazon S3",
                                               "brief archive"), "archive")]},
                  "brief.held": {"label": "queue for review", "target": ("queue", SQS, "Amazon SQS", "review queue")},
                  "brief.approved": {"label": "notify field",
                                     "target": ("notify", "vendor:aws/compute/lambda", "Lambda", "notifier"),
                                     "then": [(("ses", "vendor:aws/engagement/simple-email-service-ses", "Amazon SES",
                                                "Monday email"), "email"),
                                              (("teams", "vendor:saas/chat/teams", "Teams", "field alert"), "alert")]}},
        "dlq": {"event": "brief.approved", "node": ("dlq", SQS, "Amazon SQS", "dead letters"),
                "label": "after 3 retries"},
        "keys": {"data": "Publish", "async": "Rule to target", "control": "Schema", "feedback": "Dead letter"},
    }
    return af.events(C, th, spec, name="events", draft=DRAFT).render(s)


def state(prs, th):
    s = blank(prs)
    tag.tag(s, "decision-logic", "arch-state")
    title(s, "The brief graph stops for a reviewer when a claim has no source",
          keywords=("stops for a reviewer",), canvas=WIDE, theme=th)
    spec = {
        "graph": {"label": "brief_graph", "fact": "on ECS Fargate", "mark": "mark:langgraph-logo"},
        "path": [("start", "terminal", "START", None), ("plan", "state", "plan", "Claude"),
                 ("scan", "state", "scan_signals", "tool calls"), ("draft", "state", "draft_brief", "Claude"),
                 ("check", "state", "check_claims", "cite or flag"), ("publish", "state", "publish", "write ledger"),
                 ("end", "terminal", "END", None)],
        "path_labels": ["request", "plan", "signals", "draft", ("all cited",), "brief id"],
        "branch": [("fail", "terminal", "END", None, 1), ("retry", "state", "retry", "at most 2", 2),
                   ("review", "interrupt", "human_review", "interrupt", 4)],
        "edges": [("check", "review", "unsupported", "conditional"),
                  ("review", "publish", "approve", "conditional"),
                  ("review", "draft", "revise", "conditional"),
                  ("scan", "retry", "tool error", "conditional", None, (-0.125, None)),
                  ("retry", "scan", "retry", "data", None, (None, 0.125)),
                  ("retry", "fail", "2 failures", "conditional")],
        "above": [(("search", "vendor:aws/analytics/amazon-opensearch-service", "OpenSearch", "claim index"), 2, "scan",
                   "search_claims"),
                  (("bedrock", "vendor:aws/ml/bedrock", "Bedrock", "Claude models"), 3, "draft", "model call"),
                  (("rds", "vendor:aws/database/rds", "RDS", "brief ledger"), 5, "publish", "write")],
        "below": [(("ddb", "vendor:aws/database/dynamodb", "DynamoDB", "checkpoints"), 3, "left",
                   ("review", "ddb", "save state", "control", ("l", "t"), (0.125, None))),
                  (("sqs", "vendor:aws/integration/simple-queue-service-sqs", "Amazon SQS", "review queue"), 4, "right",
                   ("review", "sqs", "enqueue", "data", ("b", "t")), "center"),
                  (("app", "vendor:programming/framework/react", "Review app", "React"), 5, "right",
                   ("app", "review", "resume", "data"), "right")],
        "keys": {"data": "Edge", "conditional": "Conditional edge", "control": "Checkpoint"},
    }
    return af.state_graph(C, th, spec, name="graph", draft=DRAFT).render(s)


def vendor_media(deck: Path) -> dict:
    """{sha256: reason} for every embedded official vendor mark (P18 allowance; P22 verifies the bytes)."""
    import hashlib
    import zipfile
    table = json.loads((SKILL / "kit" / "assets" / "vendor_hashes.json").read_text(encoding="utf-8"))["keys"]
    keys = {h: k for k, sizes in table.items() for h in sizes.values()}
    out = {}
    with zipfile.ZipFile(deck) as z:
        for n in sorted(z.namelist()):
            if n.startswith("ppt/media/"):
                h = hashlib.sha256(z.read(n)).hexdigest()
                if h in keys:
                    out[h] = ("official vendor mark %s from the diagrams package, byte-verified by P22 against "
                              "vendor_hashes.json" % keys[h])
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(SKILL / ".cache" / "arch_flow" / "arch_flow.pptx"))
    ap.add_argument("--theme", default=str(SKILL / "themes" / "sigmoid.json"))
    ap.add_argument("--draft", action="store_true")
    a = ap.parse_args()
    global DRAFT
    DRAFT = a.draft
    th = Theme.load(a.theme)
    prs = Presentation(str(SKILL / "kit" / "assets" / "base-wide.pptx"))
    prims.bind(th)
    metrics = {"cicd": cicd(prs, th), "mlops": mlops(prs, th), "medallion": medallion(prs, th),
               "events": events(prs, th), "state": state(prs, th)}
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    sha = package.save(prs, out, title="Flow diagrams", theme=th, manifest={"numbers": {}}, overwrite=True,
                       date="2026-10-04")
    dj = {"type": "example", "canvas": "wide", "theme": a.theme, "client": None, "spelling": "US",
          "media_allow": vendor_media(out), "series": [], "deliverable": out.name,
          "build": ["{python}", str(Path(__file__).resolve()), "--theme", a.theme, "--out", "{out}"]}
    (out.parent / "deck.json").write_text(json.dumps(dj, indent=2) + "\n", encoding="utf-8")
    for k, m in metrics.items():
        print("%-9s %s" % (k, {x: m[x] for x in ("marks", "edges", "labelled", "steps", "crossings", "problems")}))
    print("OK %s sha256 %s" % (out, sha[:12]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
