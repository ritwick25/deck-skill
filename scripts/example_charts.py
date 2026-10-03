"""Five chart slides as real proposal content, the pattern for deckkit.charts (ruling 22): headline figures with a bar
chart, value horizons as a line over an area wash with a crossover annotation, commercials as a waterfall that shows
the arithmetic, a case study before and after, and a prioritisation scatter plotted from scores.

    $PY scripts/example_charts.py [--out .cache/charts/charts.pptx] [--theme themes/sigmoid.json]

Fictional client (the showcase canon: Quillmont, 1,200 reps, six regions, $180K build, $24K a month run); every
figure is illustrative. The manifest records every chart (charts.records), so gate P27 fails a chart whose data or
labels change after the build.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

SKILL = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SKILL / "kit"))

from pptx import Presentation  # noqa: E402

from deckkit import charts, compose, package, prims, tag  # noqa: E402
from deckkit.canvas import WIDE, Box  # noqa: E402
from deckkit.text import Para, Run, need_height, textbox, title  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

C = WIDE.content
GAP = 0.40


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[0])


def days(v):
    return "%g day%s" % (v, "" if v == 1 else "s")


def heading(s, box, text, th, name):
    """A chart's heading: what is plotted and in what unit (a label of the exhibit, not a subtitle)."""
    return textbox(s, box, [Para((Run(text, "ink-deep", bold=True, size=13),))], name, th, anchor="top")


def lines(s, box, items, th, name, *, head_pt=13.0, body_pt=11.0, background="paper", gap=8.0):
    """[(head, body)] stacked in one text box: a bold lead and one or two lines each."""
    paras = []
    for k, (h, b) in enumerate(items):
        paras.append(Para((Run(h, "ink-deep", bold=True, size=head_pt),), space_before=gap if k else 0.0))
        if b:
            paras.append(Para((Run(b, "ink", size=body_pt),), space_before=2.0))
    return textbox(s, box, paras, name, th, anchor="top", background=background)


def text_h(items, w, th, *, head_pt=13.0, body_pt=11.0, gap=8.0):
    paras = []
    for k, (h, b) in enumerate(items):
        paras.append(Para((Run(h, "ink-deep", bold=True, size=head_pt),), space_before=gap if k else 0.0))
        if b:
            paras.append(Para((Run(b, "ink", size=body_pt),), space_before=2.0))
    return need_height(paras, w, th)


# ------------------------------------------------------------------ 1 headline figures: bar chart

def headline(prs, th, numbers):
    s = blank(prs)
    tag.tag(s, "headline-figures", "stat-band-stage-bars")
    title(s, "A competitor move reaches reps in five days, and rewriting takes two", keywords=("five days",),
          canvas=WIDE, theme=th)
    band_h = 1.55
    names = compose.stat_band(s, Box(C.x, C.y, C.w, band_h), [
        ("5 days", "from a competitor move to the field", "illustrative"),
        ("40+", "public competitor signals a week", "illustrative"),
        ("1,200", "reps who hear it from customers first", "illustrative"),
        ("6", "regional versions of every brief", "illustrative")], name="stat", theme=th)
    numbers.update({n: {"value": v, "kind": "illustrative"} for n, v in zip(names, ("5 days", "40+", "1,200", "6"))})
    top = C.y + band_h + 0.30
    left = Box(C.x, top, 7.55, C.b - top)
    rail = Box(left.r + GAP, top, C.r - left.r - GAP, C.b - top)
    heading(s, Box(left.x, left.y, left.w, 0.30), "Days from a competitor move to the field, by stage", th, "bars.head")
    plot = charts.bar(s, Box(left.x, left.y + 0.40, left.w, left.h - 0.40), ["Collect signals", "Rewrite by region",
                      "Medical review", "Send to reps"], [1, 2, 1, 1], subject="Rewrite by region", fmt=days,
                      kind="illustrative", name="bars", theme=th)
    charts.annotate(plot, "Rewrite by region", "Six versions of one brief", side="right")
    items = [("Collect: 1 day", "Analysts sift 40+ signals by hand from registries, press and congress"),
             ("Rewrite: 2 days", "Each of six regions rewrites the same brief for its own reps"),
             ("Review: 1 day", "Medical reads every regional version, one after another"),
             ("Send: 1 day", "The Monday digest holds every brief to one weekly slot")]
    heading(s, Box(rail.x, rail.y, rail.w, 0.30), "Where each day goes", th, "why.head")
    body = Box(rail.x, rail.y + 0.40, rail.w, rail.h - 0.40)
    gap = 8.0
    h = text_h(items, body.w, th, gap=gap)
    gap += max(0.0, (body.h - h) / (len(items) - 1) * 72 - 2)       # spread the four stages down the rail
    lines(s, Box(body.x, body.y, body.w, text_h(items, body.w, th, gap=gap)), items, th, "why.body", gap=gap)


# ------------------------------------------------------------------ 2 value horizons: line, area wash, crossover

def horizons(prs, th):
    s = blank(prs)
    tag.tag(s, "value-horizons", "reach-line-horizon-cards")
    title(s, "Every rep is briefed within a day once the six regions go live", keywords=("within a day",),
          canvas=WIDE, theme=th)
    left = Box(C.x, C.y, 7.95, C.h)
    rail = Box(left.r + GAP, C.y, C.r - left.r - GAP, C.h)
    heading(s, Box(left.x, left.y, 4.0, 0.30), "Reps briefed within a day", th, "reach.head")
    plot = charts.area(s, Box(left.x, left.y + 0.40, left.w, left.h - 0.40),
                       ["Today", "Pilot, Dec 2026", "Rollout, Jan 2027", "Brand three, Sep 2027"], [0, 200, 1200, 1200],
                       series_name="Reps", fmt="{:,.0f}", kind="illustrative", name="reach", theme=th)
    charts.annotate(plot, "Rollout, Jan 2027", "1,200 reps, six regions", kind="crossover", side="above")
    cards = [("Pilot, North region", "One brand, from Dec 2026", "Prove the brief in one region first",
              "Gate: medical signs off two weeks running"),
             ("Rollout, six regions", "All 1,200 reps from 22 Jan 2027", "One brief replaces six regional versions",
              "Gate: every region reads the one brief"),
             ("Run, three brands", "The same reps by Sep 2027", "More brands on one service, no new team",
              "Gate: three brands through one review queue")]
    built = []
    for head, fact, line, gate in cards:
        paras = [Para((Run(head, "ink-deep", bold=True, size=13),)),
                 Para((Run(fact, "ink", bold=True, size=11),), space_before=4.0),
                 Para((Run(line, "ink", size=11),), space_before=2.0),
                 Para((Run(gate, "ink-muted", size=11),), space_before=6.0)]
        built.append((paras, need_height(paras, rail.w - 0.92, th)))
    pad = 0.26
    gap = (rail.h - sum(h + 2 * pad for _p, h in built)) / (len(built) - 1)   # cards hug their text, air between
    y = rail.y
    for i, (paras, h) in enumerate(built, 1):
        box = Box(rail.x, y, rail.w, h + 2 * pad)
        prims.card(s, box, fill="card", line=None, name="horizon.%d.card" % i, theme=th)
        prims.disc(s, Box(box.x + 0.20, box.y + pad, 0.40, 0.40), i, fill="accent" if i == 2 else "ink-strong",
                   name="horizon.%d.disc" % i, theme=th)
        textbox(s, Box(box.x + 0.74, box.y + pad, rail.w - 0.92, h), paras, "horizon.%d.text" % i, th,
                background="card")
        y = box.b + gap


# ------------------------------------------------------------------ 3 commercials: waterfall that shows the sum

def commercials(prs, th, numbers):
    from deckkit.read.colour import contrast
    s = blank(prs)
    tag.tag(s, "commercials", "price-bridge-panel")
    title(s, "One first-year price of $468K covers the build and a year of run", keywords=("$468K",),
          canvas=WIDE, theme=th)
    left = Box(C.x, C.y, 8.10, C.h)
    col = Box(left.r + GAP, C.y, C.r - left.r - GAP, C.h)
    heading(s, Box(left.x, left.y, left.w, 0.30), "First-year price, $K", th, "price.head")
    plot = charts.waterfall(s, Box(left.x, left.y + 0.40, left.w, left.h - 0.40),
                            [("Build, three gates", 180), ("Run, 12 x $24K", 288)], total="First-year total",
                            fmt="${:,.0f}K", kind="illustrative", name="price", theme=th)
    charts.annotate(plot, "Run, 12 x $24K", "From rollout, 22 Jan 2027", side="below")
    num = "accent-ink" if contrast(th.hex("accent-ink"), th.hex("ink-deep")) >= 3.0 else "accent-bright"
    terms = [("Build: $180K fixed", "Three gates, paid at each sign-off"),
             ("Run: $24K a month", "From 22 Jan 2027, all six regions")]
    paras = []
    for k, (h, b) in enumerate(terms):
        paras.append(Para((Run(h, "paper", bold=True, size=13),), space_before=10.0 if k else 0.0))
        paras.append(Para((Run(b, "paper", size=11),), space_before=2.0))
    inner_w = col.w - 0.72
    sh = prims.stat_height(s, inner_w, "$468K", "first-year price, build and run", kind="illustrative", basis=False,
                           theme=th)
    th_terms = need_height(paras, inner_w, th)
    panel = Box(col.x, col.y, col.w, 0.36 + sh + 0.30 + th_terms + 0.36)
    prims.block(s, panel, fill="ink-deep", name="panel", theme=th)
    prims.stat(s, Box(panel.x + 0.36, panel.y + 0.36, inner_w, sh), "$468K", "first-year price, build and run",
               kind="illustrative", role=num, background="ink-deep", label_role="paper", basis=False,
               name="price.stat", theme=th)
    numbers["price.stat"] = {"value": "$468K", "kind": "illustrative"}
    textbox(s, Box(panel.x + 0.36, panel.b - 0.36 - th_terms, inner_w, th_terms), paras, "panel.terms", th,
            background="ink-deep")
    buys = [("Six firm roles, three client roles", "One team from kickoff on 2 Nov 2026"),
            ("Three gates before rollout", "20 Nov, 18 Dec 2026 and 8 Jan 2027"),
            ("One weekly brief, six regions", "Sourced claims, medical review in the flow")]
    rest = Box(col.x, panel.b + 0.34, col.w, col.b - panel.b - 0.34)
    heading(s, Box(rest.x, rest.y, rest.w, 0.30), "What the price buys", th, "buys.head")
    body = Box(rest.x, rest.y + 0.42, rest.w, rest.h - 0.42)
    gap = 8.0 + max(0.0, (body.h - text_h(buys, body.w, th)) / (len(buys) - 1) * 72 - 2)
    lines(s, Box(body.x, body.y, body.w, text_h(buys, body.w, th, gap=gap)), buys, th, "buys.body", gap=gap)


# ------------------------------------------------------------------ 4 case study: before and after

def case(prs, th):
    s = blank(prs)
    tag.tag(s, "case-study", "need-build-before-after")
    title(s, "A specialty pharma field team cut brief lead time from four days to one",
          keywords=("four days to one",), canvas=WIDE, theme=th)
    cols = C.split_cols(3, GAP, weights=(1, 1, 1.2))
    compose.numbered_rows(s, cols[0], [
        ("Briefs reached reps late", "Four days from a competitor move to the field"),
        ("Eight regional versions", "Every region rewrote the same weekly brief"),
        ("Review came last", "Medical read briefs after reps had used them"),
        ("No source on claims", "Reps could not show where a claim came from")],
        header="The need", header_glyph="clipboard-check", header_fill="ink-strong", disc_fill="ink-strong",
        name="need", theme=th)
    compose.numbered_rows(s, cols[1], [
        ("One sourced brief", "Every claim cites its source, on one page"),
        ("Review in the flow", "Medical approves before the brief goes out"),
        ("Regions add one line", "Local notes ride on the shared brief"),
        ("Ten weeks, three gates", "Pilot in one region, then all eight")],
        header="What was built", header_glyph="layers", header_fill="ink-strong", disc_fill="ink-strong",
        name="build", theme=th)
    out = cols[2]
    hb = 0.42
    prims.block(s, Box(out.x, out.y, out.w, hb), fill="accent", radius_in=0, name="after.header", theme=th)
    from deckkit.glyphs import glyph
    glyph(s, "target", Box(out.x + 0.18, out.y + hb * 0.2, hb * 0.6, hb * 0.6), role="on-accent",
          shape_name="after.header.icon", theme=th)
    textbox(s, Box(out.x + 0.18 + hb * 0.6 + 0.12, out.y, out.w - 0.6 - hb * 0.6, hb),
            [Para((Run("The outcome", "on-accent", bold=True, size=12),))], "after.header.text", th,
            background="accent", anchor="middle")
    heading(s, Box(out.x, out.y + hb + 0.22, out.w, 0.30), "Days from a move to the field", th, "after.head")
    items = [("One brief, not eight", "The same reps, briefed the day after a move"),
             ("Every claim sourced", "Medical signs off before the brief goes out")]
    h = text_h(items, out.w, th)
    top = out.y + hb + 0.62
    plot = charts.column(s, Box(out.x, top, out.w, out.b - top - h - 0.30), ["Before", "After"], [4, 1],
                         subject="After", fmt=days, kind="illustrative", name="after", theme=th)
    charts.annotate(plot, "After", "One shared brief", side="above")
    lines(s, Box(out.x, out.b - h, out.w, h), items, th, "after.text")


# ------------------------------------------------------------------ 5 prioritisation: scatter from scores

CAPS = [  # (name, what it does, effort 1-5, value 1-5, wave)
    ("Signal intake", "registries, press, congress", 1.5, 4.6, "Wave 1"),
    ("Source ledger", "every claim to its source", 2.0, 4.0, "Wave 1"),
    ("Brief writer", "one page a week", 2.6, 4.4, "Wave 1"),
    ("Review queue", "medical sign-off", 1.2, 3.5, "Wave 1"),
    ("Regional notes", "one local line each", 3.4, 4.2, "Wave 2"),
    ("Field app card", "the brief in the app", 2.4, 2.6, "Wave 2"),
    ("Payer feeds", "formulary moves", 4.2, 3.6, "Wave 2"),
    ("Congress tracker", "abstracts and posters", 3.6, 3.2, "Wave 2"),
    ("Rep questions", "answers from the brief", 4.5, 2.4, "Later"),
    ("Brand dashboards", "trends by brand", 3.2, 1.8, "Later"),
    ("Archive search", "past briefs", 1.8, 1.5, "Later"),
    ("Translation", "local languages", 4.6, 1.4, "Later")]


def priorities(prs, th):
    s = blank(prs)
    tag.tag(s, "prioritization-matrix", "value-effort-scatter-key")
    title(s, "Four capabilities give most of the value for the least effort", keywords=("Four capabilities",),
          canvas=WIDE, theme=th)
    left = Box(C.x, C.y, 7.45, C.h)
    rail = Box(left.r + GAP, C.y, C.r - left.r - GAP, C.h)
    charts.scatter(s, left, [(n, e, v) for n, _d, e, v, _w in CAPS], x_range=(1, 5), y_range=(1, 5),
                   x_label="Effort to build", y_label="Value to the field, 1 to 5",
                   groups={n: w for n, _d, _e, _v, w in CAPS}, subject="Wave 1", splits=(3, 3),
                   quadrants=("Build first", "Plan for", "Fill in", "Hold"), numbered=True,
                   kind="illustrative", name="prio", theme=th)
    waves = [("Wave 1", "by Jan 2027", "accent", "Highest value for the least effort"),
             ("Wave 2", "by Sep 2027", "series-1", "Regional depth and new sources"),
             ("Later", "after Sep 2027", "rule-strong", "Low value for the effort today")]
    blocks = []
    for wave, when, role, why in waves:
        members = [(k + 1, n, d) for k, (n, d, _e, _v, w) in enumerate(CAPS) if w == wave]
        paras = [Para((Run("%s, %s" % (wave, when), "ink-deep", bold=True, size=13),)),
                 Para((Run(why, "ink-muted", size=11),), space_before=2.0)]
        paras += [Para((Run("%d  %s" % (k, n), "ink-deep", bold=True, size=11), Run(", %s" % d, "ink", size=11)),
                       space_before=3.0 if j else 5.0) for j, (k, n, d) in enumerate(members)]
        blocks.append((role, paras, need_height(paras, rail.w - 0.28, th)))
    gap = (rail.h - sum(h for _r, _p, h in blocks)) / (len(blocks) - 1)
    y = rail.y
    for i, (role, paras, h) in enumerate(blocks, 1):
        prims.rect(s, Box(rail.x, y + 0.06, 0.14, 0.14), fill=role, name="key.%d.swatch" % i, theme=th)
        textbox(s, Box(rail.x + 0.28, y, rail.w - 0.28, h), paras, "key.%d.text" % i, th, anchor="top")
        y += h + gap


def build(theme_path: str):
    th = Theme.load(theme_path)
    prs = Presentation(str(SKILL / "kit" / "assets" / "base-wide.pptx"))
    prims.bind(th)
    numbers: dict = {}
    headline(prs, th, numbers)
    horizons(prs, th)
    commercials(prs, th, numbers)
    case(prs, th)
    priorities(prs, th)
    return prs, th, numbers


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(SKILL / ".cache" / "charts" / "charts.pptx"))
    ap.add_argument("--theme", default=str(SKILL / "themes" / "sigmoid.json"))
    a = ap.parse_args()
    a.theme = str(Path(a.theme).resolve())          # deck.json names it from any folder
    prs, th, numbers = build(a.theme)
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    sha = package.save(prs, out, title="Native charts", theme=th,
                       manifest={"numbers": numbers, "charts": charts.records(prs)}, overwrite=True,
                       date="2026-10-04")
    dj = {"type": "example", "canvas": "wide", "theme": a.theme, "client": None, "spelling": "US", "overrides": {},
          "series": [], "deliverable": out.name,
          "build": ["{python}", str(Path(__file__).resolve()), "--theme", a.theme, "--out", "{out}"]}
    (out.parent / "deck.json").write_text(json.dumps(dj, indent=2) + "\n", encoding="utf-8")
    print("OK %s sha256 %s" % (out, sha[:12]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
