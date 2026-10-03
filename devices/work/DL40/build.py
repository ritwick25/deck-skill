"""DL40: the firm in facts. A solid brand band of five glyph-led scale figures; below it a workforce ring split by
discipline beside a near-black panel of delivery figures, a three-column delivery footprint on the right, and a
labelled strip of partner platform slots along the foot (credentials)."""
import argparse
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "scripts"))
from devwork import *  # noqa: E402,F403
from deckkit import prims  # noqa: E402
from deckkit import text as txt  # noqa: E402
from deckkit.text import Para, Run  # noqa: E402

ap = argparse.ArgumentParser()
ap.add_argument("--out")
a = ap.parse_args()

prs, s, th = open_lifted("DL40")
VARIANT = "scale-band-ring-footprint"
NUMBERS = {}
KIND = "illustrative"

# every lorem text the build redraws as a kit stat or a measured text box, the empty duplicates behind them,
# the wide empty label and the full-width bar under the region heads, and the dashed rules inside the side
# panel (each stat now carries its own hairline)
for n in list(range(14, 29)) + [34, 35, 36, 37, 38] + list(range(42, 53)):
    try:
        delete(s, "dev.%d" % n)
    except KeyError:                       # a member of a group already removed with its group
        pass


def tb(name, box, paras, bg="paper", anchor="top"):
    return txt.textbox(s, box, paras, name, th, background=bg, anchor=anchor)


def P(t, role="ink", size=11.0, bold=False, align="left", before=0):
    return Para((Run(t, role, bold=bold, size=size),), align=align, space_before=before)


def stat(name, box, value, label, *, role="accent-ink", label_role="ink", bg="paper", pt=20, label_pt=11,
         centre=False):
    g = prims.stat(s, box, value, label, kind=KIND, role=role, pt=pt, label_pt=label_pt, label_role=label_role,
                   background=bg, basis=False, name=name, theme=th)
    if centre:
        for p_ in g._element.iter(qn("a:pPr")):
            p_.set("algn", "ctr")
    NUMBERS[name] = {"value": value, "kind": KIND}
    return g


# title: one claim, two-tone, at the kit title geometry
rename(s, "dev.2", "slot.title")
move(s, "slot.title", 0.50, 0.13, 12.40, 0.40)
text(s, "slot.title", [Para(tuple(txt.title_runs(
    "A data and AI engineering partner with twelve years of enterprise delivery",
    ("twelve years of enterprise delivery",))))], th, size=21, bold=True, anchor="middle")

# the one basis label every figure on the slide shares (P23), where the source's partner badge sat
tb("slot.basis", Box(9.90, 0.64, 3.00, 0.24), [P(prims.basis_label(KIND), "ink-muted", 9, align="right")])

# the scale band: solid brand, five glyph-led figures between dashed dividers
BAND_Y, BAND_H = 0.94, 1.06
rename(s, "dev.9", "struct.band")
prims.set_fill(find(s, "struct.band"), "accent", theme=th)
move(s, "struct.band", 0.50, BAND_Y, 12.40, BAND_H)
PITCH = 12.40 / 5
BAND = [("users", "40+", "Enterprise clients"), ("repeat", "85%", "Repeat business"),
        ("award", "4.7/5", "Client satisfaction"), ("chart-column-increasing", "$90M", "Annual revenue"),
        ("trending-up", "4x", "Growth in five years")]
SH = max(prims.stat_height(s, 1.50, v, lab, kind=KIND, pt=22, label_pt=12, basis=False, theme=th)
         for _g, v, lab in BAND)
for i, (g, v, lab) in enumerate(BAND, 1):
    x = 0.50 + (i - 1) * PITCH
    icon(s, "slot.band.%d.icon" % i, g, x + 0.24, BAND_Y + (BAND_H - 0.50) / 2, 0.50, th, role="on-accent")
    stat("slot.band.%d" % i, Box(x + 0.88, BAND_Y + (BAND_H - SH) / 2, 1.50, SH), v, lab, role="on-accent",
         label_role="on-accent", bg="accent", pt=22, label_pt=12)
for k, n in enumerate(("dev.10", "dev.11", "dev.12", "dev.13"), 1):
    rename(s, n, "struct.band.rule.%d" % k)
    move(s, "struct.band.rule.%d" % k, 0.50 + k * PITCH, BAND_Y + 0.22, 0, BAND_H - 0.44)

# section heads, each with its rule running to the zone edge
HEAD_Y = 2.24
rename(s, "dev.3", "slot.left.head")
move(s, "slot.left.head", 0.50, HEAD_Y, 2.10, 0.32)
text(s, "slot.left.head", "Depth by discipline", th, size=14, bold=True, role="ink-deep", anchor="middle")
rename(s, "dev.29", "struct.left.rule")
move(s, "struct.left.rule", 2.66, HEAD_Y + 0.16, 6.80 - 2.66, 0)
rename(s, "dev.30", "slot.right.head")
move(s, "slot.right.head", 7.32, HEAD_Y, 2.10, 0.32)
text(s, "slot.right.head", "Delivery footprint", th, size=14, bold=True, role="ink-deep", anchor="middle")
rename(s, "dev.31", "struct.right.rule")
move(s, "struct.right.rule", 9.40, HEAD_Y + 0.16, 12.90 - 9.40, 0)
rename(s, "dev.5", "struct.divider")
move(s, "struct.divider", 7.06, HEAD_Y, 0, 6.02 - HEAD_Y)

# the workforce ring, drawn where the source's chart picture sat: four arcs by discipline, the total in its hole
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
CX, CY, D, THICK = 2.02, 4.36, 2.40, 0.40
RING = [("Strategy consultants", 250, "accent-deep"), ("AI engineers", 250, "ink-strong"),
        ("Data scientists", 300, "rule-strong"), ("Data engineers", 700, "accent")]
total = sum(n for _l, n, _r in RING)
start = 0.0
for i, (_lab, n, role) in enumerate(RING, 1):
    sweep = 360.0 * n / total
    sh = s.shapes.add_shape(MSO_SHAPE.BLOCK_ARC, int((CX - D / 2) * EMU), int((CY - D / 2) * EMU),
                            int(D * EMU), int(D * EMU))
    sh.name = "struct.ring.%d" % i
    a0, a1 = (270.0 + start) % 360, (270.0 + start + sweep) % 360    # OOXML angles run clockwise from 3 o'clock
    # python-pptx scales an adjustment by 100000 and OOXML angles are 60000ths of a degree, so degrees x 0.6
    sh.adjustments[0], sh.adjustments[1], sh.adjustments[2] = a0 * 0.6, a1 * 0.6, THICK / D
    prims.set_fill(sh, role, theme=th)
    prims.set_line(sh, "paper", 1.5, theme=th)
    if sh.has_text_frame:
        sh.text_frame.text = ""
    start += sweep
HOLE = D - 2 * THICK - 0.06              # a paper disc fills the hole, so the total reads on paper (P25)
prims.circle(s, Box(CX - HOLE / 2, CY - HOLE / 2, HOLE, HOLE), fill="paper", name="struct.ring.hole", theme=th)
CH = prims.stat_height(s, 1.20, "1,500+", "Data and AI specialists", kind=KIND, pt=22, label_pt=10, basis=False,
                       theme=th)
stat("slot.ring.total", Box(CX - 0.60, CY - CH / 2, 1.20, CH), "1,500+", "Data and AI specialists", pt=22,
     label_pt=10, centre=True)
LBL = {  # each discipline's figure beside its own arc
    1: Box(3.06, 2.66, 1.62, 0.66),       # top right
    2: Box(3.36, 4.00, 1.40, 0.66),       # right
    3: Box(3.20, 5.22, 1.50, 0.66),       # lower right
    4: Box(0.50, 2.66, 1.20, 0.66),       # top left, over the long arc
}
for i, (lab, n, _r) in enumerate(RING, 1):
    stat("slot.ring.%d" % i, LBL[i], "%d+" % n, lab, pt=18, label_pt=10)

# the near-black panel: four delivery figures stacked
PX, PY, PW, PH = 4.86, 2.66, 1.94, 3.38
rename(s, "dev.4", "struct.panel")
prims.set_fill(find(s, "struct.panel"), "ink-deep", theme=th)
prims.set_line(find(s, "struct.panel"), None, theme=th)
move(s, "struct.panel", PX, PY, PW, PH)
PANEL = [("24x7", "Always-on support"), ("60+", "AI products live"),
         ("25+", "Data platforms built"), ("$400M", "Revenue influenced")]
ph = max(prims.stat_height(s, PW - 0.36, v, lab, kind=KIND, pt=20, label_pt=10, basis=False, theme=th)
         for v, lab in PANEL)
gap = (PH - 0.36 - 4 * ph) / 3
assert gap > 0.12, "panel stats need %.2fin more height" % (0.12 - gap)
for i, (v, lab) in enumerate(PANEL, 1):
    stat("slot.panel.%d" % i, Box(PX + 0.18, PY + 0.18 + (i - 1) * (ph + gap), PW - 0.36, ph), v, lab,
         role="accent-bright", label_role="paper", bg="ink-deep", pt=20, label_pt=10)

# the delivery footprint: three columns, each a glyph, a head, a figure and one line
COLW, CGAP = 1.62, 0.32
FOOT = [("building", "Onsite", "180+", "Leads and architects", "Discovery, design authority and adoption"),
        ("clock", "Nearshore", "320+", "Engineers in shared hours", "Daily overlap with product owners"),
        ("globe", "Offshore", "1,000+", "Engineers in two centers", "Build, test and run around the clock")]
FH = max(prims.stat_height(s, COLW, f[2], f[3], kind=KIND, pt=20, label_pt=11, basis=False, theme=th) for f in FOOT)
for i, (g, head, v, lab, body) in enumerate(FOOT, 1):
    x = 7.32 + (i - 1) * (COLW + CGAP)
    icon(s, "slot.foot.%d.icon" % i, g, x, 2.74, 0.48, th, role="accent")
    tb("slot.foot.%d.head" % i, Box(x, 3.34, COLW, 0.34), [P(head, "accent-ink", 16, True)], anchor="middle")
    stat("slot.foot.%d" % i, Box(x, 3.86, COLW, FH), v, lab, pt=20, label_pt=11)
    tb("slot.foot.%d.body" % i, Box(x, 3.86 + FH + 0.18, COLW, 0.62), [P(body, "ink", 11)])
for k, n in enumerate(("dev.32", "dev.33"), 1):
    rename(s, n, "struct.foot.rule.%d" % k)
    move(s, "struct.foot.rule.%d" % k, 7.32 + k * (COLW + CGAP) - CGAP / 2, 2.74, 0, 6.02 - 2.74)

# the partner strip along the foot: a label, then six platform slots (a glyph a builder swaps for a vendor: mark)
rename(s, "dev.7", "struct.partners.rule")
move(s, "struct.partners.rule", 0.50, 6.13, 12.40, 0)
rename(s, "dev.6", "slot.partners.head")
move(s, "slot.partners.head", 0.50, 6.18, 2.10, 0.66)
text(s, "slot.partners.head", [P("Technology partners", "ink-deep", 14, True),
                               P("Six certified platforms", "ink", 11)], th, size=14, anchor="middle")
rename(s, "dev.8", "struct.partners.divider")
move(s, "struct.partners.divider", 2.72, 6.22, 0, 0.50)
PARTNERS = [("cloud", "Cloud platforms"), ("layers", "Lakehouse"), ("warehouse", "Cloud warehouse"),
            ("sparkles", "Model providers"), ("workflow", "Data modeling"), ("activity", "Streaming")]
PP = (12.90 - 2.92) / 6
for i, (g, lab) in enumerate(PARTNERS, 1):
    x = 2.92 + (i - 1) * PP
    icon(s, "slot.partner.%d.icon" % i, g, x, 6.25, 0.42, th, role="ink-strong")
    tb("slot.partner.%d.label" % i, Box(x + 0.50, 6.25, PP - 0.58, 0.42), [P(lab, "ink", 11)], anchor="middle")

from assemble_devices import finish  # noqa: E402  content band (P03), small tints to card (P16)
from deckkit import package, tag  # noqa: E402
finish(prs.slides[0], th)
tag.tag(prs.slides[0], "credentials", VARIANT)
for sh in prs.slides[0].shapes:
    ooxml.finish(sh)
out = WORK / "DL40" / "device.pptx"
package.save(prs, out, title="Device DL40", theme=th, manifest={"numbers": NUMBERS}, overwrite=True,
             date="2026-10-03")
if a.out:
    shutil.copyfile(out, a.out)
