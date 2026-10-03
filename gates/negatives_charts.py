"""Planters for P27 charts and for P12's chart-part schema pass (track C). Each adds one kit chart slide to the golden
copy (deckkit.charts.bar, recorded in the manifest's "charts" exactly as a build records it), so the chart passes
P27 before the plant, then breaks exactly one thing in the chart part or around it.

Discovered by run_all.py --negatives (every gates/negatives_*.py joins the list). Runnable alone, which rebuilds the
golden fixture when it is missing, proves the planted chart is clean, and runs only these planters:

    $PY gates/negatives_charts.py [--keep]
"""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path

GATES_DIR = Path(__file__).resolve().parent
if str(GATES_DIR) not in sys.path:
    sys.path.insert(0, str(GATES_DIR))

from negatives import EMU, Planter, edit_part, edit_slide, members, shape, theme, write  # noqa: E402

from lxml import etree  # noqa: E402

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
CNS = C[1:-1]
ANS = A[1:-1]
NAME = "neg.chart"


def _el(xml: str):
    return etree.fromstring('<w xmlns:c="%s" xmlns:a="%s">%s</w>' % (CNS, ANS, xml))[0]


def _record(p, prs):
    from deckkit import charts
    mp = Path(p).with_name(Path(p).stem + ".manifest.json")
    man = json.loads(mp.read_text(encoding="utf-8")) if mp.is_file() else {}
    man["charts"] = charts.records(prs)
    mp.write_text(json.dumps(man, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def add_chart(p, how="bar"):
    """A clean kit chart on a new last slide, recorded in the manifest. Returns (slide index, chart part name)."""
    from pptx import Presentation
    from deckkit import charts, prims
    from deckkit.canvas import Box
    th = theme()
    prims.bind(th)
    prs = Presentation(str(p))
    s = prs.slides.add_slide(prs.slide_layouts[0])
    if how == "ring":
        from deckkit.text import Para, Run, textbox
        textbox(s, Box(6.0, 1.0, 2.0, 0.3), [Para((Run("Illustrative figures", "ink-muted", size=9),))], "neg.basis",
                th)
        charts.ring(s, Box(0.5, 1.0, 3.0, 3.0), 3, 5, kind="illustrative", name=NAME, theme=th)
    else:
        charts.bar(s, Box(0.5, 1.0, 8.0, 3.5), ["Collect", "Rewrite", "Review", "Send"], [1, 2, 1, 1],
                   subject="Rewrite", fmt="{:g} days", kind="illustrative", name=NAME, theme=th)
    prs.save(str(p))
    _record(p, prs)
    n = len(prs.slides)
    part = sorted(k for k in members(p) if k.startswith("ppt/charts/chart") and k.endswith(".xml"))[-1]
    return n, part


def _chart(p, fn, how="bar"):
    _i, part = add_chart(p, how)
    edit_part(p, part, fn)


def _ser(root):
    return root.find(".//" + C + "ser")


def _set_fill(sppr, xml):
    sf = sppr.find(A + "solidFill")
    sf.getparent().replace(sf, _el(xml))


# ------------------------------------------------------------------ planters

def c_6pt(p, w):
    _chart(p, lambda r: r.find(C + "txPr").find(".//" + A + "defRPr").set("sz", "600"))


def c_no_size(p, w):
    def fn(r):
        d = r.find(C + "txPr").find(".//" + A + "defRPr")
        del d.attrib["sz"]
    _chart(p, fn)


def c_srgb(p, w):
    _chart(p, lambda r: _set_fill(_ser(r).find(C + "spPr"), '<a:solidFill><a:srgbClr val="FF00FF"/></a:solidFill>'))


def c_status(p, w):
    _chart(p, lambda r: _set_fill(_ser(r).find(C + "spPr"), '<a:solidFill><a:srgbClr val="F38B1F"/></a:solidFill>'))


def c_3d(p, w):
    def fn(r):
        r.find(C + "chart/" + C + "autoTitleDeleted").addnext(_el('<c:view3D><c:rotX val="15"/></c:view3D>'))
    _chart(p, fn)


def c_gradient(p, w):
    _chart(p, lambda r: _set_fill(_ser(r).find(C + "spPr"), '<a:gradFill><a:gsLst><a:gs pos="0"><a:schemeClr '
                                  'val="accent2"/></a:gs><a:gs pos="100000"><a:schemeClr val="bg1"/></a:gs></a:gsLst>'
                                  '</a:gradFill>'))


def c_shadow(p, w):
    def fn(r):
        sp = _ser(r).find(C + "spPr")
        sp.append(_el('<a:effectLst><a:outerShdw blurRad="50800" dist="38100"><a:schemeClr val="tx1"/></a:outerShdw>'
                      '</a:effectLst>'))
    _chart(p, fn)


def c_style(p, w):
    _chart(p, lambda r: r.find(C + "roundedCorners").addnext(_el('<c:style val="34"/>')))


def c_rounded(p, w):
    _chart(p, lambda r: r.find(C + "roundedCorners").set("val", "1"))


def c_title(p, w):
    def fn(r):
        r.find(C + "chart").insert(0, _el('<c:title><c:tx><c:rich><a:bodyPr/><a:p><a:pPr><a:defRPr sz="1400"/></a:pPr>'
                                          '<a:r><a:rPr sz="1400"/><a:t>Days by stage</a:t></a:r></a:p></c:rich></c:tx>'
                                          '<c:overlay val="0"/></c:title>'))
    _chart(p, fn)


def c_legend(p, w):
    _chart(p, lambda r: r.find(C + "chart/" + C + "plotArea").addnext(
        _el('<c:legend><c:legendPos val="r"/><c:overlay val="0"/></c:legend>')))


def c_unpinned(p, w):
    def fn(r):
        lay = r.find(".//" + C + "plotArea/" + C + "layout")
        lay.remove(lay.find(C + "manualLayout"))
    _chart(p, fn)


def _dpt(idx, clr):
    return _el('<c:dPt><c:idx val="%d"/><c:invertIfNegative val="0"/><c:bubble3D val="0"/><c:spPr><a:solidFill>%s'
               '</a:solidFill></c:spPr></c:dPt>' % (idx, clr))


def _put_dpt(ser, d):
    for old in ser.findall(C + "dPt"):
        if old.find(C + "idx").get("val") == d.find(C + "idx").get("val"):
            ser.remove(old)
    ser.find(C + "cat").addprevious(d)


def c_two_subjects(p, w):
    _chart(p, lambda r: _put_dpt(_ser(r), _dpt(0, '<a:schemeClr val="accent1"/>')))


def c_four_colours(p, w):
    def fn(r):
        s = _ser(r)
        _put_dpt(s, _dpt(0, '<a:schemeClr val="accent2"/>'))
        _put_dpt(s, _dpt(2, '<a:schemeClr val="accent3"/>'))
        _put_dpt(s, _dpt(3, '<a:schemeClr val="accent1"><a:lumMod val="75000"/></a:schemeClr>'))
    _chart(p, fn)


def c_truncated(p, w):
    _chart(p, lambda r: r.find(".//" + C + "valAx/" + C + "scaling/" + C + "min").set("val", "0.5"))


def c_dual_axis(p, w):
    def fn(r):
        pa = r.find(".//" + C + "plotArea")
        bar = pa.find(C + "barChart")
        g = etree.fromstring(etree.tostring(bar))
        for a, v in zip(g.findall(C + "axId"), ("60010", "60020")):
            a.set("val", v)
        bar.addnext(g)
        for ax in (pa.find(C + "catAx"), pa.find(C + "valAx")):
            cp = etree.fromstring(etree.tostring(ax))
            cp.find(C + "axId").set("val", "60010" if ax.tag == C + "catAx" else "60020")
            cp.find(C + "crossAx").set("val", "60020" if ax.tag == C + "catAx" else "60010")
            pa.find(C + "spPr").addprevious(cp)
    _chart(p, fn)


def c_banned_word(p, w):
    def fn(r):
        v = _ser(r).find(C + "cat").find(".//" + C + "pt/" + C + "v")
        v.text = "Seam" + "less handoff"
    _chart(p, fn)


def c_label_contrast(p, w):
    def fn(r):
        s = _ser(r)
        s.find(C + "cat").addprevious(_el(
            '<c:dLbls><c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="1000"><a:solidFill><a:schemeClr '
            'val="bg1"><a:lumMod val="85000"/></a:schemeClr></a:solidFill></a:defRPr></a:pPr><a:endParaRPr '
            'lang="en-US"/></a:p></c:txPr><c:dLblPos val="outEnd"/><c:showLegendKey val="0"/><c:showVal val="1"/>'
            '<c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/>'
            '</c:dLbls>'))
    _chart(p, fn)


def c_tick_contrast(p, w):
    def fn(r):
        ax = r.find(".//" + C + "catAx")
        ax.find(C + "tickLblPos").set("val", "nextTo")
        sf = ax.find(C + "txPr").find(".//" + A + "solidFill")
        sf.getparent().replace(sf, _el('<a:solidFill><a:schemeClr val="bg1"><a:lumMod val="85000"/></a:schemeClr>'
                                       '</a:solidFill>'))
    _chart(p, fn)


def c_no_kind(p, w):
    i, _part = add_chart(p)
    edit_slide(p, i, lambda r: shape(r, NAME).find(".//" + P + "cNvPr").set("descr", "Bar chart of days by stage"))


def c_no_basis(p, w):
    i, _part = add_chart(p)

    def fn(r):
        sh = shape(r, NAME + ".basis")
        sh.getparent().remove(sh)
    edit_slide(p, i, fn)


def c_data_drift(p, w):
    _chart(p, lambda r: _ser(r).find(C + "val").find(".//" + C + "pt/" + C + "v").__setattr__("text", "4"))


def c_label_drift(p, w):
    i, _part = add_chart(p)
    edit_slide(p, i, lambda r: setattr(shape(r, NAME + ".value.2").find(".//" + A + "t"), "text", "3 days"))


def c_no_record(p, w):
    add_chart(p)
    mp = Path(p).with_name(Path(p).stem + ".manifest.json")
    man = json.loads(mp.read_text(encoding="utf-8"))
    man["charts"].pop(NAME)
    mp.write_text(json.dumps(man, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def c_ring_four(p, w):
    def fn(r):
        cache = _ser(r).find(C + "val").find(".//" + C + "numCache")
        cache.find(C + "ptCount").set("val", "4")
        for k, v in ((2, "1"), (3, "1")):
            cache.append(_el('<c:pt idx="%d"><c:v>%s</c:v></c:pt>' % (k, v)))
    _chart(p, fn, "ring")


def c_bogus(p, w):
    _chart(p, lambda r: r.find(C + "chart").addprevious(_el('<c:bogus val="1"/>')))


def c_misordered(p, w):
    def fn(r):
        rc = r.find(C + "roundedCorners")
        r.remove(rc)
        r.find(C + "chart").addnext(rc)
    _chart(p, fn)


PLANTERS = [
    Planter("chart-6pt-text", "P27", "a chart's own text size set to 6pt", c_6pt, "under the 8pt floor"),
    Planter("chart-no-size", "P27", "the chart's txPr with no explicit size", c_no_size, "sets no text size"),
    Planter("chart-srgb", "P27", "a series filled srgb FF00FF", c_srgb, "neither a theme role"),
    Planter("chart-status-fill", "P27", "a series filled in the rag-amber status hex", c_status, "status identity"),
    Planter("chart-3d", "P27", "a c:view3D in the chart", c_3d, "3D"),
    Planter("chart-gradient", "P27", "a series filled with a gradient", c_gradient, "gradient"),
    Planter("chart-shadow", "P27", "an outer shadow on a series", c_shadow, "effect list"),
    Planter("chart-style", "P27", "c:style 34 on the chart", c_style, "c:style"),
    Planter("chart-rounded", "P27", "roundedCorners set to 1", c_rounded, "roundedCorners"),
    Planter("chart-title", "P27", "a chart title inside the frame", c_title, "chart title"),
    Planter("chart-legend", "P27", "a legend box over one series", c_legend, "legend box"),
    Planter("chart-unpinned", "P27", "the plot area's manualLayout removed", c_unpinned, "not pinned"),
    Planter("chart-two-subjects", "P27", "a second bar painted in accent", c_two_subjects, "one subject"),
    Planter("chart-four-colours", "P27", "four categorical colours across the bars", c_four_colours,
            "categorical colours"),
    Planter("chart-truncated-bar", "P27", "the bar value axis starting at 0.5", c_truncated, "starts at 0.5"),
    Planter("chart-dual-axis", "P27", "a second bar group on its own axis pair", c_dual_axis, "axis pairs"),
    Planter("chart-banned-word", "P27", "a banned word in a category name", c_banned_word, "fingerprint"),
    Planter("chart-label-contrast", "P27", "native data labels in hairline-soft on paper", c_label_contrast,
            "data labels"),
    Planter("chart-tick-contrast", "P27", "category tick labels shown in hairline-soft", c_tick_contrast,
            "tick labels"),
    Planter("chart-no-kind", "P27", "the frame's alt text without its chart kind", c_no_kind, "states no kind"),
    Planter("chart-no-basis", "P27", "an illustrative chart with its basis label removed", c_no_basis,
            "'Illustrative figures'"),
    Planter("chart-data-drift", "P27", "a value changed in the chart cache after the build", c_data_drift,
            "data drift"),
    Planter("chart-label-drift", "P27", "a value label retyped on glass after the build", c_label_drift,
            "label drift"),
    Planter("chart-no-record", "P27", "a chart with no manifest record", c_no_record, "no manifest['charts'] record"),
    Planter("chart-ring-four", "P27", "a share ring holding four points", c_ring_four, "of 4 points"),
    Planter("chart-bogus-element", "P12", "a <c:bogus/> before c:chart in a chart part", c_bogus, "bogus"),
    Planter("chart-misordered", "P12", "roundedCorners moved after c:chart", c_misordered, "roundedCorners"),
]


# ------------------------------------------------------------------ standalone runner

def main(argv=None) -> int:
    import argparse
    import run_all
    from _shared import Ctx, run_gate
    ap = argparse.ArgumentParser(description="run the chart planters alone")
    ap.add_argument("--keep", action="store_true")
    a = ap.parse_args(argv)
    if not run_all.GOLDEN.is_file():
        ok, out = run_all.build_golden()
        if not ok:
            print("golden build FAILED:\n" + out)
            return 1
    import importlib
    mods = {"P27": importlib.import_module("p27_charts"), "P12": importlib.import_module("p12_package")}
    bad = 0
    # the planted chart itself must be clean, or a planter's failure proves nothing
    tmp = Path(tempfile.mkdtemp(prefix="neg-chart-clean-"))
    try:
        dst = tmp / run_all.GOLDEN.name
        shutil.copy2(run_all.GOLDEN, dst)
        for sc in run_all.sidecars(run_all.GOLDEN):
            shutil.copy2(sc, tmp / sc.name)
        add_chart(dst)
        for gid in ("P27", "P12"):
            code, found = run_gate(mods[gid], str(dst), Ctx(dst))
            print("%-24s %-5s %s" % ("clean chart", gid, "pass" if code == 0 else "FAIL " + run_all.first_fail(found)))
            bad += code != 0
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    for pl in PLANTERS:
        tmp = Path(tempfile.mkdtemp(prefix="neg-%s-" % pl.name))
        try:
            dst = tmp / run_all.GOLDEN.name
            shutil.copy2(run_all.GOLDEN, dst)
            for sc in run_all.sidecars(run_all.GOLDEN):
                shutil.copy2(sc, tmp / sc.name)
            (tmp / "work").mkdir()
            pl.fn(dst, tmp / "work")
            code, found = run_gate(mods[pl.gate], str(dst), Ctx(dst))
            hit = [f for f in found if f.level == "fail" and pl.expect.lower() in (f.where + " " + f.msg).lower()]
            ok = code == 1 and bool(hit)
            bad += not ok
            print("%-24s %-5s %-14s %s" % (pl.name, pl.gate, "failed as expected" if ok else "DID NOT",
                                           ((hit[0].msg if hit else run_all.first_fail(found)) or "")[:70]))
        finally:
            if not a.keep:
                shutil.rmtree(tmp, ignore_errors=True)
    print("%d chart planters, %d misbehaved" % (len(PLANTERS), bad))
    return 1 if bad else 0


__all__ = ["PLANTERS", "add_chart"]

if __name__ == "__main__":
    sys.exit(main())
