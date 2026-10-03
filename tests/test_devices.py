"""The device library at runtime: a placed device fills every slot and leaves no sample behind."""
import json
import sys
import unittest
from pathlib import Path

SKILL = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SKILL / "kit"))

from pptx import Presentation  # noqa: E402

from deckkit import devices, prims  # noqa: E402
from deckkit.errors import GeometryError  # noqa: E402
from deckkit.theme import Theme  # noqa: E402


@unittest.skipUnless((SKILL / "devices" / "devices.json").is_file(), "device library not assembled")
class TestPlace(unittest.TestCase):
    def setUp(self):
        self.th = Theme.load(SKILL / "themes" / "sigmoid.json")
        self.prs = Presentation(str(SKILL / "kit" / "assets" / "base-wide.pptx"))
        self._use = prims.use(self.th)           # scoped: a global bind leaks into other suites
        self._use.__enter__()
        self.e = next(d for d in devices.catalogue() if d["id"] == "DL01")

    def tearDown(self):
        self._use.__exit__(None, None, None)

    def _fill(self):
        out = {}
        for s in self.e["slots"]:
            if s["kind"] == "glyph":
                out[s["name"]] = "check"
            elif s["kind"] == "title":
                out[s["name"]] = ("Two decisions start the build", ("Two decisions",))
            else:
                out[s["name"]] = "01" if (s.get("max_words") or 2) <= 1 else "Fresh text"
        return out

    def _text(self, s):
        return " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)

    def test_filled_device_carries_no_sample(self):
        s = devices.place(self.prs, "DL01", self._fill(), theme=self.th)
        text = self._text(s)
        left = [x for x in self.e["samples"] if len(x) >= 8 and x in text]
        self.assertEqual(left, [])

    def test_unfilled_slot_keeps_its_sample(self):
        fill = self._fill()
        name = next(s["name"] for s in self.e["slots"] if s["kind"] == "text" and (s.get("max_words") or 0) >= 4)
        fill.pop(name)
        s = devices.place(self.prs, "DL01", fill, theme=self.th)
        self.assertTrue(any(x in self._text(s) for x in self.e["samples"] if len(x) >= 8))

    def test_unknown_slot_refused(self):
        with self.assertRaises(GeometryError):
            devices.place(self.prs, "DL01", {"slot.nope": "x"}, theme=self.th)


if __name__ == "__main__":
    unittest.main()
