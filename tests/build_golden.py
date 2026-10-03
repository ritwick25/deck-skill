#!/usr/bin/env python3
"""The golden fixture: a seven-slide WIDE deck in the sigmoid theme, drawn only from kit primitives and
text, tagged per slide, that every static gate must pass. The negatives runner copies it and plants one
breakage per copy; a gate that does not go red on its plant has proved nothing.

    $PY tests/build_golden.py                 # writes .cache/golden.pptx, its manifest and deck.json
    $PY tests/build_golden.py --out X.pptx    # the same deck anywhere (gate P21 rebuilds it this way)

Slides (archetype:variant, family):
  1 cover:title-only            Frame     one two-tone title, nothing else
  2 offering-pillars:icon-row   Parallel columns   three cards, each a glyph, a heading, a measured body and five rows
  3 process-flow:chevron-row    Sequence  a chevron header, four step cards, three arrows card to card
  4 headline-figures:stat-row   Numbers-led   four stat cards (measured, delivered, illustrative, target), a glyph each
  5 reference-table:table      Two-dimensional field  a readiness table and two side cards
  6 from-to-shift:two-panel     Split     two zones, one arrow between them
  7 closing:line                Frame     the closing line

Every content slide fills the canvas content band (gate P24): its structure runs from just inside the band's top
to the band's foot (BOTTOM), and the air inside a card is spread evenly between its blocks (spread), never left
under its content (gate P08). Every text box is sized from deckkit.measure at build time (tight fit, no slack
beyond SLACK), so a build with SLACK patched to 0.5 produces boxes that gate P02 rejects. Copy is fictional: a regional
planning team with no client, brand or person in it. Module constants are patch points for the
builder-refusal tests; nothing here is ever edited by a test.
"""
from __future__ import annotations

import argparse
import json
import math
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL / "kit"))

from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

from deckkit import glyphs, ooxml, package, prims, tag, text  # noqa: E402
from deckkit.canvas import EMU_PER_IN, WIDE, Box  # noqa: E402
from deckkit.read.ns import NS  # noqa: E402
from deckkit.text import Para, Run  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

BASE = SKILL / "kit" / "assets" / "base-wide.pptx"
THEME_PATH = SKILL / "themes" / "sigmoid.json"
DEFAULT_OUT = SKILL / ".cache" / "golden.pptx"
TITLE = "Golden fixture"

# patch points (the refusal tests patch these, never the source)
BODY_PT = 11.0
HEAD_PT = 13.0
CELL_PT = 11.0
PAD = 0.20            # card padding
ICON = 0.45           # glyph side
GAP = 0.10            # between stacked text blocks
CELL_MY = 0.11        # a table cell's top and bottom margin

_A = "{%s}" % NS["a"]
NO_TABLE_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"     # PowerPoint's built-in "No Style, No Grid"

C = WIDE.content       # 0.50, 0.85, 12.403 x 5.99


# ------------------------------------------------------------------ measured helpers

def need(paras, w: float, theme, size: float) -> float:
    """Height the paragraphs need at width w (measured, SLACK applied), rounded UP to 0.0001 in because
    measure.fits reports need_in rounded to four places."""
    n = text.check_fit("probe", paras, Box(0, 0, w, 100.0), theme, size=size).need_in
    return math.ceil(n * 10000 + 1) / 10000.0


def stack(slide, x: float, y: float, w: float, blocks, theme, prefix: str) -> float:
    """Text blocks one under another, each box exactly its measured height. Returns the bottom."""
    for i, (paras, size) in enumerate(blocks):
        h = need(paras, w, theme, size)
        text.textbox(slide, Box(x, y, w, h), paras, "%s.%d" % (prefix, i + 1), theme, size=size)
        y += h + (GAP if i + 1 < len(blocks) else 0.0)
    return y


def blocks_height(blocks, w, theme) -> float:
    hs = [need(p, w, theme, s) for p, s in blocks]
    return sum(hs) + GAP * (len(hs) - 1)


def head(t):
    return [Para(Run(t, role="ink-deep", bold=True))]


def body(*lines):
    return [Para(Run(t, role="ink"), space_after=4.0 if i + 1 < len(lines) else 0.0) for i, t in enumerate(lines)]


def new_slide(prs, archetype, variant):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    tag.tag(s, archetype, variant)
    return s


# ------------------------------------------------------------------ the table (no table module yet)

def table(slide, x, y, col_w, rows, theme, name, height=None):
    """A plain table: header row in ink-strong with paper text, body rows on paper, one hairline under
    every row, explicit margins, every cell's text written by text.write and sized by measurement. With height,
    the body rows share whatever the measured rows leave (the cells are anchored middle)."""
    from lxml import etree
    mx, my = 0.10, CELL_MY
    heights = []
    for ri, row in enumerate(rows):
        hh = 0.0
        for ci, cell in enumerate(row):
            paras = [Para(Run(cell, role="paper" if ri == 0 else "ink", bold=ri == 0))]
            hh = max(hh, need(paras, col_w[ci] - 2 * mx, theme, CELL_PT) + 2 * my)
        heights.append(hh)
    if height is not None and height > sum(heights):
        extra = (height - sum(heights)) / (len(heights) - 1)
        heights = heights[:1] + [h + extra for h in heights[1:]]
    gf = slide.shapes.add_table(len(rows), len(col_w), *Box(x, y, sum(col_w), sum(heights)).emu)
    gf.name = name
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for k in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        tblPr.set(k, "0")
    sid = tblPr.find(_A + "tableStyleId")
    if sid is None:
        sid = etree.SubElement(tblPr, _A + "tableStyleId")
    sid.text = NO_TABLE_STYLE
    for ci, w in enumerate(col_w):
        tbl.columns[ci].width = Emu(int(round(w * EMU_PER_IN)))
    for ri, h in enumerate(heights):
        tbl.rows[ri].height = Emu(int(round(h * EMU_PER_IN)))
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = tbl.cell(ri, ci)
            paras = [Para(Run(cell_text, role="paper" if ri == 0 else "ink", bold=ri == 0))]
            text.write(cell.text_frame, paras, theme, size=CELL_PT, anchor="middle",
                       background="ink-strong" if ri == 0 else "paper")
            tcPr = cell._tc.get_or_add_tcPr()
            for c in list(tcPr):
                tcPr.remove(c)
            for k, v in (("marL", mx), ("marR", mx), ("marT", my), ("marB", my)):
                tcPr.set(k, str(int(round(v * EMU_PER_IN))))
            tcPr.set("anchor", "ctr")
            for side in ("lnL", "lnR", "lnT", "lnB"):
                ln = etree.SubElement(tcPr, _A + side)
                if side == "lnB":
                    ln.set("w", "9525")
                    etree.SubElement(ln, _A + "solidFill").append(theme.spec("hairline"))
                else:
                    ln.set("w", "0")
                    etree.SubElement(ln, _A + "noFill")
            etree.SubElement(tcPr, _A + "solidFill").append(theme.spec("ink-strong" if ri == 0 else "paper"))
    ooxml.finish(gf)
    return gf, sum(heights)


# ------------------------------------------------------------------ slides

def s_cover(prs, th):
    s = new_slide(prs, "cover", "title-only")
    paras = [Para(text.title_runs("One shared demand view for every region", ["shared demand view"]))]
    h = need(paras, C.w, th, 30.0)
    text.textbox(s, Box(C.x, 3.0, C.w, h), paras, "title", th, size=30.0, face="Arial")


TOP = C.y + 0.10      # content starts just inside the content band
BOTTOM = C.b          # ... and runs to the band's foot (gate P24)
LABEL_PT = 10.0       # the bold label over a value line
ROW_ICON = 0.45       # glyph side in a label-value row
ROW_GAP = 0.22        # between label-value rows


def label_value(label, value):
    return [Para(Run(label, role="ink-deep", bold=True, size=LABEL_PT)), Para(Run(value, role="ink"))]


def spread(heights, y0: float, y1: float, min_gap: float, what: str) -> list[float]:
    """Tops for blocks of these heights between y0 and y1, the first at y0 and the last ending at y1, the air
    shared evenly between them. Raises when the blocks need more than the room at min_gap."""
    free = (y1 - y0) - sum(heights)
    n = len(heights)
    gap = free / (n - 1) if n > 1 else 0.0
    if n > 1 and gap + 1e-9 < min_gap:
        raise ValueError("%s: %d blocks need %.3fin, the room is %.3fin" % (what, n, sum(heights) + min_gap * (n - 1),
                                                                         y1 - y0))
    tops, y = [], y0
    for h in heights:
        tops.append(y)
        y += h + gap
    return tops


def row_heights(items, w, theme) -> list[float]:
    tw = w - ROW_ICON - 0.15
    return [max(ROW_ICON, need(label_value(lab, val), tw, theme, BODY_PT)) for _g, lab, val in items]


def rows(slide, x, y, w, items, theme, prefix, bottom=None, heights=None) -> float:
    """Label-value rows, each a glyph with its label and value beside it, the glyph centred on the text block.
    With bottom, the rows are spread from y to bottom (the last row ends on it). heights, when given, are the row
    heights to lay out (a row of cards shares the tallest row at each index, so their rows line up across cards).
    Returns the bottom."""
    tx, tw = x + ROW_ICON + 0.15, w - ROW_ICON - 0.15
    hs = list(heights) if heights is not None else row_heights(items, w, theme)
    if bottom is None:
        tops = [y + sum(hs[:k]) + ROW_GAP * k for k in range(len(hs))]
    else:
        tops = spread(hs, y, bottom, ROW_GAP, prefix)
    for k, ((g, label, value), rh, ry) in enumerate(zip(items, hs, tops), start=1):
        paras = label_value(label, value)
        h = need(paras, tw, theme, BODY_PT)
        glyphs.glyph(slide, g, Box(x, ry + (rh - ROW_ICON) / 2, ROW_ICON, ROW_ICON), role="ink-strong",
                     shape_name="%s.%d.icon" % (prefix, k), theme=theme)
        text.textbox(slide, Box(tx, ry + (rh - h) / 2, tw, h), paras, "%s.%d.text" % (prefix, k), theme, size=BODY_PT)
    return tops[-1] + hs[-1]


def rows_height(items, w, theme) -> float:
    hs = row_heights(items, w, theme)
    return sum(hs) + ROW_GAP * (len(hs) - 1)


PILLARS = [
    ("database", "Shared demand signal",
     "Every region reads the same weekly demand figures, drawn from one governed store and refreshed each Monday.",
     [("user", "Owner", "Demand planning lead"), ("calendar", "Cadence", "Refreshed every Monday"),
      ("gauge", "Measure", "Forecast error by product family"), ("layers", "Scope", "Every region and channel"),
      ("file-spreadsheet", "Inputs", "Orders, shipments and promotions")]),
    ("search", "Exception-first review",
     "Planners see only the accounts that moved outside tolerance, with the reason and the size of the gap beside each.",
     [("users", "Owner", "Regional planners"), ("clock", "Cadence", "Tuesday review, one hour"),
      ("target", "Measure", "Exceptions closed in the cycle"), ("funnel", "Scope", "Lines beyond tolerance only"),
      ("triangle-alert", "Trigger", "A line outside its tolerance band")]),
    ("file-check", "Signed plan of record",
     "Each cycle closes on one versioned plan that finance, supply and sales all sign before it is released.",
     [("handshake", "Owner", "Finance, supply and sales"), ("calendar-check", "Cadence", "Signed each Friday"),
      ("archive", "Measure", "Plans signed on time"), ("lock", "Scope", "One version per cycle"),
      ("share-2", "Readers", "Every planner and the regional leads")]),
]
RULE_GAP = 0.20       # air above and below the hairline that separates a card's claim from its rows


def s_pillars(prs, th):
    s = new_slide(prs, "offering-pillars", "icon-row")
    text.title(s, "Three capabilities carry the regional planning model", ["Three capabilities"], WIDE, th)
    cols = Box(C.x, TOP, C.w, BOTTOM - TOP).split_cols(3, 0.30)
    iw = cols[0].w - 2 * PAD
    head_h = max(blocks_height([(head(h), HEAD_PT), (body(b), BODY_PT)], iw, th) for _, h, b, _r in PILLARS)
    row_hs = [max(col) for col in zip(*(row_heights(r, iw, th) for _, _, _, r in PILLARS))]
    for i, (col, (g, h, b, items)) in enumerate(zip(cols, PILLARS), start=1):
        card = col
        prims.card(s, card, fill="card", line="hairline", name="pillar.%d.card" % i, theme=th)
        glyphs.glyph(s, g, Box(card.x + PAD, card.y + PAD, ICON, ICON), role="accent",
                     shape_name="pillar.%d.icon" % i, theme=th)
        stack(s, card.x + PAD, card.y + PAD + ICON + GAP, iw, [(head(h), HEAD_PT), (body(b), BODY_PT)], th,
              "pillar.%d.text" % i)
        ry = card.y + PAD + ICON + GAP + head_h + RULE_GAP
        prims.hairline(s, (card.x + PAD, ry), (card.r - PAD, ry), name="pillar.%d.rule" % i, theme=th)
        rows(s, card.x + PAD, ry + RULE_GAP, iw, items, th, "pillar.%d.row" % i, bottom=card.b - PAD,
             heights=row_hs)


STEPS = [
    ("Draft forecast", "The model proposes a weekly forecast per region and product family.",
     [("Input", "Last week's orders and shipments"), ("Output", "A draft for every line"),
      ("Owner", "Demand planning"), ("Done when", "Every line has a draft by Monday noon")]),
    ("Review exceptions", "Planners accept or adjust the lines that moved beyond tolerance.",
     [("Input", "The ranked exception list"), ("Output", "Accepted or adjusted lines"),
      ("Owner", "Regional planners"), ("Done when", "No exception is left open on Tuesday")]),
    ("Reconcile supply", "Supply confirms capacity and flags any line it cannot meet.",
     [("Input", "The reviewed forecast"), ("Output", "Confirmed capacity"),
      ("Owner", "Supply planning"), ("Done when", "Every flagged line has a named fix")]),
    ("Sign the plan", "Finance, supply and sales sign one versioned plan of record.",
     [("Input", "The reconciled plan"), ("Output", "The signed plan of record"),
      ("Owner", "Finance, supply and sales"), ("Done when", "Three signatures by Friday")]),
]
DISC = 0.55           # the step disc; its digit is set large, an ordinal and never a figure (gate P23)


def s_process(prs, th):
    s = new_slide(prs, "process-flow", "chevron-row")
    text.title(s, "Four steps take a forecast from draft to signed plan", ["Four steps"], WIDE, th)
    prims.step_header(s, Box(C.x, TOP, C.w, 0.50), ["Draft", "Review", "Reconcile", "Sign"], name="steps",
                      theme=th)
    zpad = 0.25
    panel_y = TOP + 0.50 + 0.20
    panel = Box(C.x, panel_y, C.w, BOTTOM - panel_y)
    inner = Box(C.x + zpad, panel_y + zpad, C.w - 2 * zpad, panel.h - 2 * zpad)
    cols = inner.split_cols(4, 0.50)
    iw = cols[0].w - 2 * PAD
    head_h = max(blocks_height([(head(h), HEAD_PT), (body(b), BODY_PT)], iw, th) for h, b, _d in STEPS)
    detail_hs = [max(col) for col in zip(*([need(label_value(a, v), iw, th, BODY_PT) for a, v in d]
                                           for _h, _b, d in STEPS))]
    prims.zone(s, panel, fill="accent-tint", name="step.panel", theme=th)
    cards = []
    for i, (card, (h, b, detail)) in enumerate(zip(cols, STEPS), start=1):
        prims.card(s, card, fill="paper", line="rule-strong", name="step.%d.card" % i, theme=th)
        prims.disc(s, Box(card.x + PAD, card.y + PAD, DISC, DISC), i, fill="accent", name="step.%d.disc" % i, theme=th)
        stack(s, card.x + PAD, card.y + PAD + DISC + GAP, iw, [(head(h), HEAD_PT), (body(b), BODY_PT)], th,
              "step.%d.text" % i)
        ry = card.y + PAD + DISC + GAP + head_h + RULE_GAP
        prims.hairline(s, (card.x + PAD, ry), (card.r - PAD, ry), name="step.%d.rule" % i, theme=th)
        tops = spread(detail_hs, ry + RULE_GAP, card.b - PAD, GAP, "step.%d.detail" % i)
        for k, ((lab, val), dy) in enumerate(zip(detail, tops), start=1):
            p = label_value(lab, val)
            text.textbox(s, Box(card.x + PAD, dy, iw, need(p, iw, th, BODY_PT)), p, "step.%d.detail.%d" % (i, k), th,
                         size=BODY_PT)
        cards.append(card)
    for i, (a, b) in enumerate(zip(cards, cards[1:]), start=1):
        prims.connect(s, (a.r, a.cy), (b.x, b.cy), name="step.arrow.%d" % i, theme=th)


STATS = [
    ("trending-down", "42%", "Fewer manual adjustments per planning cycle", "measured",
     "Counted from the planning log over the last two quarters, in both pilot regions.",
     [("Window", "Last two quarters"), ("Regions", "North and Coastal")]),
    ("timer", "5 days", "From draft forecast to signed plan, down from twelve", "delivered",
     "Elapsed days between the Monday draft and the Friday signature, averaged over the quarter.",
     [("Window", "The last full quarter"), ("Regions", "North and Coastal")]),
    ("package", "18%", "Lower safety stock at the same service level", "illustrative",
     "A planning assumption from the supply model, not yet seen in any region's stock.",
     [("Window", "Modelled, not observed"), ("Regions", "Every region")]),
    ("flag", "8 wk", "To a working pilot in two regions", "target",
     "The committed date for the pilot. The clock starts when the data owners sign off.",
     [("Window", "From data owner sign-off"), ("Regions", "North and Coastal")]),
]
STAT_BOX = 1.30          # stat.3 (a label and a basis chip) needs 1.21in
STAT_ICON = 0.60


def s_stats(prs, th):
    s = new_slide(prs, "headline-figures", "stat-row")
    text.title(s, "Planning cycle time falls once regions share one view", ["cycle time"], WIDE, th)
    cols = Box(C.x, TOP, C.w, BOTTOM - TOP).split_cols(4, 0.40)
    iw = cols[0].w - 2 * PAD

    def src_blocks(src):
        return [(head("How it is counted"), BODY_PT), (body(src), BODY_PT)]

    hs = [max(col) for col in zip(*([STAT_ICON, STAT_BOX] + [need(label_value(a, b), iw, th, BODY_PT) for a, b in f]
                                    + [blocks_height(src_blocks(src), iw, th)] for *_x, src, f in STATS))]
    tops = spread(hs, cols[0].y + PAD, cols[0].b - PAD, GAP, "stat cards")      # one grid for the four cards
    numbers = {}
    for i, (card, (g, v, lab, kind, src, facts)) in enumerate(zip(cols, STATS), start=1):
        prims.card(s, card, fill="card", line="hairline", name="stat.%d.card" % i, theme=th)
        x = card.x + PAD
        fact_paras = [label_value(a, b) for a, b in facts]
        glyphs.glyph(s, g, Box(x, tops[0], STAT_ICON, STAT_ICON), role="accent", shape_name="stat.%d.icon" % i,
                     theme=th)
        prims.stat(s, Box(x, tops[1], iw, STAT_BOX), v, lab, kind=kind, background="card", name="stat.%d" % i,
                   theme=th)
        numbers["stat.%d" % i] = {"value": v, "kind": kind}
        for k, (p, fy) in enumerate(zip(fact_paras, tops[2:-1]), start=1):
            text.textbox(s, Box(x, fy, iw, need(p, iw, th, BODY_PT)), p, "stat.%d.fact.%d" % (i, k), th, size=BODY_PT,
                         background="card")
        stack(s, x, tops[-1], iw, src_blocks(src), th, "stat.%d.source.text" % i)
    return numbers


REGIONS = [("Region", "Data ready", "Planners trained"),
           ("North", "Yes", "Yes"), ("Coastal", "Yes", "In progress"), ("Inland", "By week six", "Not started"),
           ("Highland", "Yes", "Not started"), ("Lakeside", "By week four", "In progress"),
           ("Delta", "By week eight", "Not started"), ("Plains", "Yes", "In progress"),
           ("Valley", "By week six", "Not started"), ("Harbour", "Yes", "Yes"), ("Ridge", "By week ten", "Not started")]
NEXT = [("calendar-check", "Wave two", "Lakeside and Plains, once their planners finish training"),
        ("calendar", "Wave three", "Every other region, once its data owner signs the readiness test"),
        ("graduation-cap", "Training", "Two half-day sessions per region, run by the planning lead"),
        ("user-check", "Data owners", "One named owner per region signs the readiness test"),
        ("gauge", "Review", "Forecast error checked at the end of every wave")]


def s_table(prs, th):
    s = new_slide(prs, "reference-table", "table")
    text.title(s, "Two regions are ready for the pilot this quarter", ["Two regions"], WIDE, th)
    _gf, _th = table(s, C.x, TOP, [3.0, 2.6, 2.6], REGIONS, th, "readiness.table", height=BOTTOM - TOP)
    x = C.x + 8.2 + 0.40
    w = C.r - x
    iw = w - 2 * PAD
    blocks = [(head("Pilot regions"), HEAD_PT),
              (body("North and Coastal meet both readiness tests today and start first."), BODY_PT)]
    card = Box(x, TOP, w, ICON + GAP + blocks_height(blocks, iw, th) + 2 * PAD)
    prims.card(s, card, fill="card", line="hairline", name="pilot.card", theme=th)
    glyphs.glyph(s, "circle-check", Box(card.x + PAD, card.y + PAD, ICON, ICON), role="accent",
                 shape_name="pilot.icon", theme=th)
    stack(s, card.x + PAD, card.y + PAD + ICON + GAP, iw, blocks, th, "pilot.text")
    nxt = Box(x, card.b + 0.30, w, BOTTOM - card.b - 0.30)
    prims.card(s, nxt, fill="paper", line="rule-strong", name="next.card", theme=th)
    rows(s, nxt.x + PAD, nxt.y + PAD, iw, NEXT, th, "next.row", bottom=nxt.b - PAD)


PANELS = [
    ("card", "clock", "Weekly reconciliation", "Each region keeps its own spreadsheet of demand.",
     [("files", "Figures", "Five copies of the same demand, matched by hand"),
      ("hourglass", "Effort", "Two planner days before every review"),
      ("triangle-alert", "Disagreement", "Found in the meeting, not before it")]),
    ("accent-tint", "shield-check", "One governed view", "Regions read demand from one store with a named owner.",
     [("database", "Figures", "One store, refreshed each Monday"),
      ("list-ordered", "Effort", "Exceptions arrive ranked by the size of the gap"),
      ("circle-check", "Agreement", "The review starts from figures that already agree")]),
]


def s_fromto(prs, th):
    s = new_slide(prs, "from-to-shift", "two-panel")
    text.title(s, "Weekly reconciliation gives way to one governed view", ["one governed view"], WIDE, th)
    cols = Box(C.x, TOP, C.w, BOTTOM - TOP).split_cols(2, 0.90)
    pad = 0.25
    iw = cols[0].w - 2 * pad
    cw = iw - 2 * PAD
    head_h = max(blocks_height([(head(h), HEAD_PT), (body(b), BODY_PT)], iw, th) for _, _, h, b, _i in PANELS)
    item_h = max(max(ROW_ICON, need(label_value(lab, val), cw - ROW_ICON - 0.15, th, BODY_PT))
                 for *_x, items in PANELS for _g, lab, val in items)
    items_top = TOP + pad + ICON + GAP + head_h + 0.25
    card_h = (BOTTOM - pad - items_top - 2 * 0.20) / 3          # three item cards share the room to the foot
    zones = []
    for i, (col, (fill, g, h, b, items)) in enumerate(zip(cols, PANELS), start=1):
        z = col
        prims.zone(s, z, fill=fill, name="panel.%d" % i, theme=th)
        glyphs.glyph(s, g, Box(z.x + pad, z.y + pad, ICON, ICON), role="accent" if i == 2 else "ink-strong",
                     shape_name="panel.%d.icon" % i, theme=th)
        stack(s, z.x + pad, z.y + pad + ICON + GAP, iw, [(head(h), HEAD_PT), (body(b), BODY_PT)], th,
              "panel.%d.text" % i)
        y = z.y + pad + ICON + GAP + head_h + 0.25
        for k, item in enumerate(items, start=1):
            c = Box(z.x + pad, y, iw, card_h)
            prims.card(s, c, fill="paper", line="rule-strong", name="panel.%d.item.%d.card" % (i, k), theme=th)
            rh = row_heights([item], cw, th)[0]
            rows(s, c.x + PAD, c.y + (card_h - rh) / 2, cw, [item], th, "panel.%d.item.%d" % (i, k))
            y = c.b + 0.20
        zones.append(z)
    a, b = zones
    prims.connect(s, (a.r, a.cy), (b.x, b.cy), weight_pt=1.5, role="accent", name="panel.arrow", theme=th)


def s_close(prs, th):
    s = new_slide(prs, "closing", "line")
    paras = [Para(text.title_runs("Next step: confirm the pilot region by month end", ["confirm the pilot region"]))]
    h = need(paras, C.w, th, 24.0)
    text.textbox(s, Box(C.x, 3.1, C.w, h), paras, "title", th, size=24.0, face="Arial")


# ------------------------------------------------------------------ build

def build(out: Path = DEFAULT_OUT) -> str:
    """Build the golden deck at out, with its manifest and deck.json beside it. Returns the pptx sha256."""
    out = Path(out)
    th = Theme.load(THEME_PATH)
    prs = Presentation(str(BASE))
    with prims.use(th):
        s_cover(prs, th)
        s_pillars(prs, th)
        s_process(prs, th)
        numbers = s_stats(prs, th)
        s_table(prs, th)
        s_fromto(prs, th)
        s_close(prs, th)
    manifest = {"fixture": "golden", "canvas": WIDE.name, "numbers": numbers}
    sha = package.save(prs, out, title=TITLE, theme=th, manifest=manifest, overwrite=True)
    deck_json = {
        "type": "fixture", "canvas": WIDE.name, "theme": str(THEME_PATH), "client": None, "spelling": "US",
        "overrides": {}, "budgets": {}, "series": [], "deliverable": out.name, "notes_deck": False,
        "build": ["{python}", str(Path(__file__).resolve()), "--out", "{out}"],
    }
    (out.parent / "deck.json").write_text(json.dumps(deck_json, indent=2) + "\n", encoding="utf-8")
    return sha


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", default=str(DEFAULT_OUT))
    a = ap.parse_args(argv)
    sha = build(Path(a.out))
    print("%s  sha256 %s" % (a.out, sha))
    return 0


if __name__ == "__main__":
    sys.exit(main())
