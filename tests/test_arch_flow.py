"""Flow diagrams (deckkit.arch_flow): the measured board, its router and labels, the refusals, and the five type
builders on the example's Quillmont content, each held to gate P26 at its arch-* type.

Run: cd ~/.claude/skills/deck && $PY -m unittest tests.test_arch_flow -v
"""
from __future__ import annotations

import importlib.util
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(SKILL, "kit"))
sys.path.insert(0, os.path.join(SKILL, "gates"))

from pptx import Presentation  # noqa: E402

from deckkit import arch_flow as af, package, prims, tag  # noqa: E402
from deckkit.canvas import WIDE, Box  # noqa: E402
from deckkit.errors import GeometryError, Overflows  # noqa: E402
from deckkit.text import title  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

BASE = os.path.join(SKILL, "kit", "assets", "base-wide.pptx")
SIGMOID = os.path.join(SKILL, "themes", "sigmoid.json")
TH = Theme.load(SIGMOID)
C = WIDE.content
LAMBDA, S3, RDS = "vendor:aws/compute/lambda", "vendor:aws/storage/simple-storage-service-s3", "vendor:aws/database/rds"


def slide():
    prs = Presentation(BASE)
    return prs, prs.slides.add_slide(prs.slide_layouts[0])


def tearDownModule():
    prims._ACTIVE.clear()


def example():
    spec = importlib.util.spec_from_file_location("example_arch_flow", os.path.join(SKILL, "scripts",
                                                                                     "example_arch_flow.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def small(draft=False, fact="one account"):
    """Three marks in a boundary, a gate between two of them, a legend row."""
    b = af.Board(C, TH, cols=[3, 1, 3, 3], rows=[af.HEAD_LABEL, ("auto", 1), 0.30], name="t", draft=draft)
    b.frame("acct", "boundary", "Account", fact=fact, cols=(0, 3), rows=(0, 1))
    b.mark("a", LAMBDA, "Lambda", "collectors", cell=(0, 1))
    b.gate("g", "quality", "tests pass", cell=(1, 1))
    b.mark("b", S3, "Amazon S3", "raw zone", cell=(2, 1))
    b.mark("c", RDS, "RDS", "ledger", cell=(3, 1))
    b.edge("a", "g")
    b.edge("g", "b", "land", step=1)
    b.edge("b", "c", "load nightly", step=2)
    b.steps({1: "Land raw files", 2: "Load the ledger"})
    return b


class Board(unittest.TestCase):
    def test_render_routes_lands_and_labels(self):
        prs, s = slide()
        m = small().render(s)
        self.assertEqual(m["crossings"], 0)
        self.assertEqual(m["marks"], 3)
        self.assertEqual(m["gates"], 1)
        self.assertEqual(m["steps"], 2)
        self.assertEqual(m["problems"], [])
        names = {sh.name for sh in s.shapes}
        for n in ("t.f.1", "t.f.2", "t.f.2.label", "t.f.2.step", "t.n.g.gate", "t.g.acct.frame", "t.legend.1.step"):
            self.assertIn(n, names)

    def test_edges_end_on_their_nodes(self):
        prs, s = slide()
        b = small()
        b.render(s)
        for e in b.edges:
            a, z = b.nodes[e.a].body, b.nodes[e.b].body
            (x0, y0), (x1, y1) = e.pts[0], e.pts[-1]
            self.assertTrue(abs(x0 - a.r) < 1e-3 or abs(x0 - a.x) < 1e-3 or abs(y0 - a.y) < 1e-3 or
                            abs(y0 - a.b) < 1e-3, (e.a, e.pts[0], a))
            self.assertTrue(abs(x1 - z.x) < 1e-3 or abs(x1 - z.r) < 1e-3 or abs(y1 - z.y) < 1e-3 or
                            abs(y1 - z.b) < 1e-3, (e.b, e.pts[-1], z))

    def test_outer_boundary_needs_a_fact(self):
        prs, s = slide()
        b = af.Board(C, TH, cols=[1, 1], rows=[af.HEAD_LABEL, "auto"], name="t")
        b.frame("z", "boundary", "Account", cols=(0, 1), rows=(0, 1))
        b.mark("a", LAMBDA, "Lambda", cell=(0, 1))
        b.mark("b", S3, "Amazon S3", cell=(1, 1))
        b.edge("a", "b", "land")
        with self.assertRaises(GeometryError):
            b.render(s)

    def test_declarations_refuse(self):
        b = af.Board(C, TH, cols=[1, 1], rows=["auto"], name="t")
        b.mark("a", LAMBDA, "Lambda", cell=(0, 0))
        b.mark("b", S3, "Amazon S3", cell=(1, 0))
        with self.assertRaises(GeometryError):
            b.edge("a", "b", "five words is too many")
        with self.assertRaises(GeometryError):
            b.gate("g", "quality", "a criterion that rambles on", cell=(0, 0))
        with self.assertRaises(GeometryError):
            b.chip("e", "BriefDrafted", cell=(0, 0))
        with self.assertRaises(GeometryError):
            b.frame("f", "region", "", cols=(0, 1), rows=(0, 0))
        with self.assertRaises(GeometryError):
            b.mark("x", "glyph:nope-not-a-glyph", "X", cell=(0, 0))
        with self.assertRaises(GeometryError):
            b.edge("a", "b", kind="sideways")

    def test_steps_run_one_to_n(self):
        prs, s = slide()
        b = af.Board(C, TH, cols=[1, 1, 1], rows=["auto", 0.3], name="t")
        for k, mk in enumerate((LAMBDA, S3, RDS)):
            b.mark("n%d" % k, mk, "Node %d" % k, cell=(k, 0))
        b.edge("n0", "n1", "a", step=1)
        b.edge("n1", "n2", "b", step=3)
        b.steps({1: "One", 3: "Three"})
        with self.assertRaises(GeometryError):
            b.render(s)

    def test_caption_over_two_lines_raises_and_draft_lists(self):
        def build(draft):
            b = af.Board(C, TH, cols=[1, 9], rows=["auto"], name="t", draft=draft)
            b.mark("a", LAMBDA, "Lambda", "a role long enough to wrap onto more than two lines here", cell=(0, 0))
            b.mark("b", S3, "Amazon S3", cell=(1, 0))
            b.edge("a", "b", "land")
            return b
        prs, s = slide()
        with self.assertRaises(Overflows):
            build(False).render(s)
        prs, s = slide()
        m = build(True).render(s)
        self.assertTrue(any("caption" in p for p in m["problems"]))

    def test_crossing_is_refused(self):
        """Four nodes on a cross: the second edge cannot avoid the first."""
        prs, s = slide()
        b = af.Board(C, TH, cols=[1, 1, 1], rows=["auto", "auto", "auto"], name="t")
        b.mark("w", LAMBDA, "West", cell=(0, 1))
        b.mark("e", S3, "East", cell=(2, 1))
        b.mark("n", RDS, "North", cell=(1, 0), caption="right")
        b.mark("s", "vendor:aws/compute/fargate", "South", cell=(1, 2), caption="right")
        b.edge("w", "e", "across", sides=("r", "l"))
        b.edge("n", "s", "down", sides=("b", "t"))
        with self.assertRaises(GeometryError) as cm:
            b.render(s)
        self.assertIn("crossing", str(cm.exception))

    def test_blocked_edge_draws_its_glyph_and_key(self):
        prs, s = slide()
        b = af.Board(C, TH, cols=[1, 1], rows=["auto", 0.3], name="t")
        b.mark("a", "vendor:aws/management/management-console", "AWS Console", cell=(0, 0))
        b.mark("b", S3, "Amazon S3", cell=(1, 0))
        b.edge("a", "b", "manual deploy", kind="blocked")
        m = b.render(s)
        names = {sh.name for sh in s.shapes}
        self.assertIn("t.f.1.block", names)
        self.assertIn("state.t.legend.key.blocked", names)
        self.assertEqual(m["kinds"], ["blocked"])

    def test_bus_merges_feeders_into_one_head(self):
        prs, s = slide()
        b = af.Board(C, TH, cols=[2, 2], rows=["auto", "auto", "auto"], name="t")
        b.stack([("p", "glyph:megaphone", "Press", None), ("q", "glyph:landmark", "Payers", None),
                 ("r", "glyph:folder-open", "Partners", None)], cell=(0, 0), rows=(0, 2))
        b.mark("l", LAMBDA, "Lambda", "collectors", cell=(1, 1))
        b.bus(["p", "q", "r"], "l", "weekly pull")
        b.render(s)
        heads = [e for e in b.edges if getattr(e, "bus", None)]
        self.assertEqual(len(heads), 1)
        self.assertEqual(len(heads[0].bus["feeders"]), 3)
        names = [sh.name for sh in s.shapes]
        self.assertEqual(sum(1 for n in names if ".bus." in n), 4)        # three feeders and the trunk


class Types(unittest.TestCase):
    """The example's five slides, built by the type builders, gated by P26 at their arch-* types."""

    @classmethod
    def setUpClass(cls):
        cls.dir = tempfile.mkdtemp(prefix="arch-flow-")
        cls.ex = example()
        prs = Presentation(BASE)
        cls.metrics = {k: getattr(cls.ex, k)(prs, TH) for k in ("cicd", "mlops", "medallion", "events", "state")}
        cls.out = os.path.join(cls.dir, "flow.pptx")
        package.save(prs, cls.out, title="flow", theme=TH, overwrite=True)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, True)
        prims._ACTIVE.clear()

    def test_every_type_draws_clean(self):
        for k, m in self.metrics.items():
            self.assertEqual(m["problems"], [], k)
            self.assertEqual(m["crossings"], 0, k)

    def test_spec_bands(self):
        m = self.metrics
        self.assertTrue(12 <= m["cicd"]["marks"] <= 24 and m["cicd"]["gates"] == 4 and 4 <= m["cicd"]["steps"] <= 10)
        self.assertTrue(14 <= m["mlops"]["marks"] <= 28 and m["mlops"]["gates"] == 3 and 5 <= m["mlops"]["steps"] <= 10)
        self.assertTrue(20 <= m["medallion"]["marks"] <= 40)
        self.assertTrue(10 <= m["events"]["marks"] <= 22)
        self.assertEqual(m["events"]["labelled"], m["events"]["edges"])
        self.assertEqual(m["state"]["labelled"], m["state"]["edges"])
        self.assertIn("conditional", m["state"]["kinds"])

    def test_p26_passes_each_type(self):
        import p26_architecture
        from _shared import Ctx, run_gate
        code, found = run_gate(p26_architecture, self.out, Ctx(self.out, theme=SIGMOID))
        self.assertEqual(code, 0, [f.msg for f in found if f.level != "note"])
        notes = " ".join(f.msg for f in found)
        for t in ("cicd:", "mlops:", "dataflow:", "event:", "state:"):
            self.assertIn(t, notes)

    def test_cicd_refuses_a_gate_count_mismatch(self):
        spec = {"envs": [{}], "gates": [], "prod": {"waves": [()]}}
        with self.assertRaises(GeometryError):
            af.cicd(C, TH, spec)


class Layout(unittest.TestCase):
    def test_hinge_sits_on_the_divider(self):
        prs, s = slide()
        b = af.Board(C, TH, cols=[2, 1, 1, 2], rows=[af.HEAD_SECTION, "auto"], name="t")
        b.frame("ops", "segmented", "ops", cols=(0, 3), rows=(0, 1), cells=[("Build · dev", 2), ("Run · prod", 2)])
        b.mark("a", LAMBDA, "Lambda", cell=(0, 1))
        b.mark("h", "vendor:aws/ml/bedrock", "Bedrock", "prompt registry", cell=(1, 1), hinge=2)
        b.mark("c", S3, "Amazon S3", cell=(3, 1))
        b.edge("a", "h", "register")
        b.edge("h", "c", "promote")
        b.render(s)
        self.assertAlmostEqual(b.nodes["h"].body.cx, b.colx[2], places=6)
        rules = [sh for sh in s.shapes if ".cell.1.rule" in sh.name]
        self.assertTrue(rules)
        cap = b.nodes["h"].cap
        for r in rules:                                     # the divider breaks round the hinge's caption
            y0, y1 = r.top / 914400, (r.top + r.height) / 914400
            self.assertFalse(y0 < cap.cy < y1)

    def test_frames_of_one_row_share_a_foot(self):
        prs, s = slide()
        b = af.Board(C, TH, cols=[1, 1], rows=[af.HEAD_BAR, "auto", "auto"], name="t")
        b.frame("l", "region", "Left", cols=(0, 0), rows=(0, 2))
        b.frame("r", "region", "Right", cols=(1, 1), rows=(0, 2))
        b.mark("a", LAMBDA, "Lambda", cell=(0, 1))
        b.mark("b", S3, "Amazon S3", cell=(1, 2))
        b.edge("a", "b", "land", sides=("r", "l"))
        b.render(s)
        self.assertAlmostEqual(b.frames["l"].box.b, b.frames["r"].box.b, places=6)


if __name__ == "__main__":
    unittest.main()
