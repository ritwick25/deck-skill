"""Sequence, DFD and C4 diagrams (deckkit.arch_seq; rulings 20 and 21, spec 3.8, 3.10 and the C4 notes).

Run: cd ~/.claude/skills/deck && $PY -m unittest tests.test_arch_seq -v
Each diagram is drawn on a fresh WIDE slide in the sigmoid theme; checks read the drawn shapes back, and the four
example slides (scripts/example_arch_seq.py) are gated.
"""
from __future__ import annotations

import importlib.util
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(SKILL, "kit"))
sys.path.insert(0, os.path.join(SKILL, "gates"))

from pptx import Presentation  # noqa: E402

from deckkit import arch_seq, package, prims, tag  # noqa: E402
from deckkit.canvas import WIDE, Box  # noqa: E402
from deckkit.errors import Capacity, GeometryError, Overflows  # noqa: E402
from deckkit.read.ns import NS  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

A = "{%s}" % NS["a"]
BASE = os.path.join(SKILL, "kit", "assets", "base-wide.pptx")
SIGMOID = os.path.join(SKILL, "themes", "sigmoid.json")
TH = Theme.load(SIGMOID)
C = WIDE.content
EMU = 914400.0
M = {"app": "vendor:programming/framework/react", "api": "vendor:aws/network/api-gateway",
     "agent": "vendor:aws/compute/elastic-container-service-task", "model": "vendor:aws/ml/bedrock",
     "search": "vendor:aws/analytics/amazon-opensearch-service", "s3": "vendor:aws/storage/simple-storage-service-s3"}


def slide():
    prs = Presentation(BASE)
    prims.bind(TH)
    return prs, prs.slides.add_slide(prs.slide_layouts[0])


def tearDownModule():
    prims._ACTIVE.clear()


def box_of(sh) -> Box:
    return Box(sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU)


def named(s, prefix):
    return [sh for sh in s.shapes if sh.name.startswith(prefix)]


def tail(sh):
    t = sh._element.find(".//" + A + "ln/" + A + "tailEnd")
    return None if t is None else t.get("type")


def dash(sh):
    d = sh._element.find(".//" + A + "ln/" + A + "prstDash")
    return None if d is None else d.get("val")


def seq5():
    q = arch_seq.Sequence()
    q.participant("app", M["app"], "Field app", "React Native")
    q.participant("api", M["api"], "API Gateway", "brief API")
    q.participant("agent", M["agent"], "Writer agent", "LangGraph")
    q.participant("model", M["model"], "Bedrock", "Claude models")
    q.participant("search", M["search"], "OpenSearch", "claim index")
    q.band("Agent zone · private subnets", "api", "search")
    q.message("app", "api", "POST /briefs")
    q.message("api", "agent", "start run")
    q.message("agent", "model", "plan the brief")
    q.message("model", "agent", "tool_use: search", kind="reply")
    a = q.message("agent", "search", "search claims")
    b = q.message("search", "agent", "sources", kind="reply")
    c = q.message("agent", "api", "brief ready", kind="reply")
    e = q.message("agent", "model", "flag the claim", kind="async")
    d = q.message("agent", "api", "held", kind="reply")
    q.message("api", "app", "200 {brief_id}", kind="reply")
    q.fragment("loop", "each claim", a, b)
    q.fragment("alt", "all sourced", c, d, else_at=e, else_guard="a claim has no source")
    return q


class SequenceTest(unittest.TestCase):
    def test_numbered_labelled_messages_in_notation(self):
        prs, s = slide()
        q = seq5()
        m = q.render(s, C, TH)
        self.assertEqual((m["participants"], m["messages"], m["labelled"], m["fragments"], m["bands"]),
                         (5, 10, 10, 2, 1))
        labels = sorted(named(s, "seq.msg."), key=lambda sh: sh.name)
        texts = {sh.name: sh.text_frame.text for sh in labels if sh.name.endswith(".label")}
        self.assertEqual(len(texts), 10)
        for n in range(1, 11):
            self.assertTrue(texts["seq.msg.%d.label" % n].startswith("%d " % n), "numbered at the label start")
        lines = {sh.name: sh for sh in labels if not sh.name.endswith(".label")}
        self.assertEqual(tail(lines["seq.msg.1"]), "triangle")          # a call: filled head, solid
        self.assertIsNone(dash(lines["seq.msg.1"]))
        self.assertEqual((tail(lines["seq.msg.4"]), dash(lines["seq.msg.4"])), ("arrow", "sysDash"))   # a reply
        self.assertEqual((tail(lines["seq.msg.8"]), dash(lines["seq.msg.8"])), ("arrow", "dash"))      # an event

    def test_messages_land_on_activation_bars_and_stay_horizontal(self):
        prs, s = slide()
        seq5().render(s, C, TH)
        bars = [box_of(sh) for sh in named(s, "state.seq.act.")]
        self.assertGreaterEqual(len(bars), 5)
        for sh in named(s, "seq.msg."):
            if sh.name.endswith(".label"):
                continue
            x0, y0 = sh.begin_x / EMU, sh.begin_y / EMU
            x1, y1 = sh.end_x / EMU, sh.end_y / EMU
            self.assertAlmostEqual(y0, y1, places=4, msg="%s is horizontal" % sh.name)
            for x, y in ((x0, y0), (x1, y1)):
                self.assertTrue(any(abs(x - b.x) < 0.01 or abs(x - b.r) < 0.01 for b in bars
                                    if b.y - 0.01 <= y <= b.b + 0.01), "%s lands on a bar" % sh.name)

    def test_fragments_enclose_their_messages_and_sit_inside_the_band(self):
        prs, s = slide()
        q = seq5()
        q.render(s, C, TH)
        band = box_of(named(s, "seq.band")[0])
        for f in q.frags:
            fb = f.box
            self.assertTrue(band.x < fb.x and fb.r < band.r)
            for m in q.msgs[f.first - 1:f.last]:
                self.assertTrue(fb.y < m.y < fb.b and fb.x < min(m.x0, m.x1) and max(m.x0, m.x1) < fb.r)
                self.assertTrue(fb.x < m.lbox.x and m.lbox.r < fb.r)
            for m in q.msgs:
                if not f.first <= m.n <= f.last:
                    self.assertFalse(fb.y < m.y < fb.b, "message %d outside fragment %s" % (m.n, f.op))

    def test_labels_never_cover_an_activation_bar(self):
        prs, s = slide()
        q = seq5()
        q.render(s, C, TH)
        bars = [box_of(sh) for sh in named(s, "state.seq.act.")]
        for m in q.msgs:
            for b in bars:
                hit = m.lbox.x < b.r and b.x < m.lbox.r and m.lbox.y < b.b and b.y < m.lbox.b
                self.assertFalse(hit, "message %d label on a bar" % m.n)

    def test_refusals(self):
        q = arch_seq.Sequence()
        q.participant("a", M["app"], "A")
        q.participant("b", M["api"], "B")
        with self.assertRaises(Overflows):
            q.message("a", "b", "one two three four five six")
        with self.assertRaises(GeometryError):
            q.message("a", "b", "   ")
        with self.assertRaises(GeometryError):
            q.band("Agent zone", "a", "b")                          # no fact after a middot
        q.message("b", "a", "orphan reply", kind="reply")
        prs, s = slide()
        with self.assertRaises(GeometryError):                      # a reply with no open call
            q.render(s, C, TH)
        q2 = arch_seq.Sequence()
        q2.participant("a", M["app"], "A")
        q2.participant("b", M["api"], "B")
        q2.message("a", "b", "call, no reply")
        prs, s = slide()
        with self.assertRaises(GeometryError):                      # a call needs its reply ...
            q2.render(s, C, TH)
        q2.fragment("break", "policy denied", 1, 1)
        prs, s = slide()
        q2.render(s, C, TH)                                         # ... unless it sits in a break
        q3 = arch_seq.Sequence()
        for k in range(8):
            q3.participant("p%d" % k, M["app"], "P%d" % k)
        with self.assertRaises(Capacity):
            q3.participant("p8", M["app"], "P8")

    def test_fragment_across_the_band_is_refused(self):
        q = arch_seq.Sequence()
        q.participant("app", M["app"], "Field app")
        q.participant("api", M["api"], "API Gateway")
        q.participant("agent", M["agent"], "Writer agent")
        q.band("Agent zone · private", "api", "agent")
        q.message("app", "api", "request")
        q.message("api", "app", "response", kind="reply")
        q.fragment("opt", "cached", 1, 2)
        prs, s = slide()
        with self.assertRaises(GeometryError):
            q.render(s, C, TH)

    def test_nested_fragments_and_overlap_refused(self):
        q = seq5()
        with self.assertRaises(GeometryError):
            q.fragment("opt", "partial", 6, 7)                      # neither nested in nor apart from the loop
        with self.assertRaises(GeometryError):
            q.fragment("opt", "x", 1, 3, else_at=2, else_guard="y")  # else only on alt or par

    def test_overflow_raises_never_shrinks(self):
        q = arch_seq.Sequence()
        q.participant("a", M["app"], "A")
        q.participant("b", M["api"], "B")
        for k in range(8):
            q.message("a", "b", "call %d" % k)
            q.message("b", "a", "reply %d" % k, kind="reply")
        prs, s = slide()
        with self.assertRaises(Overflows):
            q.render(s, Box(C.x, C.y, C.w, 3.0), TH)


def dfd3():
    d = arch_seq.DFD(grid=(4, 2), weights=[2.0, 3.0, 3.0, 2.0])
    d.boundary("acct", "AWS account · production", at=(1, 0, 2, 2))
    d.entity("web", "Public web", "press", at=(0, 0), mark="glyph:globe")
    d.entity("rep", "Field rep", "reads briefs", at=(3, 0), mark="glyph:user")
    d.process("collect", "1", "Collect signals", at=(1, 0), tech="AWS Lambda")
    d.process("draft", "2", "Draft the brief", at=(2, 0), tech="Writer agent")
    d.store("lake", "D1", "Signal lake", at=(1.5, 1), mark=M["s3"])
    d.flow("web", "collect", "web pages", control="TLS 1.3")
    d.flow("collect", "lake", "raw signals")
    d.flow("lake", "draft", "sources")
    d.flow("draft", "rep", "brief", control="OIDC")
    return d


class DFDTest(unittest.TestCase):
    def test_numbered_flows_markers_and_chips(self):
        prs, s = slide()
        m = dfd3().render(s, C, TH)
        self.assertEqual((m["flows"], m["labelled"], m["steps"], m["crossing_flows"], m["chips"]), (4, 4, 4, 2, 2))
        self.assertEqual(m["markers"], 2)
        steps = sorted(int(sh.text_frame.text) for sh in named(s, "dfd.step."))
        self.assertEqual(steps, [1, 2, 3, 4])
        bound = box_of(named(s, "dfd.bound.acct")[0])
        for mk in named(s, "state.dfd.cross."):
            b = box_of(mk)
            on_edge = min(abs(b.cx - bound.x), abs(b.cx - bound.r), abs(b.cy - bound.y), abs(b.cy - bound.b))
            self.assertLess(on_edge, 0.01, "a crossing marker sits on the boundary edge")
        for chip in named(s, "dfd.control."):
            cb = box_of(chip)
            inside = bound.x < cb.x and cb.r < bound.r and bound.y < cb.y and cb.b < bound.b
            apart = cb.r < bound.x or cb.x > bound.r or cb.b < bound.y or cb.y > bound.b
            self.assertTrue(inside or apart, "a chip never straddles a boundary")

    def test_dfd_rules_refuse(self):
        d = dfd3()
        d.flow("web", "rep", "gossip")
        prs, s = slide()
        with self.assertRaises(GeometryError):                       # entity to entity
            d.render(s, C, TH)
        d = dfd3()
        d.flows[0].control = None
        prs, s = slide()
        with self.assertRaises(GeometryError):                       # a crossing without its control
            d.render(s, C, TH)
        d = dfd3()
        d.flow("collect", "draft", "signals", control="HTTPS")       # a control on a flow that crosses nothing
        prs, s = slide()
        with self.assertRaises(GeometryError):
            d.render(s, C, TH)
        with self.assertRaises(GeometryError):
            dfd3().boundary("z", "Agent zone", at=(1, 1, 1, 1))      # a bare dashed zone
        with self.assertRaises(GeometryError):
            dfd3().process("p", "x", "Bad id", at=(1, 1))
        c = arch_seq.DFD(grid=(3, 1), level="context")
        c.entity("a", "A", at=(0, 0))
        c.process("p", "0", "System", at=(1, 0))
        c.store("s", "D1", "Store", at=(2, 0))
        prs, s = slide()
        with self.assertRaises(GeometryError):                       # a context DFD carries no store
            c.render(s, C, TH)

    def test_draft_lists_problems_instead_of_raising(self):
        d = dfd3()
        d.flows[0].label = "a label far too long to sit"
        prs, s = slide()
        m = d.render(s, Box(C.x, C.y, 7.0, 3.0), TH, draft=True)
        self.assertTrue(m["problems"])


def c4ctx():
    c = arch_seq.C4("context", grid=(3, 2), weights=[3.0, 4.0, 3.0])
    c.boundary("ent", "Quillmont · enterprise", at=(1, 0, 2, 2), kind="enterprise")
    c.person("rep", "Field rep", "Reads the brief", at=(2, 0))
    c.system("svc", "Brief service", "Drafts a weekly brief", at=(1, 1))
    c.external("web", "Public web", "Press, trials", at=(0, 1))
    c.rel("rep", "svc", "reads · HTTPS")
    c.rel("svc", "web", "collects")
    return c


class C4Test(unittest.TestCase):
    def test_context_view(self):
        prs, s = slide()
        m = c4ctx().render(s, C, TH, name="c4")
        self.assertEqual((m["flows"], m["labelled"], m["markers"], m["chips"], m["steps"]), (2, 2, 0, 0, 0))
        self.assertTrue(named(s, "c4.node.svc"))
        self.assertTrue(named(s, "state.c4.legend.key.person"))
        self.assertTrue(named(s, "c4.legend.key.person.text"))

    def test_c4_rules_refuse(self):
        with self.assertRaises(GeometryError):
            c4ctx().container("x", "X", "Lambda", "does x", at=(1, 0))   # containers are one level down
        c = arch_seq.C4("context", grid=(2, 1))
        c.person("p", "P", "uses", at=(0, 0))
        prs, s = slide()
        with self.assertRaises(GeometryError):                           # one system in focus
            c.render(s, C, TH)
        c = arch_seq.C4("container", grid=(2, 1))
        with self.assertRaises(GeometryError):
            c.container("x", "X", None, "does x", at=(0, 0))            # technology is required
        c.container("x", "X", "Lambda", "does x", at=(0, 0))
        prs, s = slide()
        with self.assertRaises(GeometryError):                           # the system boundary is required
            c.render(s, C, TH)


def _example():
    spec = importlib.util.spec_from_file_location("example_arch_seq", os.path.join(SKILL, "scripts",
                                                                                   "example_arch_seq.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class ExampleGatesTest(unittest.TestCase):
    """The four example slides pass the layout, style and colour gates; P26 is held on the sequence and the DFD
    (its C4 rules are gate track B/E1's to settle, see docs/notes/E3.md)."""

    @classmethod
    def setUpClass(cls):
        cls.dir = tempfile.mkdtemp(prefix="arch-seq-")
        ex = _example()
        prs = Presentation(BASE)
        prims.bind(TH)
        cls.metrics = {"seq": ex.sequence(prs, TH), "dfd": ex.dfd(prs, TH), "c4ctx": ex.c4_context(prs, TH),
                       "c4ctr": ex.c4_container(prs, TH)}
        cls.deck = os.path.join(cls.dir, "all.pptx")
        package.save(prs, cls.deck, title="all", theme=TH, overwrite=True, date="2026-10-04")
        prs2 = Presentation(BASE)
        ex.sequence(prs2, TH)
        ex.dfd(prs2, TH)
        cls.deck2 = os.path.join(cls.dir, "two.pptx")
        package.save(prs2, cls.deck2, title="two", theme=TH, overwrite=True, date="2026-10-04")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, ignore_errors=True)

    def _gate(self, name, deck):
        from _shared import FAIL, Ctx, run_gate
        mod = __import__(name)
        ctx = Ctx(deck, theme=SIGMOID)
        try:
            _code, found = run_gate(mod, deck, ctx)
        finally:
            ctx.close()
        return [f.line() for f in found if f.level == FAIL]

    def test_layout_and_style_gates(self):
        for g in ("p01_floor", "p03_bounds", "p04_padding", "p05_overlap", "p06_arrows", "p07_row", "p08_balance",
                  "p09_chrome", "p10_density", "p11_style", "p14_titles", "p16_colour", "p20_notes", "p22_icons",
                  "p24_fill", "p25_contrast"):
            self.assertEqual(self._gate(g, self.deck), [], g)

    def test_p26_on_sequence_and_dfd(self):
        self.assertEqual(self._gate("p26_architecture", self.deck2), [])

    def test_title_counts_come_from_the_drawing(self):
        self.assertEqual(self.metrics["dfd"]["crossing_flows"], 6)
        self.assertEqual(self.metrics["seq"]["labelled"], self.metrics["seq"]["messages"])


if __name__ == "__main__":
    unittest.main()
