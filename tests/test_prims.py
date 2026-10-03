"""K3: prims, glyphs, icons.

Run: cd ~/.claude/skills/deck && $PY -m unittest tests.test_prims -v
The fixture is the primitive sheet (scripts/render_prims.py build), one slide with every creator, built
on both themes into a temp dir. Refusals assert the typed error; each structural check is also shown
failing once on a planted copy so it is not tautological.
"""
from __future__ import annotations

import copy
import importlib.util
import io
import math
import os
import shutil
import sys
import tempfile
import unittest
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(SKILL, "kit"))

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402

from deckkit import glyphs, icons, package, prims  # noqa: E402
from deckkit.canvas import Box  # noqa: E402
from deckkit.errors import (BelowFloor, GeometryError, Overflows, ThemeInvariant,  # noqa: E402
                            UnsupportedPart)
from deckkit.read.ns import NS  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

A = "{%s}" % NS["a"]
P = "{%s}" % NS["p"]
BASE = os.path.join(SKILL, "kit", "assets", "base-wide.pptx")
SIGMOID = Theme.load(os.path.join(SKILL, "themes", "sigmoid.json"))
NEUTRAL = Theme.load(os.path.join(SKILL, "themes", "neutral.json"))
SHAPES = (P + "sp", P + "cxnSp", P + "pic")
EFFECTS = {"outerShdw", "innerShdw", "prstShdw", "glow", "softEdge", "reflection", "blur", "effectDag"}

_spec = importlib.util.spec_from_file_location("render_prims", os.path.join(SKILL, "scripts", "render_prims.py"))
render_prims = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(render_prims)


def new_slide():
    prs = Presentation(BASE)
    return prs, prs.slides.add_slide(prs.slide_layouts[0])


def slide_xml(pptx_path):
    with zipfile.ZipFile(pptx_path) as z:
        return etree.fromstring(z.read("ppt/slides/slide1.xml"))


# ------------------------------------------------------------------ structural checks (also used on plants)

def style_defects(root):
    return [s for s in root.iter(*SHAPES) if s.find(P + "style") is not None]


def effect_defects(root):
    out = []
    for s in root.iter(*SHAPES):
        sp = s.find(P + "spPr")
        lst = [c for c in sp if etree.QName(c).localname in ("effectLst", "effectDag")]
        if len(lst) != 1 or etree.QName(lst[0]).localname != "effectLst" or len(lst[0]):
            out.append(s)
    out += [e for e in root.iter() if isinstance(e.tag, str) and etree.QName(e).localname in EFFECTS]
    return out


class Fixture(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="k3-")
        cls.paths = {}
        for th in (SIGMOID, NEUTRAL):
            p = os.path.join(cls.tmp, "prims-%s.pptx" % th.name)
            render_prims.build(th, __import__("pathlib").Path(p))
            cls.paths[th.name] = p
        cls.root = slide_xml(cls.paths["sigmoid"])

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)


class TestSheet(Fixture):
    def test_no_shape_carries_p_style(self):
        self.assertEqual(style_defects(self.root), [])
        n = len(list(self.root.iter(*SHAPES)))
        self.assertGreater(n, 60)

    def test_p_style_check_fails_on_a_plant(self):
        r = copy.deepcopy(self.root)
        sp = next(r.iter(P + "sp"))
        sp.append(etree.Element(P + "style"))
        self.assertEqual(len(style_defects(r)), 1)

    def test_only_one_empty_effect_list(self):
        self.assertEqual(effect_defects(self.root), [])

    def test_effect_check_fails_on_plants(self):
        r = copy.deepcopy(self.root)
        sps = list(r.iter(P + "sp"))
        sps[0].find(P + "spPr").append(etree.Element(A + "effectLst"))        # doubled
        lst = sps[1].find(P + "spPr").find(A + "effectLst")
        etree.SubElement(lst, A + "outerShdw", blurRad="40000")                # blurred shadow
        self.assertGreaterEqual(len(effect_defects(r)), 2)

    def test_no_gradient_no_autofit_no_srgb_on_kit_shapes(self):
        tags = {etree.QName(e).localname for e in self.root.iter() if isinstance(e.tag, str)}
        self.assertNotIn("gradFill", tags)
        self.assertNotIn("normAutofit", tags)
        self.assertNotIn("spAutoFit", tags)
        self.assertNotIn("srgbClr", tags)           # every kit colour is a schemeClr role
        for body in self.root.iter(P + "txBody"):
            self.assertIsNotNone(body.find(A + "bodyPr/" + A + "noAutofit"))

    def test_every_run_explicit_size_colour_face_at_or_above_floor(self):
        runs = list(self.root.iter(A + "r"))
        self.assertGreater(len(runs), 20)
        for r in runs:
            rpr = r.find(A + "rPr")
            self.assertIsNotNone(rpr.get("sz"))
            self.assertGreaterEqual(int(rpr.get("sz")), 800)
            self.assertEqual(int(rpr.get("sz")) % 50, 0)
            self.assertIsNotNone(rpr.find(A + "solidFill/" + A + "schemeClr"))
            self.assertIsNotNone(rpr.find(A + "latin"))

    def test_every_shape_named_and_no_mono(self):
        for nv in self.root.iter(P + "cNvPr"):
            if etree.QName(nv.getparent().getparent()).localname == "spTree":
                continue                                   # the slide tree's own record, not a shape
            nm = nv.get("name")
            self.assertTrue(nm, "unnamed shape")
            self.assertNotRegex(nm, r"^(Rectangle|Oval|Straight Connector|Freeform|TextBox|Picture|Group) \d+$")
        faces = {e.get("typeface") for e in self.root.iter(A + "latin")}
        self.assertFalse(faces & {"Consolas", "Courier New", "JetBrains Mono", "Menlo"})

    def test_only_state_shapes_are_bars(self):
        """No strip <= prims.BAR_MAX_IN (0.10in) thin and 6x long filled or outlined in a non-neutral role
        unless named state.* (the quote bar)."""
        neutral = {SIGMOID.spec(r).get("val") + repr([(c.tag, c.get("val")) for c in SIGMOID.spec(r)])
                   for r in prims.NEUTRAL_STRIP}
        bars = []
        for s in self.root.iter(P + "sp"):
            ext = s.find(P + "spPr/" + A + "xfrm/" + A + "ext")
            colours = s.findall(P + "spPr/" + A + "solidFill/" + A + "schemeClr") + \
                s.findall(P + "spPr/" + A + "ln/" + A + "solidFill/" + A + "schemeClr")
            if ext is None or not colours:
                continue
            keys = {c.get("val") + repr([(m.tag, m.get("val")) for m in c]) for c in colours}
            w, h = int(ext.get("cx")) / 914400, int(ext.get("cy")) / 914400
            if min(w, h) <= prims.BAR_MAX_IN + 1e-6 and max(w, h) >= 6 * min(w, h) and keys - neutral:
                bars.append(s.find(".//" + P + "cNvPr").get("name"))
        self.assertEqual(prims.BAR_MAX_IN, 0.10)
        self.assertEqual(bars, ["state.quote"])

    def test_cards_rounded_with_physical_radius(self):
        prs = Presentation(self.paths["sigmoid"])
        from deckkit.geom import find
        sh = find(prs.slides[0], "row1.card.filled")
        self.assertAlmostEqual(sh.adjustments[0] * min(sh.width, sh.height) / 914400, 0.05, places=3)

    def test_xsd_both_themes(self):
        for name, p in self.paths.items():
            ok, msg = package.validate_xsd(p)
            self.assertTrue(ok, "%s: %s" % (name, msg[:400]))

    def test_theme_swap_changes_only_the_theme(self):
        a, b = slide_xml(self.paths["sigmoid"]), slide_xml(self.paths["neutral"])
        strip = lambda r: etree.tostring(r)  # noqa: E731
        self.assertEqual(strip(a), strip(b))  # same slide XML: the colour lives in the theme part


# ------------------------------------------------------------------ refusals

class TestRefusals(unittest.TestCase):
    def setUp(self):
        self.prs, self.s = new_slide()
        self.ctx = prims.use(SIGMOID)
        self.ctx.__enter__()

    def tearDown(self):
        self.ctx.__exit__(None, None, None)

    def test_radius_on_rect_raises(self):
        with self.assertRaises(GeometryError):
            prims.rect(self.s, Box(1, 1, 2, 1), fill="card", radius_in=0.1, name="r")

    def test_card_radius_becomes_normalised_adjustment(self):
        c = prims.card(self.s, Box(1, 1, 3, 1), radius_in=0.25, name="c")
        self.assertAlmostEqual(c.adjustments[0], 0.25, places=4)
        c = prims.card(self.s, Box(1, 1, 3, 0.4), radius_in=1.0, name="c2")
        self.assertAlmostEqual(c.adjustments[0], 0.5, places=4)

    def test_outline_card_on_white_needs_rule_strong(self):
        with self.assertRaises(ThemeInvariant):
            prims.card(self.s, Box(1, 1, 2, 1), fill="paper", line="hairline", name="c")
        prims.card(self.s, Box(1, 1, 2, 1), fill="paper", line="rule-strong", name="c")

    def test_accent_bar_refused_unless_state(self):
        with self.assertRaises(GeometryError):
            prims.rect(self.s, Box(1, 1, 0.05, 1.5), fill="accent", name="card.edge")
        with self.assertRaises(GeometryError):
            prims.rect(self.s, Box(1, 1, 2.0, 0.04), fill="accent", name="card.top")
        prims.rect(self.s, Box(1, 1, 0.05, 1.5), fill="accent", name="state.active")

    def test_below_floor_raises(self):
        with self.assertRaises(BelowFloor):
            prims.caps_label(self.s, Box(1, 1, 2, 0.3), "small", pt=7.5, name="l")
        with self.assertRaises(BelowFloor):
            prims.disc(self.s, Box(1, 1, 0.2, 0.2), 1, name="d")       # derived size 6pt

    def test_off_ladder_size_raises(self):
        with self.assertRaises(GeometryError):
            prims.caps_label(self.s, Box(1, 1, 2, 0.3), "x", pt=9.3, name="l")

    def test_overflow_raises_never_shrinks(self):
        with self.assertRaises(Overflows):
            prims.pill(self.s, Box(1, 1, 0.5, 0.25), "A far too long tag for this pill", name="p")
        with self.assertRaises(Overflows):
            prims.stat(self.s, Box(1, 1, 1.5, 0.6), "42%", "A label", kind="measured", name="s")

    def test_step_depth_band(self):
        for d in (0.25, 0.4):
            with self.assertRaises(GeometryError):
                prims.step_header(self.s, Box(1, 1, 6, 0.5), ["A", "B"], depth=d, name="st")
        g = prims.step_header(self.s, Box(1, 1, 6, 0.5), ["A", "B", "C"], depth=0.31, name="st")
        kids = [c for c in g._element if etree.QName(c).localname == "sp"]
        prst = [k.find(".//" + A + "prstGeom").get("prst") for k in kids]
        self.assertEqual(prst, ["homePlate", "chevron", "chevron"])
        for k in kids:
            ext = k.find(".//" + A + "ext")
            w, h = int(ext.get("cx")), int(ext.get("cy"))
            adj = int(k.find(".//" + A + "gd").get("fmla").split()[1])
            depth = adj / 100000 * min(w, h) / h
            self.assertTrue(0.31 - 1e-3 <= depth <= 0.35 + 1e-3, depth)

    def test_stat_needs_kind_and_labels_illustrative(self):
        with self.assertRaises(GeometryError):
            prims.stat(self.s, Box(1, 1, 2, 1.3), "3x", "Faster", kind="guess", name="s")
        g = prims.stat(self.s, Box(1, 1, 2, 1.3), "3x", "Faster", kind="illustrative", name="s")
        texts = [t.text for t in g._element.iter(A + "t")]
        self.assertIn("Illustrative", texts)
        names = [n.get("name") for n in g._element.iter(P + "cNvPr")]
        self.assertEqual(names[0], "s")
        # the rule sits between the numeral and the label: above the label
        from deckkit.geom import find
        rule, label = find(self.s, "s.rule"), find(self.s, "s.label")
        self.assertLessEqual(rule.top, label.top)

    def test_disc_needs_square(self):
        with self.assertRaises(GeometryError):
            prims.disc(self.s, Box(1, 1, 0.5, 0.4), 1, name="d")

    def test_connect_unglued_with_tail(self):
        c = prims.connect(self.s, (1, 1), (3, 1), name="a")
        self.assertIsNone(c._element.find(".//" + A + "stCxn"))
        self.assertIsNotNone(c._element.find(".//" + A + "tailEnd"))
        with self.assertRaises(GeometryError):
            prims.connect(self.s, (1, 1), (1, 1), name="z")

    def test_elbow_orthogonal_only(self):
        with self.assertRaises(GeometryError):
            prims.elbow(self.s, [(1, 1), (2, 2)], name="e")
        e = prims.elbow(self.s, [(1, 1), (1, 2), (3, 2)], name="e")
        self.assertEqual(len(e._element.findall(".//" + A + "lnTo")), 2)

    def test_unknown_role_and_unbound_theme(self):
        with self.assertRaises(ThemeInvariant):
            prims.rect(self.s, Box(1, 1, 1, 1), fill="purple", name="r")
        self.ctx.__exit__(None, None, None)
        try:
            with self.assertRaises(ThemeInvariant):
                prims.rect(self.s, Box(1, 1, 1, 1), fill="card", name="r")
        finally:
            self.ctx = prims.use(SIGMOID)
            self.ctx.__enter__()

    def test_name_required(self):
        with self.assertRaises(GeometryError):
            prims.card(self.s, Box(1, 1, 1, 1), name="")


# ------------------------------------------------------------------ glyphs

class TestGlyphs(unittest.TestCase):
    def test_catalogue_size_and_licence(self):
        n = len(glyphs.names())
        self.assertTrue(140 <= n <= 220, n)
        lic = os.path.join(SKILL, "kit", "assets", "LICENSE-lucide.txt")
        with open(lic, encoding="utf-8") as f:
            self.assertIn("ISC License", f.read())

    def test_every_glyph_produces_a_valid_custgeom(self):
        for name in glyphs.names():
            groups = glyphs.paths(name)
            self.assertEqual(glyphs.validate_paths(groups), [], name)
            g = glyphs.custgeom(name)
            paths = g.findall(A + "pathLst/" + A + "path")
            self.assertTrue(paths, name)
            for p, grp in zip(paths, groups):
                self.assertEqual((p.get("w"), p.get("h")), ("24000", "24000"))
                self.assertEqual(p.get("fill") == "none", not grp["fill"], name)
                kids = [etree.QName(c).localname for c in p]
                self.assertEqual(kids[0], "moveTo", name)
                self.assertTrue(set(kids) <= {"moveTo", "lnTo", "cubicBezTo", "close"}, name)
                self.assertEqual(kids.count("moveTo"), len(grp["d"]), name)
                for c in p.iter(A + "cubicBezTo"):
                    self.assertEqual(len(c), 3)
                for pt in p.iter(A + "pt"):
                    for k in ("x", "y"):
                        self.assertTrue(-2000 <= int(pt.get(k)) <= 26000, (name, pt.get(k)))

    def test_validator_fails_on_plants(self):
        self.assertTrue(glyphs.validate_paths([{"fill": False, "d": [[["L", 1, 1]]]}]))
        self.assertTrue(glyphs.validate_paths([{"fill": False, "d": [[["M", 1, 1], ["A", 1, 1, 0, 0, 1, 2, 2]]]}]))
        self.assertTrue(glyphs.validate_paths([{"fill": False, "d": [[["M", 1, 1], ["L", float("nan"), 1]]]}]))
        self.assertTrue(glyphs.validate_paths([]))

    def test_all_glyphs_on_one_slide_pass_xsd(self):
        prs, s = new_slide()
        with prims.use(NEUTRAL):
            for i, n in enumerate(glyphs.names()):
                glyphs.glyph(s, n, Box(0.3 + (i % 26) * 0.49, 0.3 + (i // 26) * 0.6, 0.36, 0.36))
        root = s._element
        self.assertEqual(style_defects(root), [])
        self.assertEqual(effect_defects(root), [])
        for sp in root.iter(P + "sp"):
            ln = sp.find(P + "spPr/" + A + "ln")
            self.assertEqual(ln.get("cap"), "rnd")
            self.assertIsNotNone(ln.find(A + "round"))
            self.assertIsNone(sp.find(P + "txBody"))
        with tempfile.TemporaryDirectory() as td:
            p = os.path.join(td, "glyphs.pptx")
            package.save(prs, p, title="Glyphs", theme=NEUTRAL, overwrite=True)
            ok, msg = package.validate_xsd(p)
            self.assertTrue(ok, msg[:400])

    def test_aliases_and_unknown(self):
        self.assertEqual(glyphs.resolve("alert-triangle"), "triangle-alert")
        self.assertEqual(glyphs.resolve("bar-chart-2"), "chart-no-axes-column")
        self.assertEqual(glyphs.resolve("home"), "house")
        with self.assertRaises(glyphs.UnknownGlyph) as cm:
            glyphs.resolve("databse")
        self.assertIn("database", str(cm.exception))
        self.assertIsInstance(cm.exception, UnsupportedPart)

    def test_stroke_scales_with_side(self):
        self.assertAlmostEqual(glyphs.stroke_for(0.4), 2.4)
        self.assertAlmostEqual(glyphs.stroke_for(0.2), 1.2)
        prs, s = new_slide()
        with prims.use(SIGMOID):
            g = glyphs.glyph(s, "database", Box(1, 1, 0.6, 0.3))
            self.assertEqual(g.width, g.height)                    # square, centred in the box
            self.assertEqual(int(g._element.find(".//" + A + "ln").get("w")), int(round(1.8 * 12700)))
            with self.assertRaises(GeometryError):
                glyphs.glyph(s, "database", Box(1, 1, 0.1, 0.1))

    def test_path_parser(self):
        subs = glyphs.parse_path("M2 2h4v4H2z m10 0 l2 2 2-2")
        self.assertEqual(subs[0], [("M", 2, 2), ("L", 6, 2), ("L", 6, 6), ("L", 2, 6), ("Z",)])
        self.assertEqual(subs[1], [("M", 12, 2), ("L", 14, 4), ("L", 16, 2)])
        subs = glyphs.parse_path("M1 1a1 1 0 012 0")                # compact flags
        self.assertEqual(subs[0][-1][-2:], (3.0, 1.0))
        subs = glyphs.parse_path("M0 0Q2 4 4 0T8 0")                # T reflects Q's control point
        self.assertEqual([o[0] for o in subs[0]], ["M", "C", "C"])
        self.assertAlmostEqual(subs[0][2][2], -8 / 3)

    def test_arc_quarter_circle_is_accurate(self):
        ops = glyphs.arc_to_cubics(10, 0, 10, 10, 0, 0, 1, 0, 10)   # centre (0,0) clockwise in y-down
        self.assertEqual(len(ops), 1)
        _, x1, y1, x2, y2, x, y = ops[0]
        t = 0.5                                                    # midpoint must lie on the circle
        mx = (1 - t) ** 3 * 10 + 3 * (1 - t) ** 2 * t * x1 + 3 * (1 - t) * t ** 2 * x2 + t ** 3 * x
        my = 3 * (1 - t) ** 2 * t * y1 + 3 * (1 - t) * t ** 2 * y2 + t ** 3 * y
        self.assertAlmostEqual(math.hypot(mx, my), 10, delta=0.003)
        self.assertEqual(len(glyphs.arc_to_cubics(0, 0, 5, 5, 0, 1, 1, 0.01, 0)), 4)   # large arc


# ------------------------------------------------------------------ icons

class TestIcons(unittest.TestCase):
    def setUp(self):
        self.prs, self.s = new_slide()

    def test_vendor_keeps_hue_alt_and_raster_floor(self):
        p = icons.vendor(self.s, "aws/compute/ec2", Box(1, 1, 0.5, 0.5))
        self.assertEqual(p._element.find(".//" + P + "cNvPr").get("descr"), "Ec2 (vendor mark aws/compute/ec2)")
        self.assertEqual(p._element.find(".//" + P + "cNvPr").get("name"), "vendor.aws/compute/ec2")
        self.assertEqual(effect_defects(self.s._element), [])
        with self.assertRaises(GeometryError):
            icons.vendor(self.s, "aws/compute/ec2", Box(1, 1, 2.5, 2.5))
        with self.assertRaises(UnsupportedPart):
            icons.vendor(self.s, "aws/compute/ec3", Box(1, 1, 0.5, 0.5))

    def test_svg_picture_has_both_blips_and_survives_save(self):
        with self.assertRaises(UnsupportedPart):
            icons.svg_picture(self.s, render_prims.SAMPLE_SVG, Box(1, 1, 0.5, 0.5), None, name="i", alt="x")
        pic = icons.svg_picture(self.s, render_prims.SAMPLE_SVG, Box(1, 1, 0.5, 0.5), "auto", name="i", alt="Mark")
        blip = pic._element.find(".//" + A + "blip")
        R = "{%s}" % NS["r"]
        self.assertTrue(blip.get(R + "embed"))
        sb = blip.find(".//{%s}svgBlip" % icons.ASVG_NS)
        self.assertTrue(sb.get(R + "embed"))
        png = self.s.part.related_part(blip.get(R + "embed")).blob
        self.assertGreaterEqual(icons._png_size(png)[0], 0.5 * 96 * 4)
        with tempfile.TemporaryDirectory() as td:
            out = os.path.join(td, "svg.pptx")
            package.save(self.prs, out, title="Svg", theme=SIGMOID, overwrite=True)
            self.assertEqual(package.structural_check(out), [])
            with zipfile.ZipFile(out) as z:
                self.assertIn(b'Extension="svg"', z.read("[Content_Types].xml"))
            # fix_svg_ct repairs a package that lost the Default, and reports a no-op otherwise
            self.assertFalse(icons.fix_svg_ct(out))
            with zipfile.ZipFile(out) as z:
                m = {i.filename: z.read(i.filename) for i in z.infolist()}
            m["[Content_Types].xml"] = m["[Content_Types].xml"].replace(
                b'<Default Extension="svg" ContentType="image/svg+xml"/>', b"")
            package.write_members(m, out)
            self.assertTrue(icons.fix_svg_ct(out))

    def test_monogram(self):
        with prims.use(SIGMOID):
            m = icons.monogram(self.s, Box(1, 1, 0.5, 0.5), "Dx", "accent", name="m")
            self.assertEqual(m._element.find(".//" + A + "prstGeom").get("prst"), "roundRect")
            with self.assertRaises(GeometryError):
                icons.monogram(self.s, Box(1, 1, 0.5, 0.5), "ABCD", name="m2")
            with self.assertRaises(BelowFloor):
                icons.monogram(self.s, Box(1, 1, 0.2, 0.2), "AB", name="m3")

    def test_rasterise_keeps_aspect(self):
        png = icons.rasterise(b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 40 20"><rect width="40" '
                              b'height="20" fill="#000"/></svg>', 200)
        self.assertEqual(icons._png_size(png), (200, 100))


if __name__ == "__main__":
    unittest.main()
