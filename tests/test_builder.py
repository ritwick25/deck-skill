"""Builder fixes: transactional creators, one text writer, one floor, finish via OxmlElement, the strip rule,
the WIDE title box and the docProps date.

Run: cd ~/.claude/skills/deck && $PY -m unittest tests.test_builder -v
Every refusal test compares the whole slide tree (serialised) and the slide part's relationships before
and after, so "unchanged" means byte-for-byte, not a shape count.
"""
from __future__ import annotations

import dataclasses
import json
import os
import re
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(SKILL, "kit"))

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.oxml.xmlchemy import OxmlElement  # noqa: E402

from deckkit import canvas, glyphs, icons, measure, ooxml, package, prims, text  # noqa: E402
from deckkit.canvas import STD, WIDE, Box  # noqa: E402
from deckkit.errors import BelowFloor, DeckkitError, GeometryError, Overflows, ThemeInvariant  # noqa: E402
from deckkit.read.ns import NS  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

A = "{%s}" % NS["a"]
P = "{%s}" % NS["p"]
KIT = os.path.join(SKILL, "kit", "deckkit")
BASE = os.path.join(SKILL, "kit", "assets", "base-wide.pptx")
SIGMOID = Theme.load(os.path.join(SKILL, "themes", "sigmoid.json"))
SVG = (b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 48 48"><rect x="4" y="4" width="40" height="40" '
       b'fill="#4B5960"/></svg>')
LONG = "A label far too long for the narrow box it was given on this slide"


def new_slide():
    prs = Presentation(BASE)
    return prs, prs.slides.add_slide(prs.slide_layouts[0])


def state(slide):
    return etree.tostring(slide.shapes._spTree), tuple(sorted(slide.part.rels.keys()))


class Base(unittest.TestCase):
    def setUp(self):
        self.prs, self.s = new_slide()
        self.ctx = prims.use(SIGMOID)
        self.ctx.__enter__()

    def tearDown(self):
        self.ctx.__exit__(None, None, None)

    def unchanged(self, err, fn, *a, **k):
        """fn(*a, **k) raises err and leaves the slide tree and its relationships exactly as they were."""
        before = state(self.s)
        n = len(self.s.shapes)
        with self.assertRaises(err):
            fn(*a, **k)
        self.assertEqual(len(self.s.shapes), n, "a refused creator left a shape behind")
        self.assertEqual(state(self.s), before, "a refused creator changed the slide")


def _raise_floor(*a, **k):
    raise BelowFloor("forced", need=8.0, have=7.5)


def _raise_overflow(*a, **k):
    raise Overflows("forced", need=1.0, have=0.5)


def _nth(n, real, err=_raise_overflow):
    """A stand-in that runs `real` until its n-th call, which raises (a failure late in a compound)."""
    calls = {"k": 0}

    def f(*a, **k):
        calls["k"] += 1
        if calls["k"] >= n:
            err()
        return real(*a, **k)
    return f


# ------------------------------------------------------------------ 1. transactional creators

class Transactional(Base):
    CREATORS = [(prims, n) for n in ("rect", "card", "tab", "pill", "step_header", "disc", "person", "hairline",
                                     "connect", "elbow", "caps_label", "stat", "quote", "zone")] + \
               [(text, "textbox"), (text, "title"), (icons, "vendor"), (icons, "svg_picture"),
                (icons, "monogram"), (glyphs, "glyph")]

    def test_every_creator_is_wrapped_in_a_transaction(self):
        for mod, name in self.CREATORS:
            fn = getattr(mod, name)
            label = "%s.%s" % (mod.__name__, name)
            self.assertTrue(hasattr(fn, "__wrapped__"), label)
            self.assertIs(fn.__code__, ooxml.transactional(lambda s: s).__code__, label)

    # real refusals: the text is measured before the shape exists
    def test_text_creators_refuse_on_real_input(self):
        U = self.unchanged
        U(Overflows, prims.tab, self.s, Box(1, 1, 0.8, 0.3), LONG, name="t")
        U(BelowFloor, prims.tab, self.s, Box(1, 1, 2, 0.3), "Tab", pt=7.5, name="t")
        U(Overflows, prims.pill, self.s, Box(1, 1, 0.5, 0.25), LONG, name="p")
        U(BelowFloor, prims.pill, self.s, Box(1, 1, 2, 0.3), "Tag", pt=7.5, name="p")
        U(BelowFloor, prims.disc, self.s, Box(1, 1, 0.2, 0.2), 1, name="d")
        U(Overflows, prims.step_header, self.s, Box(1, 1, 6, 0.5), ["Scope", "Build", LONG], name="st")
        U(BelowFloor, prims.step_header, self.s, Box(1, 1, 6, 0.5), ["A", "B"], pt=7.5, name="st")
        U(Overflows, prims.caps_label, self.s, Box(1, 1, 1.0, 0.2), LONG, name="l")
        U(BelowFloor, prims.caps_label, self.s, Box(1, 1, 2, 0.3), "small", pt=7.5, name="l")
        U(Overflows, prims.stat, self.s, Box(1, 1, 1.5, 0.6), "42%", "A label", kind="measured", name="s")
        U(BelowFloor, prims.stat, self.s, Box(1, 1, 2, 1.5), "42%", "A label", kind="measured", label_pt=7.5,
          name="s")
        U(Overflows, prims.quote, self.s, Box(1, 1, 2, 0.4), LONG, "Someone", name="q")
        U(Overflows, prims.zone, self.s, Box(1, 1, 1.0, 1.0), label=LONG, name="z")
        U(BelowFloor, text.textbox, self.s, Box(1, 1, 4, 1), [text.Run("x", size=7.5)], "tb", SIGMOID)
        U(Overflows, text.textbox, self.s, Box(1, 1, 0.6, 0.15), LONG, "tb", SIGMOID, size=11)
        U(Overflows, text.title, self.s, " ".join(["Reconciliation"] * 12), (), WIDE, SIGMOID)
        U(BelowFloor, icons.monogram, self.s, Box(1, 1, 0.2, 0.2), "AB", name="m")

    # forced refusals after the first shape is already on the slide
    def test_forced_failure_after_the_shape_is_added(self):
        U = self.unchanged
        with mock.patch.object(prims, "set_line", _raise_floor):
            U(BelowFloor, prims.rect, self.s, Box(1, 1, 2, 1), fill="card", name="r")
            U(BelowFloor, prims.card, self.s, Box(1, 1, 2, 1), name="c")
            U(BelowFloor, prims.card, self.s, Box(1, 1, 2, 1), radius_in=0, name="c0")
            U(BelowFloor, prims.zone, self.s, Box(1, 1, 2, 1), name="z")
            U(BelowFloor, prims.person, self.s, Box(1, 1, 0.5, 0.7), name="pe")
            U(BelowFloor, prims.hairline, self.s, (1, 1), (3, 1), name="h")
            U(BelowFloor, prims.connect, self.s, (1, 1), (3, 1), name="a")
            U(BelowFloor, prims.elbow, self.s, [(1, 1), (1, 2), (3, 2)], name="e")
            U(BelowFloor, prims.tab, self.s, Box(1, 1, 2, 0.3), "Tab", name="t")
            U(BelowFloor, prims.pill, self.s, Box(1, 1, 1.2, 0.3), "Tag", name="p")
            U(BelowFloor, prims.disc, self.s, Box(1, 1, 0.4, 0.4), 1, name="d")
            U(BelowFloor, icons.monogram, self.s, Box(1, 1, 0.5, 0.5), "QA", name="m")
            U(BelowFloor, glyphs.glyph, self.s, "database", Box(1, 1, 0.4, 0.4))
        with mock.patch.object(text, "write", _nth(3, text.write)):           # third chevron fails
            U(Overflows, prims.step_header, self.s, Box(1, 1, 6, 0.5), ["A", "B", "C"], name="st")
        with mock.patch.object(prims, "pill", _nth(1, prims.pill, _raise_floor)):     # the basis chip, last
            U(BelowFloor, prims.stat, self.s, Box(1, 1, 2, 1.5), "3x", "Faster", kind="target", name="s")
        with mock.patch.object(text, "textbox", _nth(2, text.textbox)):         # the attribution box
            U(Overflows, prims.quote, self.s, Box(1, 1, 3.4, 1.1), "Short passage.", "Someone", name="q")
        with mock.patch.object(prims, "caps_label", _nth(1, prims.caps_label)):  # the zone's label
            U(Overflows, prims.zone, self.s, Box(1, 1, 2, 1), label="Group", name="z")
        with mock.patch.object(text, "_write_prepped", _raise_overflow):
            U(Overflows, text.textbox, self.s, Box(1, 1, 4, 1), "copy", "tb", SIGMOID, size=11)
            U(Overflows, text.title, self.s, "One short title", (), WIDE, SIGMOID)

    def test_picture_rollback_drops_the_relationships_too(self):
        U = self.unchanged
        with mock.patch.object(icons, "_set_nv", _raise_floor):
            U(BelowFloor, icons.vendor, self.s, "aws/compute/ec2", Box(1, 1, 0.5, 0.5))
            U(BelowFloor, icons.svg_picture, self.s, SVG, Box(1, 1, 0.5, 0.5), "auto", name="i", alt="Mark")
        U(GeometryError, icons.svg_picture, self.s, SVG, Box(1, 1, 0.5, 0.5), "auto", name="i", alt="")

    def test_nested_compound_rolls_back_whole(self):
        before = state(self.s)
        with self.assertRaises(Overflows):
            with ooxml.transaction(self.s):
                prims.card(self.s, Box(1, 1, 2, 1), name="c")
                prims.stat(self.s, Box(1, 3, 2, 1.5), "3x", "Faster", kind="measured", name="s")
                prims.pill(self.s, Box(1, 5, 0.5, 0.25), LONG, name="p")
        self.assertEqual(state(self.s), before)

    def test_success_finishes_every_shape_it_added(self):
        from pptx.enum.shapes import MSO_SHAPE
        with ooxml.transaction(self.s):
            raw = self.s.shapes.add_shape(MSO_SHAPE.OVAL, 0, 0, 914400, 914400)     # python-pptx injects p:style
            self.assertIsNotNone(raw._element.find(P + "style"))
        self.assertIsNone(raw._element.find(P + "style"))
        self.assertEqual(len(raw._element.spPr.findall(A + "effectLst")), 1)
        g = prims.stat(self.s, Box(1, 1, 2, 1.5), "3x", "Faster", kind="target", name="s")
        for el in g._element.iter(*ooxml.SHAPE_TAGS):
            self.assertIsNone(el.find(P + "style"), el.find(".//" + P + "cNvPr").get("name"))

    def test_successful_creators_leave_no_p_style_anywhere(self):
        prims.rect(self.s, Box(1, 1, 1, 1), fill="card", name="r")
        prims.step_header(self.s, Box(1, 2.5, 6, 0.5), ["A", "B", "C"], active=1, name="st")
        prims.quote(self.s, Box(1, 3.5, 3.4, 1.1), "Short passage.", "Someone", name="q")
        icons.monogram(self.s, Box(1, 5, 0.5, 0.5), "QA", name="m")
        icons.vendor(self.s, "aws/compute/ec2", Box(2, 5, 0.5, 0.5))
        tree = self.s.shapes._spTree
        self.assertEqual([e for e in tree.iter(*ooxml.SHAPE_TAGS) if e.find(P + "style") is not None], [])
        self.assertEqual(ooxml.once_only_violations(self.s), [])


# ------------------------------------------------------------------ 2. one text writer

class OneWriter(Base):
    def test_prims_private_writer_is_gone(self):
        for attr in ("R", "write_text", "textbox", "_mparas", "_label_paras"):
            self.assertFalse(hasattr(prims, attr), "prims.%s still exists" % attr)

    def test_only_text_py_builds_runs(self):
        """No writer module but text.py constructs a:r / a:rPr / a:endParaRPr."""
        pat = re.compile(r'''(Element|SubElement)\(.*(_A \+ "|"a:)(r|rPr|endParaRPr)"''')
        for f in sorted(Path(KIT).glob("*.py")):
            if f.name in ("text.py",):
                continue
            hits = [i + 1 for i, line in enumerate(f.read_text().splitlines()) if pat.search(line)]
            self.assertEqual(hits, [], "%s builds runs at lines %s" % (f.name, hits))

    def test_role_guards_apply_in_prims_and_icons(self):
        U = self.unchanged
        for role in ("accent", "accent-tint", "series-2", "hairline", "rule-strong", "card"):
            U(ThemeInvariant, prims.pill, self.s, Box(1, 1, 1.5, 0.3), "Tag", role=role, name="p")
            U(ThemeInvariant, prims.caps_label, self.s, Box(1, 1, 2, 0.3), "Label", role=role, name="l")
            U(ThemeInvariant, prims.stat, self.s, Box(1, 1, 2, 1.5), "3x", "Faster", kind="measured", role=role,
              name="s")
        with self.assertRaises(ThemeInvariant) as cm:
            prims.caps_label(self.s, Box(1, 1, 2, 0.3), "Label", role="accent", name="l")
        self.assertIn("accent-ink", cm.exception.hint)

    def test_prims_runs_carry_the_text_writer_signature(self):
        tab = prims.tab(self.s, Box(1, 1, 2, 0.34), "Header", fill="accent", name="t")
        g = prims.stat(self.s, Box(4, 1, 2, 1.5), "3x", "Faster", kind="target", name="s")
        m = icons.monogram(self.s, Box(7, 1, 0.5, 0.5), "QA", name="m")
        for el in (tab._element, g._element, m._element):
            for rpr in el.iter(A + "rPr"):
                self.assertEqual(rpr.get("kern"), "0")
                self.assertIsNotNone(rpr.find(A + "ea"))
                self.assertIsNotNone(rpr.find(A + "solidFill/" + A + "schemeClr"))
        self.assertEqual(tab._element.find(".//" + A + "rPr/" + A + "solidFill/" + A + "schemeClr").get("val"),
                         SIGMOID.spec("on-accent").get("val"))
        def expr(sc):
            return [sc.get("val")] + [(etree.QName(c).localname, c.get("val")) for c in sc]
        val = g._element.find(".//" + A + "r/" + A + "rPr/" + A + "solidFill/" + A + "schemeClr")
        self.assertEqual(expr(val), expr(SIGMOID.spec("accent-ink")))                          # the numeral

    def test_title_keywords_are_accent_ink(self):
        roles = {r.role for r in text.title_runs("Three capabilities carry the plan", ["Three capabilities"])}
        self.assertEqual(roles, {"accent-ink", "ink-deep"})


# ------------------------------------------------------------------ 3. one floor

class OneFloor(unittest.TestCase):
    def test_one_object(self):
        from deckkit import text as T
        self.assertIs(canvas.FLOOR, measure.FLOOR)
        self.assertIs(T.FLOOR, measure.FLOOR)
        self.assertIs(prims.FLOOR, measure.FLOOR)
        self.assertEqual(measure.FLOOR, 8.0)

    def test_no_second_definition(self):
        pat = re.compile(r"^\s*FLOOR\w*\s*=", re.M)
        defs = []
        for root in (Path(SKILL) / "kit", Path(SKILL) / "gates", Path(SKILL) / "scripts"):
            for f in root.rglob("*.py"):
                if pat.search(f.read_text()):
                    defs.append(str(f.relative_to(SKILL)))
        self.assertEqual(defs, ["kit/deckkit/measure.py"])

    def test_read_layer_uses_the_same_floor(self):
        from deckkit.read import walk
        self.assertFalse(hasattr(walk, "FLOOR_PT"))
        with mock.patch.object(measure, "FLOOR", 9.0):
            self.assertEqual(walk._floor(), 9.0)


# ------------------------------------------------------------------ 4. finish builds effectLst with OxmlElement

class FinishOxml(Base):
    def test_effectlst_is_python_pptx_built_and_accessors_work(self):
        c = prims.card(self.s, Box(1, 1, 2, 1), name="c")
        sp = c._element.spPr
        lst = sp.effectLst
        self.assertIsNotNone(lst)
        self.assertIs(type(lst), type(OxmlElement("a:effectLst")))
        self.assertEqual(len(lst), 0)
        self.assertFalse(c.shadow.inherit)
        c.shadow.inherit = True                                  # python-pptx removes it through the proxy
        self.assertIsNone(sp.effectLst)
        c.shadow.inherit = False
        self.assertEqual(len(sp.findall(A + "effectLst")), 1)
        ooxml.finish(c)
        self.assertEqual(len(sp.findall(A + "effectLst")), 1)
        self.assertFalse(c.shadow.inherit)
        self.assertEqual(c.line.width, int(round(WIDE.hairline_pt * 12700)))
        self.assertEqual(int(c.fill.type), 1)                     # MSO_FILL.SOLID
        self.assertIsNotNone(sp.prstGeom)
        self.assertIsNotNone(sp.get_or_add_xfrm())
        names = [etree.QName(e).localname for e in sp]
        self.assertLess(names.index("ln"), names.index("effectLst"))
        self.assertEqual(ooxml.once_only_violations(self.s), [])

    def test_group_and_connector_accessors_after_finish(self):
        g = prims.person(self.s, Box(1, 1, 0.5, 0.7), name="pe")
        self.assertFalse(g.shadow.inherit)
        h = prims.hairline(self.s, (1, 3), (3, 3), name="h")
        self.assertIsNotNone(h._element.spPr.effectLst)
        self.assertGreater(h.line.width, 0)


# ------------------------------------------------------------------ 5. accent strips

class Strips(Base):
    def test_non_neutral_strips_refused(self):
        U = self.unchanged
        for role in ("accent", "accent-tint", "accent-deep", "accent-wash", "series-1", "series-5", "navy",
                     "ink-strong", "ink-deep"):
            U(GeometryError, prims.rect, self.s, Box(1, 1, 0.10, 1.0), fill=role, name="edge")      # exactly 0.10
            U(GeometryError, prims.rect, self.s, Box(1, 1, 2.0, 0.06), fill=role, name="top")
        U(GeometryError, prims.rect, self.s, Box(1, 1, 0.05, 1.5), line="accent", name="edge.line")   # line only
        U(GeometryError, prims.card, self.s, Box(1, 1, 0.08, 1.2), fill="accent-tint", line=None, name="c")
        U(GeometryError, prims.card, self.s, Box(1, 1, 2.0, 0.09), fill="accent-tint", line=None, radius_in=0,
          name="c")
        U(GeometryError, prims.zone, self.s, Box(1, 1, 3.0, 0.08), name="z")

    def test_neutral_state_and_non_strips_pass(self):
        for role in prims.NEUTRAL_STRIP:
            prims.rect(self.s, Box(1, 1, 0.05, 2.0), fill=role, name="divider.%s" % role)
        prims.rect(self.s, Box(1, 1, 3.0, 0.02), line="hairline", name="rule.line")
        prims.card(self.s, Box(1, 1, 3.0, 0.10), fill="card", line="hairline", name="card.thin")
        prims.rect(self.s, Box(1, 1, 0.05, 1.5), fill="accent", name="state.active")
        prims.rect(self.s, Box(1, 1, 0.11, 2.0), fill="accent", name="block.wide")              # thicker than 0.10
        prims.rect(self.s, Box(1, 1, 0.10, 0.59), fill="accent", name="block.short")            # under 6x long
        self.assertEqual(prims.BAR_MAX_IN, 0.10)


# ------------------------------------------------------------------ 6. WIDE title box

class TitleBox(Base):
    def test_wide_title_box(self):
        tb = WIDE.title_box
        self.assertEqual((tb.x, tb.y, tb.w, tb.h), (0.50, 0.13, 12.40, 0.40))
        self.assertLessEqual(tb.b, WIDE.content.y)

    def test_one_21pt_bold_line_fits_with_no_tolerance(self):
        face = measure.face_for("Arial", True)
        self.assertGreater(measure.line_in(21.0, face), 0.35)     # why 0.35 needed a tolerance
        f = measure.fits([measure.MPara((measure.MRun("Planning", 21.0, face),))], WIDE.title_box)
        self.assertTrue(f.ok)
        self.assertLessEqual(f.need_in, WIDE.title_box.h)
        sh = text.title(self.s, "Four steps take a forecast from draft to signed plan", ["Four steps"], WIDE,
                        SIGMOID)
        self.assertEqual(sh.height, int(round(0.40 * 914400)))
        self.assertFalse(hasattr(text, "FIT_TOL_IN"))

    def test_title_refuses_a_box_shorter_than_its_line(self):
        tight = dataclasses.replace(WIDE, title_box=Box(0.50, 0.13, 12.40, 0.35))
        self.unchanged(Overflows, text.title, self.s, "One short title", (), tight, SIGMOID)

    def test_std_two_lines_fit_outright(self):
        face = measure.face_for("Arial", True)
        self.assertLessEqual(2 * measure.line_in(STD.title_pt, face), STD.title_box.h)


# ------------------------------------------------------------------ 7. docProps date

def core(path):
    with zipfile.ZipFile(path) as z:
        return z.read("docProps/core.xml").decode()


def small():
    prs, s = new_slide()
    with prims.use(SIGMOID):
        prims.card(s, Box(1, 1, 2, 1), name="c")
    return prs


class DocPropsDate(unittest.TestCase):
    def setUp(self):
        self.td = tempfile.TemporaryDirectory()
        self.dir = Path(self.td.name)

    def tearDown(self):
        self.td.cleanup()

    def stamps(self, path):
        return re.findall(r"<dcterms:(created|modified)[^>]*>([^<]+)<", core(path))

    def test_date_from_deck_json(self):
        (self.dir / "deck.json").write_text(json.dumps({"date": "2026-10-03"}))
        date = json.loads((self.dir / "deck.json").read_text())["date"]
        a, b = self.dir / "a.pptx", self.dir / "b" / "a.pptx"
        sa = package.save(small(), a, title="Sample", theme=SIGMOID, date=date)
        sb = package.save(small(), b, title="Sample", theme=SIGMOID, date=date)
        self.assertEqual(sa, sb)                                         # deterministic with a date
        self.assertEqual(a.read_bytes(), b.read_bytes())
        self.assertEqual(self.stamps(a), [("created", "2026-10-03T00:00:00Z"), ("modified", "2026-10-03T00:00:00Z")])
        self.assertEqual(json.loads(package.manifest_path(a).read_text())["date"], "2026-10-03T00:00:00Z")
        with zipfile.ZipFile(a) as z:
            self.assertEqual({i.date_time for i in z.infolist()}, {package.FIXED_TIME})

    def test_fallback_is_1980_and_deterministic(self):
        a, b = self.dir / "a.pptx", self.dir / "b.pptx"
        self.assertEqual(package.save(small(), a, title="Sample", theme=SIGMOID),
                         package.save(small(), b, title="Sample", theme=SIGMOID))
        self.assertEqual({v for _, v in self.stamps(a)}, {package.FIXED_DATE})
        self.assertNotIn("date", json.loads(package.manifest_path(a).read_text()))

    def test_date_changes_only_the_stamps(self):
        a, b = self.dir / "a.pptx", self.dir / "b.pptx"
        package.save(small(), a, title="Sample", theme=SIGMOID, date="2026-10-03")
        package.save(small(), b, title="Sample", theme=SIGMOID)
        with zipfile.ZipFile(a) as za, zipfile.ZipFile(b) as zb:
            diff = [n for n in za.namelist() if za.read(n) != zb.read(n)]
        self.assertEqual(diff, ["docProps/core.xml"])

    def test_bad_date_refused_before_anything_is_written(self):
        out = self.dir / "a.pptx"
        for bad in ("2026-13-01", "3 Oct 2026", 20261003):
            with self.assertRaises(DeckkitError):
                package.save(small(), out, title="Sample", theme=SIGMOID, date=bad)
        self.assertEqual(sorted(os.listdir(self.dir)), [])


if __name__ == "__main__":
    unittest.main()
