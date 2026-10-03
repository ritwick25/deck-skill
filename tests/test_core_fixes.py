"""Track A1 kit-core fixes (ruling 22): accent-bright, one basis label per stat band, cards sized to content, the
device title normalised to the kit's title geometry, archetype= required on a multi-archetype device, and the
new_deck scaffold (template_build.py) that builds, is refused while it holds placeholders, and passes every gate
once filled.

Run: cd ~/.claude/skills/deck && $PY -m unittest tests.test_core_fixes -v
Fictional content only.
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL / "kit"))

from pptx import Presentation  # noqa: E402

from deckkit import brand, compose, devices, prims  # noqa: E402
from deckkit.canvas import WIDE, Box  # noqa: E402
from deckkit.errors import GeometryError, Overflows  # noqa: E402
from deckkit.geom import find  # noqa: E402
from deckkit.read.colour import contrast  # noqa: E402
from deckkit.theme import BRIGHT_AGAINST, BRIGHT_MIN, Theme  # noqa: E402

PY = sys.executable
BASE = SKILL / "kit" / "assets" / "base-wide.pptx"
EMU = 914400
SIGMOID = SKILL / "themes" / "sigmoid.json"
NEUTRAL = SKILL / "themes" / "neutral.json"


def _green():
    """The brand agent's blocker: a dark client primary (#0B6E4F) through deckkit.brand.make_theme."""
    data, _notes = brand.make_theme("Greenco", "0B6E4F", house=SIGMOID, date="2026-10-04")
    return Theme(data)


def _walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from _walk(sh.shapes)


def _box(sh):
    return Box(sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU)


def _hit(a: Box, b: Box) -> bool:
    return a.x < b.r - 1e-6 and b.x < a.r - 1e-6 and a.y < b.b - 1e-6 and b.y < a.b - 1e-6


class _Slide(unittest.TestCase):
    theme_path = SIGMOID

    def setUp(self):
        self.th = Theme.load(self.theme_path)
        self.prs = Presentation(str(BASE))
        self._use = prims.use(self.th)
        self._use.__enter__()
        self.s = self.prs.slides.add_slide(self.prs.slide_layouts[0])

    def tearDown(self):
        self._use.__exit__(None, None, None)

    def named(self, prefix):
        return [sh for sh in _walk(self.s.shapes) if sh.name.startswith(prefix)]


class TestAccentBright(unittest.TestCase):
    def test_every_theme_solves_it_against_the_dark_fills(self):
        for th in (Theme.load(SIGMOID), Theme.load(NEUTRAL), _green()):
            hx = th.hex("accent-bright")
            for fill in BRIGHT_AGAINST:
                self.assertGreaterEqual(contrast(hx, th.hex(fill)) + 1e-9, BRIGHT_MIN, (th.name, fill, hx))
            rows = [r for r in th.contrast_table() if r["fg"] == "accent-bright"]
            self.assertEqual(sorted(r["bg"] for r in rows), sorted(BRIGHT_AGAINST))
            self.assertTrue(all(r["status"] == "PASS" for r in rows), th.name)
            self.assertEqual(th.solved["accent-bright"]["method"], "lighten")

    def test_a_brand_that_already_clears_is_kept(self):
        data = json.loads(SIGMOID.read_text())
        data["name"] = "lightbrand"
        data["slots"]["accent1"] = "F2A900"
        data["roles"].pop("accent-ink", None)
        data["roles"].pop("on-accent", None)
        th = Theme(data)
        self.assertEqual(th.expr("accent-bright"), "accent1")

    def test_hue_is_kept(self):
        import colorsys
        th = _green()
        h0 = colorsys.rgb_to_hls(*[int(th.hex("accent")[i:i + 2], 16) / 255 for i in (0, 2, 4)])[0]
        h1 = colorsys.rgb_to_hls(*[int(th.hex("accent-bright")[i:i + 2], 16) / 255 for i in (0, 2, 4)])[0]
        self.assertAlmostEqual(h0, h1, delta=0.02)


class TestStatBasis(_Slide):
    def test_stat_without_its_chip(self):
        g = prims.stat(self.s, Box(1, 1, 3, 2), "40+", "signals a week", kind="illustrative", basis=False, name="st")
        names = [sh.name for sh in g.shapes]
        self.assertNotIn("st.basis", names)
        self.assertEqual(g._element.find(".//{*}cNvPr").get("descr"), "stat kind illustrative")
        with_chip = prims.stat_height(self.s, 3, "40+", "signals a week", kind="illustrative")
        without = prims.stat_height(self.s, 3, "40+", "signals a week", kind="illustrative", basis=False)
        self.assertLess(without, with_chip)

    def test_stat_band_draws_one_label(self):
        names = compose.stat_band(self.s, Box(0.5, 4.5, 12.4, 1.8), [
            ("40+", "signals a week", "illustrative"), ("5 days", "move to field", "illustrative"),
            ("6", "versions of a brief", "illustrative")], name="band")
        self.assertEqual(names, ["band.1", "band.2", "band.3"])
        labels = self.named("band.basis")
        self.assertEqual(len(labels), 1)
        self.assertEqual(labels[0].text_frame.text, "Illustrative figures")
        self.assertEqual([sh for sh in _walk(self.s.shapes) if sh.name.endswith(".basis") and sh.name != "band.basis"],
                         [])
        band = self.named("band.band")[0]
        self.assertGreaterEqual(band.top, labels[0].top + labels[0].height)

    def test_measured_band_has_no_label(self):
        compose.stat_band(self.s, Box(0.5, 4.5, 12.4, 1.8), [("3", "agents live", "delivered"),
                                                            ("1", "gate", "measured")], name="band")
        self.assertEqual(self.named("band.basis"), [])

    def test_mixed_basis_refused(self):
        with self.assertRaises(GeometryError):
            compose.stat_band(self.s, Box(0.5, 4.5, 12.4, 1.8), [("3", "agents live", "delivered"),
                                                                ("40+", "signals", "illustrative")], name="band")
        with self.assertRaises(GeometryError):
            compose.stat_band(self.s, Box(0.5, 4.5, 12.4, 1.8), [("3", "agents", "target"),
                                                                ("40+", "signals", "illustrative")], name="band")


class TestStatBandDarkPrimary(unittest.TestCase):
    def test_dark_primary_builds_and_reads(self):
        for th in (_green(), Theme.load(NEUTRAL), Theme.load(SIGMOID)):
            prs = Presentation(str(BASE))
            with prims.use(th):
                s = prs.slides.add_slide(prs.slide_layouts[0])
                compose.stat_band(s, Box(0.5, 4.5, 12.4, 1.8), [("40+", "signals a week", "illustrative"),
                                                               ("6", "versions", "illustrative")], name="band")
            val = next(sh for sh in _walk(s.shapes) if sh.name == "band.1.value")
            clr = val._element.find(".//{*}rPr/{*}solidFill")[0]
            role = next(r for r in ("accent-ink", "accent-bright") if _same(th, r, clr))
            self.assertGreaterEqual(contrast(th.hex(role), th.hex("ink-deep")) + 1e-9, BRIGHT_MIN, th.name)
            if contrast(th.hex("accent-ink"), th.hex("ink-deep")) < BRIGHT_MIN:
                self.assertEqual(role, "accent-bright", th.name)


def _same(th, role, clr) -> bool:
    spec = th.spec(role)
    key = lambda e: (e.get("val"), tuple(sorted((c.tag.split("}")[-1], c.get("val")) for c in e)))  # noqa: E731
    return key(spec) == key(clr)


class TestCards(_Slide):
    ASK = [("Approve", "check", "Endorse the regions", "Three regions and two lines first"),
           ("Name", "user-check", "One owner", "A planning lead who signs off accuracy"),
           ("Open", "database", "Sales data", "Two years of history by week two")]
    ICONS = [("search", "Scan", "Read the feeds every weekday"), ("bot", "Draft", "One page with sources"),
             ("user-check", "Review", "A named reviewer signs")]

    def test_ask_glyph_clear_of_text(self):
        compose.ask_cards(self.s, Box(0.5, 1.0, 12.4, 4.0), self.ASK, name="ask")
        for i in (1, 2, 3):
            icon = _box(self.named("ask.%d.icon" % i)[0])
            for t in self.named("ask.%d.mark" % i) + self.named("ask.%d.text" % i):
                self.assertFalse(_hit(icon, _box(t)), "ask %d glyph overlaps %s" % (i, t.name))

    def test_cards_sized_to_content(self):
        foot = compose.ask_cards(self.s, Box(0.5, 1.0, 12.4, 4.5), self.ASK, name="ask")
        card = _box(self.named("ask.1.card")[0])
        self.assertLess(card.h, 4.5)
        self.assertAlmostEqual(foot, card.b, places=3)
        foot = compose.icon_cards(self.s, Box(0.5, 1.0, 12.4, 4.5), self.ICONS, name="ic")
        card = _box(self.named("ic.1.card")[0])
        self.assertLess(card.h, 3.0)
        self.assertAlmostEqual(foot, card.b, places=3)

    def test_spread_fills_or_refuses(self):
        need = compose.icon_cards(self.s, Box(0.5, 1.0, 12.4, 4.0), self.ICONS, name="ic0") - 1.0
        compose.icon_cards(self.s, Box(0.5, 1.0, 12.4, need * 1.2), self.ICONS, name="ic", spread=True)
        self.assertAlmostEqual(_box(self.named("ic.1.card")[0]).h, need * 1.2, places=3)
        with self.assertRaises(GeometryError):
            compose.icon_cards(self.s, Box(0.5, 1.0, 12.4, 5.5), self.ICONS, name="ic2", spread=True)
        with self.assertRaises(GeometryError):
            compose.ask_cards(self.s, Box(0.5, 1.0, 12.4, 5.5), self.ASK, name="ask2", spread=True)


@unittest.skipUnless(devices.JSON.is_file(), "device library not assembled")
class TestDevicePlace(unittest.TestCase):
    def setUp(self):
        self.th = Theme.load(SIGMOID)
        self.prs = Presentation(str(BASE))
        self._use = prims.use(self.th)
        self._use.__enter__()

    def tearDown(self):
        self._use.__exit__(None, None, None)

    def _plain(self, dev):
        e = next(d for d in devices.catalogue() if d["id"] == dev)
        lib = Presentation(str(devices.PPTX))
        src = lib.slides[e["index"] - 1]
        fill = {}
        for sl in e["slots"]:
            sh = find(src, sl["name"])
            if sl["kind"] == "glyph":
                fill[sl["name"]] = "check"
            elif sl["kind"] == "title":
                fill[sl["name"]] = ("Two decisions start the build", ("Two decisions",))
            elif sh.shape_type == 6:
                kids = [x for x in _walk(sh.shapes) if x.shape_type != 6 and x.has_text_frame and
                        x.text_frame.text.strip()]
                fill[sl["name"]] = ["1"] * len(kids)
            elif getattr(sh, "has_table", False) and sh.has_table:
                fill[sl["name"]] = [["1"] * len(sh.table.columns) for _ in sh.table.rows]
            else:
                fill[sl["name"]] = "1" if (sl.get("max_words") or 2) <= 1 else "Ab"
        return e, fill

    def test_title_takes_the_kit_geometry(self):
        e, fill = self._plain("DL07")
        s = devices.place(self.prs, "DL07", fill, theme=self.th, archetype="process-flow")
        t = find(s, "slot.title")
        tb = WIDE.title_box
        b = _box(t)
        for got, want in zip((b.x, b.y, b.w, b.h), (tb.x, tb.y, tb.w, tb.h)):
            self.assertAlmostEqual(got, want, places=2)
        sizes = {r.font.size.pt for p in t.text_frame.paragraphs for r in p.runs}
        self.assertEqual(sizes, {WIDE.title_pt})

    def test_title_clears_a_device_mark_beside_it(self):
        e, fill = self._plain("DL28")
        s = devices.place(self.prs, "DL28", fill, theme=self.th, archetype=e["archetypes"][0])
        t = _box(find(s, "slot.title"))
        legend = min(_box(sh).x for sh in s.shapes if sh.name.startswith("struct.legend"))
        self.assertLessEqual(t.r, legend - 0.19)

    def test_long_title_refused(self):
        e, fill = self._plain("DL05")
        fill["slot.title"] = ("Every signal passes five automated checks and two independent reviews before planners",
                              ())
        with self.assertRaises(Overflows):
            devices.place(self.prs, "DL05", fill, theme=self.th)

    def test_multi_archetype_needs_archetype(self):
        e, fill = self._plain("DL07")
        n = len(self.prs.slides)
        with self.assertRaises(GeometryError) as cm:
            devices.place(self.prs, "DL07", fill, theme=self.th)
        self.assertIn("archetype", str(cm.exception))
        self.assertEqual(len(self.prs.slides), n)              # refused before a slide is added
        s = devices.place(self.prs, "DL05", self._plain("DL05")[1], theme=self.th)   # one archetype: optional
        self.assertIsNotNone(s)


FILLED = '''COPY = {
    "summary": {
        "title": ("A shared forecast cuts stockouts before the peak season", ("cuts stockouts",)),
        "left": ("What changes", [("One forecast for every store", "Planners, buyers and finance read one number"),
                                  ("Exceptions come first", "Planners review only the items the model flags"),
                                  ("Every change is recorded", "Each override keeps its reason and its author"),
                                  ("Accuracy is checked weekly", "A holdout test runs before each release")]),
        "right": ("What it takes", [("Two years of sales history", "Exported by the data team before week two"),
                                    ("One planning owner", "Signs off accuracy targets and the gate"),
                                    ("Three pilot regions", "Chosen by the merchandising lead in week one"),
                                    ("A weekly review slot", "Thirty minutes with planners every Monday")]),
        "stats": [("12%", "fewer stockouts in the pilot regions", "illustrative"),
                  ("4 hours", "saved per planner each week", "illustrative"),
                  ("14 weeks", "from kickoff to the first live forecast", "illustrative")],
    },
    "platform": {
        "title": ("Four platform layers carry every forecast from raw sales to planners", ("Four platform layers",)),
        "sources": ("Sources · internal and partner", [
            ("vendor:aws/general/internet-alt1", "Public data"), ("glyph:file-search", "Sales feeds"),
            ("glyph:landmark", "Finance"), ("glyph:megaphone", "Promotions"), ("glyph:folder-open", "Supplier files")]),
        "boundary": "Forecast platform · one AWS account, eu-west-1",
        "rail": ("Identity", [("vendor:aws/security/cognito", "Cognito"),
                              ("vendor:aws/security/identity-and-access-management-iam", "AWS IAM"),
                              ("vendor:aws/security/key-management-service", "AWS KMS"),
                              ("vendor:aws/security/secrets-manager", "Secrets")]),
        "ingest": ("Ingest", [("vendor:aws/integration/eventbridge-scheduler", "EventBridge"),
                              ("vendor:aws/compute/lambda", "Lambda"), ("vendor:aws/ml/textract", "Textract"),
                              ("vendor:aws/integration/simple-queue-service-sqs", "Amazon SQS")]),
        "store": ("Store", [("vendor:aws/storage/simple-storage-service-s3", "Amazon S3"),
                            ("vendor:aws/analytics/amazon-opensearch-service", "OpenSearch"),
                            ("vendor:aws/database/rds", "RDS"), ("vendor:aws/analytics/glue", "AWS Glue")]),
        "agents": ("Agents", [("vendor:aws/ml/bedrock", "Bedrock"), ("mark:claude-spark", "Claude"),
                              ("mark:langgraph-logo", "LangGraph"),
                              ("vendor:aws/compute/elastic-container-service-task", "Forecast agent"),
                              ("vendor:aws/compute/elastic-container-service-task", "Review agent")]),
        "serve": ("Serve", [("vendor:aws/network/api-gateway", "API Gateway"),
                            ("vendor:programming/framework/react", "Planner app"),
                            ("vendor:aws/database/dynamodb", "DynamoDB"),
                            ("vendor:aws/engagement/simple-email-service-ses", "Amazon SES")]),
        "band": ("Delivery and operations", [("vendor:onprem/ci/github-actions", "GitHub Actions"),
                                             ("vendor:aws/compute/ec2-container-registry", "Amazon ECR"),
                                             ("vendor:onprem/iac/terraform", "Terraform"),
                                             ("mark:langfuse-icon", "Langfuse"),
                                             ("vendor:aws/management/cloudwatch", "CloudWatch")]),
    },
    "ask": {
        "title": ("Four decisions this month put the pilot live by week fourteen", ("Four decisions",)),
        "cards": [("check", "Approve the regions", "Three regions and two product lines first"),
                  ("user-check", "Name one owner", "A planning lead who signs off accuracy"),
                  ("database", "Open the sales data", "Two years of sales history by week two"),
                  ("shield-check", "Set the go-live test", "Live after four weeks inside the error band")],
        "rail": "What is true when",
        "milestones": [("Week 2", "Scope and owner confirmed"), ("Week 6", "First forecasts beside the current plan"),
                       ("Week 10", "Accuracy checked against the gate"), ("Week 14", "Planners act on the forecast")],
    },
}
'''


def _gates(deck: Path, *only) -> tuple[int, str]:
    cmd = [PY, str(SKILL / "gates" / "run_all.py"), "--deck", str(deck)] + (["--only", ",".join(only)] if only else [])
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    return r.returncode, r.stdout + r.stderr


class TestScaffold(unittest.TestCase):
    """new_deck.py scaffolds template_build.py: it builds, P19 refuses it while placeholders remain, and once its COPY
    is filled the deck passes every gate."""

    def test_scaffold_builds_refuses_then_passes(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / "pilot"
            r = subprocess.run([PY, str(SKILL / "scripts" / "new_deck.py"), str(folder), "--name", "Forecast Pilot"],
                               capture_output=True, text=True)
            self.assertEqual(r.returncode, 0, r.stderr)
            dj = json.loads((folder / "deck.json").read_text())
            self.assertEqual(dj["type"], "deck")
            src = (folder / "build.py").read_text()
            self.assertTrue(src.startswith('"""Forecast Pilot: '))
            self.assertNotIn("__NAME__", src)
            for banned in ("pharma", "medical", "rep ", "brief"):
                self.assertNotIn(banned, src.split("# --- COPY ---")[1].split("# --- END COPY ---")[0].lower())
            deck = folder / "Forecast Pilot.pptx"
            r = subprocess.run([PY, str(folder / "build.py")], capture_output=True, text=True, cwd=folder)
            self.assertEqual(r.returncode, 0, r.stderr)
            code, out = _gates(deck, "P19")
            self.assertNotEqual(code, 0)
            self.assertIn("unfilled slot token", out)
            a = src.index("# --- COPY ---\n") + len("# --- COPY ---\n")
            b = src.index("# --- END COPY ---")
            (folder / "build.py").write_text(src[:a] + FILLED + src[b:])
            r = subprocess.run([PY, str(folder / "build.py")], capture_output=True, text=True, cwd=folder)
            self.assertEqual(r.returncode, 0, r.stderr)
            code, out = _gates(deck)
            self.assertEqual(code, 0, out[-3000:])


class TestEpicOnEveryTheme(unittest.TestCase):
    """scripts/example_epic.py builds on the house, neutral and a dark-primary client theme, and holds P25 and P26."""

    def test_epic_builds_and_reads(self):
        with tempfile.TemporaryDirectory() as tmp:
            green = Path(tmp) / "green.json"
            data, _ = brand.make_theme("Greenco", "0B6E4F", house=SIGMOID, date="2026-10-04")
            green.write_text(json.dumps(data))
            for th in (SIGMOID, NEUTRAL, green):
                out = Path(tmp) / th.stem / "epic.pptx"
                r = subprocess.run([PY, str(SKILL / "scripts" / "example_epic.py"), "--theme", str(th), "--out",
                                    str(out)], capture_output=True, text=True)
                self.assertEqual(r.returncode, 0, r.stderr)
                code, txt = _gates(out, "P23", "P25", "P26")
                self.assertEqual(code, 0, txt[-2000:])


if __name__ == "__main__":
    unittest.main()
