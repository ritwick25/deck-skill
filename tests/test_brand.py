"""Client brand intake (ruling 23): deckkit.brand and scripts/brand.py.

Run: cd ~/.claude/skills/deck && $PY -m unittest tests.test_brand -v
Fictional brands only. The CSS read goes through a file:// page, so no test touches the network.
"""
from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(SKILL, "kit"))
sys.path.insert(0, os.path.join(SKILL, "scripts"))

from pptx import Presentation  # noqa: E402

from deckkit import brand, package, prims, tag  # noqa: E402
from deckkit.canvas import Box  # noqa: E402
from deckkit.errors import GeometryError  # noqa: E402
from deckkit.icons import ASVG_NS  # noqa: E402
from deckkit.read.colour import contrast  # noqa: E402
from deckkit.read.ns import NS  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

import brand as brand_cli  # noqa: E402  (scripts/brand.py)

A = "{%s}" % NS["a"]
R = "{%s}" % NS["r"]
BASE = os.path.join(SKILL, "kit", "assets", "base-wide.pptx")
SIGMOID = json.loads(Path(SKILL, "themes", "sigmoid.json").read_text(encoding="utf-8"))
PY = sys.executable

LOGO_SVG = b"""<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 300 80" width="300" height="80">
<circle cx="40" cy="40" r="32" fill="#0B6E4F"/><circle cx="40" cy="40" r="14" fill="#F2A900"/>
<rect x="88" y="26" width="196" height="28" rx="4" fill="#0B6E4F"/></svg>"""
FIRM_SVG = b"""<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 160 80"><rect x="4" y="16" width="48" height="48"
fill="#3A3A3A"/><rect x="64" y="30" width="90" height="20" fill="#3A3A3A"/></svg>"""


def _png(w, h):
    from PIL import Image
    b = io.BytesIO()
    Image.new("RGB", (w, h), (11, 110, 79)).save(b, "PNG")
    return b.getvalue()


def _theme(primary="#0B6E4F", secondary="#F2A900", **kw):
    data, _ = brand.make_theme("Quillmoor Bio", primary, secondary, date="2026-10-03", **kw)
    return data


class ThemeFromHex(unittest.TestCase):
    def test_primary_theme_passes_theme_check(self):
        data = _theme()
        text, rc = brand.check_text(data)
        self.assertEqual(rc, 0, text)
        self.assertIn("PASS: every gated pair", text)
        self.assertEqual(data["name"], "quillmoor-bio")
        self.assertEqual(data["slots"]["accent1"], "0B6E4F")
        self.assertEqual(data["slots"]["accent2"], "F2A900")          # the secondary is series-1
        self.assertIn("accent2", data["fill_only"])                     # F2A900 is under 4.5 on white
        self.assertGreaterEqual(contrast(data["slots"]["dk2"], "FFFFFF"), brand.DK2_TARGET)
        self.assertEqual(data["identity"], SIGMOID["identity"])         # notation colours survive the swap
        self.assertEqual(data["vendors"], SIGMOID["vendors"])
        th = Theme(json.loads(json.dumps(data)))                        # loads with its records, as written
        self.assertEqual(th.expr("accent-ink"), "accent1")
        self.assertEqual(th.expr("on-accent"), "bg1")

    def test_low_contrast_primary_gets_solved_accent_ink(self):
        data = _theme(primary="#F2A900", secondary=None)
        self.assertEqual(brand.check_text(data)[1], 0)
        rec = data["roles"]["accent-ink"]["record"]
        self.assertNotEqual(rec, "accent1")                              # lowered for text
        self.assertTrue(rec.startswith("accent1 lumMod "))
        th = Theme(json.loads(json.dumps(data)))
        for bg in ("paper", "card"):
            self.assertGreaterEqual(contrast(th.hex("accent-ink"), th.hex(bg)), 4.5)
        self.assertEqual(th.hex("accent"), "F2A900")                     # fills keep the brand colour
        self.assertEqual(th.expr("on-accent"), "tx1")                    # black text on the yellow

    def test_mid_luminance_primary_refused_with_reason(self):
        with self.assertRaises(brand.BrandRefused) as cm:
            _theme(primary="#E4572E", secondary=None)
        msg = str(cm.exception)
        self.assertIn("on-accent on accent", msg)
        self.assertIn("mid luminance", msg)
        deeper = brand.deeper("E4572E")
        self.assertIn(deeper, msg)
        self.assertEqual(brand.check_text(_theme(primary=deeper, secondary="#E4572E"))[1], 0)

    def test_series_padding_and_limit(self):
        data = _theme(accents=["#2A6FDB"])
        self.assertEqual([data["slots"]["accent%d" % i] for i in (2, 3)], ["F2A900", "2A6FDB"])
        self.assertEqual(data["slots"]["accent4"], brand.HOUSE_SERIES[0])
        with self.assertRaises(brand.BrandRefused):
            _theme(accents=["#111111", "#222222", "#333333", "#444444", "#555555"])
        with self.assertRaises(brand.BrandRefused):
            _theme(primary="green")

    def test_unknown_font_falls_back_to_arial(self):
        face, note = brand.font_choice("Quillmoor Sans Display")
        self.assertEqual(face, "Arial")
        self.assertIn("not installed", note)
        data = _theme(font="Quillmoor Sans Display")
        self.assertEqual((data["fonts"]["major"], data["fonts"]["minor"]), ("Arial", "Arial"))
        self.assertEqual(data["brand"]["font"]["requested"], "Quillmoor Sans Display")


class Candidates(unittest.TestCase):
    CSS = """:root{--brand-primary:#0b6e4f;--brand-gold:#F2A900;--grey:#777;--white:#fff}
    a{color:#0B6E4F} .btn{background:rgb(11,110,79)} .x{color:#2a6fdb} .y{color:#2A6FDB}
    .z{color:#2a6fdb} .shadow{color:rgba(200,30,30,0.1)} body{color:#222;background:#fafafa}"""

    def test_ranked_custom_properties_first_neutrals_out(self):
        c = brand.css_candidates([self.CSS])
        hexes = [r["hex"] for r in c]
        self.assertEqual(hexes[:2], ["0B6E4F", "F2A900"])               # named by properties
        self.assertEqual(c[0]["count"], 3)                              # #0b6e4f, #0B6E4F, rgb(11,110,79)
        self.assertIn("--brand-primary", c[0]["properties"])
        self.assertIn("2A6FDB", hexes)
        for grey in ("777777", "FFFFFF", "222222", "FAFAFA", "C81E1E"):  # neutrals and a 10% alpha colour
            self.assertNotIn(grey, hexes)

    def test_fetch_from_file_url_and_cli_waits_for_choice(self):
        with tempfile.TemporaryDirectory() as d:
            Path(d, "site.css").write_text(self.CSS, encoding="utf-8")
            Path(d, "index.html").write_text('<html><head><link rel="stylesheet" href="site.css"><style>'
                                             'h1{color:#8A1538}</style></head><body></body></html>',
                                             encoding="utf-8")
            url = Path(d, "index.html").as_uri()
            texts, read = brand.fetch_css(url)
            self.assertEqual(len(read), 2)
            self.assertIn("8A1538", [r["hex"] for r in brand.css_candidates(texts)])
            out = Path(d, "brand")
            with redirect_stdout(io.StringIO()) as buf:
                rc = brand_cli.main(["--name", "Quillmoor Bio", "--from-url", url, "--out", str(out),
                                     "--date", "2026-10-03"])
            self.assertEqual(rc, 3, buf.getvalue())                     # printed, nothing picked
            self.assertFalse((out / "theme.json").exists())
            cand = json.loads((out / "brand-candidates.json").read_text())
            self.assertEqual((cand["url"], cand["fetched"]), (url, "2026-10-03"))
            self.assertEqual(cand["candidates"][0]["hex"], "0B6E4F")
            with redirect_stdout(io.StringIO()) as buf:
                rc = brand_cli.main(["--name", "Quillmoor Bio", "--from-url", url, "--primary", "#0B6E4F",
                                     "--secondary", "#F2A900", "--out", str(out), "--date", "2026-10-03"])
            self.assertEqual(rc, 0, buf.getvalue())
            data = json.loads((out / "theme.json").read_text())
            self.assertEqual(data["brand"]["verified"]["url"], url)
            self.assertIn("fetched 2026-10-03", data["source"])


class Lockup(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.TemporaryDirectory()
        self.dir = Path(self.d.name)
        (self.dir / "logo.svg").write_bytes(LOGO_SVG)
        (self.dir / "firm.svg").write_bytes(FIRM_SVG)
        self.th = Theme(_theme())
        self.prs = Presentation(BASE)
        self.active = list(prims._ACTIVE)          # the lockup takes theme=; nothing is bound here

    def tearDown(self):
        self.d.cleanup()
        self.assertEqual(prims._ACTIVE, self.active)

    def slide(self, arch, variant="title-stack"):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[0])
        if arch:
            tag.tag(s, arch, variant)
        return s

    def test_refused_on_content_and_untagged_slides(self):
        for arch in ("process-flow", None):
            s = self.slide(arch, "numbered-rows")
            with self.assertRaises(GeometryError) as cm:
                brand.lockup(s, Box(0.9, 0.75, 5.0, 0.56), self.dir / "logo.svg", theme=self.th, name="Quillmoor Bio")
            self.assertIn("cover or closing", str(cm.exception))
            self.assertEqual(len(s.shapes), 0)                          # nothing left behind

    def test_svg_logo_embeds_with_fallback_and_firm_after_divider(self):
        s = self.slide("cover")
        box = Box(0.9, 0.75, 5.0, 0.56)
        out = brand.lockup(s, box, self.dir / "logo.svg", self.dir / "firm.svg", theme=self.th,
                           name="Quillmoor  Bio", firm_name="Example Firm")
        names = {sh.name: sh for sh in s.shapes}
        self.assertEqual(set(names), {"logo.client", "logo.divider", "logo.firm"})
        pic = names["logo.client"]
        self.assertEqual(pic._element.find(".//{*}cNvPr").get("descr"), "Quillmoor Bio logo")
        self.assertEqual(names["logo.firm"]._element.find(".//{*}cNvPr").get("descr"), "Example Firm logo")
        self.assertAlmostEqual(pic.width / pic.height, 300 / 80, places=2)          # aspect kept
        blip = pic._element.find(".//" + A + "blip")
        fallback = s.part.related_part(blip.get(R + "embed"))
        self.assertEqual(fallback.content_type, "image/png")
        self.assertGreaterEqual(_png_w(fallback.blob), pic.width / 914400 * 96 * 2)  # P22: 2x its size
        svgb = blip.find(".//{%s}svgBlip" % ASVG_NS)
        self.assertIsNotNone(svgb)
        self.assertEqual(s.part.related_part(svgb.get(R + "embed")).blob, LOGO_SVG)  # byte for byte
        # left to right: client, hairline divider, firm, all inside the box
        cl, dv, fm = names["logo.client"], names["logo.divider"], names["logo.firm"]
        self.assertLess(cl.left + cl.width, dv.left)
        self.assertLess(dv.left, fm.left)
        self.assertLessEqual(fm.left + fm.width, int(box.r * 914400) + 2)
        self.assertLessEqual(out["w"], box.w + 1e-6)

    def test_png_as_is_and_refused_when_too_small(self):
        (self.dir / "big.png").write_bytes(_png(1200, 300))
        (self.dir / "small.png").write_bytes(_png(120, 30))
        s = self.slide("closing", "one-line")
        brand.lockup(s, Box(0.9, 0.75, 5.0, 0.56), self.dir / "big.png", theme=self.th, name="Quillmoor Bio")
        pic = s.shapes[0]
        self.assertEqual(pic.image.blob, (self.dir / "big.png").read_bytes())
        s2 = self.slide("cover")
        with self.assertRaises(GeometryError):
            brand.lockup(s2, Box(0.9, 0.75, 5.0, 0.56), self.dir / "small.png", theme=self.th, name="Quillmoor Bio")

    def test_cli_logo_record_and_gates_on_saved_cover(self):
        out = self.dir / "brand"
        with redirect_stdout(io.StringIO()) as buf:
            rc = brand_cli.main(["--name", "Quillmoor Bio", "--primary", "#0B6E4F", "--secondary", "#F2A900",
                                 "--logo", str(self.dir / "logo.svg"), "--out", str(out), "--date", "2026-10-03"])
        self.assertEqual(rc, 0, buf.getvalue())
        th = Theme.load(out / "theme.json")
        lp = brand.logo_path(th)
        self.assertEqual(lp.read_bytes(), LOGO_SVG)
        self.assertEqual(th.data["brand"]["logo"]["alt"], "Quillmoor Bio logo")
        prs = Presentation(BASE)
        s = prs.slides.add_slide(prs.slide_layouts[0])
        tag.tag(s, "cover", "title-stack")
        brand.lockup(s, Box(0.9, 0.75, 5.0, 0.56), lp, theme=th, name=th.data["brand"]["client"])
        deck = self.dir / "deck" / "cover.pptx"
        deck.parent.mkdir()
        package.save(prs, deck, title="Cover", theme=th, overwrite=True, date="2026-10-03")
        (deck.parent / "deck.json").write_text(json.dumps({"theme": str(out / "theme.json")}))
        r = subprocess.run([PY, os.path.join(SKILL, "gates", "run_all.py"), "--deck", str(deck), "--only", "P09,P22"],
                           capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        lp.write_bytes(LOGO_SVG.replace(b"#F2A900", b"#FFFFFF"))         # the logo changed after intake
        with self.assertRaises(brand.BrandRefused):
            brand.logo_path(th)


def _png_w(blob):
    from PIL import Image
    with Image.open(io.BytesIO(blob)) as im:
        return im.size[0]


if __name__ == "__main__":
    unittest.main()
