"""deckkit.text: runs with explicit size and role colour, measured boxes, the two-tone title, retext.

Run: cd ~/.claude/skills/deck && $PY -m unittest tests.test_text -v
"""
from __future__ import annotations

import io
import os
import sys
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
KIT = os.path.join(SKILL, "kit")
sys.path.insert(0, KIT)

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402

from deckkit import text as T  # noqa: E402
from deckkit.canvas import STD, WIDE, Box  # noqa: E402
from deckkit.errors import BelowFloor, Capacity, DeckkitError, OverBudget, Overflows, ThemeInvariant  # noqa: E402
from deckkit.read.ns import NS  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

A = "{%s}" % NS["a"]
P = "{%s}" % NS["p"]
SIGMOID = Theme.load(os.path.join(SKILL, "themes", "sigmoid.json"))
BASE = {"wide": os.path.join(KIT, "assets", "base-wide.pptx"), "std": os.path.join(KIT, "assets", "base-std.pptx")}
BODY = ("Regional teams review the quarterly plan and agree which accounts move first, then publish "
        "one shared view of demand and supply so every market reads the same numbers.")


def new_slide(kind="wide"):
    prs = Presentation(BASE[kind])
    return prs, prs.slides.add_slide(prs.slide_layouts[0])


def scheme(el) -> str:
    """'val mod=v ...' of the schemeClr inside a run's solidFill, comparable to Theme.expr()."""
    sc = el.find(A + "solidFill/" + A + "schemeClr")
    mods = ["%s %s" % (etree.QName(c).localname, c.get("val")) for c in sc]
    return " ".join([sc.get("val")] + mods)


def runs(shape):
    return list(shape._element.iter(A + "r"))


def run_text(r):
    return r.find(A + "t").text


def roundtrip(prs):
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return Presentation(buf)


class Floor(unittest.TestCase):
    def test_7_5pt_raises_and_nothing_is_written(self):
        prs, s = new_slide()
        before = len(s.shapes)
        with self.assertRaises(BelowFloor) as cm:
            T.textbox(s, Box(1, 1, 4, 1), [T.Para([T.Run("small print", size=7.5)])], "note", SIGMOID)
        self.assertEqual(cm.exception.have, 7.5)
        self.assertEqual(cm.exception.need, 8.0)
        self.assertEqual(len(s.shapes), before, "a refused box must not leave a shape behind")

    def test_7_5pt_raises_in_write_and_from_the_default_size(self):
        prs, s = new_slide()
        tb = T.textbox(s, Box(1, 1, 4, 1), "placeholder copy", "card.1.body", SIGMOID, size=11)
        with self.assertRaises(BelowFloor):
            T.write(tb.text_frame, "tiny", SIGMOID, size=7.5)
        self.assertEqual(run_text(runs(tb)[0]), "placeholder copy", "a refused write leaves the old text")

    def test_8pt_is_allowed(self):
        prs, s = new_slide()
        tb = T.textbox(s, Box(1, 1, 4, 1), [T.Run("source line", size=8.0)], "src", SIGMOID)
        self.assertEqual(runs(tb)[0].find(A + "rPr").get("sz"), "800")


class Explicit(unittest.TestCase):
    def test_every_run_has_size_typeface_and_role_colour(self):
        prs, s = new_slide()
        paras = [T.Para([T.Run("Lead ", role="ink-strong", bold=True), T.Run("then the body copy")]),
                 T.Para([T.Run("Second line", role="ink-muted", italic=True)], space_before=4)]
        tb = T.textbox(s, Box(1, 1, 5, 1.2), paras, "card.1.body", SIGMOID, size=11)
        rs = runs(tb)
        self.assertEqual(len(rs), 3)
        for r, role in zip(rs, ("ink-strong", "ink", "ink-muted")):
            rpr = r.find(A + "rPr")
            self.assertEqual(rpr.get("sz"), "1100")
            self.assertEqual(rpr.find(A + "latin").get("typeface"), "Arial")
            self.assertEqual(scheme(rpr), SIGMOID.expr(role))
            self.assertIsNone(rpr.find(A + "solidFill/" + A + "srgbClr"))
        body = tb._element.find(".//" + A + "bodyPr")
        self.assertEqual(body.get("wrap"), "square")
        self.assertIsNotNone(body.find(A + "noAutofit"))
        self.assertIsNone(body.find(A + "spAutoFit"))
        self.assertEqual([body.get(k) for k in ("lIns", "tIns", "rIns", "bIns")], ["0"] * 4)
        ppr = tb._element.find(".//" + A + "pPr")
        self.assertIsNotNone(ppr.find(A + "buNone"))
        self.assertIsNotNone(ppr.find(A + "lnSpc"))
        self.assertIsNone(tb._element.find(P + "style"))
        self.assertEqual(len(tb._element.spPr.findall(A + "effectLst")), 1)

    def test_insets_are_written_in_emu(self):
        prs, s = new_slide()
        from deckkit.measure import Insets
        tb = T.textbox(s, Box(1, 1, 5, 1), "inset copy", "box", SIGMOID, size=11,
                       insets=Insets(0.1, 0.05, 0.1, 0.05))
        body = tb._element.find(".//" + A + "bodyPr")
        self.assertEqual([body.get(k) for k in ("lIns", "tIns", "rIns", "bIns")],
                         ["91440", "45720", "91440", "45720"])

    def test_a_run_with_no_size_is_refused(self):
        prs, s = new_slide()
        with self.assertRaises(ValueError):
            T.textbox(s, Box(1, 1, 4, 1), "no size anywhere", "x", SIGMOID)

    def test_fill_role_on_text_is_refused_with_a_hint(self):
        prs, s = new_slide()
        with self.assertRaises(ThemeInvariant) as cm:
            T.textbox(s, Box(1, 1, 4, 1), [T.Run("brand", role="accent", size=11)], "x", SIGMOID)
        self.assertIn("accent-ink", cm.exception.hint)
        with self.assertRaises(ThemeInvariant):
            T.textbox(s, Box(1, 1, 4, 1), [T.Run("series", role="series-2", size=11)], "x", SIGMOID)

    def test_bullets_hang_and_round_trip(self):
        prs, s = new_slide()
        tb = T.textbox(s, Box(1, 1, 5, 1.5), [T.Para("first point", bullet="-"), T.Para("second", bullet="-")],
                       "card.1.points", SIGMOID, size=11)
        ppr = tb._element.find(".//" + A + "pPr")
        self.assertEqual(ppr.find(A + "buChar").get("char"), "-")
        self.assertEqual(int(ppr.get("indent")), -int(ppr.get("marL")))
        prs2 = roundtrip(prs)
        self.assertEqual([p.text for p in prs2.slides[0].shapes[0].text_frame.paragraphs], ["first point", "second"])


class Fit(unittest.TestCase):
    def test_overflow_raises_with_a_hint_before_any_shape(self):
        prs, s = new_slide()
        before = len(s.shapes)
        with self.assertRaises(Overflows) as cm:
            T.textbox(s, Box(1, 1, 3, 0.4), [T.Para(BODY)], "card.2.body", SIGMOID, size=11)
        e = cm.exception
        self.assertEqual(e.what, "card.2.body")
        self.assertGreater(e.need, e.have)
        self.assertRegex(e.hint, r"cut \d+ of \d+ words")
        self.assertEqual(len(s.shapes), before)

    def test_the_same_text_fits_a_taller_box(self):
        prs, s = new_slide()
        T.textbox(s, Box(1, 1, 3, 1.5), [T.Para(BODY)], "card.2.body", SIGMOID, size=11)

    def test_require_fit_false_writes_without_measuring(self):
        prs, s = new_slide()
        tb = T.textbox(s, Box(1, 1, 3, 0.4), [T.Para(BODY)], "card.2.body", SIGMOID, require_fit=False, size=11)
        self.assertEqual("".join(run_text(r) for r in runs(tb)), BODY)

    def test_indented_overflow_also_raises(self):
        prs, s = new_slide()
        with self.assertRaises(Overflows) as cm:
            T.textbox(s, Box(1, 1, 3, 0.4), [T.Para(BODY, bullet="-")], "card.3.points", SIGMOID, size=11)
        self.assertTrue(cm.exception.hint)


class Title(unittest.TestCase):
    TEXT = "Shared demand view cuts planning cycle time by a third"

    def test_two_tone_writes_exactly_the_keyword_runs_in_accent(self):
        prs, s = new_slide()
        kws = ("planning cycle time", "a third")
        sh = T.title(s, self.TEXT, kws, theme=SIGMOID)
        self.assertEqual(sh.name, "title")
        accent = SIGMOID.expr("accent-ink")
        deep = SIGMOID.expr("ink-deep")
        rs = runs(sh)
        hot = [run_text(r) for r in rs if scheme(r.find(A + "rPr")) == accent]
        cold = [run_text(r) for r in rs if scheme(r.find(A + "rPr")) == deep]
        self.assertEqual(hot, list(kws))
        self.assertEqual(len(hot) + len(cold), len(rs))
        self.assertEqual("".join(run_text(r) for r in rs), self.TEXT)
        self.assertEqual(SIGMOID.hex("accent-ink"), SIGMOID.hex("accent"))
        for r in rs:
            rpr = r.find(A + "rPr")
            self.assertEqual(rpr.get("sz"), str(int(WIDE.title_pt * 100)))
            self.assertEqual(rpr.get("b"), "1")
        self.assertEqual((sh.left, sh.top, sh.width, sh.height), WIDE.title_box.emu)

    def test_single_colour_title_is_valid(self):
        prs, s = new_slide()
        sh = T.title(s, self.TEXT, theme=SIGMOID)
        self.assertEqual({scheme(r.find(A + "rPr")) for r in runs(sh)}, {SIGMOID.expr("ink-deep")})

    def test_four_keywords_raise(self):
        prs, s = new_slide()
        with self.assertRaises(Capacity) as cm:
            T.title(s, self.TEXT, ("Shared", "demand", "planning", "third"), theme=SIGMOID)
        self.assertEqual((cm.exception.need, cm.exception.have), (4, 3))
        self.assertEqual(len(s.shapes), 0)

    def test_keyword_must_be_a_whole_word_substring(self):
        prs, s = new_slide()
        with self.assertRaises(DeckkitError):
            T.title(s, self.TEXT, ("Planning",), theme=SIGMOID)        # case differs
        with self.assertRaises(DeckkitError):
            T.title(s, self.TEXT, ("plan",), theme=SIGMOID)            # inside "planning"
        with self.assertRaises(DeckkitError):
            T.title(s, self.TEXT, ("cycle time", "time by"), theme=SIGMOID)   # overlap

    def test_fifteen_words_raise(self):
        prs, s = new_slide()
        with self.assertRaises(OverBudget):
            T.title(s, "one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen",
                    theme=SIGMOID)

    def test_wide_title_past_one_line_raises(self):
        prs, s = new_slide()
        long = ("Regional commercial operations consolidate fragmented quarterly demand forecasting "
                "into accountable governance")
        with self.assertRaises(Overflows) as cm:
            T.title(s, long, theme=SIGMOID)
        self.assertEqual(cm.exception.unit, "lines")

    def test_std_allows_two_lines(self):
        prs, s = new_slide("std")
        long = ("Regional commercial operations consolidate fragmented quarterly demand forecasting "
                "into accountable governance")
        sh = T.title(s, long, theme=SIGMOID)
        self.assertEqual((sh.left, sh.top, sh.width, sh.height), STD.title_box.emu)


class Retext(unittest.TestCase):
    def _two_tone(self):
        prs, s = new_slide()
        paras = [T.Para([T.Run("Bold lead ", role="ink-strong", bold=True), T.Run("muted tail", role="ink-muted")]),
                 T.Para([T.Run("second paragraph", role="accent-ink", italic=True)])]
        return prs, T.textbox(s, Box(1, 1, 5, 1), paras, "device.body", SIGMOID, size=12)

    def test_retext_keeps_bold_and_colour(self):
        prs, sh = self._two_tone()
        T.retext(sh, [["New lead ", "new tail"], "replacement line"])
        rs = runs(sh)
        self.assertEqual([run_text(r) for r in rs], ["New lead ", "new tail", "replacement line"])
        r0, r1, r2 = (r.find(A + "rPr") for r in rs)
        self.assertEqual((r0.get("b"), scheme(r0)), ("1", SIGMOID.expr("ink-strong")))
        self.assertEqual((r1.get("b"), scheme(r1)), ("0", SIGMOID.expr("ink-muted")))
        self.assertEqual((r2.get("i"), scheme(r2)), ("1", SIGMOID.expr("accent-ink")))
        self.assertEqual({r.get("sz") for r in (r0, r1, r2)}, {"1200"})

    def test_string_chunk_takes_the_first_run_formatting(self):
        prs, sh = self._two_tone()
        T.retext(sh, ["Whole new line", "x"])
        p0 = sh._element.findall(".//" + A + "p")[0]
        rs = p0.findall(A + "r")
        self.assertEqual(len(rs), 1)
        self.assertEqual(rs[0].find(A + "rPr").get("b"), "1")
        self.assertEqual(scheme(rs[0].find(A + "rPr")), SIGMOID.expr("ink-strong"))

    def test_paragraph_count_must_match_and_none_removes(self):
        prs, sh = self._two_tone()
        with self.assertRaises(Capacity):
            T.retext(sh, ["only one"])
        with self.assertRaises(Capacity):
            T.retext(sh, ["a", "b", "c"])
        with self.assertRaises(Capacity):
            T.retext(sh, [["one", "two", "three"], "b"])
        self.assertEqual(run_text(runs(sh)[0]), "Bold lead ", "a refused retext changes nothing")
        T.retext(sh, ["kept", None])
        self.assertEqual(len(sh._element.findall(".//" + A + "p")), 1)

    def test_sub_floor_source_run_is_refused(self):
        prs, sh = self._two_tone()
        runs(sh)[2].find(A + "rPr").set("sz", "700")
        with self.assertRaises(BelowFloor):
            T.retext(sh, [["a ", "b"], "c"])

    def test_round_trip_reopens(self):
        prs, sh = self._two_tone()
        T.retext(sh, [["New lead ", "new tail"], "replacement line"])
        prs2 = roundtrip(prs)
        self.assertEqual(prs2.slides[0].shapes[0].text_frame.text, "New lead new tail\nreplacement line")


if __name__ == "__main__":
    unittest.main()
