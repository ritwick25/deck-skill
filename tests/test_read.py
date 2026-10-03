"""Fixtures for deckkit.read against the real reference decks (read-only), plus synthetic unit checks.

Run: cd ~/.claude/skills/deck && $PY -m unittest tests.test_read -v
Sources are opened read-only; the module asserts every source sha256 is unchanged at the end.
Expected values follow library/design/pipeline.md 3.8 as adjudicated in this suite (see each docstring).
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import sys
import tempfile
import time
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "kit"))

from deckkit.read import deck_info, open_deck, para_text, read_slide, words  # noqa: E402
from deckkit.read.colour import apply_mods, contrast  # noqa: E402
from deckkit.read.deckinfo import font_words, status_ranges  # noqa: E402
from deckkit.read.ns import NS, A  # noqa: E402
from deckkit.read.xform import EMU_IN, IDENT, group_child_to_parent, leaf_geometry, read_xfrm  # noqa: E402
from deckkit.read.zones import slide_zones  # noqa: E402

LIB = "/Users/Nikunj/Codes/Decks-For-Visuals/library"
MANIFEST = json.load(open(os.path.join(LIB, "manifest.json")))
DECKS = {k: v for k, v in MANIFEST["decks"].items() if k <= "d18"}   # the v1 corpus these counts are pinned on
REFERENCE = [k for k, v in DECKS.items() if v["tier"] == "reference"]

_D01 = {}


def d01():
    """The flagship opened once, every slide read once (timed for the performance check)."""
    if not _D01:
        deck = open_deck(DECKS["d01"]["path"])
        t = time.time()
        slides = [read_slide(deck, i) for i in range(1, deck.n_slides + 1)]
        _D01.update(deck=deck, slides=slides, seconds=time.time() - t)
    return _D01


def runs(slide, chrome=False):
    for L in slide["shapes"] + (slide["chrome"] if chrome else []):
        if L.get("text"):
            for p in L["text"]["paragraphs"]:
                for r in p["runs"]:
                    if r["t"].strip():
                        yield L, r


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ---------------------------------------------------------------------------------------------- unit

class TestTokenizer(unittest.TestCase):
    def test_words(self):
        self.assertEqual(words("Long Range Supply & Capacity"), 4)        # '&' alone is not a word
        self.assertEqual(words("AI-native  pods, 5 core + 2 fractional"), 6)
        self.assertEqual(words(""), 0)
        self.assertEqual(words(" ".join(chr(c) for c in (0x2014, 0x2022)) + " ..."), 0)

    def test_mid_word_run_split(self):
        from lxml import etree
        p = etree.fromstring(
            '<a:p xmlns:a="%s"><a:r><a:t>L</a:t></a:r><a:r><a:t>eader board</a:t></a:r><a:br/>'
            '<a:fld type="slidenum"><a:t>7</a:t></a:fld></a:p>' % NS["a"])
        self.assertEqual(para_text(p), "Leader board\n7")
        self.assertEqual(words(para_text(p)), 3)


class TestColour(unittest.TestCase):
    def test_contrast(self):
        self.assertAlmostEqual(contrast("000000", "FFFFFF"), 21.0, places=2)
        self.assertAlmostEqual(contrast("FFFFFF", "FFFFFF"), 1.0, places=3)

    def test_epic_modifiers(self):
        """pipeline.md 3.7: bg1 lumMod 75000 = BFBFBF; accent1/accent2 at lumMod 20000 lumOff 80000."""
        deck = d01()["deck"]
        th = deck.master(deck.master_parts[0]).theme["colors"]
        self.assertEqual(apply_mods(th["lt1"], [("lumMod", 75000)]), "BFBFBF")
        self.assertEqual(apply_mods(th["accent1"], [("lumMod", 20000), ("lumOff", 80000)]), "FFCACC")
        self.assertEqual(apply_mods(th["accent2"], [("lumMod", 20000), ("lumOff", 80000)]), "FAE6D8")


class TestXform(unittest.TestCase):
    def test_rotated_nonuniform_group(self):
        """A 90-degree group with x scale 0.5: a child square becomes a 0.5 x 1.0 frame turned on its side."""
        g = dict(x=0, y=0, w=457200, h=914400, rot=90.0, fh=False, fv=False, chx=0, chy=0, chw=914400, chh=914400)
        M = group_child_to_parent(g)
        child = dict(x=0, y=0, w=914400, h=914400, rot=0.0, fh=False, fv=False)
        geo = leaf_geometry(child, M)
        bx, by, bw, bh = geo["bbox_in"]
        self.assertAlmostEqual(bw, 1.0, places=6)
        self.assertAlmostEqual(bh, 0.5, places=6)
        self.assertEqual(geo["group_scale"], (0.5, 1.0))
        self.assertAlmostEqual(geo["frame"]["rot_deg"], 90.0)

    def test_sheared_child(self):
        """A rotated child inside a non-uniformly scaled group is sheared in slide space."""
        g = dict(x=0, y=0, w=457200, h=914400, rot=0.0, fh=False, fv=False, chx=0, chy=0, chw=914400, chh=914400)
        child = dict(x=0, y=0, w=914400, h=914400, rot=45.0, fh=False, fv=False)
        self.assertTrue(leaf_geometry(child, group_child_to_parent(g))["frame"]["sheared"])

    def test_flip_endpoints(self):
        r = dict(x=0, y=0, w=914400, h=914400, rot=0.0, fh=True, fv=False)
        geo = leaf_geometry(r, IDENT)
        self.assertEqual(tuple(round(v, 6) for v in geo["start_in"]), (1.0, 0.0))
        self.assertEqual(tuple(round(v, 6) for v in geo["end_in"]), (0.0, 1.0))
        self.assertTrue(geo["frame"]["mirrored"])


class TestSyntheticPlaceholders(unittest.TestCase):
    """The corpus has no xfrm-less placeholders; python-pptx's default template writes them, so it is the fixture."""

    def test_inherited_geometry_and_txstyles(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            path = os.path.join(td, "ph.pptx")
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[0])
            s.shapes.title.text = "Quarterly review"
            s.placeholders[1].text = "Subtitle words here"
            prs.save(path)
            with open_deck(path) as deck:
                sl = read_slide(deck, 1)
                lay = deck.layout(deck.layout_part_of(deck.slide_part(1)))
                lt = next(e for info, e in lay.ph.items if info["type"] == "ctrTitle")
                r = read_xfrm(lt.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
                              .find(A("xfrm")))
        title = next(L for L in sl["shapes"] if (L["placeholder"] or {}).get("type") == "ctrTitle")
        self.assertEqual(title["placeholder"]["geom_from"], "layout")
        self.assertAlmostEqual(title["bbox"]["x_in"], round(r["x"] / EMU_IN, 3), places=3)
        self.assertAlmostEqual(title["bbox"]["w_in"], round(r["w"] / EMU_IN, 3), places=3)
        run = title["text"]["paragraphs"][0]["runs"][0]
        self.assertEqual(run["size_pt"], 44.0)                 # master titleStyle lvl1 sz 4400
        self.assertEqual(run["size_src"], "master_txStyles")
        self.assertEqual(run["font"], "Calibri")              # +mj-lt through the master theme
        self.assertEqual(run["font_requested"], "+mj-lt")
        self.assertEqual(title["role_guess"]["role"], "title")
        sub = next(L for L in sl["shapes"] if (L["placeholder"] or {}).get("type") == "subTitle")
        self.assertEqual(sub["role_guess"]["role"], "subtitle")


    def test_floor_reads_effective_size(self):
        """Planted negative: 10pt text shrunk by normAutofit fontScale 70% must count as a sub-8pt run."""
        from lxml import etree
        from pptx import Presentation
        from pptx.util import Inches, Pt
        with tempfile.TemporaryDirectory() as td:
            path = os.path.join(td, "fs.pptx")
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[6])
            tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
            tb.text_frame.text = "Shrunk caption text"
            tb.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
            bp = tb.text_frame._txBody.find(A("bodyPr"))
            for c in list(bp):
                bp.remove(c)
            etree.SubElement(bp, A("normAutofit")).set("fontScale", "70000")
            prs.save(path)
            with open_deck(path) as deck:
                sl = read_slide(deck, 1)
        r = sl["shapes"][0]["text"]["paragraphs"][0]["runs"][0]
        self.assertEqual((r["size_pt"], r["size_rendered_pt"]), (10.0, 7.0))
        self.assertEqual(sl["text"]["runs_below_8pt"], 1)
        self.assertIn("sub-8pt", [h["kind"] for h in sl["hygiene"]])
        self.assertEqual(sl["fidelity"]["autofit_shrunk"], 1)


class TestSyntheticIcons(unittest.TestCase):
    """Planted cases for the refined icon definition: a lone custGeom glyph counts, a lone preset disc or chip
    never does, an out-of-range glyph never does, a texted glyph never does, a glyph inside an icon-sized
    group is counted once as the group."""

    def test_definition(self):
        from lxml import etree
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as td:
            path = os.path.join(td, "ic.pptx")
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[6])

            def glyph(shapes, x, y, w, h, text=None):
                fb = shapes.build_freeform(Inches(x), Inches(y))
                fb.add_line_segments([(Inches(x + w), Inches(y)), (Inches(x + w / 2), Inches(y + h))])
                sp = fb.convert_to_shape()
                if text:
                    sp.text_frame.text = text
                return sp

            g1 = glyph(s.shapes, 1, 1, 0.5, 0.5)                                    # counts: vector
            disc = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2), Inches(1), Inches(0.4), Inches(0.4))
            chip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3), Inches(1), Inches(0.6), Inches(0.3))
            big = glyph(s.shapes, 4, 1, 1.5, 1.5)                                    # too large
            thin = glyph(s.shapes, 6, 1, 0.7, 0.1)                                   # aspect 7
            texted = glyph(s.shapes, 7, 1, 0.5, 0.5, text="A")                       # carries text
            grp = s.shapes.add_group_shape()
            glyph(grp.shapes, 1, 3, 0.3, 0.3)
            glyph(grp.shapes, 1.3, 3, 0.3, 0.3)                                      # group 0.6 x 0.3: one unit
            gx = grp._element.find(".//" + A("xfrm"))                                # python-pptx leaves it 0 x 0
            for tag, a, b in (("off", "x", "y"), ("chOff", "x", "y"), ("ext", "cx", "cy"), ("chExt", "cx", "cy")):
                v = (Inches(1), Inches(3)) if tag.endswith("ff") else (Inches(0.6), Inches(0.3))
                gx.find(A(tag)).set(a, str(int(v[0])))
                gx.find(A(tag)).set(b, str(int(v[1])))
            prs.save(path)
            want_vec = g1.shape_id
            gid = grp.shape_id
            with open_deck(path) as deck:
                sl = read_slide(deck, 1)
        units = {(u["kind"], u["id"]) for u in sl["icons"]}
        self.assertEqual(units, {("vector", want_vec), ("group", gid)})
        ids = {u["id"] for u in sl["icons"]}
        for sp in (disc, chip, big, thin, texted):
            self.assertNotIn(sp.shape_id, ids)


# ---------------------------------------------------------------------------------------------- fixtures

class TestCorpusTotals(unittest.TestCase):
    def test_slides_and_hidden(self):
        """583 slides, 79 hidden (d02 21, d09 7, d12 31, d13 20); every deck agrees with the manifest."""
        total = hidden = 0
        per = {}
        for k, v in DECKS.items():
            with open_deck(v["path"]) as deck:
                n = deck.n_slides
                h = [i for i in range(1, n + 1) if deck.xml(deck.slide_part(i)).get("show") in ("0", "false")]
            self.assertEqual(n, v["slides"], k)
            self.assertEqual(h, v["hidden"], k)
            total += n
            hidden += len(h)
            per[k] = len(h)
        self.assertEqual(total, 583)
        self.assertEqual(hidden, 79)
        self.assertEqual({k: n for k, n in per.items() if n}, {"d02": 21, "d09": 7, "d12": 31, "d13": 20})

    def test_every_slide_reads(self):
        """read_slide succeeds on every slide of every deck and emits the shapes@1 top-level keys."""
        keys = {"schema", "slide_id", "deck_id", "index", "part", "canvas", "hidden", "status", "layout",
                "show_master_shapes", "background", "shapes", "chrome", "groups", "text", "icons", "fidelity",
                "hygiene"}
        for k, v in DECKS.items():
            if k == "d01":
                continue
            with open_deck(v["path"]) as deck:
                for i in range(1, deck.n_slides + 1):
                    s = read_slide(deck, i)
                    self.assertEqual(set(s), keys, (k, i))
                    self.assertEqual(s["hidden"], i in v["hidden"], (k, i))
                    ids = [L["id"] for L in s["shapes"]]
                    self.assertEqual(len(ids), len(set(ids)), (k, i, "duplicate leaf ids"))


class TestD01(unittest.TestCase):
    def slide(self, i):
        return d01()["slides"][i - 1]

    def test_performance(self):
        """read_slide over all 135 d01 slides, cold (including the deck-wide font pass), in under 120 s."""
        sys.stderr.write("[d01 135 slides: %.1f s] " % d01()["seconds"])
        self.assertLess(d01()["seconds"], 120.0)

    def test_s15_title(self):
        """Shape 560 at 0.50, 0.13 in, 12.40 x 0.35 in; 21pt bold; two-tone with an accent1 run."""
        L = next(L for L in self.slide(15)["shapes"] if L["id"] == 560)
        b = L["bbox"]
        for got, want in ((b["x_in"], 0.50), (b["y_in"], 0.13), (b["w_in"], 12.40), (b["h_in"], 0.35)):
            self.assertLessEqual(abs(got - want), 0.01)
        rs = [r for p in L["text"]["paragraphs"] for r in p["runs"]]
        self.assertTrue(all(r["size_rendered_pt"] == 21.0 and r["bold"] for r in rs))
        self.assertTrue(L["text"]["two_tone"])
        self.assertTrue(any("scheme:accent1" in r["colour_src"] for r in rs))
        self.assertEqual(L["role_guess"]["role"], "title")

    def test_s57_pres_default(self):
        """A 14pt run on slide 57 whose size comes from the presentation default."""
        hits = [(L["id"], r) for L, r in runs(self.slide(57)) if r["size_src"] == "pres_default"]
        self.assertTrue(hits)
        self.assertTrue(all(r["size_pt"] == 14.0 for _, r in hits))
        self.assertIn(4, [i for i, _ in hits])

    def test_s50_group(self):
        """'Group 651' (cNvPr id 652): rot 90, own scale 0.4846 x 1.0, bbox computed by hand from off/ext."""
        g = next(g for g in self.slide(50)["groups"] if g["name"] == "Group 651")
        self.assertEqual(g["id"], 652)
        self.assertEqual(g["rot_deg"], 90.0)
        self.assertAlmostEqual(g["scale"][0], 0.4846, places=4)
        self.assertAlmostEqual(g["scale"][1], 1.0, places=4)
        # off (11025982, 4488532) ext (437161, 207358) turned 90 degrees about its centre
        cx, cy = 11025982 + 437161 / 2, 4488532 + 207358 / 2
        want = ((cx - 207358 / 2) / EMU_IN, (cy - 437161 / 2) / EMU_IN, 207358 / EMU_IN, 437161 / EMU_IN)
        got = (g["bbox"]["x_in"], g["bbox"]["y_in"], g["bbox"]["w_in"], g["bbox"]["h_in"])
        for a, b in zip(got, want):
            self.assertLessEqual(abs(a - b), 0.001)
        self.assertEqual(g["children"], [650, 651])

    def test_sub8pt_runs(self):
        """Effective size below 8pt: 1,275 of 6,566 non-blank content runs over 1-135; 799 of 4,449 over 1-104."""
        for lo, hi, want_b, want_n in ((1, 135, 1275, 6566), (1, 104, 799, 4449)):
            n = b = 0
            for s in d01()["slides"][lo - 1:hi]:
                for _L, r in runs(s):
                    n += 1
                    b += r["size_rendered_pt"] < 8.0
            self.assertLessEqual(abs(b - want_b), 0.01 * want_b, (lo, hi, b))
            self.assertLessEqual(abs(n - want_n), 0.01 * want_n, (lo, hi, n))
            self.assertEqual(sum(s["text"]["runs_below_8pt"] for s in d01()["slides"][lo - 1:hi]), b)

    def test_pictures(self):
        """22 svgBlip parts; 105 image layers (92 artisticPhotocopy) on 41 slides; ChatGPT title on 72, 118, 127, 131."""
        svg = layers = photocopy = 0
        layer_slides, chatgpt = set(), set()
        for s in d01()["slides"]:
            for L in s["shapes"]:
                p = L.get("picture")
                if not p:
                    continue
                svg += bool(p["svg_part"])
                if p["hdphoto_part"]:
                    layers += 1
                    layer_slides.add(s["index"])
                photocopy += p["artistic"].count("artisticPhotocopy")
                if re.search("chatgpt", L["title"], re.I):
                    chatgpt.add(s["index"])
                    self.assertEqual(p["ai_hint"].lower(), "chatgpt")
        self.assertEqual(svg, 22)
        self.assertEqual(layers, 105)
        self.assertEqual(photocopy, 92)
        self.assertEqual(len(layer_slides), 41)
        self.assertEqual(chatgpt, {72, 118, 127, 131})

    def test_connectors(self):
        """1,360 cxnSp; 328 glued at both ends. The '424' survey figure is glued at one end or more (377 + 375 - 328)."""
        cxn = [L for s in d01()["slides"] for L in s["shapes"] if L["kind"] == "cxn"]
        g = [L["connector"]["glue"] for L in cxn]
        self.assertEqual(len(cxn), 1360)
        self.assertEqual(sum(1 for x in g if x["start"] and x["end"]), 328)
        self.assertEqual(sum(1 for x in g if x["start"] or x["end"]), 424)
        self.assertEqual(sum(1 for x in g if x["start"]), 377)
        self.assertEqual(sum(1 for x in g if x["end"]), 375)
        # every connector carries a touch verdict at both ends
        self.assertTrue(all(L["connector"]["touch"]["start"] is not None for L in cxn))

    def test_status_ranges(self):
        """d01 1-104 is the narrative (dividers 103 APPENDIX and 104 GRAVEYARD included); 105-135 the graveyard."""
        sr = status_ranges(d01()["deck"])
        self.assertEqual(sr["main"], [[1, 104]])
        self.assertEqual(sr["graveyard"], [[105, 135]])
        self.assertEqual(self.slide(105)["status"], "graveyard")
        self.assertEqual(self.slide(104)["status"], "main")

    def test_zones_and_icons(self):
        z = slide_zones(self.slide(15))["zones"]
        self.assertGreaterEqual(len(z), 2)
        self.assertEqual(z[0]["id"], "Z1")
        self.assertIn(560, z[0]["ids"])            # the title opens the reading order

    def test_icons(self):
        """Refined icon definition (library DECISIONS): outermost icon-sized text-free groups, icon-sized
        text-free pictures, and lone text-free custGeom glyphs. 1,055 over 1-104 (240 groups, 786 pictures,
        29 vector glyphs); 1,780 over 1-135 (480, 1,256, 44). An open stroked path that is straight or headed is
        a line, not a glyph: 30 (60 over 1-135) agent-mesh edges counted as glyphs before open custGeom paths got
        connector blocks (1,085 and 1,840 then). The pre-refinement figures were 1,026 and 1,736."""
        for hi, want, kinds in ((104, 1055, {"group": 240, "pic": 786, "vector": 29}),
                                (135, 1780, {"group": 480, "pic": 1256, "vector": 44})):
            got = {}
            for s in d01()["slides"][:hi]:
                for u in s["icons"]:
                    got[u["kind"]] = got.get(u["kind"], 0) + 1
            self.assertEqual(got, kinds, hi)
            self.assertEqual(sum(got.values()), want, hi)
        for s in d01()["slides"]:
            ids = [u["id"] for u in s["icons"]]
            self.assertEqual(len(ids), len(set(ids)), s["index"])
            leaves = {L["id"]: L for L in s["shapes"]}
            for u in s["icons"]:
                if u["kind"] == "vector":
                    L = leaves[u["id"]]
                    self.assertTrue(L["geom"]["cust"] and not L.get("icon_group")
                                    and not (L["text"] and L["text"]["text"].strip()))
                    self.assertEqual(L["role_guess"]["role"], "icon")

    def test_deck_info(self):
        info = deck_info(d01()["deck"], deck_id="d01", with_fonts=False)
        self.assertEqual(info["default_text_style"]["lvl1"], {"size_pt": 14.0, "font": "Arial", "bold": False,
                                                               "hex": "000000"})
        self.assertEqual(len(info["masters"]), 5)
        self.assertEqual(info["masters"][4]["theme"]["fonts"]["minor"], "Calibri")
        self.assertTrue(all(m["theme"]["fonts"]["minor"] == "Arial" for m in info["masters"][:4]))
        self.assertEqual(info["status_ranges"]["graveyard"], [[105, 135]])
        self.assertEqual(info["origin"]["guess"], "google-export")


class TestFontWords(unittest.TestCase):
    """Words by requested font. The survey figures (Calibri 9,106; Space Grotesk 5,858 +327 Light +49 Medium)
    were measured over the 14 reference decks with str.split() per a:r run and explicit rPr latin faces only.
    Under the ruled word definition, paragraph-joined, with the inheritance chain, the true figures differ;
    both are asserted so the adjudication stays reproducible."""

    @classmethod
    def setUpClass(cls):
        cls.ruled14, cls.ruled18 = {}, {}
        for k, v in DECKS.items():
            with open_deck(v["path"]) as deck:
                fw = font_words(deck)
            for f, n in fw.items():
                cls.ruled18[f] = cls.ruled18.get(f, 0) + n
                if k in REFERENCE:
                    cls.ruled14[f] = cls.ruled14.get(f, 0) + n

    def test_ruled_definition(self):
        r14, r18 = self.ruled14, self.ruled18
        self.assertEqual(r14.get("Calibri"), 8212)
        self.assertEqual(r14.get("Space Grotesk"), 5502)
        self.assertEqual(r14.get("Space Grotesk Light"), 328)
        self.assertEqual(r14.get("Space Grotesk Medium"), 272)    # 227 of these inherit from a layout placeholder
        self.assertEqual(r18.get("Calibri"), 10105)
        self.assertEqual(next(iter(sorted(r18, key=lambda f: -r18[f]))), "Arial")

    def test_legacy_survey_method_reproduces(self):
        """Characteriser, not a tokenizer used anywhere: re-runs the superseded survey method."""
        legacy = {}
        for k in REFERENCE:
            with open_deck(DECKS[k]["path"]) as deck:
                for i in range(1, deck.n_slides + 1):
                    for r in deck.xml(deck.slide_part(i)).iter(A("r")):
                        rpr = r.find(A("rPr"))
                        lat = rpr.find(A("latin")) if rpr is not None else None
                        t = r.find(A("t"))
                        f = lat.get("typeface") if lat is not None else None
                        legacy[f] = legacy.get(f, 0) + len(((t.text or "") if t is not None else "").split())
        self.assertEqual(legacy["Calibri"], 9106)
        self.assertEqual(legacy["Space Grotesk"], 5858)
        self.assertEqual(legacy["Space Grotesk Light"], 327)
        self.assertEqual(legacy["Space Grotesk Medium"], 49)


# ---------------------------------------------------------------------------------------------- titles, lines, hygiene

def _leaf(lid, x, y, w, h, kind="sp", prst="rect", text=None, pt=12.0, fill="none", line="none", ph=None, **kw):
    """A synthetic shapes@1 leaf: enough of the record for pick_title, connect and layout_hygiene."""
    L = {"id": lid, "kind": kind, "z": lid, "layer": "content", "on_canvas": "in", "line_like": False,
         "group_path": [], "placeholder": {"type": ph} if ph else None, "geom": {"prst": prst, "cust": False},
         "bbox": {"x_in": x, "y_in": y, "w_in": w, "h_in": h},
         "frame": {"w_in": w, "h_in": h, "rot_deg": 0.0, "mirrored": False},
         "xfrm": {"flipH": False, "flipV": False},
         "corners_in": [[x, y], [x + w, y], [x + w, y + h], [x, y + h]],
         "fill": {"type": fill}, "line": {"type": line}, "grad": None, "effects": None, "text": None,
         "connector": None, "container_of": [], "icon_like": False}
    if text is not None:
        L["text"] = {"text": text, "words": words(text), "size": {"dominant_pt": pt}, "bold_share": 0.0,
                     "body": {"vert": "horz", "anchor": "t", "autofit": "none", "insets_in": [0.1, 0.05, 0.1, 0.05]},
                     "occupancy": None, "paragraphs": [{"text": text, "algn": "l",
                                                        "runs": [{"t": text, "size_rendered_pt": pt}]}]}
    L.update(kw)
    return L


def _line(lid, s, e, prst="line", head="none", tail="none", glue=None, kind="sp"):
    from deckkit.read import connect as cx
    x0, y0 = min(s[0], e[0]), min(s[1], e[1])
    L = _leaf(lid, x0, y0, abs(e[0] - s[0]), abs(e[1] - s[1]), kind=kind, prst=prst, line="solid", line_like=True)
    L["connector"] = cx.connector_block(prst, s, e, {"head": {"type": head}, "tail": {"type": tail}}, glue)
    return L


def _classify(leaves, canvas=(13.333, 7.5)):
    from deckkit.read import connect as cx
    area = canvas[0] * canvas[1]
    legend = cx.legend_lines(leaves, area)
    out = {}
    for L in leaves:
        c = L.get("connector")
        if c is not None:
            c["contact"] = cx.contact(tuple(c["start_in"]), tuple(c["end_in"]), leaves, L["id"], area,
                                      straight=c["orientation"] in ("h", "v", "diag"))
            c["class"] = out[L["id"]] = cx.classify(L, legend)
    return out


class TestTitleRule(unittest.TestCase):
    """walk.pick_title tiers on planted slides (13.333 x 7.5 in)."""
    W, H = 13.333, 7.5

    def pick(self, leaves, words_content=None):
        from deckkit.read.walk import pick_title
        L, tier = pick_title(leaves, self.W, self.H, words_content)
        return (L["id"] if L else None), tier

    def body(self, start, n, pt=12.0, y=3.0):
        return [_leaf(start + i, 1 + i * 2, y + (i % 3) * 0.6, 1.8, 0.4, text="body copy line %d here" % i, pt=pt)
                for i in range(n)]

    def test_placeholder_wins(self):
        leaves = [_leaf(1, 0.5, 0.3, 12, 0.6, text="Big headline", pt=40), _leaf(2, 0.5, 5, 6, 0.6,
                  text="Placeholder title", pt=18, ph="title")] + self.body(10, 4)
        self.assertEqual(self.pick(leaves), (2, "placeholder"))

    def test_empty_placeholder_skipped(self):
        leaves = [_leaf(2, 0.5, 5, 6, 0.6, text="", ph="title"), _leaf(1, 0.5, 0.3, 12, 0.6, text="Real title", pt=28)]
        leaves += self.body(10, 4)
        self.assertEqual(self.pick(leaves), (1, "top-band"))

    def test_top_band_ratio(self):
        """The top band (22%) and 1.3x the median body size: 16pt over a 12pt median is a top-band title; 15.5pt
        is not, and is found only as the slide's largest text; below a larger body line it is no title at all."""
        for pt, want in ((16.0, (1, "top-band")), (15.5, (1, "top-largest"))):
            leaves = [_leaf(1, 0.5, 1.5, 12, 0.5, text="Top band heading words", pt=pt)] + self.body(10, 5)
            self.assertEqual(self.pick(leaves), want, pt)
        leaves = [_leaf(1, 0.5, 1.5, 12, 0.5, text="Top band heading words", pt=15.5),
                  _leaf(2, 1, 6, 3, 0.4, text="A larger body line set low", pt=16)] + self.body(10, 5)
        self.assertEqual(self.pick(leaves), (None, None))
        leaves = [_leaf(1, 0.5, 1.7, 12, 0.5, text="Below the band", pt=30)] + self.body(10, 5, pt=12)
        self.assertNotEqual(self.pick(leaves)[1], "top-band")       # top at 22.7%: outside the band

    def test_few_words_divider(self):
        leaves = [_leaf(1, 0.5, 3.5, 12, 0.8, text="Section two", pt=28), _leaf(2, 0.5, 5, 4, 0.4, text="Notes", pt=40)]
        self.assertEqual(self.pick(leaves), (2, "few-words"))      # 3 words: the largest text anywhere

    def test_figures_never_titles(self):
        """A '02' divider numeral and a '5 to 10%' stat are figures, not headings."""
        leaves = [_leaf(1, 0.5, 0.2, 2, 1, text="02", pt=66), _leaf(2, 0.5, 0.4, 3, 0.6, text="5 to 10%", pt=60),
                  _leaf(3, 0.5, 1.2, 12, 0.5, text="The benchmark comes from three sources", pt=19)]
        leaves += self.body(10, 6, pt=9)
        self.assertEqual(self.pick(leaves), (3, "top-band"))

    def test_dominant_cover(self):
        """A cover title mid-canvas, 1.2x every other text, over 12 words."""
        leaves = [_leaf(1, 0.5, 3.0, 12, 0.8, text="Data Migration and Rationalization", pt=48),
                  _leaf(2, 0.5, 4.5, 12, 0.5, text="A proposal for the data and reporting migration programme", pt=20),
                  _leaf(3, 0.5, 6.5, 4, 0.3, text="September 2026", pt=14)]
        self.assertEqual(self.pick(leaves), (1, "dominant"))
        leaves[1]["text"]["size"]["dominant_pt"] = 44.0             # no longer dominant: no title
        self.assertEqual(self.pick(leaves), (None, None))

    def test_top_largest(self):
        leaves = [_leaf(1, 0.5, 0.2, 12, 0.5, text="The framework accommodates adjacent capabilities", pt=21.3),
                  _leaf(2, 5, 4, 2, 0.4, text="INTERMEDIATE", pt=21.1), _leaf(3, 8, 4, 2, 0.4, text="HIGH", pt=21.1)]
        leaves += self.body(10, 6, pt=17)
        self.assertEqual(self.pick(leaves), (1, "top-largest"))

    def test_chrome_never_titled(self):
        from deckkit.read.walk import guess_roles
        leaves = [_leaf(1, 0.5, 0.2, 12, 0.5, text="Layout banner text", pt=30)]
        guess_roles(leaves, self.W, self.H, find_title=False)
        self.assertNotEqual(leaves[0]["role_guess"]["role"], "title")
        guess_roles(leaves, self.W, self.H)
        self.assertEqual(leaves[0]["role_guess"]["role"], "title")


class TestLineClasses(unittest.TestCase):
    """connect.classify: degenerate, stroke, connector, rule; contact on two distinct shapes."""

    def boxes(self):
        return [_leaf(1, 1, 1, 2, 1, fill="solid"), _leaf(2, 5, 1, 2, 1, fill="solid")]

    def test_degenerate_and_short(self):
        got = _classify(self.boxes() + [_line(10, (3, 1.5), (3.04, 1.5)), _line(11, (3, 1.5), (3.1, 1.5), head="triangle")])
        self.assertEqual(got, {10: "degenerate", 11: "stroke"})

    def test_contact_two_shapes(self):
        """A bare line from one box edge to another is a connector; touching_both needs two distinct shapes."""
        leaves = self.boxes() + [_line(10, (3.0, 1.5), (5.05, 1.5))]
        self.assertEqual(_classify(leaves), {10: "connector"})
        ct = leaves[-1]["connector"]["contact"]
        self.assertEqual((ct["start"]["id"], ct["end"]["id"], ct["distinct"]), (1, 2, True))

    def test_head_and_glue(self):
        leaves = self.boxes() + [_line(10, (3.5, 3), (4.5, 3), tail="triangle"),
                                 _line(11, (3.5, 4), (4.5, 4), kind="cxn", prst="straightConnector1",
                                       glue={"start": {"id": 1, "idx": 0}, "end": None}),
                                 _line(12, (3.5, 5), (4.5, 5))]
        self.assertEqual(_classify(leaves), {10: "connector", 11: "connector", 12: "rule"})

    def test_rules(self):
        """A separator inside a card, a line running along a card's edge, a divider across a band: rules."""
        card = _leaf(1, 1, 1, 6, 3, fill="solid")
        inner_a, inner_b = _leaf(2, 1.5, 1.5, 2, 0.5, text="Left label"), _leaf(3, 4.5, 1.5, 2, 0.5, text="Right")
        band = _leaf(4, 0, 5, 13.333, 0.5, fill="solid")
        above, below = _leaf(5, 0, 4.5, 13.333, 0.5, fill="solid"), _leaf(6, 0, 5.5, 13.333, 0.5, fill="solid")
        leaves = [card, inner_a, inner_b, band, above, below,
                  _line(10, (1.2, 3.0), (6.8, 3.0)),          # separator inside the card, near nothing
                  _line(11, (1.0, 4.0), (7.0, 4.0)),          # along the card's bottom edge
                  _line(12, (3.0, 5.0), (3.0, 5.5))]          # a tick across the band, ends on its outline
        got = _classify(leaves)
        self.assertEqual(got, {10: "rule", 11: "rule", 12: "rule"})
        self.assertEqual(leaves[-1]["connector"]["contact"]["across"], 4)

    def test_icon_part_and_legend(self):
        legend = _leaf(1, 8, 5, 3, 1.5, fill="solid")
        leaves = [legend, _leaf(2, 8.1, 5.05, 1.5, 0.3, text="Legend"),
                  _line(10, (8.2, 5.6), (8.7, 5.6), tail="triangle"), _leaf(3, 8.8, 5.45, 2, 0.3, text="Data flow"),
                  _line(11, (8.2, 6.1), (8.7, 6.1)), _leaf(4, 8.8, 5.95, 2, 0.3, text="Private link"),
                  _line(12, (1, 1), (2, 1), tail="triangle")]
        leaves[-1]["icon_group"] = 99
        self.assertEqual(_classify(leaves), {10: "stroke", 11: "stroke", 12: "stroke"})

    def test_card_with_dividers_is_not_a_legend(self):
        """Vertical dividers between labels inside a card are rules, never legend samples."""
        leaves = [_leaf(1, 1, 1, 6, 1, line="solid"), _leaf(2, 1.1, 1.2, 1.5, 0.5, text="Planners validate"),
                  _line(10, (2.8, 1.2), (2.8, 1.8)), _leaf(3, 2.9, 1.2, 1.5, 0.5, text="Leaders approve")]
        self.assertEqual(_classify(leaves), {10: "rule"})


class TestAiHint(unittest.TestCase):
    def test_markers(self):
        from deckkit.read.pics import ai_hint
        for s in ("ChatGPT Image Aug 24, 2026, 08_33_15 PM.png", "made with DALL-E 3", "dalle render",
                  "Midjourney v6 hero", "Adobe Firefly fill", "Stable Diffusion XL", "Imagen 3 output",
                  "Gemini image of a plant", "Generated by AI using Midjourney"):
            self.assertIsNotNone(ai_hint("Picture 4", s, ""), s)
        for s in ("A logo with red and black text\n\nDescription automatically generated",
                  "Text\n\nDescription automatically generated with low confidence",
                  "A person smiling\n\nAI-generated content may be incorrect.",
                  "AI-generated monochrome distribution centre", "Generated CPG Supply Chain Background",
                  "Dallas skyline", "Fireflies at dusk"):
            self.assertIsNone(ai_hint("Picture 4", s, ""), s)


class TestLayoutHygiene(unittest.TestCase):
    """walk.layout_hygiene kinds, each planted once and each with its ids."""

    def run_(self, leaves, canvas=(13.333, 7.5)):
        from deckkit.read.walk import layout_hygiene
        flags = {}

        def add(kind, sid, detail=""):
            flags.setdefault(kind, []).append(sid)
        layout_hygiene(leaves, canvas[0], canvas[1], add)
        return flags

    def test_overlap(self):
        a = _leaf(1, 1, 1, 3, 0.4, text="First label text", pt=14)
        b = _leaf(2, 1.2, 1.05, 3, 0.4, text="Second label text", pt=14)
        self.assertEqual(sorted(self.run_([a, b])["overlap"]), [1, 2])
        card = _leaf(3, 5, 1, 4, 2, text="Card heading", fill="solid")
        lab = _leaf(4, 5.2, 1.2, 2, 0.4, text="Inside label")
        self.assertNotIn("overlap", self.run_([card, lab]))           # a label on a drawn card is inside it
        far = _leaf(5, 1, 1, 12, 0.5, text="Short title", pt=20)     # a wide frame, a short line of ink
        tab = _leaf(6, 10, 1, 0.6, 0.4, text="01", pt=12, fill="solid")
        self.assertNotIn("overlap", self.run_([far, tab]))

    def test_canvas_shadow_gradient_rotation_degenerate(self):
        off = _leaf(1, -0.2, 1, 2, 1, fill="solid", on_canvas="partial")
        edge = _leaf(2, -0.03, 2, 2, 1, fill="solid", on_canvas="partial")
        soft = _leaf(3, 3, 1, 2, 1, fill="solid",
                     effects={"shadow": {"blur_pt": 4.0, "alpha": 0.35}, "soft_shadow": True})
        crisp = _leaf(4, 3, 3, 2, 1, fill="solid", effects={"shadow": {"blur_pt": 0.5, "alpha": 0.35}})
        grad = _leaf(5, 6, 1, 2, 1, fill="solid", grad={"kind": "lin", "stops": []})
        rot = _leaf(6, 6, 3, 2, 0.4, text="Rotated label")
        rot["frame"]["rot_deg"] = 90.0
        rot["text"]["body"]["autofit"] = "shape"
        vert = _leaf(7, 9, 3, 0.4, 2, text="Vertical")
        vert["text"]["body"]["vert"] = "vert270"
        flat = _leaf(8, 9, 1, 2, 0.4, text="Plain spAutoFit")
        flat["text"]["body"]["autofit"] = "shape"
        dot = _line(9, (11, 6), (11.02, 6))
        dot["connector"]["class"] = "degenerate"
        f = self.run_([off, edge, soft, crisp, grad, rot, vert, flat, dot])
        self.assertEqual(f["partial-off-canvas"], [1])
        self.assertEqual(f["soft-shadow"], [3])
        self.assertEqual(f["gradient"], [5])
        self.assertEqual(sorted(f["render-divergence-risk"]), [6, 7])
        self.assertEqual(f["degenerate-connector"], [9])

    def test_read_slide_emits_kinds(self):
        """Through read_slide: the new kinds reach the hygiene list with ids."""
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as td:
            path = os.path.join(td, "hy.pptx")
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[6])
            a = s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(0.5))
            a.text_frame.text = "Overlapping words one"
            b = s.shapes.add_textbox(Inches(1.1), Inches(1.05), Inches(3), Inches(0.5))
            b.text_frame.text = "Overlapping words two"
            c = s.shapes.add_textbox(Inches(-0.5), Inches(4), Inches(2), Inches(0.5))
            c.text_frame.text = "Off the edge"
            prs.save(path)
            with open_deck(path) as deck:
                sl = read_slide(deck, 1)
        kinds = {h["kind"]: h for h in sl["hygiene"]}
        self.assertEqual(sorted(kinds["overlap"]["ids"]), sorted([a.shape_id, b.shape_id]))
        self.assertEqual(kinds["partial-off-canvas"]["ids"], [c.shape_id])


class TestOpenPaths(unittest.TestCase):
    """Open custGeom paths (stroked, unfilled) carry a connector block: ends at the path's first and last points
    through the shape's flips, arrowheads from headEnd/tailEnd, classified like a cxnSp. The kit's elbow is one."""

    def build(self, td):
        from pptx import Presentation
        from pptx.util import Inches
        from deckkit import prims
        from deckkit.theme import Theme
        skill = os.path.dirname(HERE)
        theme = Theme.load(os.path.join(skill, "themes", "sigmoid.json"))
        prs = Presentation(os.path.join(skill, "kit", "assets", "base-wide.pptx"))
        s = prs.slides.add_slide(prs.slide_layouts[0])
        with prims.use(theme):
            a = prims.rect(s, prims.Box(1, 1, 2, 1), fill="card", name="a")
            b = prims.rect(s, prims.Box(6, 3, 2, 1), fill="card", name="b")
            e = prims.elbow(s, [(3, 1.5), (4.5, 1.5), (4.5, 3.5), (6, 3.5)], name="route")
            flip = prims.elbow(s, [(1, 6), (3, 6), (3, 5)], head="end", name="flipped")
        flip._element.find(".//" + A("xfrm")).set("flipH", "1")     # mirrored: the head moves to x = 1
        arc = s.shapes.add_shape(1, Inches(9), Inches(4), Inches(2), Inches(2))
        filled = s.shapes.add_shape(1, Inches(9), Inches(1), Inches(1), Inches(1))
        from lxml import etree
        for sh, body, fill in ((arc, '<a:moveTo><a:pt x="0" y="100"/></a:moveTo><a:arcTo wR="50" hR="50" '
                                'stAng="10800000" swAng="16200000"/>', "none"),
                               (filled, '<a:moveTo><a:pt x="0" y="0"/></a:moveTo><a:lnTo><a:pt x="100" y="0"/>'
                                '</a:lnTo><a:lnTo><a:pt x="50" y="100"/></a:lnTo>', None)):
            sp = sh._element.spPr
            sp.remove(sp.find(A("prstGeom")))
            g = etree.fromstring('<a:custGeom xmlns:a="%s"><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>'
                                 '<a:rect l="0" t="0" r="r" b="b"/><a:pathLst><a:path w="100" h="100"%s>%s'
                                 '</a:path></a:pathLst></a:custGeom>'
                                 % (NS["a"], ' fill="none"' if fill else "", body))
            sp.insert(1, g)
            if fill:
                sh.fill.background()
            sh.line.width = 12700
        path = os.path.join(td, "open.pptx")
        prs.save(path)
        with open_deck(path) as deck:
            sl = read_slide(deck, 1)
        return {L["name"]: L for L in sl["shapes"]}, (a, b, e, flip, arc, filled)

    def test_elbow_and_arc(self):
        with tempfile.TemporaryDirectory() as td:
            by, (a, b, e, flip, arc, filled) = self.build(td)
        k = by["route"]["connector"]
        self.assertIsNotNone(k)
        self.assertEqual((k["prst"], k["orientation"], k["points"]), ("custom", "elbow", 4))
        self.assertEqual(k["start_in"], [3.0, 1.5])
        self.assertEqual(k["end_in"], [6.0, 3.5])
        self.assertAlmostEqual(k["path_len_in"], 1.5 + 2.0 + 1.5, places=2)
        self.assertNotEqual(k["tail"]["type"], "none")                # elbow's default head sits at the last point
        self.assertEqual(k["class"], "connector")
        self.assertTrue(k["touch"]["start"]["ok"] and k["touch"]["end"]["ok"])
        self.assertEqual({k["touch"]["start"]["id"], k["touch"]["end"]["id"]}, {a.shape_id, b.shape_id})
        self.assertTrue(by["route"]["line_like"])
        f = by["flipped"]["connector"]
        self.assertEqual(f["start_in"], [3.0, 6.0])                   # flipH mirrors the path in its box
        self.assertEqual(f["end_in"], [1.0, 5.0])
        c = by[arc.name]["connector"]
        self.assertEqual(c["orientation"], "curve")
        self.assertAlmostEqual(c["start_in"][0], 9.0, places=2)       # (0, 100) of a 100-unit path in a 2 in box
        self.assertAlmostEqual(c["start_in"][1], 6.0, places=2)
        self.assertGreater(c["path_len_in"], c["length_in"])
        self.assertIn(c["class"], ("rule", "stroke"))
        self.assertIsNone(by[filled.name]["connector"])               # an open path with a fill paints an area
        self.assertTrue(by[filled.name]["geom"]["open"])

    def test_loop_is_a_stroke(self):
        from deckkit.read import connect as cx
        L = _line(1, (1.0, 1.0), (1.5, 1.0), prst="custom")
        L["connector"]["path_len_in"] = 3.0
        L["connector"]["contact"] = {"distinct": False, "across": None}
        self.assertEqual(cx.classify(L), "stroke")                    # a ring drawn as one open stroke
        L["connector"]["head"] = {"type": "triangle"}
        self.assertEqual(cx.classify(L), "connector")                 # a cycle arrow
        L["connector"]["length_in"] = 0.01
        L["connector"]["head"] = {"type": "none"}
        self.assertEqual(cx.classify(L), "stroke")                    # chord tiny, drawn length real: not degenerate


class TestParagraphProps(unittest.TestCase):
    def test_margin_indent_end_size(self):
        from pptx import Presentation
        from pptx.util import Inches
        from lxml import etree
        with tempfile.TemporaryDirectory() as td:
            path = os.path.join(td, "para.pptx")
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[6])
            tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
            tf = tb.text_frame
            tf.text = "Bulleted line"
            p0 = tf.paragraphs[0]._p
            p0.get_or_add_pPr().set("marL", "228600")
            p0.get_or_add_pPr().set("indent", "-114300")
            p1 = tf.add_paragraph()._p
            etree.SubElement(p1, A("endParaRPr")).set("sz", "1400")
            prs.save(path)
            with open_deck(path) as deck:
                sl = read_slide(deck, 1)
        ps = sl["shapes"][0]["text"]["paragraphs"]
        self.assertEqual((ps[0]["mar_l_in"], ps[0]["indent_in"], ps[0]["end_pt"]), (0.25, -0.125, None))
        self.assertEqual((ps[1]["mar_l_in"], ps[1]["indent_in"], ps[1]["end_pt"]), (0.0, 0.0, 14.0))


class TestCoverTitle(unittest.TestCase):
    W, H = 13.333, 7.5

    def test_cover_tier(self):
        from deckkit.read.walk import cover_like, pick_title
        leaves = [_leaf(1, 1, 4.5, 6, 0.8, text="Programme planning and execution with agents", pt=24),
                  _leaf(2, 1, 5.2, 6, 0.5, text="Executive proposal for the planning team", pt=22),
                  _leaf(3, 1, 6.0, 4, 0.4, text="25 August 2026", pt=23)]
        leaves[0]["text"]["bold_share"] = 1.0
        self.assertEqual(pick_title(leaves, self.W, self.H, 17), (None, None))     # no tier separates them
        self.assertTrue(cover_like(leaves, self.W, self.H, 17, index1=1))
        L, tier = pick_title(leaves, self.W, self.H, 17, cover=True)
        self.assertEqual((L["id"], tier), (1, "cover"))
        self.assertFalse(cover_like(leaves, self.W, self.H, 17, index1=2))
        pic = _leaf(9, 0, 0, self.W, self.H, kind="pic", prst=None)
        self.assertTrue(cover_like(leaves[:1] + [pic], self.W, self.H, 8, index1=4))   # few words over a picture
        self.assertFalse(cover_like(leaves[:1] + [pic], self.W, self.H, 13, index1=4))
        small = [dict(leaves[0], text=dict(leaves[0]["text"], size={"dominant_pt": 12.0}))] + leaves[1:]
        self.assertNotEqual(pick_title(small, self.W, self.H, 17, cover=True)[1], "cover")  # bold but small
        head = _leaf(5, 0.3, 0.1, 6, 0.5, text="Data marketplace walkthrough", pt=21)
        label = _leaf(6, 4, 2.5, 1, 0.4, text="LIVE", pt=18)
        label["text"]["bold_share"] = 1.0
        self.assertEqual(pick_title([head, label, pic], self.W, self.H, 4, cover=True), (head, "few-words"))


_CORPUS = {}


def corpus():
    """Every slide of every deck, read once (d01 from the timed cache)."""
    if not _CORPUS:
        for k, v in DECKS.items():
            if k == "d01":
                _CORPUS[k] = d01()["slides"]
                continue
            with open_deck(v["path"]) as deck:
                _CORPUS[k] = [read_slide(deck, i) for i in range(1, deck.n_slides + 1)]
    return _CORPUS


class TestCorpusTitlesLines(unittest.TestCase):
    """Corpus facts for the title rule, the line classes and the ai-image hint (library DECISIONS 2026-10-03)."""

    def test_untitled_main_slides(self):
        """Visible main slides with words and no title: three profile pages set wholly at 8 to 9pt. The two
        covers whose 24pt bold title sits beside 23pt and 22pt lines (d08 and d14, slide 1) take the cover tier:
        the largest bold text on a deck's first slide."""
        missing = [(k, s["index"]) for k, ss in corpus().items() for s in ss
                   if s["status"] == "main" and not s["hidden"] and s["text"]["words_content"]
                   and not any((L["role_guess"] or {}).get("role") == "title" for L in s["shapes"])]
        self.assertEqual(sorted(missing), [("d09", 51), ("d09", 52), ("d09", 53)])
        for k in ("d08", "d14"):
            t = [L for L in corpus()[k][0]["shapes"] if (L["role_guess"] or {}).get("role") == "title"]
            self.assertEqual([(L["id"], L["text"]["size"]["dominant_pt"]) for L in t], [(545, 24.0)], k)
            self.assertIn("cover", t[0]["role_guess"]["why"])

    def test_line_classes(self):
        """cxnSp, line presets and block arrows: six lines whose two "shapes" were open custGeom strokes (lines
        themselves, now line-like) moved from connector to rule (2,039 and 1,667 distinct before). Open custGeom
        paths, stroked and unfilled, carry a connector block too: 806 of them, none classified before."""
        got, d1, both, cust = {}, {}, 0, {}
        for k, ss in corpus().items():
            for s in ss:
                for L in s["shapes"]:
                    c = L.get("connector")
                    if c is None:
                        g = L.get("geom") or {}
                        if g.get("open") and (L.get("fill") or {}).get("type") == "none" and \
                                (L.get("line") or {}).get("type") not in (None, "none"):
                            self.fail("open stroked path without a connector block: %s %s" % (k, L["id"]))
                        continue
                    self.assertIn(c["class"], ("connector", "rule", "stroke", "degenerate"))
                    if c["prst"] == "custom":
                        cust[c["class"]] = cust.get(c["class"], 0) + 1
                        continue
                    got[c["class"]] = got.get(c["class"], 0) + 1
                    if k == "d01":
                        d1[c["class"]] = d1.get(c["class"], 0) + 1
                    both += c["class"] == "connector" and c["contact"]["distinct"]
        self.assertEqual(got, {"connector": 2033, "rule": 2254, "stroke": 316, "degenerate": 32})
        self.assertEqual(d1, {"connector": 744, "rule": 687, "stroke": 37, "degenerate": 7})
        self.assertEqual(both, 1661)
        self.assertEqual(cust, {"connector": 261, "rule": 52, "stroke": 68, "degenerate": 425})

    def test_ai_hint_explicit_only(self):
        """126 pictures carried a hint under the old pattern (119 of them Office alt text); 6 name ChatGPT."""
        hits = sorted((k, s["index"], L["picture"]["ai_hint"]) for k, ss in corpus().items() for s in ss
                      for L in s["shapes"] + s["chrome"] if (L.get("picture") or {}).get("ai_hint"))
        self.assertEqual(hits, [("d01", 72, "ChatGPT"), ("d01", 118, "ChatGPT"), ("d01", 127, "ChatGPT"),
                                ("d01", 131, "ChatGPT"), ("d08", 5, "ChatGPT"), ("d14", 5, "ChatGPT")])


class TestZSourcesUntouched(unittest.TestCase):
    """A1: every source sha256 equals the manifest after the suite (runs last by name)."""

    def test_sha256(self):
        for k, v in DECKS.items():
            self.assertEqual(sha256(v["path"]), v["sha256"], k)


if __name__ == "__main__":
    unittest.main()
