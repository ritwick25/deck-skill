"""Native charts (deckkit.charts), the chart-part schema pass in package.part_xsd, chart frames in clone.copy_shapes,
the series slots against the RAG identities, and gate P27 (ruling 22, track C).

Run: cd ~/.claude/skills/deck && $PY -m unittest tests.test_charts -v
"""
from __future__ import annotations

import hashlib
import importlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(SKILL, "kit"))
sys.path.insert(0, os.path.join(SKILL, "gates"))

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402

from deckkit import charts, clone, package, prims  # noqa: E402
from deckkit.canvas import Box  # noqa: E402
from deckkit.errors import GeometryError, Overflows  # noqa: E402
from deckkit.theme import MACHADO, Theme, _simulate, delta_e00  # noqa: E402

PY = sys.executable
BASE = os.path.join(SKILL, "kit", "assets", "base-wide.pptx")
SIGMOID = os.path.join(SKILL, "themes", "sigmoid.json")
NEUTRAL = os.path.join(SKILL, "themes", "neutral.json")
TH = Theme.load(SIGMOID)
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
FRAME = Box(0.5, 1.0, 8.0, 4.0)
STAGES = (["Collect", "Rewrite", "Review", "Send"], [1, 2, 1, 1])


def sha(path) -> str:
    with open(path, "rb") as fh:
        return hashlib.sha256(fh.read()).hexdigest()


def new_deck():
    prims.bind(TH)
    prs = Presentation(BASE)
    return prs, prs.slides.add_slide(prs.slide_layouts[0])


def every_chart(s):
    """One chart of every type on fresh slides of s's deck; returns the plots."""
    prs = s.part.package.presentation_part.presentation
    out = [charts.bar(s, FRAME, *STAGES, subject="Rewrite", kind="illustrative", name="bar")]
    s2 = prs.slides.add_slide(prs.slide_layouts[0])
    out.append(charts.column(s2, FRAME, ["Before", "After"], [5, 1], subject="After", kind="illustrative",
                             name="col"))
    s3 = prs.slides.add_slide(prs.slide_layouts[0])
    out.append(charts.line(s3, FRAME, ["Q1", "Q2", "Q3", "Q4"], {"Briefed": [0, 200, 1200, 1200],
                                                                  "Target": [0, 400, 900, 1200]},
                           subject="Briefed", kind="target", name="line"))
    s4 = prs.slides.add_slide(prs.slide_layouts[0])
    out.append(charts.area(s4, FRAME, ["Q1", "Q2", "Q3"], [1, 3, 4], series_name="Reach", kind="illustrative",
                           name="area"))
    s5 = prs.slides.add_slide(prs.slide_layouts[0])
    out.append(charts.waterfall(s5, FRAME, [("Build", 180), ("Run", 288), ("Credit", -30)], total="Total",
                                kind="illustrative", name="fall"))
    s6 = prs.slides.add_slide(prs.slide_layouts[0])
    out.append(charts.stacked(s6, FRAME, ["Today", "Target"], {"Collect": [1, 0.5], "Rewrite": [2, 0.5]},
                              subject="Rewrite", kind="illustrative", name="stack"))
    s7 = prs.slides.add_slide(prs.slide_layouts[0])
    out.append(charts.scatter(s7, Box(0.5, 1.0, 8.0, 5.5), [("A", 1.5, 4.5), ("B", 4, 2), ("C", 2.5, 3.5)],
                              x_range=(1, 5), y_range=(1, 5), x_label="Effort", y_label="Value", subject="A",
                              kind="illustrative", name="dots"))
    s8 = prs.slides.add_slide(prs.slide_layouts[0])
    out.append(charts.scatter(s8, Box(0.5, 1.0, 8.0, 5.5), [("A", 1.5, 4.5), ("B", 4, 2)], x_range=(1, 5),
                              y_range=(1, 5), x_label="Effort", y_label="Value", sizes=[3, 1], kind="illustrative",
                              name="bubbles"))
    return out


class ChartBuild(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="charts-")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def save(self, prs, name="t.pptx"):
        out = os.path.join(self.tmp, name)
        package.save(prs, out, title="charts", theme=TH, manifest={"numbers": {}, "charts": charts.records(prs)},
                     overwrite=True, date="2026-10-04")
        return out

    def test_every_type_is_schema_valid_and_house_styled(self):
        prs, s = new_deck()
        plots = every_chart(s)
        out = self.save(prs)
        self.assertEqual(package.part_xsd(out), [])
        self.assertEqual(package.structural_check(out), [])
        roles = {(e.slot, tuple(e.mods)) for e in TH.roles.values()}
        with zipfile.ZipFile(out) as z:
            parts = sorted(n for n in z.namelist() if n.startswith("ppt/charts/chart"))
            self.assertEqual(len(parts), len(plots))
            for n in parts:
                r = etree.fromstring(z.read(n))
                self.assertEqual(r.find(C + "roundedCorners").get("val"), "0", n)
                self.assertEqual(r.find(C + "chart/" + C + "autoTitleDeleted").get("val"), "1", n)
                self.assertIsNone(r.find(C + "chart/" + C + "legend"), n)
                self.assertIsNone(r.find(C + "chart/" + C + "title"), n)
                self.assertIsNotNone(r.find(".//" + C + "manualLayout"), n)
                self.assertFalse(list(r.iter(A + "srgbClr")), n)
                self.assertFalse(list(r.iter(A + "gradFill")), n)
                for e in r.iter(A + "defRPr"):
                    self.assertGreaterEqual(int(e.get("sz")), 800, n)
                for e in r.iter(A + "schemeClr"):
                    mods = tuple((etree.QName(m).localname, int(m.get("val"))) for m in e)
                    self.assertIn((e.get("val"), mods), roles, n)

    def test_labels_land_on_data_coordinates(self):
        prs, s = new_deck()
        plot = charts.bar(s, FRAME, *STAGES, subject="Rewrite", kind="illustrative", name="b")
        shapes = {sh.name: sh for sh in s.shapes}
        for i, v in enumerate(STAGES[1]):
            lab = shapes["b.value.%d" % (i + 1)]
            self.assertAlmostEqual(lab.left / 914400, plot.val(v) + 0.08, places=2)
            self.assertAlmostEqual((lab.top + lab.height / 2) / 914400, plot.cat(i), places=2)
        self.assertEqual(shapes["b.value.2"].text_frame.paragraphs[0].runs[0].font.bold, True)
        # the pinned plot area in the part is the Plot's box
        ml = s.shapes[[sh.name for sh in s.shapes].index("b")].chart._chartSpace.find(".//" + C + "manualLayout")
        x = float(ml.find(C + "x").get("val"))
        self.assertAlmostEqual(FRAME.x + FRAME.w * x, plot.box.x, places=3)

    def test_one_basis_label_per_slide_and_kind_in_alt_text(self):
        prs, s = new_deck()
        charts.bar(s, Box(0.5, 1.0, 6.0, 3.0), *STAGES, kind="illustrative", name="a")
        charts.bar(s, Box(6.8, 1.0, 6.0, 3.0), *STAGES, kind="illustrative", name="b")
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        self.assertEqual(texts.count("Illustrative figures"), 1)
        descr = [sh._element.find(".//{*}cNvPr").get("descr") for sh in s.shapes if getattr(sh, "has_chart", False)]
        self.assertEqual([charts.frame_kind(d) for d in descr], ["illustrative", "illustrative"])
        _p, s2 = new_deck()
        charts.bar(s2, FRAME, *STAGES, kind="delivered", name="c")
        self.assertNotIn("Illustrative figures", [sh.text_frame.text for sh in s2.shapes if sh.has_text_frame])

    def test_refusals(self):
        prs, s = new_deck()
        with self.assertRaises(charts.ChartError):
            charts.bar(s, FRAME, ["A", "B"], [1, -2], kind="illustrative", name="neg")
        with self.assertRaises(charts.ChartError):
            charts.bar(s, FRAME, ["A", "B"], [1, 2], subject="C", kind="illustrative", name="nosubj")
        with self.assertRaises(charts.ChartError):
            charts.line(s, FRAME, ["a", "b"], {k: [1, 2] for k in "wxyz"}, kind="illustrative", name="four")
        with self.assertRaises(charts.ChartError):
            charts.line(s, FRAME, ["a", "b"], {"x": [1, 2], "y": [2, 1]}, area=True, kind="illustrative", name="sa")
        with self.assertRaises(charts.ChartError):
            charts.bar(s, FRAME, ["A", "B"], [1, 2], vmax=1.5, kind="illustrative", name="clip")
        with self.assertRaises(GeometryError):
            charts.bar(s, FRAME, ["A", "B"], [1, 2], kind="guess", name="kind")
        with self.assertRaises(charts.ChartError):
            charts.stacked(s, FRAME, ["a"], {k: [1] for k in "wxyz"}, kind="illustrative", name="st4")
        with self.assertRaises(Overflows):
            charts.column(s, Box(0.5, 1.0, 2.0, 3.0), ["A very long category name indeed"] * 6, [1] * 6,
                          kind="illustrative", name="crowd")
        with self.assertRaises(charts.ChartError):
            charts.annotate(charts.bar(s, FRAME, ["A", "B"], [1, 2], kind="illustrative", name="ok"), "A", "x",
                            kind="glow")
        # a refused chart leaves nothing behind: no shape, no chart part relationship
        self.assertEqual([sh.name for sh in s.shapes if not sh.name.startswith("ok")], [])

    def test_bullet_progress_and_ring(self):
        prs, s = new_deck()
        b = charts.bullet(s, Box(0.5, 1.0, 7.0, 1.8), [("Regions live", 4, 6), ("Brands covered", 1, 3)],
                          scale_max=6, subject="Regions live", kind="target", name="kpi")
        names = {sh.name: sh for sh in s.shapes}
        self.assertEqual(names["kpi.1.value"].text_frame.text, "4")
        self.assertAlmostEqual(names["kpi.1.bar"].height / 914400, charts.TRACK_IN, places=3)
        self.assertGreater(charts.TRACK_IN, 0.10)                      # never the accent-strip thickness
        self.assertAlmostEqual(names["kpi.1.target"].left / 914400, b.x(6), places=2)
        self.assertIn("Target figures", [sh.text_frame.text for sh in s.shapes if sh.has_text_frame])
        charts.progress(s, Box(0.5, 3.2, 7.0, 0.6), 1200, 1200, label="Reps briefed", kind="target", name="prog")
        with self.assertRaises(charts.ChartError):
            charts.bullet(s, Box(0.5, 4.0, 7.0, 0.6), [("Over", 9)], scale_max=6, kind="target", name="over")
        _p, s2 = new_deck()
        with self.assertRaises(charts.ChartError):                     # no basis label on the slide yet
            charts.ring(s2, Box(1, 1, 2, 2), 3, 5, kind="illustrative", name="r")
        charts.ring(s2, Box(1, 1, 2, 2), 3, 5, kind="delivered", name="r")
        with self.assertRaises(charts.ChartError):
            charts.ring(s2, Box(4, 1, 2, 2), 5, 5, kind="delivered", name="r2")

    def test_waterfall_arithmetic_and_carry_lines(self):
        prs, s = new_deck()
        p = charts.waterfall(s, FRAME, [("Build", 180), ("Run", 288)], total="Total", fmt="${:,.0f}K",
                             kind="illustrative", name="w")
        self.assertEqual(p.series["base"], [0.0, 180.0, 0.0])
        self.assertEqual(p.series["value"], [180.0, 288.0, 468.0])
        names = {sh.name: sh for sh in s.shapes}
        self.assertEqual(names["w.value.3"].text_frame.text, "$468K")
        self.assertIn("w.carry.1", names)
        self.assertIn("w.carry.2", names)
        x = names["w.carry.1"].left / 914400
        self.assertAlmostEqual(x, p.val(180), places=2)

    def test_annotations(self):
        prs, s = new_deck()
        p = charts.line(s, Box(0.5, 1.0, 9.0, 4.5), ["a", "b", "c", "d"], {"Reps": [0, 200, 1200, 1200]},
                        end_labels=False, kind="illustrative", name="l")
        out = charts.annotate(p, "c", "1,200 reps, six regions", kind="crossover")
        names = [o.name for o in out]
        self.assertEqual(names, ["l.note.1.drop", "l.note.1.dot", "l.note.1.text"])
        dot = out[1]
        cx = (dot.left + dot.width / 2) / 914400
        self.assertAlmostEqual(cx, p.point(2, 1200)[0], places=2)
        band = charts.annotate(p, 300, "Target", kind="band", to=700)[0]
        order = [sh.name for sh in s.shapes]
        self.assertLess(order.index(band.name), order.index("l"))      # behind the chart frame
        with self.assertRaises(Overflows):
            q = charts.line(s, Box(0.5, 1.0, 9.0, 4.5), ["a", "b"], {"R": [0, 10]}, kind="illustrative", name="m")
            charts.annotate(q, 6, "", kind="band", to=10)                # the end label straddles its top edge

    def test_records_round_trip_and_p27(self):
        import p27_charts
        from _shared import Ctx, run_gate
        prs, s = new_deck()
        every_chart(s)
        out = self.save(prs)
        with open(os.path.join(self.tmp, "deck.json"), "w") as fh:
            json.dump({"theme": SIGMOID}, fh)
        code, found = run_gate(p27_charts, out, Ctx(out))
        self.assertEqual(code, 0, [f.line() for f in found if f.level == "fail"])
        with open(out.replace(".pptx", ".manifest.json")) as fh:
            rec = json.load(fh)["charts"]
        self.assertEqual(sorted(rec), sorted(["bar", "col", "line", "area", "fall", "stack", "dots", "bubbles"]))
        self.assertEqual(rec["fall"]["series"][1]["val"], [180.0, 288.0, 30.0, 438.0])

    def test_build_is_deterministic(self):
        a = os.path.join(self.tmp, "a", "x.pptx")
        b = os.path.join(self.tmp, "b", "x.pptx")
        for o in (a, b):
            r = subprocess.run([PY, os.path.join(SKILL, "scripts", "example_charts.py"), "--out", o],
                               capture_output=True, text=True, timeout=300)
            self.assertEqual(r.returncode, 0, r.stderr[-2000:])
        self.assertEqual(sha(a), sha(b))

    def test_clone_carries_chart_and_workbook(self):
        prs, s = new_deck()
        charts.bar(s, FRAME, *STAGES, subject="Rewrite", kind="illustrative", name="bar")
        src = self.save(prs, "src.pptx")
        sp = Presentation(src)
        dst = Presentation(BASE)
        for _ in range(2):                         # twice into one package: two chart parts, never shared
            d = dst.slides.add_slide(dst.slide_layouts[0])
            m = clone.copy_shapes(sp.slides[0], d, strip=clone.CLEAN_STRIP)
            self.assertTrue(any("/ppt/charts/chart" in n for n in m.parts_added))
        for k, sl in enumerate(dst.slides):
            for sh in sl.shapes:
                if sh.name.startswith("bar"):
                    sh.name = sh.name.replace("bar", "bar%d" % k, 1)
        out = os.path.join(self.tmp, "dst.pptx")
        package.save(dst, out, title="c", theme=TH, manifest={"numbers": {}, "charts": charts.records(dst)},
                     overwrite=True)
        self.assertEqual(package.structural_check(out), [])
        self.assertEqual(package.part_xsd(out), [])
        with zipfile.ZipFile(out) as z:
            names = z.namelist()
            self.assertEqual(len([n for n in names if n.startswith("ppt/charts/chart")]), 2)
            books = [n for n in names if n.startswith("ppt/embeddings/")]
            self.assertEqual(len(books), 2)
            zipfile.ZipFile(io.BytesIO(z.read(books[0]))).testzip()
        back = Presentation(out)
        gf = [sh for sh in back.slides[1].shapes if getattr(sh, "has_chart", False)][0]
        self.assertEqual(gf.chart.plots[0].categories[1], "Rewrite")
        self.assertEqual(charts.frame_kind(gf._element.find(".//{*}cNvPr").get("descr")), "illustrative")

    def test_extended_chart_still_refused(self):
        from deckkit.errors import UnsupportedPart
        prs, s = new_deck()
        charts.bar(s, FRAME, *STAGES, kind="illustrative", name="bar")
        gd = [sh for sh in s.shapes if getattr(sh, "has_chart", False)][0]._element.find(".//" + A + "graphicData")
        gd.set("uri", "http://schemas.microsoft.com/office/drawing/2014/chartex")
        d = Presentation(BASE).slides.add_slide(Presentation(BASE).slide_layouts[0])
        with self.assertRaises(UnsupportedPart):
            clone.copy_shapes(s, d)


class SchemaAndTheme(unittest.TestCase):
    def test_part_xsd_catches_a_bogus_chart_element(self):
        prs, s = new_deck()
        charts.bar(s, FRAME, *STAGES, kind="delivered", name="bar")
        tmp = tempfile.mkdtemp(prefix="charts-xsd-")
        try:
            out = os.path.join(tmp, "x.pptx")
            prs.save(out)
            with zipfile.ZipFile(out) as z:
                m = {n: z.read(n) for n in z.namelist()}
            self.assertEqual(package.part_xsd(out), [])
            root = etree.fromstring(m["ppt/charts/chart1.xml"])
            root.find(C + "chart").addprevious(etree.Element(C + "bogus"))
            m["ppt/charts/chart1.xml"] = etree.tostring(root)
            package.write_members(m, out)
            errs = package.part_xsd(out)
            self.assertTrue(errs and "bogus" in errs[0], errs)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_python_pptx_signed_axis_ids_are_known_not_errors(self):
        self.assertIsNotNone(package.known_value(
            "Element '{http://schemas.openxmlformats.org/drawingml/2006/chart}axId', attribute 'val': '-2068027336' "
            "is not a valid value of the atomic type 'xs:unsignedInt'."))

    def test_series_slots_never_read_as_status(self):
        for path in (SIGMOID, NEUTRAL):
            th = Theme.load(path)
            rag = [v["hex"] for k, v in th.identity.items() if k.startswith("rag-")]
            self.assertEqual(len(rag), 3)
            for i in range(1, 6):
                h = th.hex("series-%d" % i)
                for r in rag:
                    worst = min([delta_e00(h, r)] + [delta_e00(_simulate(h, k), _simulate(r, k)) for k in MACHADO])
                    self.assertGreaterEqual(worst, 10.0, "%s series-%d %s vs %s" % (th.name, i, h, r))

    def test_chart_order_is_separable_on_both_themes(self):
        for path in (SIGMOID, NEUTRAL):
            self.assertEqual(charts.check_theme(Theme.load(path)), [], path)

    def test_nice_ceiling(self):
        self.assertEqual(charts.nice_ceiling(1200, 4, headroom=0.08), (1500, 500))
        top, unit = charts.nice_ceiling(468, 4)
        self.assertGreaterEqual(top, 468)
        self.assertAlmostEqual(top / unit, round(top / unit))

    def test_basis_text_matches_the_stat_band(self):
        if hasattr(prims, "basis_label"):
            for k, v in charts.BASIS.items():
                self.assertEqual(v, prims.basis_label(k))


class Planters(unittest.TestCase):
    def test_every_chart_planter_fails_its_gate(self):
        r = subprocess.run([PY, "-B", os.path.join(SKILL, "gates", "negatives_charts.py")], capture_output=True,
                           text=True, timeout=600)
        self.assertEqual(r.returncode, 0, r.stdout[-3000:] + r.stderr[-1500:])
        import negatives_charts
        self.assertIn("%d chart planters, 0 misbehaved" % len(negatives_charts.PLANTERS), r.stdout)

    def test_p27_is_registered(self):
        with open(os.path.join(SKILL, "gates", "registry.json")) as fh:
            reg = json.load(fh)["gates"]
        g = [x for x in reg if x["id"] == "P27"]
        self.assertEqual(len(g), 1)
        self.assertEqual(g[0]["module"], "p27_charts")
        importlib.import_module("p27_charts")


if __name__ == "__main__":
    unittest.main()
