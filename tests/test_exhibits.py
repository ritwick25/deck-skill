"""The exhibit components (deckkit.exhibits, deckkit.tables) and scripts/capture.py (ruling 22, track A2).

Run: cd ~/.claude/skills/deck && $PY -m unittest tests.test_exhibits -v
Every component is drawn on a fresh WIDE deck in the sigmoid theme; checks read the drawn shapes back. The last
test builds scripts/example_exhibits.py into a scratch folder and runs the geometry and colour gates on it.
"""
from __future__ import annotations

import hashlib
import io
import json
import math
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL / "kit"))
sys.path.insert(0, str(SKILL / "scripts"))

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402

from deckkit import exhibits, prims, tables  # noqa: E402
from deckkit.canvas import EMU_PER_IN, WIDE, Box  # noqa: E402
from deckkit.errors import BelowFloor, Capacity, GeometryError, Overflows, UnsupportedPart  # noqa: E402
from deckkit.read.ns import NS  # noqa: E402
from deckkit.theme import Theme  # noqa: E402

A = "{%s}" % NS["a"]
P = "{%s}" % NS["p"]
PY = "/Library/Frameworks/Python.framework/Versions/3.14/bin/python3"
BASE = SKILL / "kit" / "assets" / "base-wide.pptx"
TH = Theme.load(SKILL / "themes" / "sigmoid.json")
C = WIDE.content
STAGES6 = [("scan-search", "Scan", "Approved feeds"), ("list-ordered", "Rank", "By territory"),
           ("file-text", "Draft", "A source per claim"), ("user-check", "Review", "Medical signs"),
           ("send", "Release", "Monday"), ("refresh-cw", "Learn", "Ratings retune")]


def slide():
    prs = Presentation(str(BASE))
    return prs, prs.slides.add_slide(prs.slide_layouts[0])


def setUpModule():
    prims.bind(TH)


def tearDownModule():
    prims._ACTIVE.clear()


def named(s):
    return {sh.name: sh for sh in s.shapes}


def box_of(sh):
    return Box(sh.left / EMU_PER_IN, sh.top / EMU_PER_IN, sh.width / EMU_PER_IN, sh.height / EMU_PER_IN)


def overlaps(a, b, tol=0.0):
    return a.x < b.r - tol and b.x < a.r - tol and a.y < b.b - tol and b.y < a.b - tol


def png(w, h, colour=(180, 185, 190)):
    buf = io.BytesIO()
    Image.new("RGB", (w, h), colour).save(buf, "PNG")
    return buf.getvalue()


class CycleRing(unittest.TestCase):
    def test_wedges_one_ring_six_paths_labels_outside(self):
        _prs, s = slide()
        out = exhibits.cycle_ring(s, C, STAGES6, hub=("Weekly brief", "One cycle"), name="loop")
        sh = named(s)
        ring = sh["loop.ring"]
        self.assertEqual(len(ring._element.findall(".//" + A + "path")), 6)       # one closed path per stage
        self.assertEqual(sum(1 for n in sh if n.startswith("loop.") and n.endswith(".icon")), 6)
        sq = out["ring"]
        for i, lb in enumerate(out["labels"], 1):
            self.assertFalse(overlaps(lb, sq), "label %d crosses the ring's square" % i)
            self.assertTrue(C.contains(lb, 1e-6), "label %d leaves the box" % i)
            for j, other in enumerate(out["labels"][i:], i + 1):
                self.assertFalse(overlaps(lb, other), "labels %d and %d overlap" % (i, j))
        self.assertTrue(box_of(ring).contains(out["hub"], 0.01))
        self.assertTrue(box_of(ring).contains(box_of(sh["loop.hub"]), 0.01))

    def test_segments_point_clockwise_and_are_equal(self):
        _prs, s = slide()
        exhibits.cycle_ring(s, C, STAGES6, name="loop")
        paths = named(s)["loop.ring"]._element.findall(".//" + A + "path")
        counts = {len(p) for p in paths}
        self.assertEqual(len(counts), 1, "every segment has the same construction")

    def test_active_wedge_shares_the_ring_frame(self):
        _prs, s = slide()
        exhibits.cycle_ring(s, C, STAGES6, active=2, name="loop")
        sh = named(s)
        self.assertEqual(box_of(sh["loop.ring"]), box_of(sh["loop.3.active"]))

    def test_nodes_arrows_land_on_node_outlines(self):
        _prs, s = slide()
        exhibits.cycle_ring(s, C, STAGES6[:5], hub="Quarterly loop", style="nodes", name="ring")
        sh = named(s)
        nodes = [box_of(sh["ring.%d.node" % i]) for i in range(1, 6)]
        for i in range(1, 6):
            arrow = sh["ring.%d.next" % i]
            ln = arrow._element.find(".//" + A + "ln")
            self.assertIsNotNone(ln.find(A + "tailEnd"), "arrow %d has a head" % i)
            path = arrow._element.find(".//" + A + "path")
            self.assertEqual(path.get("fill"), "none")
            pts = path.findall(".//" + A + "pt")
            fx, fy = arrow.left / EMU_PER_IN, arrow.top / EMU_PER_IN
            start = (fx + int(pts[0].get("x")) / EMU_PER_IN, fy + int(pts[0].get("y")) / EMU_PER_IN)
            end = (fx + int(pts[-1].get("x")) / EMU_PER_IN, fy + int(pts[-1].get("y")) / EMU_PER_IN)
            for pt, node in ((start, nodes[i - 1]), (end, nodes[i % 5])):
                d = math.hypot(pt[0] - node.cx, pt[1] - node.cy)
                self.assertAlmostEqual(d, node.w / 2, delta=0.01)

    def test_refusals_leave_the_slide_unchanged(self):
        _prs, s = slide()
        n0 = len(s.shapes)
        with self.assertRaises(Capacity):
            exhibits.cycle_ring(s, C, STAGES6[:2], name="loop")
        with self.assertRaises(Capacity):
            exhibits.cycle_ring(s, C, STAGES6 + STAGES6[:3], name="loop")
        with self.assertRaises(Overflows):
            exhibits.cycle_ring(s, Box(C.x, C.y, 4.0, 2.0), STAGES6, name="loop")
        with self.assertRaises(Exception):
            exhibits.cycle_ring(s, C, STAGES6, hub_fill="card", name="loop")
        self.assertEqual(len(s.shapes), n0)


class MarkWall(unittest.TestCase):
    GROUPS = [("Cloud", [("vendor:aws/ml/bedrock", "Amazon Bedrock"), ("vendor:gcp/ml/vertex-ai", "Vertex AI")]),
              ("Agents", [("mark:langfuse-icon", "Langfuse"), ("mark:claude-spark", "Claude"),
                          ("mono:QX", "Queue service")]),
              ("Delivery", [("vendor:programming/language/python", "Python")])]

    def test_marks_captions_tabs_no_pills(self):
        _prs, s = slide()
        out = exhibits.mark_wall(s, C, self.GROUPS, name="wall")
        sh = named(s)
        pics = [x for x in s.shapes if x.shape_type == 13]
        self.assertEqual(len(pics), 5)
        for p in pics:
            descr = p._element.find(".//" + P + "cNvPr").get("descr")
            self.assertTrue("(vendor mark " in descr or "(brand mark " in descr, descr)
        caps = [n for n in sh if n.endswith(".caption")]
        self.assertEqual(len(caps), 6)
        self.assertEqual(sum(1 for n in sh if n.endswith(".tab")), 3)
        for x in s.shapes:
            g = x._element.find(".//" + A + "prstGeom")
            if g is not None and g.get("prst") == "roundRect":
                adj = g.find(".//" + A + "gd")
                self.assertFalse(adj is not None and adj.get("fmla") == "val 50000", "a pill: %s" % x.name)
        widths = [b.w for b in out["groups"]]
        self.assertGreater(max(widths), min(widths) + 0.3, "columns follow their marks, not an even tile grid")

    def test_first_marks_share_one_line(self):
        _prs, s = slide()
        exhibits.mark_wall(s, C, self.GROUPS, name="wall")
        tops = {round(x.top / EMU_PER_IN + x.height / EMU_PER_IN / 2, 2) for x in s.shapes
                if x.name.startswith(("vendor.", "mark.")) and x.name.endswith(("-1.1", "-2.1", "-3.1"))}
        self.assertEqual(len(tops), 1, tops)

    def test_refusals(self):
        _prs, s = slide()
        n0 = len(s.shapes)
        with self.assertRaises(UnsupportedPart):
            exhibits.mark_wall(s, C, [("X", [("vendor:aws/nope/none", "Nope")])], name="wall")
        with self.assertRaises(UnsupportedPart):
            exhibits.mark_wall(s, C, [("X", [("logo:aws", "Nope")])], name="wall")
        with self.assertRaises(GeometryError):
            exhibits.mark_wall(s, C, [("X", [("vendor:aws/ml/bedrock", "")])], name="wall")
        with self.assertRaises(Overflows):
            exhibits.mark_wall(s, Box(C.x, C.y, 3.0, 1.0), self.GROUPS, name="wall")
        self.assertEqual(len(s.shapes), n0)


class Portraits(unittest.TestCase):
    def test_never_a_silhouette(self):
        _prs, s = slide()
        with self.assertRaises(GeometryError):
            exhibits.portrait(s, Box(1, 1, 1, 1), name="p")
        self.assertEqual(len(s.shapes), 0)

    def test_initials_disc(self):
        _prs, s = slide()
        exhibits.portrait(s, Box(1, 1, 1, 1), initials="EL", name="p")
        sh = named(s)
        self.assertEqual(sh["p"]._element.find(".//" + A + "prstGeom").get("prst"), "ellipse")
        self.assertEqual(sh["p.initials"].text_frame.text, "EL")

    def test_photo_clipped_to_circle_and_square_cropped(self):
        _prs, s = slide()
        pic = exhibits.portrait(s, Box(1, 1, 1, 1), image=png(400, 600), alt="Programme sponsor", name="p")
        self.assertEqual(pic._element.find(".//" + A + "prstGeom").get("prst"), "ellipse")
        self.assertAlmostEqual(pic.crop_top + pic.crop_bottom, 1 - 400 / 600, places=3)
        self.assertIn("Programme sponsor", pic._element.find(".//" + P + "cNvPr").get("descr"))

    def test_photo_refusals(self):
        _prs, s = slide()
        with self.assertRaises(GeometryError):
            exhibits.portrait(s, Box(1, 1, 1, 1), image=png(120, 120), alt="Lead", name="p")    # under 192 dpi
        with self.assertRaises(GeometryError):
            exhibits.portrait(s, Box(1, 1, 1, 1), image=png(400, 400), name="p")                # no alt
        self.assertEqual(len(s.shapes), 0)

    def test_quote_portrait_panel_hugs_its_words(self):
        _prs, s = slide()
        used = exhibits.quote_portrait(s, Box(C.x, C.y, C.w, 3.0), "Every brief carries its sources.",
                                       "Engagement lead", initials="EL", name="q")
        sh = named(s)
        panel = box_of(sh["q.quote"])
        self.assertLess(panel.h, 1.5)
        self.assertLessEqual(used.h, 3.0)
        self.assertIn("q.portrait", sh)


class ScreenFrame(unittest.TestCase):
    def capture(self, tmp, w=2560, h=1180, css_w=1280, min_px=14):
        p = Path(tmp) / "shot.png"
        p.write_bytes(png(w, h, (240, 242, 245)))
        rec = {"css_width": css_w, "css_height": h // 2, "scale": 2.0, "min_font_px": min_px,
               "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
        (Path(tmp) / "shot.capture.json").write_text(json.dumps(rec))
        return p

    def test_browser_fills_height_and_records_type(self):
        with tempfile.TemporaryDirectory() as tmp:
            _prs, s = slide()
            out = exhibits.screen_frame(s, C, self.capture(tmp), name="scr")
            self.assertGreater(out["window"].h, 0.95 * C.h)
            self.assertAlmostEqual(out["type_pt"], 14 * out["screen"].w * 72 / 1280, places=1)
            pic = named(s)["scr.screen"]
            self.assertIn("smallest type", pic._element.find(".//" + P + "cNvPr").get("descr"))

    def test_tall_capture_is_cropped_not_shrunk(self):
        with tempfile.TemporaryDirectory() as tmp:
            _prs, s = slide()
            out = exhibits.screen_frame(s, C, self.capture(tmp, w=2560, h=3000), name="scr")
            pic = named(s)["scr.screen"]
            self.assertGreater(pic.crop_bottom, 0.3)
            self.assertAlmostEqual(out["window"].w, C.w, delta=0.01)

    def test_floor_refusal(self):
        with tempfile.TemporaryDirectory() as tmp:
            _prs, s = slide()
            with self.assertRaises(BelowFloor):
                exhibits.screen_frame(s, C, self.capture(tmp, min_px=9), name="scr")
            self.assertEqual(len(s.shapes), 0)

    def test_needs_a_record_of_its_type(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "bare.png"
            p.write_bytes(png(2560, 1180))
            _prs, s = slide()
            with self.assertRaises(UnsupportedPart):
                exhibits.screen_frame(s, C, p, name="scr")
            exhibits.screen_frame(s, C, p, chrome=None, type_pt=0, name="scr")     # a capture with no text
            p2 = self.capture(tmp)
            p2.write_bytes(png(2560, 1180, (1, 2, 3)))                              # sidecar now stale
            with self.assertRaises(UnsupportedPart):
                exhibits.screen_frame(s, C, p2, name="scr2")

    def test_band_bleeds_and_carries_the_pill(self):
        with tempfile.TemporaryDirectory() as tmp:
            _prs, s = slide()
            exhibits.screen_frame(s, C, self.capture(tmp, w=1600, h=1180, css_w=800), chrome="laptop",
                                  band="ink-deep", pill="Week 41", name="scr")
            sh = named(s)
            band = box_of(sh["bleed.band"])
            self.assertAlmostEqual(band.x, 0.0, places=3)
            self.assertAlmostEqual(band.w, WIDE.w, places=2)
            self.assertIn("scr.pill", sh)
            self.assertGreaterEqual(box_of(sh["scr.base"]).x, C.x - 1e-3)
            with self.assertRaises(GeometryError):
                exhibits.screen_frame(s, C, self.capture(tmp), pill="Week 41", name="scr2")


class Table(unittest.TestCase):
    ROWS = [["Time to first brief", "6 weeks", "4 weeks"], ["Approved sources", "Yes", "No"]]

    def test_header_marked_column_and_glyph_cells(self):
        _prs, s = slide()
        gf, h, tops, lefts = tables.table(s, Box(C.x, C.y, 8.0, 3.0), ["Criterion", "Option A", "Option B"],
                                          self.ROWS, marked_col=2,
                                          cell_glyphs={"Yes": "circle-check", "No": ("circle-x", "ink-muted")},
                                          name="opt")
        self.assertAlmostEqual(h, 3.0, places=3)                     # the rows share the box's height
        tbl = gf.table
        fill = lambda r, c: tbl.cell(r, c)._tc.tcPr.find(A + "solidFill")[0].get("val")   # noqa: E731
        self.assertEqual(fill(0, 2), TH.spec("accent").get("val"))
        self.assertEqual(fill(0, 1), TH.spec("ink-strong").get("val"))
        self.assertEqual(fill(1, 2), TH.spec("card").get("val"))
        icons = sorted(n for n in named(s) if n.endswith(".icon"))
        self.assertEqual(icons, ["opt.r2.c2.icon", "opt.r2.c3.icon"])
        self.assertEqual(tbl._tbl.tblPr.find(A + "tableStyleId").text, tables.NO_TABLE_STYLE)

    def test_overflow_refused(self):
        _prs, s = slide()
        with self.assertRaises(Overflows):
            tables.table(s, Box(C.x, C.y, 8.0, 1.0), ["A", "B"], [["x", "y"]] * 12, name="t")
        with self.assertRaises(GeometryError):
            tables.table(s, C, ["A", "B"], [["x"]], name="t")
        self.assertEqual(len(s.shapes), 0)


class CellGrid(unittest.TestCase):
    LEGEND = {"live": ("circle-check", "Live", "accent"), "pilot": ("clock", "Pilot", "ink-strong"),
              "next": ("calendar", "Next wave", "card"), "out": ("circle-x", "Not planned", "paper")}

    def test_every_cell_glyph_plus_word(self):
        _prs, s = slide()
        cells = [["live", "pilot", ("next", "Wave 2")], ["out", "live", "pilot"]]
        out = exhibits.cell_grid(s, C, ["Scan", ("Review", "Medical signs")], ["US", "DE", "JP"], cells,
                                 self.LEGEND, name="g")
        sh = named(s)
        self.assertEqual(len(out["cells"]), 6)
        for r in range(1, 3):
            for c in range(1, 4):
                self.assertIn("g.r%d.c%d.icon" % (r, c), sh)
                self.assertTrue(sh["g.r%d.c%d.text" % (r, c)].text_frame.text.strip())
        self.assertEqual(sh["g.r1.c3.text"].text_frame.text, "Wave 2")
        self.assertIn("g.key.next", sh)                                    # an override word brings the key

    def test_refusals(self):
        _prs, s = slide()
        with self.assertRaises(GeometryError):
            exhibits.cell_grid(s, C, ["A"], ["X", "Y"], [["live"]], self.LEGEND, name="g")
        with self.assertRaises(GeometryError):
            exhibits.cell_grid(s, C, ["A"], ["X"], [["maybe"]], self.LEGEND, name="g")
        with self.assertRaises(Overflows):
            exhibits.cell_grid(s, C, ["A"], ["X"] * 14, [["out"] * 14], self.LEGEND, name="g")
        self.assertEqual(len(s.shapes), 0)


def _playwright():
    try:
        import playwright.sync_api  # noqa: F401
        return True
    except Exception:  # noqa: BLE001
        return False


@unittest.skipUnless(_playwright(), "Playwright is not installed on this interpreter")
class Capture(unittest.TestCase):
    PAGE = ("<!doctype html><html><head><meta charset='utf-8'><style>body{margin:0;width:640px;height:300px;"
            "font-family:Arial}p{font-size:%dpx}</style></head><body><h1>Queue</h1><p>small words</p></body></html>")

    def test_records_the_smallest_type(self):
        import capture
        with tempfile.TemporaryDirectory() as tmp:
            html = Path(tmp) / "p.html"
            html.write_text(self.PAGE % 14)
            rec = capture.capture(html, Path(tmp) / "p.png", width=640, height=300, placed_width_in=6.0)
            self.assertEqual(rec["min_font_px"], 14)
            self.assertAlmostEqual(rec["min_pt"], 14 * 6.0 * 72 / 640, places=2)
            side = json.loads((Path(tmp) / "p.capture.json").read_text())
            self.assertEqual(side["sha256"], hashlib.sha256((Path(tmp) / "p.png").read_bytes()).hexdigest())
            with Image.open(Path(tmp) / "p.png") as im:
                self.assertEqual(im.size, (1280, 600))

    def test_refuses_under_the_floor_and_writes_nothing(self):
        import capture
        with tempfile.TemporaryDirectory() as tmp:
            html = Path(tmp) / "p.html"
            html.write_text(self.PAGE % 9)
            with self.assertRaises(SystemExit):
                capture.capture(html, Path(tmp) / "p.png", width=640, height=300, placed_width_in=6.0)
            self.assertFalse((Path(tmp) / "p.png").exists())

    def test_refuses_external_requests(self):
        import capture
        with tempfile.TemporaryDirectory() as tmp:
            html = Path(tmp) / "p.html"
            html.write_text("<html><head><link href='https://example.com/a.css' rel='stylesheet'></head></html>")
            with self.assertRaises(SystemExit):
                capture.capture(html, Path(tmp) / "p.png")


ASSET = SKILL / ".cache" / "exhibits" / "assets" / "review-queue.png"


@unittest.skipUnless(ASSET.exists() or _playwright(), "no capture asset and no Playwright to make one")
class ExampleDeck(unittest.TestCase):
    def test_example_passes_the_geometry_and_colour_gates(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "exhibits.pptx"
            r = subprocess.run([PY, str(SKILL / "scripts" / "example_exhibits.py"), "--out", str(out)],
                               capture_output=True, text=True, timeout=300)
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            g = subprocess.run([PY, str(SKILL / "gates" / "run_all.py"), "--deck", str(out), "--only",
                                "P01,P02,P03,P04,P05,P06,P07,P08,P09,P10,P11,P13,P14,P15,P16,P17,P20,P22,P23,P24,P25"],
                               capture_output=True, text=True, timeout=300)
            self.assertEqual(g.returncode, 0, g.stdout[-3000:] + g.stderr[-2000:])


if __name__ == "__main__":
    unittest.main()
