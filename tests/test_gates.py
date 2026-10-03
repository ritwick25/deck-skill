"""K5: the static gates, the golden fixture, the negatives runner, archetype tags.

Run: cd ~/.claude/skills/deck && $PY -m unittest tests.test_gates -v

Builder refusals are proved in-process with unittest.mock.patch; no source file is ever edited:
  * text below the floor: the golden builder with its body size patched to 7.5pt raises BelowFloor;
  * a radius on a rect: the builder with prims.card routed to prims.rect(radius_in=...) raises GeometryError;
  * SLACK 0.5: the builder with deckkit.measure.SLACK patched to 0.5 builds a deck, and gate P02 (which holds
    its own SLACK 1.15) rejects it, while the unpatched golden passes P02.
The runner is shown failing on its own rows: a planter caught by the wrong rule, a planter that changes
nothing, a gate with no planter.
"""
from __future__ import annotations

import contextlib
import importlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL / "kit"))
sys.path.insert(0, str(SKILL / "gates"))
sys.path.insert(0, str(HERE))

import build_golden  # noqa: E402
from deckkit import measure, package, prims, tag  # noqa: E402
from deckkit.errors import BelowFloor, DeckkitError, GeometryError  # noqa: E402

import _shared  # noqa: E402
import negatives  # noqa: E402
import run_all  # noqa: E402

PY = sys.executable


def gate(gid):
    reg = {g["id"]: g for g in run_all.registry()}
    return importlib.import_module(reg[gid]["module"])


def run(gid, deck, **extra):
    ctx = _shared.Ctx(deck, extra=extra)
    try:
        return _shared.run_gate(gate(gid), str(deck), ctx)
    finally:
        ctx.close()


class Golden(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="k5-"))
        cls.deck = cls.tmp / "golden.pptx"
        build_golden.build(cls.deck)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def copy(self, name="copy"):
        d = Path(tempfile.mkdtemp(prefix="k5-%s-" % name, dir=self.tmp))
        dst = d / self.deck.name
        shutil.copy2(self.deck, dst)
        for sc in run_all.sidecars(self.deck):
            shutil.copy2(sc, d / sc.name)
        return dst


class TestGolden(Golden):
    def test_golden_passes_every_implemented_gate(self):
        ctx = _shared.Ctx(self.deck)
        failed = {}
        for g in run_all.implemented(run_all.registry()):
            code, found = _shared.run_gate(run_all.load(g), str(self.deck), ctx)
            if code:
                failed[g["id"]] = [f.line() for f in found if f.level == "fail"][:3]
        ctx.close()
        self.assertEqual(failed, {})

    def test_build_is_deterministic(self):
        other = self.tmp / "again" / "golden.pptx"
        self.assertEqual(build_golden.build(other), package.sha256_file(self.deck))

    def test_seven_slides_tagged_in_both_homes(self):
        rows = tag.read(str(self.deck))
        self.assertEqual([r["tag"] for r in rows],
                         ["cover:title-only", "pillars:icon-row", "process-flow:chevron-row", "stat-hero:stat-row",
                          "matrix:table", "from-to:two-panel", "closing:line"])
        self.assertTrue(all(r["source"] == "custom" for r in rows))
        # without docProps/custom.xml every tag still reads from the slide's cSld name
        bare = self.copy("bare")
        m = negatives.members(bare)
        m.pop("docProps/custom.xml")
        negatives.write(bare, m)
        rows2 = tag.read(str(bare))
        self.assertEqual([r["tag"] for r in rows2], [r["tag"] for r in rows])
        self.assertTrue(all(r["source"] == "name" for r in rows2))

    def test_tag_refuses_bad_ids(self):
        from pptx import Presentation
        prs = Presentation(str(build_golden.BASE))
        s = prs.slides.add_slide(prs.slide_layouts[0])
        for bad in (("Pillars", "row"), ("pillars", ""), ("pillars:x", "row")):
            with self.assertRaises(DeckkitError):
                tag.tag(s, *bad)


class TestBuilderRefusals(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="k5-ref-"))

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_text_at_7_5pt_refused(self):
        with mock.patch.object(build_golden, "BODY_PT", 7.5):
            with self.assertRaises(BelowFloor):
                build_golden.build(self.tmp / "x.pptx")
        self.assertFalse((self.tmp / "x.pptx").exists())

    def test_radius_on_rect_refused(self):
        real_rect = prims.rect

        def card_as_rect(slide, box, *, fill="card", line="hairline", radius_in=None, name, theme=None):
            return real_rect(slide, box, fill=fill, line=line, radius_in=0.05, name=name, theme=theme)

        with mock.patch.object(prims, "card", card_as_rect):
            with self.assertRaises(GeometryError) as cm:
                build_golden.build(self.tmp / "x.pptx")
        self.assertIn("radius", str(cm.exception))

    def test_slack_half_builds_a_deck_p02_rejects(self):
        deck = self.tmp / "slack.pptx"
        with mock.patch.object(measure, "SLACK", 0.5):
            build_golden.build(deck)
        self.assertEqual(measure.SLACK, 1.15)
        code, found = run("P02", deck)
        self.assertEqual(code, 1, [f.line() for f in found])
        self.assertTrue(any("needs" in f.msg for f in found if f.level == "fail"))
        control = self.tmp / "control" / "golden.pptx"
        build_golden.build(control)
        self.assertEqual(run("P02", control)[0], 0)


class TestInputs(Golden):
    def test_missing_validator_fails_p12(self):
        with mock.patch.object(package, "VALIDATOR_GLOB", str(self.tmp / "nowhere" / "*" / "validate.py")):
            code, found = run("P12", self.deck)
        self.assertEqual(code, 1)
        self.assertTrue(any("validator unavailable" in f.msg for f in found))

    def test_missing_refs_folder_fails_p18(self):
        code, found = run("P18", self.deck, refs_dir=str(self.tmp / "no-refs"))
        self.assertEqual(code, 2)
        self.assertIn("missing", found[0].msg)

    def test_missing_build_command_fails_p21(self):
        d = self.copy("nobuild")
        dj = json.loads((d.parent / "deck.json").read_text())
        dj.pop("build")
        (d.parent / "deck.json").write_text(json.dumps(dj))
        self.assertEqual(run("P21", d)[0], 2)

    def test_untagged_slide_fails_p15(self):
        d = self.copy("untagged")
        m = negatives.members(d)
        m.pop("docProps/custom.xml")
        negatives.write(d, m)
        negatives.edit_slide(d, 4, lambda r: r.find("{*}cSld").attrib.pop("name"))
        code, found = run("P15", d)
        self.assertEqual(code, 1)
        self.assertTrue(any("untagged" in f.msg for f in found))

    def test_gate_cli_exit_codes(self):
        ok = subprocess.run([PY, "-B", str(SKILL / "gates" / "p01_floor.py"), str(self.deck), "-q"],
                            capture_output=True, text=True)
        self.assertEqual(ok.returncode, 0, ok.stdout + ok.stderr)
        self.assertIn("P01 floor: PASS", ok.stdout)
        gone = subprocess.run([PY, "-B", str(SKILL / "gates" / "p01_floor.py"), str(self.tmp / "absent.pptx")],
                              capture_output=True, text=True)
        self.assertEqual(gone.returncode, 2)


class TestAuthority(Golden):
    def test_override_needs_a_source(self):
        d = self.copy("nosrc")
        dj = json.loads((d.parent / "deck.json").read_text())
        dj["overrides"] = {"R-NOSUB": ""}
        (d.parent / "deck.json").write_text(json.dumps(dj))
        self.assertEqual(run("P09", d)[0], 2)

    def test_subtitle_override_honoured_in_its_folder_only(self):
        d = self.copy("nosub")
        negatives.p09_subtitle(d, d.parent)
        self.assertEqual(run("P09", d)[0], 1)
        dj = json.loads((d.parent / "deck.json").read_text())
        dj["overrides"] = {"R-NOSUB": "DESIGN.md#R36"}
        (d.parent / "deck.json").write_text(json.dumps(dj))
        code, found = run("P09", d)
        self.assertEqual(code, 0, [f.line() for f in found])
        self.assertTrue(any("DESIGN.md#R36" in f.msg for f in found))
        other = d.parent.parent / "elsewhere"
        other.mkdir()
        shutil.copy2(d, other / d.name)
        for sc in run_all.sidecars(self.deck):
            shutil.copy2(sc, other / sc.name)
        self.assertEqual(run("P09", other / d.name)[0], 1)

    def test_quotation_and_thank_you_pass_p17(self):
        d = self.copy("voice")
        th = negatives.theme()
        from deckkit.canvas import Box

        def fn(prs):
            with prims.use(th):
                prims.quote(prs.slides[6], Box(0.5, 4.5, 6.0, 1.2), "We trusted the numbers once they agreed.",
                            "Regional planning lead, fictional", name="quote")
        negatives.with_prs(d, fn)
        negatives.edit_slide(d, 7, lambda r: _set_title(r, "Thank you"))
        code, found = run("P17", d)
        self.assertEqual(code, 0, [f.line() for f in found])
        negatives.edit_slide(d, 7, lambda r: _set_title(r, "Thank you, we are ready"))
        self.assertEqual(run("P17", d)[0], 1)


def _set_title(root, text):
    sh = negatives.shape(root, "title")
    runs = negatives.runs_of(sh)
    runs[0].find("{*}t").text = text
    for r in runs[1:]:
        r.getparent().remove(r)


class TestRunner(unittest.TestCase):
    """The runner's own failure rows, each shown firing once."""

    def _negatives(self, planters, only):
        args = mock.Mock(only=only, keep=False)
        buf = io.StringIO()
        with mock.patch.object(negatives, "PLANTERS", planters), contextlib.redirect_stdout(buf):
            code = run_all.run_negatives(args, run_all.registry())
        return code, buf.getvalue()

    def test_wrong_rule_is_not_a_pass(self):
        P = negatives.Planter
        code, out = self._negatives([P("x", "P01", "7.5pt", negatives.p01_75, "no such words")], "P01")
        self.assertEqual(code, 1)
        self.assertIn("FAILED ON ANOTHER RULE", out)

    def test_plant_that_changes_nothing_is_an_error(self):
        P = negatives.Planter
        code, out = self._negatives([P("noop", "P01", "nothing", lambda p, w: None, "")], "P01")
        self.assertEqual(code, 1)
        self.assertIn("PLANT ERROR", out)

    def test_gate_without_a_planter_fails_the_run(self):
        code, out = self._negatives([], "P03")
        self.assertEqual(code, 1)
        self.assertIn("NO PLANTER", out)

    def test_full_negatives_run(self):
        r = subprocess.run([PY, "-B", str(SKILL / "gates" / "run_all.py"), "--negatives"], capture_output=True,
                           text=True, timeout=900)
        self.assertEqual(r.returncode, 0, r.stdout[-3000:] + r.stderr[-2000:])
        n = len(negatives.PLANTERS)
        self.assertIn("%d negatives, %d behaved as expected, 0 did not" % (n, n), r.stdout)
        self.assertIn("25/25 implemented gates covered", r.stdout)

    def test_missing_pptx_gate_is_missing_not_pending(self):
        reg = run_all.registry() + [{"id": "P99", "module": "p99_absent", "title": "absent", "kind": "pptx",
                                     "needs": ["deck"]}]
        args = mock.Mock(deck=str(run_all.GOLDEN), only="P99,R01", static=False, refs=None, fingerprints=None,
                         twin=None, verbose=False)
        if not run_all.GOLDEN.is_file():
            build_golden.build(run_all.GOLDEN)
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            code = run_all.run_deck(args, reg)
        out = buf.getvalue()
        self.assertEqual(code, 1)
        self.assertRegex(out, r"P99 .*MISSING")
        self.assertRegex(out, r"R01 .*PENDING")


if __name__ == "__main__":
    unittest.main()
