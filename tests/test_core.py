"""K1 core: errors, canvas, ooxml, theme, package, base templates, themes.

Run: cd ~/.claude/skills/deck && $PY -m unittest tests.test_core -v
Every refusal test asserts the typed error; every negative plants one defect in a copy and expects
the check to go red, so no check here is trusted without having been seen to fail.
"""
from __future__ import annotations

import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from unittest import mock

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
KIT = os.path.join(SKILL, "kit")
sys.path.insert(0, KIT)

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

from deckkit import ooxml, package  # noqa: E402
from deckkit.canvas import STD, WIDE, Box, canvas_of  # noqa: E402
from deckkit.errors import (BelowFloor, DeckkitError, GeometryError, HandEditDrift,  # noqa: E402
                            ThemeInvariant, ValidatorMissing)
from deckkit.read.colour import contrast  # noqa: E402
from deckkit.read.ns import NS  # noqa: E402
from deckkit.theme import Theme, apply_theme  # noqa: E402

PY = sys.executable
THEMES = os.path.join(SKILL, "themes")
SIGMOID = os.path.join(THEMES, "sigmoid.json")
NEUTRAL = os.path.join(THEMES, "neutral.json")
BASE_WIDE = os.path.join(KIT, "assets", "base-wide.pptx")
BASE_STD = os.path.join(KIT, "assets", "base-std.pptx")
A = "{%s}" % NS["a"]
P = "{%s}" % NS["p"]


def rb(path) -> bytes:
    with open(path, "rb") as f:
        return f.read()


def jload(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def theme_dict(**slots):
    d = jload(SIGMOID)
    d["slots"].update(slots)
    return d


def cli(*args):
    env = dict(os.environ, PYTHONPATH=KIT)
    return subprocess.run([PY, "-m", "deckkit"] + list(args), capture_output=True, text=True, env=env, cwd=SKILL)


def small_deck(base=BASE_WIDE):
    prs = Presentation(base)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(2))
    ooxml.finish(sh)
    tb = s.shapes.add_textbox(Inches(1), Inches(3.5), Inches(4), Inches(1))
    ooxml.finish(tb)
    ooxml.no_autofit(tb.text_frame)
    tb.text_frame.text = "Fictional sample text"
    return prs


def members(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def write(m, path):
    with open(path, "wb") as f:
        f.write(package.zip_bytes(m))


class TmpCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="deckkit-k1-")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def p(self, *parts):
        return os.path.join(self.tmp, *parts)


# ------------------------------------------------------------------ errors and canvas

class TestErrors(unittest.TestCase):
    def test_fields_carried(self):
        e = ThemeInvariant("slot dk1", need="000000", have="111111", hint="fixed")
        self.assertEqual((e.what, e.need, e.have, e.hint), ("slot dk1", "000000", "111111", "fixed"))
        self.assertIsInstance(e, DeckkitError)
        self.assertIn("slot dk1", str(e))


class TestCanvas(unittest.TestCase):
    def test_profiles(self):
        self.assertEqual(WIDE.emu, (12192000, 6858000))
        self.assertEqual(STD.emu, (9144000, 5143500))
        self.assertAlmostEqual(WIDE.title_box.r, WIDE.w - WIDE.margin_r, places=2)
        self.assertAlmostEqual(STD.title_box.r, STD.w - STD.margin_r, places=2)
        self.assertEqual((WIDE.content.y, round(WIDE.content.b, 2)), (0.85, 6.84))
        self.assertEqual((STD.content.y, round(STD.content.b, 2)), (0.99, 5.25))
        self.assertEqual((WIDE.k, STD.k), (1.0, 0.75))
        for c in (WIDE, STD):
            self.assertGreaterEqual(min(c.ramp.values()), 8.0)

    def test_split(self):
        cols = WIDE.content.split_cols(4, WIDE.gutter)
        self.assertEqual(len(cols), 4)
        self.assertAlmostEqual(cols[-1].r, WIDE.content.r, places=9)
        self.assertAlmostEqual(sum(c.w for c in cols) + 3 * WIDE.gutter, WIDE.content.w, places=9)
        rows = Box(0, 0, 2, 3).split_rows(3, 0.15, weights=[1, 2, 1])
        self.assertAlmostEqual(rows[1].h, 2 * rows[0].h, places=9)
        with self.assertRaises(GeometryError):
            Box(0, 0, 1, 1).split_cols(12, 0.2)
        with self.assertRaises(GeometryError):
            Box(0, 0, 0, 1)

    def test_ramp_floor_refused(self):
        import dataclasses
        with self.assertRaises(BelowFloor):
            dataclasses.replace(WIDE, ramp=dict(WIDE.ramp, small=7.5))

    def test_canvas_of(self):
        self.assertIs(canvas_of(Presentation(BASE_WIDE)), WIDE)
        self.assertIs(canvas_of(Presentation(BASE_STD)), STD)
        with self.assertRaises(GeometryError):
            canvas_of((9144000, 6858000))           # 4:3


# ------------------------------------------------------------------ ooxml

class TestOoxml(unittest.TestCase):
    def test_finish_strips_style_and_writes_one_effectlst(self):
        prs = Presentation(BASE_WIDE)
        s = prs.slides.add_slide(prs.slide_layouts[0])
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
        self.assertIsNotNone(sh._element.find(P + "style"))            # python-pptx injects it
        sh._element.spPr.append(etree.Element(A + "effectLst"))
        ooxml.finish(sh)
        ooxml.finish(sh)                                                 # idempotent
        self.assertIsNone(sh._element.find(P + "style"))
        self.assertEqual(len(sh._element.spPr.findall(A + "effectLst")), 1)
        self.assertEqual(len(sh._element.spPr.find(A + "effectLst")), 0)
        sh.fill.solid()
        names = [etree.QName(c).localname for c in sh._element.spPr]
        self.assertLess(names.index("solidFill"), names.index("effectLst"))
        self.assertEqual(ooxml.once_only_violations(s), [])
        with self.assertRaises(GeometryError):
            ooxml.finish(sh, shadow="soft")
        c = s.shapes.add_connector(1, Inches(1), Inches(3), Inches(3), Inches(3))
        ooxml.finish(c)
        self.assertIsNone(c._element.find(P + "style"))

    def test_no_autofit(self):
        prs = Presentation(BASE_WIDE)
        tb = prs.slides.add_slide(prs.slide_layouts[0]).shapes.add_textbox(0, 0, Inches(2), Inches(1))
        body = ooxml.no_autofit(tb.text_frame)
        kids = [etree.QName(c).localname for c in body]
        self.assertEqual(kids.count("noAutofit"), 1)
        self.assertNotIn("spAutoFit", kids)
        self.assertEqual(body.get("wrap"), "square")

    def test_once_only_and_next_id(self):
        prs = Presentation(BASE_WIDE)
        s = prs.slides.add_slide(prs.slide_layouts[0])
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(1), Inches(1))
        ooxml.finish(sh)
        sh._element.spPr.append(etree.Element(A + "effectLst"))
        v = ooxml.once_only_violations(s)
        self.assertEqual(len(v), 1)
        self.assertEqual(v[0][1:3], ("effectLst", 2))
        self.assertEqual(ooxml.next_id(s), sh.shape_id + 1)


# ------------------------------------------------------------------ theme

class TestThemeInvariants(unittest.TestCase):
    def test_client_theme_touching_dk1_refused(self):
        with self.assertRaises(ThemeInvariant):
            Theme(theme_dict(dk1="111111"))

    def test_lt1_and_unknown_slots_refused(self):
        with self.assertRaises(ThemeInvariant):
            Theme(theme_dict(lt1="FAFAFA"))
        with self.assertRaises(ThemeInvariant):
            Theme(theme_dict(accent7="123456"))
        d = theme_dict()
        del d["slots"]["accent3"]
        with self.assertRaises(ThemeInvariant):
            Theme(d)
        Theme(theme_dict(dk1="000000", lt1="ffffff"))                 # restating the fixed value is allowed

    def test_unknown_role_and_bad_expression_refused(self):
        d = theme_dict()
        d["roles"]["glow"] = "accent1"
        with self.assertRaises(ThemeInvariant):
            Theme(d)
        d = theme_dict()
        d["roles"]["ink"] = "dk1 lumMod 80000"                          # dk1 is not a mapped slot name
        with self.assertRaises(ThemeInvariant):
            Theme(d)

    def test_record_drift_refused(self):
        d = theme_dict(accent1="F2A900")                               # amber brand: accent-ink must be lowered
        with self.assertRaises(ThemeInvariant):
            Theme(d)                                                    # record still says accent1
        d["roles"]["accent-ink"].pop("record")
        d["roles"]["on-accent"].pop("record")
        t = Theme(d)
        self.assertIn("lumMod", t.expr("accent-ink"))
        self.assertGreaterEqual(contrast(t.hex("accent-ink"), t.hex("card")), 4.5)
        self.assertEqual(t.expr("on-accent"), "tx1")                    # black on amber


class TestThemeRoles(unittest.TestCase):
    def test_role_values_match_engineering_table(self):
        t = Theme.load(SIGMOID)
        want = {"paper": "FFFFFF", "card": "F2F2F2", "hairline-soft": "D9D9D9", "hairline": "BFBFBF",
                "rule-strong": "8C8C8C", "ink-deep": "1A1A1A", "ink": "333333", "ink-strong": "404040",
                "ink-muted": "595959", "accent": "D91E18", "accent-tint": "FAD0CF", "accent-ink": "D91E18",
                "on-accent": "FFFFFF", "navy": "4B5960", "series-1": "1E3C7D", "series-5": "6B7983"}
        self.assertEqual({k: t.hex(k) for k in want}, want)

    def test_spec_is_scheme_colour(self):
        t = Theme.load(SIGMOID)
        el = t.spec("ink-muted")
        self.assertEqual(el.tag, A + "schemeClr")
        self.assertEqual(el.get("val"), "tx1")
        self.assertEqual([(etree.QName(c).localname, c.get("val")) for c in el], [("lumMod", "65000"), ("lumOff", "35000")])
        self.assertEqual(t.spec_xml("accent"), '<a:schemeClr val="accent1"/>')

    def test_fmt_scheme_flat(self):
        x = etree.fromstring(Theme.load(SIGMOID).fmt_scheme_xml())
        fills = x.find(A + "fillStyleLst")
        self.assertEqual([etree.QName(c).localname for c in fills], ["solidFill"] * 3)
        self.assertEqual([etree.QName(c).localname for c in x.find(A + "bgFillStyleLst")], ["solidFill"] * 3)
        effects = x.find(A + "effectStyleLst")
        self.assertEqual(len(effects), 3)
        for es in effects:
            self.assertEqual(len(es), 1)
            self.assertEqual(len(es.find(A + "effectLst")), 0)
        self.assertEqual([ln.get("w") for ln in x.find(A + "lnStyleLst")], ["9525", "12700", "19050"])
        self.assertEqual(x.findall(".//" + A + "gradFill"), [])

    def test_contrast_tables_pass(self):
        for path in (SIGMOID, NEUTRAL):
            t = Theme.load(path)
            rows = t.contrast_table()
            self.assertEqual([r for r in rows if r["status"] == "FAIL"], [], path)
            exempt = {r["fg"] for r in rows if r["status"] == "EXEMPT"}
            self.assertEqual(exempt, {"hairline", "hairline-soft"})
            self.assertTrue(all(r["reason"] for r in rows if r["status"] == "EXEMPT"))
            gated = {(r["fg"], r["bg"]): r["min"] for r in rows if r["min"]}
            self.assertEqual(gated[("ink", "paper")], 7.0)
            self.assertEqual(gated[("on-accent", "accent")], 4.5)
            self.assertEqual(gated[("rule-strong", "paper")], 3.0)

    def test_cli_check_exit_zero_on_both_themes(self):
        for path in (SIGMOID, NEUTRAL):
            r = cli("theme", "check", path)
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            self.assertIn("PASS: every gated pair", r.stdout)

    def test_planted_failing_pair_exits_one(self):
        with tempfile.TemporaryDirectory() as d:
            bad = theme_dict(accent1="FFD700")                         # gold: white text on it fails 4.5
            bad["roles"]["accent-ink"] = "accent1"
            bad["roles"]["on-accent"] = "bg1"
            path = os.path.join(d, "planted.json")
            with open(path, "w", encoding="utf-8") as f:
                json.dump(bad, f)
            r = cli("theme", "check", path)
            self.assertEqual(r.returncode, 1, r.stdout)
            self.assertIn("FAIL", r.stdout)
            self.assertRegex(r.stdout, r"on-accent\s+accent\s+FFFFFF\s+FFD700\s+[\d.]+\s+4\.5\s+FAIL")

    def test_fill_only_slot_as_text_fails(self):
        d = jload(NEUTRAL)
        d["roles"]["navy"] = "accent2"                                  # the sky slot, fill only
        t = Theme(d)
        self.assertTrue(any("fill-only" in r["reason"] for r in t.failures()))

    def test_solve_lowers_lum_mod(self):
        d = theme_dict(accent1="00A0E9")
        d["roles"]["accent-ink"] = "accent1"
        d["roles"]["on-accent"] = "tx1"
        t = Theme(d)
        expr, ratio = t.solve("accent-ink", ["paper"], 4.5)
        self.assertRegex(expr, r"^accent1 lumMod \d+$")
        self.assertGreaterEqual(ratio, 4.5)
        lum = int(expr.split()[-1])
        self.assertLess(lum, 100000)
        d["roles"]["accent-ink"] = "accent1 lumMod %d" % (lum + 500)   # one step lighter must fail
        self.assertLess(contrast(Theme(d).hex("accent-ink"), "FFFFFF"), 4.5)
        r = cli("theme", "solve", SIGMOID, "accent-ink", "paper", "card")
        self.assertEqual(r.returncode, 0, r.stderr)

    def test_cvd_report(self):
        self.assertEqual(Theme.load(NEUTRAL).cvd(), [])
        col = Theme.load(SIGMOID).cvd()
        self.assertTrue(any(c["kind"] == "protan" and {c["a_hex"], c["b_hex"]} == {"F38B1F", "4EA347"} for c in col))
        grey = Theme(theme_dict(accent2="595959", accent3="5A5A5A"))
        self.assertTrue(any(c["kind"] == "normal" for c in grey.cvd(["series-1", "series-2"])))


# ------------------------------------------------------------------ package

class TestPackage(TmpCase):
    def test_double_build_byte_identical(self):
        t = Theme.load(SIGMOID)
        a, b = self.p("a", "deck.pptx"), self.p("b", "deck.pptx")
        sha1 = package.save(small_deck(), a, title="Sample", theme=t, manifest={"slides_meta": [1]})
        sha2 = package.save(small_deck(), a, title="Sample", theme=t, manifest={"slides_meta": [1]})
        sha3 = package.save(small_deck(), b, title="Sample", theme=t, manifest={"slides_meta": [1]})
        self.assertEqual(sha1, sha2)
        self.assertEqual(sha1, sha3)
        self.assertEqual(rb(a), rb(b))
        self.assertEqual(rb(package.manifest_path(a)), rb(package.manifest_path(b)))
        self.assertEqual(jload(package.manifest_path(a))["sha256"], sha1)

    def test_saved_package_normalised(self):
        out = self.p("deck.pptx")
        package.save(small_deck(), out, title="Sample", theme=Theme.load(NEUTRAL), manifest=None)
        m = members(out)
        core = m["docProps/core.xml"].decode()
        self.assertIn("<dc:creator>Sigmoid</dc:creator>", core)
        self.assertIn("<cp:lastModifiedBy>Sigmoid</cp:lastModifiedBy>", core)
        self.assertEqual(set(re.findall(r"W3CDTF\">([^<]+)<", core)), {package.FIXED_DATE})
        self.assertNotIn("TitlesOfParts", m["docProps/app.xml"].decode())
        self.assertFalse(any("thumbnail" in n for n in m))
        self.assertIn('Extension="svg" ContentType="image/svg+xml"', m["[Content_Types].xml"].decode())
        self.assertFalse(any(n.startswith("ppt/fonts/") for n in m))
        theme = m["ppt/theme/theme1.xml"].decode()
        self.assertIn('<a:accent1><a:srgbClr val="1E3C7D"/>', theme)
        self.assertNotIn("gradFill", theme)
        with zipfile.ZipFile(out) as z:
            self.assertEqual(z.infolist()[0].filename, "[Content_Types].xml")
            self.assertEqual({i.date_time for i in z.infolist()}, {package.FIXED_TIME})

    def test_validate_and_structure_pass(self):
        out = self.p("deck.pptx")
        package.save(small_deck(), out, title="Sample", theme=Theme.load(SIGMOID), manifest=None)
        for path in (out, BASE_WIDE, BASE_STD):
            self.assertEqual(package.structural_check(path), [], path)
            ok, detail = package.validate_xsd(path)
            self.assertTrue(ok, detail)

    def test_xsd_catches_planted_element_in_master(self):
        m = members(BASE_WIDE)
        k = "ppt/slideMasters/slideMaster1.xml"
        m[k] = m[k].replace(b"<p:txStyles>", b"<p:bogus/><p:txStyles>")
        write(m, self.p("bad.pptx"))
        ok, detail = package.validate_xsd(self.p("bad.pptx"))
        self.assertFalse(ok)
        self.assertIn("bogus", detail)

    def test_missing_validator_raises(self):
        with mock.patch.object(package, "VALIDATOR_GLOB", self.p("nowhere", "*", "validate.py")):
            with self.assertRaises(ValidatorMissing):
                package.validate_xsd(BASE_WIDE)

    def test_hand_edit_drift_refused(self):
        t = Theme.load(SIGMOID)
        out = self.p("deck.pptx")
        package.save(small_deck(), out, title="Sample", theme=t, manifest=None)
        with open(out, "ab") as f:
            f.write(b"\0")                                              # a hand edit, as far as sha can tell
        with self.assertRaises(HandEditDrift):
            package.save(small_deck(), out, title="Sample", theme=t, manifest=None)
        package.save(small_deck(), out, title="Sample", theme=t, manifest=None, overwrite=True)
        open(package.lock_path(out), "w").close()
        with self.assertRaises(HandEditDrift):
            package.save(small_deck(), out, title="Sample", theme=t, manifest=None)
        os.remove(package.lock_path(out))
        stranger = self.p("stranger.pptx")
        shutil.copy(BASE_WIDE, stranger)                                # exists, never built by the kit
        with self.assertRaises(HandEditDrift):
            package.save(small_deck(), stranger, title="Sample", theme=t, manifest=None)

    def test_apply_theme_swaps_every_theme_part(self):
        out = self.p("deck.pptx")
        package.save(small_deck(), out, title="Sample", theme=Theme.load(SIGMOID), manifest=None)
        parts = apply_theme(out, Theme.load(NEUTRAL), self.p("neutral.pptx"))
        self.assertEqual(parts, ["ppt/theme/theme1.xml"])
        th = members(self.p("neutral.pptx"))["ppt/theme/theme1.xml"].decode()
        self.assertIn('<a:accent1><a:srgbClr val="1E3C7D"/>', th)
        self.assertIn('name="neutral"', th)
        self.assertEqual(package.structural_check(self.p("neutral.pptx")), [])
        r = cli("theme", "apply", SIGMOID, self.p("neutral.pptx"))
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn('<a:accent1><a:srgbClr val="D91E18"/>', members(self.p("neutral.pptx"))["ppt/theme/theme1.xml"].decode())


class TestStructuralNegatives(TmpCase):
    """Each plant must turn structural_check red; the clean source must be green."""

    def built(self):
        out = self.p("src.pptx")
        prs = small_deck()
        s = prs.slides[0]
        c = s.shapes.add_connector(1, Inches(1), Inches(5), Inches(4), Inches(5))
        ooxml.finish(c)
        package.save(prs, out, title="Sample", theme=Theme.load(SIGMOID), manifest=None)
        self.assertEqual(package.structural_check(out), [])
        return members(out)

    def plant(self, fn, expect):
        m = self.built()
        fn(m)
        write(m, self.p("planted.pptx"))
        defects = package.structural_check(self.p("planted.pptx"))
        self.assertTrue(any(re.search(expect, d) for d in defects), "%s not in %s" % (expect, defects))

    S = "ppt/slides/slide1.xml"

    def sub(self, m, old, new, part=None):
        part = part or self.S
        self.assertIn(old, m[part])
        m[part] = m[part].replace(old, new, 1)

    def test_double_effectlst(self):
        self.plant(lambda m: self.sub(m, b"<a:effectLst/>", b"<a:effectLst/><a:effectLst/>"), r"<effectLst> x2")

    def test_dangling_rid(self):
        def f(m):
            root = etree.fromstring(m[self.S])
            sp = root.find(".//" + P + "sp")
            sp.find(P + "spPr").insert(2, etree.fromstring(
                '<a:blipFill xmlns:a="%s" xmlns:r="%s"><a:blip r:embed="rId999"/></a:blipFill>' % (NS["a"], NS["r"])))
            m[self.S] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        self.plant(f, r"r:embed=rId999 has no relationship")

    def test_svg_content_type_dropped(self):
        def f(m):
            m["ppt/media/image1.svg"] = b'<svg xmlns="http://www.w3.org/2000/svg"/>'
            m["[Content_Types].xml"] = m["[Content_Types].xml"].replace(
                b'<Default Extension="svg" ContentType="image/svg+xml"/>', b"")
        self.plant(f, r"image1\.svg: no content type")

    def test_duplicate_shape_id(self):
        def f(m):
            ids = re.findall(rb'<p:cNvPr id="(\d+)"', m[self.S])
            m[self.S] = m[self.S].replace(b'<p:cNvPr id="%s"' % ids[2], b'<p:cNvPr id="%s"' % ids[1], 1)
        self.plant(f, r"cNvPr id \d+ used 2 times")

    def test_sld_id_below_256(self):
        self.plant(lambda m: self.sub(m, b'<p:sldId id="256"', b'<p:sldId id="255"', "ppt/presentation.xml"),
                   r"sldId 255")

    def test_negative_extent(self):
        def f(m):
            m[self.S] = re.sub(rb'<a:ext cx="\d+"', b'<a:ext cx="-5"', m[self.S], count=1)
        self.plant(f, r"negative extent cx=-5")

    def test_glue_to_missing_shape(self):
        def f(m):
            self.sub(m, b"<p:cNvCxnSpPr/>", b'<p:cNvCxnSpPr><a:stCxn id="999" idx="0"/></p:cNvCxnSpPr>')
        self.plant(f, r"stCxn glued to missing shape id 999")

    def test_missing_rel_target(self):
        def f(m):
            del m["ppt/slideLayouts/slideLayout1.xml"]
        self.plant(f, r"targets missing part ppt/slideLayouts/slideLayout1\.xml")

    def test_control_character(self):
        self.plant(lambda m: self.sub(m, b"Fictional", b"Fic\x0btional"), r"control character|not well-formed")


# ------------------------------------------------------------------ base templates

class TestBase(TmpCase):
    def test_make_base_deterministic_and_fresh(self):
        r = subprocess.run([PY, os.path.join(SKILL, "scripts", "make_base.py"), "--out", self.tmp],
                           capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        for name, ref in (("base-wide.pptx", BASE_WIDE), ("base-std.pptx", BASE_STD)):
            self.assertEqual(rb(self.p(name)), rb(ref),
                             "%s in kit/assets is stale against make_base.py" % name)

    def test_base_shape(self):
        for path, canvas in ((BASE_WIDE, WIDE), (BASE_STD, STD)):
            prs = Presentation(path)
            self.assertEqual((prs.slide_width, prs.slide_height), canvas.emu)
            self.assertEqual(len(prs.slide_masters), 1)
            self.assertEqual(len(prs.slide_layouts), 1)
            self.assertEqual(prs.slide_layouts[0].name, "Blank")
            self.assertEqual(len(prs.slide_layouts[0].placeholders), 0)
            self.assertEqual(len(prs.slide_masters[0].placeholders), 0)
            self.assertEqual(len(prs.slides), 0)
            m = members(path)
            self.assertFalse(any("thumbnail" in n or n.startswith("ppt/fonts/") for n in m))
            self.assertNotIn(b"embedTrueTypeFonts", m["ppt/presentation.xml"])
            cm = etree.fromstring(m["ppt/slideMasters/slideMaster1.xml"]).find(P + "clrMap")
            self.assertEqual({k: cm.get(k) for k in ("bg1", "tx1", "bg2", "tx2")},
                             {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"})
            th = m["ppt/theme/theme1.xml"].decode()
            self.assertIn('<a:accent1><a:srgbClr val="D91E18"/>', th)
            self.assertIn('<a:latin typeface="Arial"/>', th)
            self.assertNotIn("Calibri", th)
            self.assertNotIn("gradFill", th)


if __name__ == "__main__":
    unittest.main()
