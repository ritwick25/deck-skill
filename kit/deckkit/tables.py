"""Tables: a native PowerPoint table under the house rules (promoted from tests/build_golden.py table()).

    table(slide, box, header, rows, marked_col=None, cell_glyphs=None, ...)

No table style (PowerPoint's "No Style, No Grid"), no banding flags, every cell's margins, anchor, fill and borders
written explicitly, one hairline under each row and none between columns. The header row is the table's one solid
fill (ink-strong, paper bold labels); the first column is set apart by weight, so each row is found by its name.
Every cell is written by deckkit.text.write and measured before the table exists: rows are as tall as their
fullest cell, then share whatever height the box has left so the table fills it; Overflows when the rows need
more than the box.

marked_col marks the recommended column of an options table: its header in the accent fill, its body cells bold
on the card fill. It is a state, not decoration, so it is allowed once per table.

cell_glyphs turns a word into glyph plus word: {"Yes": "circle-check", "No": ("circle-x", "ink-muted")} draws the
glyph at the left of every body cell whose whole text is that word (case-insensitive), in the brand role unless a
role is given, and keeps the word beside it. The glyph is a kit shape laid over its cell (nested in the table's
frame, so P05 sees it inside its table).
"""
from __future__ import annotations

from lxml import etree
from pptx.util import Emu

from . import measure, ooxml, prims
from .canvas import EMU_PER_IN, Box
from .errors import GeometryError, Overflows, ThemeInvariant
from .glyphs import glyph
from .prims import canvas_for, theme_of
from .read.ns import NS
from .text import Para, Run, need_height, write

_A = "{%s}" % NS["a"]
NO_TABLE_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"     # PowerPoint's built-in "No Style, No Grid"
MX, MY = 0.10, 0.09        # cell margins (in, x k)
GLYPH_IN = 0.20            # a cell glyph's side (x k)
HEADER_FILLS = ("ink-strong", "ink-deep", "navy", "accent")


def _glyph_spec(spec):
    if isinstance(spec, str):
        return spec, "accent"
    g, role = spec
    return g, role


def _cell_paras(text, role, bold, pt, align):
    return [Para((Run(str(text), role, bold=bold, size=pt),), align=align)]


@ooxml.transactional
def table(slide, box: Box, header, rows, marked_col=None, cell_glyphs=None, *, widths=None, pt=11.0,
          header_fill="ink-strong", first_col_bold=True, fill_height=True, align=None, name: str, theme=None):
    """A measured native table in box. header: the column labels; rows: lists of cell strings, one per column.

    widths: relative column weights (default: each column as wide as its longest word and its average text ask).
    pt: body and header size (11pt; never under the floor). fill_height: rows share the box's spare height (the
    table fills its zone); False keeps them at their measured height. align: per-column "left" / "center" /
    "right" (default left). Returns (graphicFrame, height_in, [row tops in inches], [column lefts in inches])."""
    th = theme_of(theme)
    k = canvas_for(slide).k
    header = [str(h) for h in header]
    rows = [[str(c) for c in r] for r in rows]
    nc = len(header)
    if nc < 1 or not rows:
        raise GeometryError("table %s" % name, need="a header and at least one row", have=(nc, len(rows)))
    if any(len(r) != nc for r in rows):
        raise GeometryError("table %s rows" % name, need="%d cells per row" % nc, have=[len(r) for r in rows])
    if header_fill not in HEADER_FILLS:
        raise ThemeInvariant("table %s header_fill" % name, need=" / ".join(HEADER_FILLS), have=header_fill)
    if marked_col is not None and not 0 <= marked_col < nc:
        raise GeometryError("table %s marked_col" % name, need="0..%d" % (nc - 1), have=marked_col)
    prims.check_pt(pt, "table %s" % name)
    align = list(align) if align else ["left"] * nc
    glyphs_by_word = {str(w).strip().lower(): _glyph_spec(v) for w, v in (cell_glyphs or {}).items()}
    mx, my, gs = MX * k, MY * k, GLYPH_IN * k
    face = measure.face_for(prims._face(th))
    bold_face = measure.face_for(prims._face(th), True)

    # column widths: weights from content, then every column at least its longest word
    if widths is None:
        weights = []
        for ci in range(nc):
            texts = [header[ci]] + [r[ci] for r in rows]
            longest = max(measure.width_in(w, pt, bold_face) for t in texts for w in t.split() or [""])
            avg = sum(measure.width_in(t, pt, face) for t in texts) / len(texts)
            weights.append(max(longest, min(avg, 3.6 * k)) + 0.3 * k)
    else:
        if len(widths) != nc or any(w <= 0 for w in widths):
            raise GeometryError("table %s widths" % name, need="%d positive weights" % nc, have=widths)
        weights = list(widths)
    tot = float(sum(weights))
    col_w = [box.w * w / tot for w in weights]

    def fill_of(ri, ci):
        if ri == 0:
            return "accent" if ci == marked_col else header_fill
        return "card" if ci == marked_col else "paper"

    def role_of(ri, ci):
        if ri == 0:
            return prims.ON_FILL[fill_of(0, ci)]
        return "ink-deep" if (ci == 0 and first_col_bold) or ci == marked_col else "ink"

    def bold_of(ri, ci):
        return ri == 0 or (ci == 0 and first_col_bold) or ci == marked_col

    grid = [header] + rows
    paras, heights = {}, []
    for ri, r in enumerate(grid):
        hh = 0.0
        for ci, t in enumerate(r):
            glyphed = ri > 0 and t.strip().lower() in glyphs_by_word
            p = _cell_paras(t, role_of(ri, ci), bold_of(ri, ci), pt, align[ci])
            lm = mx + (gs + 0.06 * k if glyphed else 0.0)
            avail = col_w[ci] - lm - mx
            longest = max(measure.width_in(w, pt, bold_face if bold_of(ri, ci) else face) for w in t.split() or [""])
            if longest * measure.SLACK > avail + 1e-6:
                raise Overflows("table %s r%d c%d" % (name, ri, ci + 1), need=round(longest * measure.SLACK + lm + mx, 3),
                                have=round(col_w[ci], 3), hint="widen the column (widths=) or shorten the word")
            h = need_height(p, avail, th) + 2 * my
            if glyphed:
                h = max(h, gs + 2 * my)
            paras[ri, ci] = (p, lm, glyphed)
            hh = max(hh, h)
        heights.append(hh)
    need = sum(heights)
    if need > box.h + 1e-6:
        raise Overflows("table %s" % name, need=round(need, 3), have=round(box.h, 3),
                        hint="fewer rows, fewer words per cell, or split the list into two slides with the same columns")
    if fill_height and len(heights) > 1:
        extra = (box.h - need) / (len(heights) - 1)
        heights = heights[:1] + [h + extra for h in heights[1:]]
    total_h = sum(heights)
    gf = slide.shapes.add_table(len(grid), nc, *Box(box.x, box.y, box.w, total_h).emu)
    gf.name = name
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for kk in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        tblPr.set(kk, "0")
    sid = tblPr.find(_A + "tableStyleId")
    if sid is None:
        sid = etree.SubElement(tblPr, _A + "tableStyleId")
    sid.text = NO_TABLE_STYLE
    for ci, w in enumerate(col_w):
        tbl.columns[ci].width = Emu(int(round(w * EMU_PER_IN)))
    for ri, h in enumerate(heights):
        tbl.rows[ri].height = Emu(int(round(h * EMU_PER_IN)))
    for ri, r in enumerate(grid):
        for ci in range(nc):
            cell = tbl.cell(ri, ci)
            p, lm, _g = paras[ri, ci]
            f = fill_of(ri, ci)
            write(cell.text_frame, p, th, anchor="middle", background=f)
            tcPr = cell._tc.get_or_add_tcPr()
            for c in list(tcPr):
                tcPr.remove(c)
            for kk, v in (("marL", lm), ("marR", mx), ("marT", my), ("marB", my)):
                tcPr.set(kk, str(int(round(v * EMU_PER_IN))))
            tcPr.set("anchor", "ctr")
            for side in ("lnL", "lnR", "lnT", "lnB"):
                ln = etree.SubElement(tcPr, _A + side)
                if side == "lnB" and ri > 0:
                    ln.set("w", "9525")
                    etree.SubElement(ln, _A + "solidFill").append(th.spec("hairline"))
                else:
                    ln.set("w", "0")
                    etree.SubElement(ln, _A + "noFill")
            etree.SubElement(tcPr, _A + "solidFill").append(th.spec(f))
    ooxml.finish(gf)
    tops, y = [], box.y
    for h in heights:
        tops.append(y)
        y += h
    lefts, x = [], box.x
    for w in col_w:
        lefts.append(x)
        x += w
    for (ri, ci), (_p, _lm, glyphed) in sorted(paras.items()):
        if not glyphed:
            continue
        g, role = glyphs_by_word[grid[ri][ci].strip().lower()]
        cy = tops[ri] + heights[ri] / 2.0
        glyph(slide, g, Box(lefts[ci] + mx, cy - gs / 2, gs, gs), role=role,
              shape_name="%s.r%d.c%d.icon" % (name, ri, ci + 1), theme=th)
    return gf, total_h, tops, lefts


__all__ = ["table", "NO_TABLE_STYLE"]
