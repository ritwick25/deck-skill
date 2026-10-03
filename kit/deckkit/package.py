"""Package output: a deterministic, normalised save, the XSD validator by reference, and the structural
checks that catch what makes PowerPoint offer a repair while every other reader passes the file.

save() lifts the PMR `_normalise` ideas and adds what the kit needs:
  * fixed zip timestamps and member order ([Content_Types].xml first), so a rebuild is byte-identical
  * docProps rewritten: creator and lastModifiedBy "Sigmoid", created and modified set to the deck's
    date (deck.json "date", passed by the caller as date=) or 1980-01-01 when none is given, so the
    save is deterministic either way; no TitlesOfParts, no HeadingPairs (app.xml would otherwise list
    every slide title), no thumbnail part
  * every theme part carries the Theme's clrScheme, fontScheme and flat fmtScheme
  * the svg Default content type is present (python-pptx drops it, orphaning every icon)
  * embedded workbooks flattened (each chart's xlsx carries its own timestamps)
  * a reopen round-trip before anything touches disk, then an atomic os.replace
  * refusal (HandEditDrift) when a ~$ lock exists beside the target, or when the target on disk no
    longer matches the sha its last manifest recorded, unless overwrite=True

The XSD validator is referenced in place, never copied (its skill is all rights reserved). If it cannot
be found or cannot run, validate_xsd raises ValidatorMissing: a missing validator is a failure.
"""
from __future__ import annotations

import collections
import datetime as _dt
import glob
import hashlib
import io
import json
import os
import posixpath
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from lxml import etree

from .errors import DeckkitError, HandEditDrift, UnsupportedPart, ValidatorMissing
from .ooxml import once_only_violations
from .read.ns import NS
from .theme import Theme, theme_parts

FIXED_TIME = (1980, 1, 1, 0, 0, 0)       # zip member timestamps: always fixed
FIXED_DATE = "1980-01-01T00:00:00Z"      # docProps created/modified when the caller passes no date
AUTHOR = "Sigmoid"
VALIDATOR_GLOB = os.path.expanduser("~/.claude/skills/synced/*/pptx/scripts/office/validate.py")

CT_NS = NS["ct"]
PR_NS = NS["pr"]
R_NS = NS["r"]
_P = "{%s}" % NS["p"]
_A = "{%s}" % NS["a"]
RT_THUMB = "http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail"
SVG_CT = "image/svg+xml"

CORE_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
    'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
    '<dc:title>{title}</dc:title><dc:creator>{author}</dc:creator>'
    '<cp:lastModifiedBy>{author}</cp:lastModifiedBy><cp:revision>1</cp:revision>'
    '<dcterms:created xsi:type="dcterms:W3CDTF">{date}</dcterms:created>'
    '<dcterms:modified xsi:type="dcterms:W3CDTF">{date}</dcterms:modified>'
    '</cp:coreProperties>')

APP_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    '<TotalTime>0</TotalTime><Application>Microsoft Office PowerPoint</Application>'
    '<PresentationFormat>Custom</PresentationFormat><Slides>{slides}</Slides>'
    '<Company>{author}</Company><AppVersion>16.0000</AppVersion></Properties>')


def _esc(s: str) -> str:
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


_W3CDTF = re.compile(r"^(\d{4})-(\d{2})-(\d{2})(?:T(\d{2}):(\d{2}):(\d{2})Z)?$")


def w3cdtf(date=None) -> str:
    """The docProps timestamp for a deck date: None gives FIXED_DATE (1980-01-01T00:00:00Z); a
    datetime.date, a datetime (naive, or aware and converted to UTC) or a 'YYYY-MM-DD' /
    'YYYY-MM-DDTHH:MM:SSZ' string gives that moment in UTC. Anything else raises DeckkitError."""
    if date is None:
        return FIXED_DATE
    if isinstance(date, _dt.datetime):
        d = date.astimezone(_dt.timezone.utc) if date.tzinfo else date
        return d.strftime("%Y-%m-%dT%H:%M:%SZ")
    if isinstance(date, _dt.date):
        return date.strftime("%Y-%m-%dT00:00:00Z")
    m = _W3CDTF.match(date) if isinstance(date, str) else None
    if m:
        y, mo, d, hh, mm, ss = (int(g) if g else 0 for g in m.groups())
        try:
            return _dt.datetime(y, mo, d, hh, mm, ss).strftime("%Y-%m-%dT%H:%M:%SZ")
        except ValueError:
            pass
    raise DeckkitError("docProps date", need="YYYY-MM-DD, YYYY-MM-DDTHH:MM:SSZ or a datetime.date",
                       have=repr(date), hint="pass deck.json's date unchanged, or omit it for the 1980 fallback")


def core_xml(title: str, date=None) -> bytes:
    return CORE_XML.format(title=_esc(title), author=AUTHOR, date=w3cdtf(date)).encode("utf-8")


def app_xml(n_slides: int) -> bytes:
    return APP_XML.format(slides=int(n_slides), author=AUTHOR).encode("utf-8")


def sha256_bytes(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def sha256_file(path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ------------------------------------------------------------------ deterministic zip

def zip_bytes(members: dict) -> bytes:
    """Members as a zip with fixed metadata: [Content_Types].xml first, then sorted by name."""
    order = sorted(members, key=lambda n: (n != "[Content_Types].xml", n))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name in order:
            zi = zipfile.ZipInfo(name, date_time=FIXED_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.create_system = 0
            zi.external_attr = 0o600 << 16
            z.writestr(zi, members[name])
    return buf.getvalue()


def _atomic_write(data: bytes, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix="." + path.name + ".", suffix=".tmp", dir=str(path.parent))
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(data)
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp, path)
    except BaseException:
        if os.path.exists(tmp):
            os.unlink(tmp)
        raise


def write_members(members: dict, path) -> str:
    """Write members deterministically and atomically. Returns the sha256 of the written file."""
    data = zip_bytes(members)
    _atomic_write(data, Path(path))
    return sha256_bytes(data)


def _flatten_embedded(blob: bytes) -> bytes:
    """An embedded Office package rewritten with fixed timestamps and no authorship (PMR `_flatten`)."""
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        members = {i.filename: z.read(i.filename) for i in z.infolist()}
    for name in ("docProps/core.xml", "docProps/app.xml"):
        if name in members:
            d = members[name]
            d = re.sub(rb"<dcterms:(created|modified)[^>]*>[^<]*</dcterms:\1>", b"", d)
            d = re.sub(rb"<dc:creator>[^<]*</dc:creator>", b"<dc:creator></dc:creator>", d)
            d = re.sub(rb"<cp:lastModifiedBy>[^<]*</cp:lastModifiedBy>", b"", d)
            members[name] = d
    return zip_bytes(members)


# ------------------------------------------------------------------ normalise

def _rels_name(part: str) -> str:
    d, b = posixpath.split(part)
    return posixpath.join(d, "_rels", b + ".rels")


def _source_of(rels: str) -> str:
    """'ppt/slides/_rels/slide1.xml.rels' to 'ppt/slides/slide1.xml'; '_rels/.rels' to ''."""
    d, b = posixpath.split(rels)
    return posixpath.join(posixpath.dirname(d), b[:-len(".rels")]) if b != ".rels" else ""


def _resolve(source: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(source), target))


def _ext(name: str) -> str:
    """Lower-case extension by OPC rules: '_rels/.rels' is 'rels' (splitext treats it as a stem)."""
    b = posixpath.basename(name)
    return b.rsplit(".", 1)[1].lower() if "." in b else ""


def _xml(b: bytes):
    return etree.fromstring(b, etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=True))


def _ser(root) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def normalise(members: dict, *, title: str, theme: Theme, date=None) -> dict:
    """The normalised member map. Pure: same input, same output. date as for w3cdtf()."""
    m = dict(members)
    # thumbnail part and its package relationship
    rels = _xml(m["_rels/.rels"])
    for rel in list(rels):
        if rel.get("Type") == RT_THUMB:
            m.pop(_resolve("", rel.get("Target")), None)
            rels.remove(rel)
    m["_rels/.rels"] = _ser(rels)
    # docProps
    n_slides = sum(1 for n in m if re.fullmatch(r"ppt/slides/slide\d+\.xml", n))
    m["docProps/core.xml"] = core_xml(title, date)
    m["docProps/app.xml"] = app_xml(n_slides)
    # theme parts
    for n in theme_parts(list(m), m.get("[Content_Types].xml")):
        m[n] = theme.retheme_part(m[n])
    # embedded workbooks
    for n in list(m):
        if n.startswith("ppt/embeddings/") and n.endswith(".xlsx"):
            m[n] = _flatten_embedded(m[n])
    # content types: svg Default, no Override pointing at a part that is gone
    ct = _xml(m["[Content_Types].xml"])
    q = "{%s}" % CT_NS
    defaults = {d.get("Extension", "").lower(): d for d in ct.findall(q + "Default")}
    if "svg" not in defaults:
        el = etree.Element(q + "Default", Extension="svg", ContentType=SVG_CT)
        first_override = ct.find(q + "Override")
        if first_override is not None:
            first_override.addprevious(el)
        else:
            ct.append(el)
    else:
        defaults["svg"].set("ContentType", SVG_CT)
    for o in ct.findall(q + "Override"):
        if o.get("PartName", "").lstrip("/") not in m:
            ct.remove(o)
    ext_used = {_ext(n) for n in m}
    for ext, d in defaults.items():
        if ext not in ext_used and ext not in ("rels", "xml", "svg"):
            ct.remove(d)
    m["[Content_Types].xml"] = _ser(ct)
    return m


# ------------------------------------------------------------------ save

def manifest_path(path) -> Path:
    p = Path(path)
    return p.with_name(p.stem + ".manifest.json")


def lock_path(path) -> Path:
    p = Path(path)
    return p.with_name("~$" + p.name)


def save(prs, path, *, title: str, theme: Theme, manifest: dict | None = None, overwrite: bool = False,
         date=None) -> str:
    """Save prs to path, normalised and atomic. Writes <stem>.manifest.json beside it. Returns sha256.

    date is the deck's date for docProps created/modified (pass deck.json "date"); None falls back to
    1980-01-01. Either way the same inputs give the same bytes. A bad date raises before anything."""
    from pptx import Presentation
    stamp = w3cdtf(date)
    path = Path(path)
    if lock_path(path).exists():
        raise HandEditDrift(str(path), need="no lock file", have=lock_path(path).name,
                            hint="the deck is open in PowerPoint; close it before rebuilding")
    if path.exists() and not overwrite:
        mp = manifest_path(path)
        have = sha256_file(path)
        try:
            prev = json.loads(mp.read_text(encoding="utf-8")).get("sha256")
        except (FileNotFoundError, ValueError):
            prev = None
        if prev != have:
            raise HandEditDrift(str(path), need=prev or "a manifest from the last build", have=have,
                                hint="the file changed since the kit wrote it; keep the hand edits (accept) "
                                     "or rebuild with overwrite=True")
    buf = io.BytesIO()
    prs.save(buf)
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as z:
        members = {i.filename: z.read(i.filename) for i in z.infolist() if not i.is_dir()}
    data = zip_bytes(normalise(members, title=title, theme=theme, date=stamp))
    back = Presentation(io.BytesIO(data))                      # round-trip before touching disk
    if len(back.slides) != len(prs.slides):
        raise UnsupportedPart("normalised package %s" % path.name, need=len(prs.slides), have=len(back.slides),
                              hint="slides were lost on reopen; a part the normaliser does not know was dropped")
    _atomic_write(data, path)
    sha = sha256_bytes(data)
    man = dict(manifest or {})
    man.update({"pptx": path.name, "sha256": sha, "title": title, "theme": theme.name,
                "slides": len(prs.slides)})
    if date is not None:
        man["date"] = stamp
    _atomic_write((json.dumps(man, indent=2, sort_keys=True, ensure_ascii=False) + "\n").encode("utf-8"),
                  manifest_path(path))
    return sha


# ------------------------------------------------------------------ validation

def find_validator() -> Path:
    hits = sorted(glob.glob(VALIDATOR_GLOB))
    if not hits:
        raise ValidatorMissing("XSD validator", need=VALIDATOR_GLOB, have="no match",
                               hint="the synced pptx skill is missing; validation fails until it is back")
    return Path(hits[0])


_PML = {}
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_PART_XSD = re.compile(r"ppt/(slides|slideLayouts|slideMasters|notesSlides|notesMasters|handoutMasters)/[^/]+\.xml$")


def _pml_schema(validator: Path):
    """The validator bundle's own pml.xsd, loaded in place (never copied)."""
    key = str(validator)
    if key not in _PML:
        xsd = validator.parent / "schemas" / "ISO-IEC29500-4_2016" / "pml.xsd"
        if not xsd.is_file():
            raise ValidatorMissing("pml.xsd", need=str(xsd), have="missing",
                                   hint="the validator's schema bundle moved; slide parts cannot be checked")
        _PML[key] = etree.XMLSchema(etree.parse(str(xsd)))
    return _PML[key]


def _strip_mc(root):
    """mc:AlternateContent replaced by its Fallback (or dropped), as a reader without the extension would."""
    for ac in list(root.iter("{%s}AlternateContent" % _MC)):
        parent = ac.getparent()
        fb = ac.find("{%s}Fallback" % _MC)
        i = parent.index(ac)
        for c in (list(fb) if fb is not None else []):
            parent.insert(i, c)
            i += 1
        parent.remove(ac)
    ign = root.get("{%s}Ignorable" % _MC)
    if ign is not None:
        del root.attrib["{%s}Ignorable" % _MC]
    return root


# PowerPoint writes ST_TextBulletSizePercent in its transitional integer form (100000 for 100%); the strict schema
# only accepts "25%".."400%". Within that range it is a known PowerPoint value, reported, never an error.
_KNOWN = (re.compile(r"buSzPct', attribute 'val': .*The value '(\d+)' is not accepted"),)
_KNOWN_RANGE = (25000, 400000)


def known_value(message: str) -> str | None:
    """A short label when a schema message is a value PowerPoint writes itself (see _KNOWN), else None."""
    for rx in _KNOWN:
        m = rx.search(message)
        if m and _KNOWN_RANGE[0] <= int(m.group(1)) <= _KNOWN_RANGE[1]:
            return "buSzPct %s (PowerPoint's transitional percent)" % m.group(1)
    return None


def part_xsd(path, validator: Path | None = None, known: list | None = None) -> list[str]:
    """Schema errors in slide, layout, master and notes parts.

    The referenced validator maps a schema only for parts directly under ppt/, so every slide, layout
    and master is reported "Skipped (no schema)" and a foreign element in a master passes (measured
    2026-10-03 with a planted <p:bogus/>). This pass closes that gap with the bundle's own pml.xsd. A known
    PowerPoint value (known_value) is appended to `known` when given, never reported as an error.
    """
    schema = _pml_schema(validator or find_validator())
    errs = []
    with zipfile.ZipFile(path) as z:
        for n in sorted(z.namelist()):
            if not _PART_XSD.match(n):
                continue
            root = _strip_mc(_xml(z.read(n)))
            if not schema.validate(root):
                shown = 0
                for e in schema.error_log:
                    k = known_value(e.message)
                    if k:
                        if known is not None:
                            known.append("%s:%d: %s" % (n, e.line, k))
                        continue
                    if shown < 3:
                        errs.append("%s:%d: %s" % (n, e.line, e.message[:240]))
                    shown += 1
    return errs


_NS_COVER = {}


def xsd_namespaces(validator: Path) -> set:
    """Every targetNamespace the validator's schema bundle declares."""
    key = str(validator)
    if key not in _NS_COVER:
        out = set()
        for x in (validator.parent / "schemas").rglob("*.xsd"):
            try:
                tn = etree.parse(str(x)).getroot().get("targetNamespace")
            except etree.XMLSyntaxError:
                continue
            if tn:
                out.add(tn)
        if not out:
            raise ValidatorMissing("schema bundle", need=str(validator.parent / "schemas"), have="no targetNamespace",
                                   hint="the validator's schema bundle is empty or moved")
        _NS_COVER[key] = out
    return _NS_COVER[key]


def xsd_scope(path, validator: Path) -> list[tuple[str, str]]:
    """[(part, root namespace)] for every XML part whose root namespace the bundle does not cover (PowerPoint's
    own extension parts: authors.xml, revisionInfo.xml, changesInfos/, modern comments, web extensions). Those
    parts are skipped by the XSD pass; every other part is validated."""
    cover = xsd_namespaces(validator)
    out = []
    with zipfile.ZipFile(path) as z:
        for n in sorted(z.namelist()):
            if not n.endswith(".xml") or n == "[Content_Types].xml":
                continue
            try:
                ns = etree.QName(_xml(z.read(n))).namespace or ""
            except (etree.XMLSyntaxError, ValueError):
                continue                  # not well-formed: the validator reports it, so it is not skipped
            if ns not in cover:
                out.append((n, ns))
    return out


def validate_xsd(path, report: dict | None = None) -> tuple[bool, str]:
    """(ok, detail): the referenced ISO 29500 validator, plus part_xsd over slide-level parts.

    Only parts whose root namespace the validator's bundle covers count (xsd_scope): the validator's errors on the
    others are dropped and the parts listed in report["skipped"] as (part, namespace). Every other check the
    validator runs (relationship ids, slide defects) still applies to the whole package. Known PowerPoint values (known_value) are listed in
    report["known"] and never fail. Raises ValidatorMissing when the validator or its schema bundle cannot be found
    or cannot run.
    """
    v = find_validator()
    skipped = xsd_scope(path, v)
    known: list = []
    r = subprocess.run([sys.executable, str(v), os.path.abspath(str(path))], cwd=str(v.parent),
                       env=dict(os.environ, PYTHONPATH=str(v.parent)), capture_output=True, text=True,
                       timeout=600)
    out = (r.stdout + r.stderr).strip()
    if r.returncode not in (0, 1) or "Traceback" in out or (r.returncode == 0 and "PASSED" not in out):
        raise ValidatorMissing("XSD validator run", need="exit 0 or 1", have="exit %d" % r.returncode,
                               hint=(out.splitlines() or [""])[-1][:300])
    vfail = r.returncode != 0
    if vfail:
        out, vfail = _drop_known(out, known, {n for n, _ in skipped})
    parts = part_xsd(path, v, known)
    ok = not vfail and not parts
    detail = [out.splitlines()[-1] if (ok and out) else out] if vfail or ok else []
    if parts:
        detail.append("slide-level parts against pml.xsd:")
        detail += ["  " + e for e in parts]
    if report is not None:
        report["skipped"] = skipped
        report["known"] = known
    return ok, "\n".join(d for d in detail if d)


def _drop_known(out: str, known: list, skipped: set = frozenset()) -> tuple[str, bool]:
    """The validator's XSD report without the blocks of skipped parts (no schema for their namespace) and without
    known PowerPoint values, and whether anything still fails. A block whose every listed error is known is dropped
    only when the validator listed all of them (it shows three). When nothing was dropped the validator's failure
    stands unchanged, so a failure this parser does not recognise can never turn into a pass."""
    lines = out.splitlines()
    keep, i, dropped = [], 0, 0
    head = re.compile(r"^  (\S.*): (\d+) new error\(s\)$")
    while i < len(lines):
        m = head.match(lines[i])
        if not m:
            keep.append(lines[i])
            i += 1
            continue
        j = i + 1
        errs = []
        while j < len(lines) and lines[j].startswith("    - "):
            errs.append(lines[j])
            j += 1
        rest = [e for e in errs if not known_value(e)]
        if m.group(1) in skipped:
            dropped += 1
        elif errs and not rest and len(errs) == int(m.group(2)):
            known.extend("%s: %s" % (m.group(1), known_value(e)) for e in errs)
            dropped += 1
        else:
            keep.append(lines[i])
            keep.extend(rest if len(errs) == int(m.group(2)) else errs)
        i = j
    if not dropped:                       # nothing known was removed: the validator's failure stands as it is
        return out, True
    still = any(head.match(ln) for ln in keep) or any(("FAILED" in ln or "rror" in ln) and "NEW validation errors"
                                                       not in ln for ln in keep)
    if not still:
        keep = [ln for ln in keep if "Found NEW validation errors" not in ln]
    return "\n".join(keep), still


_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_SHAPE_PARTS = re.compile(r"ppt/(slides|slideLayouts|slideMasters|notesSlides|notesMasters|handoutMasters)/\w+\d+\.xml$")


def structural_check(path) -> list[str]:
    """Package defects as 'part: message' strings. Empty list is a pass.

    Once-only duplicates, dangling r:ids, missing relationship targets, content types (svg included),
    unique cNvPr ids per part, glue ids that exist, sldId >= 256, no control characters, no negative
    extents, and a python-pptx reopen.
    """
    out = []
    with zipfile.ZipFile(path) as z:
        names = [i.filename for i in z.infolist() if not i.is_dir()]
        blobs = {n: z.read(n) for n in names}
    nameset = set(names)
    if "[Content_Types].xml" not in nameset:
        return ["[Content_Types].xml: missing"]
    q = "{%s}" % CT_NS
    ct = _xml(blobs["[Content_Types].xml"])
    defaults = {d.get("Extension", "").lower(): d.get("ContentType") for d in ct.findall(q + "Default")}
    overrides = {o.get("PartName", "").lstrip("/"): o.get("ContentType") for o in ct.findall(q + "Override")}
    for ext in ("rels", "xml"):
        if ext not in defaults:
            out.append("[Content_Types].xml: no Default for .%s" % ext)
    for pn in overrides:
        if pn not in nameset:
            out.append("[Content_Types].xml: Override for a missing part /%s" % pn)
    for n in names:
        if n == "[Content_Types].xml":
            continue
        ext = _ext(n)
        t = overrides.get(n) or defaults.get(ext)
        if t is None:
            out.append("%s: no content type" % n)
        elif ext == "svg" and t != SVG_CT:
            out.append("%s: svg typed %s" % (n, t))
    # relationships
    rel_ids = collections.defaultdict(set)
    for n in names:
        if not n.endswith(".rels"):
            continue
        src = _source_of(n)
        if src and src not in nameset:
            out.append("%s: rels for a missing part %s" % (n, src))
        try:
            root = _xml(blobs[n])
        except etree.XMLSyntaxError as e:
            out.append("%s: not well-formed (%s)" % (n, e))
            continue
        for rel in root:
            rid = rel.get("Id")
            if rid in rel_ids[src]:
                out.append("%s: duplicate relationship id %s" % (n, rid))
            rel_ids[src].add(rid)
            if rel.get("TargetMode") == "External":
                continue
            tgt = _resolve(src, rel.get("Target", ""))
            if tgt not in nameset:
                out.append("%s: %s targets missing part %s" % (n, rid, tgt))
    # XML parts
    for n in names:
        if not (n.endswith(".xml") or n.endswith(".rels")):
            continue
        text = blobs[n].decode("utf-8", "replace")
        if _CTRL.search(text):
            out.append("%s: control character in XML" % n)
        if n.endswith(".rels"):
            continue
        try:
            root = _xml(blobs[n])
        except etree.XMLSyntaxError as e:
            out.append("%s: not well-formed (%s)" % (n, e))
            continue
        rids = rel_ids.get(n, set())
        for el in root.iter():
            if not isinstance(el.tag, str):
                continue
            for k, v in el.attrib.items():
                if k.startswith("{%s}" % R_NS) and v and v not in rids:
                    out.append("%s: r:%s=%s has no relationship" % (n, etree.QName(k).localname, v))
            if etree.QName(el).localname in ("ext", "chExt") and el.tag.startswith(_A):
                for a in ("cx", "cy"):
                    if el.get(a) is not None and int(el.get(a)) < 0:
                        out.append("%s: negative extent %s=%s" % (n, a, el.get(a)))
        if _SHAPE_PARTS.match(n):
            for container, child, k, sid in once_only_violations(root):
                out.append("%s: <%s> holds <%s> x%d (shape id %s)" % (n, container, child, k, sid))
            ids = [e.get("id") for e in root.iter(_P + "cNvPr")]
            for i, k in collections.Counter(ids).items():
                if k > 1:
                    out.append("%s: cNvPr id %s used %d times" % (n, i, k))
            idset = set(ids)
            for tag in ("stCxn", "endCxn"):
                for e in root.iter(_A + tag):
                    if e.get("id") not in idset:
                        out.append("%s: %s glued to missing shape id %s" % (n, tag, e.get("id")))
        if n == "ppt/presentation.xml":
            for e in root.iter(_P + "sldId"):
                v = int(e.get("id", "0"))
                if not 256 <= v <= 2147483647:
                    out.append("%s: sldId %d outside 256..2147483647" % (n, v))
            for tag in ("sldMasterId", "sldLayoutId"):
                for e in root.iter(_P + tag):
                    if int(e.get("id", "0")) < 2147483648:
                        out.append("%s: %s %s below 2147483648" % (n, tag, e.get("id")))
        if n.startswith("ppt/slideMasters/"):
            for e in root.iter(_P + "sldLayoutId"):
                if int(e.get("id", "0")) < 2147483648:
                    out.append("%s: sldLayoutId %s below 2147483648" % (n, e.get("id")))
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        for s in prs.slides:
            list(s.shapes)
    except Exception as e:                                  # any reopen failure is a defect
        out.append("reopen: %s: %s" % (type(e).__name__, e))
    return out


__all__ = ["save", "w3cdtf", "validate_xsd", "part_xsd", "xsd_scope", "xsd_namespaces", "known_value", "structural_check", "find_validator", "normalise", "write_members",
           "zip_bytes", "core_xml", "app_xml", "sha256_file", "sha256_bytes", "manifest_path", "lock_path",
           "FIXED_TIME", "FIXED_DATE", "AUTHOR", "VALIDATOR_GLOB"]
