"""Exhibits: the components a proposal deck shows rather than describes (ruling 22, track A2).

Each draws a whole exhibit inside one box, measures every word before the first shape exists, and raises
(Overflows, Capacity, GeometryError, BelowFloor) instead of clamping. The whole call is one transaction: a
refusal leaves the slide as it was. Colours are theme roles; vendor and brand marks keep their own artwork.

    cycle_ring      a closed loop of N equal stages around a hub disc, direction drawn on the loop itself:
                    style "wedges" (one ring of arrow-pointed segments) or "nodes" (discs joined by curved
                    arrows); each stage's head and line sit outside the ring (lifecycle-loop)
    mark_wall       tabbed groups of vendor and brand marks, each mark over its caption (credentials, why-us,
                    tech stacks); never pills
    screen_frame    a product capture in browser chrome, a laptop or no frame, optionally straddling a band, its
                    smallest type checked against the 8pt floor at the size it is placed (product-screen)
    portrait        a photograph clipped to a circle, or an initials disc; never a generic silhouette
    quote_portrait  a quotation panel beside its speaker's portrait
    cell_grid       a matrix whose cells carry their state as glyph plus word, rows on a dark rail, columns under
                    solid header tabs (coverage-matrix)

Geometry notes the gates rely on:
  * the wedge ring is ONE custom-geometry shape with one closed path per stage, so its stages never overlap as
    boxes (P05) and the hub, the glyphs and the active stage nest inside it; labels sit outside its square;
  * node arrows are open custom paths (curved), headed at the next node, both ends on the node outlines (P06);
  * a capture's sidecar (scripts/capture.py) carries its smallest CSS font; screen_frame turns it into points at
    the placed width and refuses under the floor, and records the figure in the picture's alt text.
"""
from __future__ import annotations

import hashlib
import io
import json
import math
from pathlib import Path

from lxml import etree

from . import icons, measure, ooxml, prims
from .canvas import EMU_PER_IN, Box
from .errors import BelowFloor, Capacity, GeometryError, Overflows, ThemeInvariant, UnsupportedPart
from .glyphs import glyph
from .prims import ON_FILL, canvas_for, theme_of
from .read.ns import NS
from .text import Para, Run, need_height, textbox

_A = "{%s}" % NS["a"]
_P = "{%s}" % NS["p"]
_R = "{%s}" % NS["r"]

HEAD_PT, BODY_PT, CAPTION_PT = 12.0, 10.5, 10.0
STRONG = ("accent", "accent-deep", "ink-deep", "ink-strong", "navy")      # solid fills (ruling 12)
LIGHT = ("card", "paper")
CAPTURE_SUFFIX = ".capture.json"


# ------------------------------------------------------------------ small helpers

def _k(slide) -> float:
    return canvas_for(slide).k


def _e(v: float) -> int:
    return int(round(v * EMU_PER_IN))


def _stack(head, body, head_role="ink-deep", body_role="ink", *, align="left", head_pt=HEAD_PT, body_pt=BODY_PT,
           gap=2.0):
    ps = []
    if head:
        ps.append(Para((Run(str(head), head_role, bold=True, size=head_pt),), align=align))
    if body:
        ps.append(Para((Run(str(body), body_role, size=body_pt),), align=align, space_before=gap if head else 0.0))
    if not ps:
        raise GeometryError("text block", need="a head or a line", have="neither")
    return ps


def _need(paras, w, th) -> float:
    return need_height(paras, w, th)


def _face(th, bold=False):
    return measure.face_for(prims._face(th), bold)


def _text_role_on(fill: str, *, light_role="ink") -> str:
    return ON_FILL[fill] if fill in STRONG else light_role


def _glyph_role_on(fill: str) -> str:
    return ON_FILL[fill] if fill in STRONG else "accent"


def _arc_cubics(cx, cy, r, a0, a1):
    """Cubic Bezier segments approximating the circular arc from angle a0 to a1 (degrees, screen space: y down, so
    increasing angle runs clockwise). Returns [((x1, y1), (x2, y2), (x, y))], each segment at most 45 degrees."""
    n = max(1, int(math.ceil(abs(a1 - a0) / 45.0 - 1e-9)))
    out = []
    step = (a1 - a0) / n
    for i in range(n):
        s = math.radians(a0 + i * step)
        t = math.radians(a0 + (i + 1) * step)
        kk = 4.0 / 3.0 * math.tan((t - s) / 4.0)
        p0 = (cx + r * math.cos(s), cy + r * math.sin(s))
        p3 = (cx + r * math.cos(t), cy + r * math.sin(t))
        c1 = (p0[0] - kk * r * math.sin(s), p0[1] + kk * r * math.cos(s))
        c2 = (p3[0] + kk * r * math.sin(t), p3[1] - kk * r * math.cos(t))
        out.append((c1, c2, p3))
    return out


def _polar(cx, cy, r, a):
    return (cx + r * math.cos(math.radians(a)), cy + r * math.sin(math.radians(a)))


def _path_shape(slide, paths, name: str, *, fill, line=None, weight_pt=None, head=None, th, closed=True, frame=None):
    """One custom-geometry shape from paths in slide inches. Each path is a list of ops: ("M", p), ("L", p),
    ("C", c1, c2, p), closed with Z when closed. The frame is the bounding box of every point (for arcs of at most
    45 degrees the control points lie inside the arc's own box), or `frame` (a Box holding every point) when given,
    so a part drawn over a whole shares its frame and nests in it. Open paths carry fill="none" (a drawn line)."""
    pts = [p for path in paths for op in path for p in op[1:]]
    x0, y0 = min(p[0] for p in pts), min(p[1] for p in pts)
    x1, y1 = max(p[0] for p in pts), max(p[1] for p in pts)
    if frame is not None:
        if x0 < frame.x - 1e-6 or y0 < frame.y - 1e-6 or x1 > frame.r + 1e-6 or y1 > frame.b + 1e-6:
            raise GeometryError("path %s" % name, need="points inside its frame", have=(x0, y0, x1, y1))
        x0, y0, x1, y1 = frame.x, frame.y, frame.r, frame.b
    w, h = max(1, _e(x1) - _e(x0)), max(1, _e(y1) - _e(y0))
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _e(x0), _e(y0), w, h)
    sh.name = name
    sp = prims._sppr(sh)
    g = etree.Element(_A + "custGeom")
    for t in ("avLst", "gdLst", "ahLst", "cxnLst"):
        etree.SubElement(g, _A + t)
    etree.SubElement(g, _A + "rect", l="l", t="t", r="r", b="b")
    pl = etree.SubElement(g, _A + "pathLst")
    for path in paths:
        pe = etree.SubElement(pl, _A + "path", w=str(w), h=str(h))
        if not closed:
            pe.set("fill", "none")
        for op in path:
            tag = {"M": "moveTo", "L": "lnTo", "C": "cubicBezTo"}[op[0]]
            el = etree.SubElement(pe, _A + tag)
            for (px, py) in op[1:]:
                etree.SubElement(el, _A + "pt", x=str(_e(px) - _e(x0)), y=str(_e(py) - _e(y0)))
        if closed:
            etree.SubElement(pe, _A + "close")
    sp.replace(sp.find(_A + "prstGeom"), g)
    prims.set_fill(sh, fill if closed else None, theme=th)
    if line is None:
        prims.set_line(sh, None, theme=th)
    else:
        prims.set_line(sh, line, weight_pt, theme=th, head=head, cap="rnd" if not closed else None)
    prims._drop_text(sh)
    ooxml.finish(sh)
    return sh


# ------------------------------------------------------------------ cycle_ring

def _segment(cx, cy, r_out, r_in, a0, a1, phi):
    """One arrow-pointed annular segment: its outer and inner arcs run from a0 to a1 (degrees), its leading edge
    comes to a point at a1 + phi on the mid radius, and its trailing edge is notched to a0 + phi, so the point of
    the segment before nests in the notch with an even gap."""
    rm = (r_out + r_in) / 2.0
    ops = [("M", _polar(cx, cy, r_out, a0))]
    ops += [("C",) + seg for seg in _arc_cubics(cx, cy, r_out, a0, a1)]
    ops.append(("L", _polar(cx, cy, rm, a1 + phi)))
    ops.append(("L", _polar(cx, cy, r_in, a1)))
    ops += [("C",) + seg for seg in _arc_cubics(cx, cy, r_in, a1, a0)]
    ops.append(("L", _polar(cx, cy, rm, a0 + phi)))
    return ops


def _hub_paras(hub, role):
    if hub is None:
        return None
    head, line = (hub, None) if isinstance(hub, str) else (tuple(hub) + (None,))[:2]
    return _stack(head, line, role, role, align="center", head_pt=13.0, body_pt=10.5)


@ooxml.transactional
def cycle_ring(slide, box: Box, stages, hub=None, *, style="wedges", fill=None, hub_fill=None, active=None,
               start=None, label_w=None, name: str, theme=None) -> dict:
    """A closed loop of N equal stages (3 to 8) around a hub disc, direction drawn on the loop, each stage's head
    and line set outside the ring at its own angle.

    stages: [(glyph, head, line)] in loop order, clockwise; line may be None. hub: None, "label" or
    ("head", "line"), set in the hub disc. style "wedges": one ring of arrow-pointed segments, the point of each
    nesting in the notch of the next (fill default card, glyphs in the brand; hub default accent). style "nodes":
    N discs on the ring joined by curved brand arrows (fill default accent, hub default ink-deep). active: the
    index of the stage that is current, drawn in the accent fill (a state). start: the angle of stage 0's centre in
    degrees, 0 at three o'clock, clockwise; by default no stage sits on the vertical axis when N is even, and stage 0
    sits at twelve o'clock when N is odd. label_w caps the label width (labels otherwise take the room the ring
    leaves). The ring takes the largest radius at which every label fits; Overflows when none does.

    Returns {"ring": Box (the ring's square), "hub": Box, "centre": (x, y), "radius": r, "labels": [Box]}."""
    th = theme_of(theme)
    k = _k(slide)
    stages = list(stages)
    n = len(stages)
    if not 3 <= n <= 8:
        raise Capacity("cycle_ring %s" % name, need="3 to 8 stages", have=n,
                       hint="a loop of more than eight stages no longer reads in one glance; merge stages")
    if style not in ("wedges", "nodes"):
        raise GeometryError("cycle_ring %s style" % name, need="wedges or nodes", have=style)
    fill = fill or ("card" if style == "wedges" else "accent")
    hub_fill = hub_fill or ("accent" if style == "wedges" else "ink-deep")
    for what, f in (("fill", fill), ("hub_fill", hub_fill)):
        if f not in STRONG + LIGHT:
            raise ThemeInvariant("cycle_ring %s %s" % (name, what), need=" / ".join(STRONG + LIGHT), have=f)
    if hub_fill not in STRONG:
        raise ThemeInvariant("cycle_ring %s hub_fill" % name, need="a solid fill: " + " / ".join(STRONG),
                             have=hub_fill, hint="the hub names what the cycle serves; it is the one solid disc")
    if active is not None and not 0 <= active < n:
        raise GeometryError("cycle_ring %s active" % name, need="0..%d" % (n - 1), have=active)
    span = 360.0 / n
    if start is None:
        start = -90.0 + (span / 2.0 if n % 2 == 0 else 0.0)
    thetas = [start + i * span for i in range(n)]
    gap = 0.12 * k                                   # ring to label
    lab_gap = 0.10 * k                               # between labels on one side

    # where each label goes: side (right / left) or axis (top / bottom)
    place = []
    for t in thetas:
        c, s = math.cos(math.radians(t)), math.sin(math.radians(t))
        place.append(("right" if c > 0 else "left") if abs(c) >= 0.30 else ("bottom" if s > 0 else "top"))
    stacks = []
    for (g, head, line), p in zip(stages, place):
        align = {"right": "left", "left": "right"}.get(p, "center")
        stacks.append(_stack(head, line, align=align))
    hub_paras = _hub_paras(hub, ON_FILL[hub_fill])

    def layout(R):
        """Label boxes at outer radius R, or None when they do not fit."""
        cy_top = box.y
        tops = [i for i, p in enumerate(place) if p == "top"]
        bots = [i for i, p in enumerate(place) if p == "bottom"]
        w_axis = min(box.w * 0.34, 3.2 * k, 2 * R)          # within the ring's width, clear of side labels
        h_top = max([_need(stacks[i], w_axis, th) for i in tops] or [0.0])
        h_bot = max([_need(stacks[i], w_axis, th) for i in bots] or [0.0])
        room_y = box.h - (h_top + gap if tops else 0.0) - (h_bot + gap if bots else 0.0)
        if 2 * R > room_y + 1e-6:
            return None
        cy = cy_top + (h_top + gap if tops else 0.0) + room_y / 2.0
        cx = box.cx
        side_w = box.w / 2.0 - R - gap
        if label_w is not None:
            side_w = min(side_w, label_w)
        if side_w < 1.5 * k:
            return None
        out = [None] * n
        for side in ("right", "left"):
            idx = sorted((i for i, p in enumerate(place) if p == side),
                         key=lambda i: math.sin(math.radians(thetas[i])))
            boxes = []
            for i in idx:
                h = _need(stacks[i], side_w, th)
                yc = cy + R * math.sin(math.radians(thetas[i]))
                boxes.append([i, yc - h / 2.0, h])
            # keep inside the box, then push apart top to bottom and back up from the foot
            for b in boxes:
                b[1] = min(max(b[1], box.y), box.b - b[2])
            for j in range(1, len(boxes)):
                prev = boxes[j - 1]
                boxes[j][1] = max(boxes[j][1], prev[1] + prev[2] + lab_gap)
            for j in range(len(boxes) - 1, -1, -1):
                lim = box.b if j == len(boxes) - 1 else boxes[j + 1][1] - lab_gap
                boxes[j][1] = min(boxes[j][1], lim - boxes[j][2])
            if boxes and (boxes[0][1] < box.y - 1e-6):
                return None
            x = cx + R + gap if side == "right" else cx - R - gap - side_w
            for i, y, h in boxes:
                out[i] = Box(x, y, side_w, h)
        for i in tops:
            h = _need(stacks[i], w_axis, th)
            out[i] = Box(cx - w_axis / 2.0, cy - R - gap - h, w_axis, h)
        for i in bots:
            h = _need(stacks[i], w_axis, th)
            out[i] = Box(cx - w_axis / 2.0, cy + R + gap, w_axis, h)
        return cx, cy, out

    r_max = min(box.h / 2.0, box.w / 2.0 - gap - 1.5 * k)
    R, fit = r_max, None
    while R >= 1.0 * k:
        fit = layout(R)
        if fit is not None:
            break
        R -= 0.05 * k
    if fit is None:
        raise Overflows("cycle_ring %s" % name, need="a ring of 1.0in radius with every label beside it",
                        have="%.2f x %.2fin" % (box.w, box.h),
                        hint="a wider box, label_w, or fewer words per stage")
    cx, cy, label_boxes = fit
    ring = Box(cx - R, cy - R, 2 * R, 2 * R)

    if style == "wedges":
        r_out, r_in = R, R * 0.50
        rm = (r_out + r_in) / 2.0
        phi = math.degrees(min(0.24 * k, 0.20 * (r_out - r_in)) / rm)
        half_gap = math.degrees(0.035 * k / rm)
        hub_r = r_in - 0.10 * k
        gs = min(0.56 * k, 0.40 * (r_out - r_in))
    else:
        dn = min(1.10 * k, 0.55 * 2 * R * math.sin(math.pi / n), R * 0.50)
        r_ring = R - dn / 2.0
        hub_r = r_ring - dn / 2.0 - 0.22 * k
        # the hub's square must clear every node's square on one axis (P05 compares frames)
        for t in thetas:
            px, py = _polar(0, 0, r_ring, t)
            clear = max(abs(px), abs(py)) - dn / 2.0 - 0.06 * k
            hub_r = min(hub_r, clear)
        gs = dn * 0.48
    if hub_paras is not None:
        side = 2 * hub_r * 0.70
        if hub_r <= 0 or _need(hub_paras, side, th) > side:
            raise Overflows("cycle_ring %s hub" % name, need=round(_need(hub_paras, max(side, 0.1), th), 3),
                            have=round(max(side, 0.0), 3), hint="a shorter hub label or a larger box")

    # draw: the ring first, then its marks, the hub, the labels
    shapes = {}
    if style == "wedges":
        bounds = [t - span / 2.0 for t in thetas]
        paths = []
        for i in range(n):
            a0, a1 = bounds[i] + half_gap, bounds[i] + span - half_gap
            paths.append(_segment(cx, cy, r_out, r_in, a0, a1, phi))
        shapes["ring"] = _path_shape(slide, paths, "%s.ring" % name, fill=fill, th=th)
        if active is not None:              # drawn over its own segment on the ring's frame, so P05 nests the two
            ring_frame = Box(shapes["ring"].left / EMU_PER_IN, shapes["ring"].top / EMU_PER_IN,
                             shapes["ring"].width / EMU_PER_IN, shapes["ring"].height / EMU_PER_IN)
            _path_shape(slide, [paths[active]], "%s.%d.active" % (name, active + 1), fill="accent", th=th,
                        frame=ring_frame)
        for i, t in enumerate(thetas):
            f = "accent" if i == active else fill
            gx, gy = _polar(cx, cy, rm, t + phi / 2.0)
            glyph(slide, stages[i][0], Box(gx - gs / 2, gy - gs / 2, gs, gs), role=_glyph_role_on(f),
                  shape_name="%s.%d.icon" % (name, i + 1), theme=th)
    else:
        centres = [_polar(cx, cy, r_ring, t) for t in thetas]
        for i, (px, py) in enumerate(centres):
            f = "accent" if i == active else fill
            if active is not None and i != active and f == "accent":
                f = "ink-strong"
            prims.circle(slide, Box(px - dn / 2, py - dn / 2, dn, dn), fill=f if f in ON_FILL else "card",
                         name="%s.%d.node" % (name, i + 1), theme=th)
            glyph(slide, stages[i][0], Box(px - gs / 2, py - gs / 2, gs, gs), role=_glyph_role_on(f),
                  shape_name="%s.%d.icon" % (name, i + 1), theme=th)
        alpha = math.degrees(2 * math.asin(min(1.0, (dn / 2.0) / (2 * r_ring))))
        for i, t in enumerate(thetas):
            a0, a1 = t + alpha, t + span - alpha
            if a1 - a0 < 4.0:
                raise Overflows("cycle_ring %s arrows" % name, need="room between the nodes", have="%.1f deg" % (a1 - a0),
                                hint="fewer stages or a larger box")
            path = [("M", _polar(cx, cy, r_ring, a0))] + [("C",) + s for s in _arc_cubics(cx, cy, r_ring, a0, a1)]
            _path_shape(slide, [path], "%s.%d.next" % (name, i + 1), fill=None, line="accent",
                        weight_pt=1.25, head="end", th=th, closed=False)
    hub_box = Box(cx - hub_r, cy - hub_r, 2 * hub_r, 2 * hub_r)
    prims.circle(slide, hub_box, fill=hub_fill, name="%s.hub" % name, theme=th)
    if hub_paras is not None:
        side = 2 * hub_r * 0.70
        h = _need(hub_paras, side, th)
        textbox(slide, Box(cx - side / 2, cy - h / 2, side, h), hub_paras, "%s.hub.text" % name, th,
                background=hub_fill)
    for i, (b, paras) in enumerate(zip(label_boxes, stacks)):
        textbox(slide, b, paras, "%s.%d.text" % (name, i + 1), th)
    return {"ring": ring, "hub": hub_box, "centre": (cx, cy), "radius": R, "labels": label_boxes}


# ------------------------------------------------------------------ mark_wall

def _mark_kind(key: str):
    if not isinstance(key, str) or ":" not in key:
        raise UnsupportedPart("mark_wall key %r" % (key,), need="vendor:<key>, mark:<key> or mono:<letters>",
                              have=key)
    kind, val = key.split(":", 1)
    if kind not in ("vendor", "mark", "mono"):
        raise UnsupportedPart("mark_wall key %r" % key, need="vendor:, mark: or mono:", have=kind)
    return kind, val


def _check_mark(key, th):
    kind, val = _mark_kind(key)
    if kind == "vendor":
        why = icons.vendor_key_refusal(val)
        if why or not icons.vendor_resolves(val, th):
            raise UnsupportedPart("mark_wall vendor %r" % val, need="a vendor key the theme accepts",
                                  have=why or "unresolved", hint="icons.vendor_keys() lists them")
    elif kind == "mark":
        if val not in icons.mark_registry():
            raise UnsupportedPart("mark_wall mark %r" % val, need="a key in kit/assets/marks/registry.json",
                                  have="missing")
    return kind, val


SQUARE_SHARE = 0.72          # a square mark takes at most this share of its column's width (the rest is air)
MARK_MARGIN = 0.24           # in, x k: the least air in a column beside a wide mark, both sides together
WIDE_CAP = 2.4               # a mark is never drawn wider than this many mark sizes


_INK_CACHE: dict = {}


def _mark_ink(key: str) -> tuple[float, tuple | None]:
    """(aspect, crop) of a mark's ink, the part of its artwork that paints anything.

    aspect is the ink's width over its height; crop is (left, top, right, bottom), the share of the image on each
    side that is fully transparent margin, or None when there is none to speak of. Vendor marks and monograms are
    square tiles drawn whole: (1.0, None). A registry mark is rasterised once (rsvg-convert for an SVG) and its alpha
    bounding box read, so a wordmark published on a wide empty canvas is measured by its letters, not its canvas.
    Nothing is altered: the crop trims only pixels that paint nothing."""
    kind, val = _mark_kind(key)
    if kind != "mark":
        return 1.0, None
    entry = icons.mark_registry()[val]
    blob = (icons.marks_dir() / entry["file"]).read_bytes()
    ck = (val, icons.blob_sha256(blob))
    if ck not in _INK_CACHE:
        from PIL import Image
        if entry["file"].lower().endswith(".svg"):
            sw, sh = icons.svg_size(blob)
            raster = icons.rasterise(blob, 600)
        else:
            raster = blob
            with Image.open(io.BytesIO(blob)) as im0:
                sw, sh = im0.size
        with Image.open(io.BytesIO(raster)) as im:
            im = im.convert("RGBA")
            W, H = im.size
            bb = im.getchannel("A").getbbox()
        if bb is None:
            raise GeometryError("mark_wall mark %s" % val, need="artwork that paints something", have="blank")
        crop = [bb[0] / W, bb[1] / H, 1.0 - bb[2] / W, 1.0 - bb[3] / H]
        crop = [0.0 if c < 0.01 else c for c in crop]          # anti-aliasing noise at the canvas edge is not margin
        aspect = ((1.0 - crop[0] - crop[2]) * sw) / ((1.0 - crop[1] - crop[3]) * sh)
        _INK_CACHE[ck] = (aspect, tuple(crop) if any(crop) else None)
    return _INK_CACHE[ck]


def _mark_size(aspect: float, side: float, cap_w: float) -> tuple[float, float]:
    """The drawn (width, height) of a mark's ink at the wall's mark size `side`.

    Equal visual weight: a mark wider than tall keeps the ink area of a square mark (side x side), so a 4.7:1
    wordmark is drawn 2.16 sides wide and 0.46 sides tall instead of being squeezed into a square's width; a square
    mark is side x side, unchanged; a tall mark fits the side x side square (its height sets the caption line). The
    width is capped at WIDE_CAP sides and by the column (cap_w), the height following the mark's own aspect."""
    if aspect >= 1.0:
        w = side * min(math.sqrt(aspect), WIDE_CAP)
        h = w / aspect
    else:
        w, h = side * aspect, side
    if w > cap_w:
        w, h = cap_w, cap_w / aspect
    return w, h


@ooxml.transactional
def mark_wall(slide, box: Box, groups, *, tab_fill="ink-strong", fill="card", caption_pt=CAPTION_PT,
              max_mark=0.80, name: str, theme=None) -> dict:
    """Tabbed groups of marks: each group a solid tab naming it over a card of marks, each mark over its caption.

    groups: [(label, [(key, caption)])]; a key is "vendor:<diagrams key>" (icons.vendor, byte-verified),
    "mark:<registry key>" (icons.mark, byte-verified) or "mono:<letters>" (a monogram tile for a product with no
    offline mark). Every mark gets a caption: a wall of unlabelled logos is decoration. Columns are as wide as their
    marks need (a group of six is wider than a group of two, never an even tile grid); every group shares one row
    count and one mark size, chosen as the largest that fits. Raises Overflows when the marks would fall under
    0.34in or a caption needs more than two lines.

    Visual weight is normalised. Every mark is drawn at the ink area of a square mark (_mark_size), measured on its
    ink rather than its canvas (_mark_ink, which trims a canvas's empty margin), so a wide wordmark gets a wider
    column instead of reading as small text beside full-size squares. When a wall holds a wide mark, each column is
    as wide as its widest need (a square at SQUARE_SHARE of the column, a wide mark's drawn width plus MARK_MARGIN,
    a caption in two lines), the mark size is the largest at which every column fits, and spare width is shared out
    in proportion to need. A square mark is side x side as before, and a wall with no wide mark lays out exactly as
    it did. Every mark stays inside its column; the captions in a row share one top and one height.

    Returns {"groups": [Box], "mark": side_in, "rows": rows}."""
    th = theme_of(theme)
    k = _k(slide)
    groups = [(str(lab), list(items)) for lab, items in groups]
    if not groups or any(not items for _l, items in groups):
        raise Capacity("mark_wall %s" % name, need="at least one group, each with a mark", have=len(groups))
    if fill not in ("card", "paper"):
        raise ThemeInvariant("mark_wall %s fill" % name, need="card or paper", have=fill,
                             hint="marks keep their own colours, so they sit on a light ground")
    for _l, items in groups:
        for key, cap in items:
            _check_mark(key, th)
            if not str(cap or "").strip():
                raise GeometryError("mark_wall %s caption for %s" % (name, key), need="a caption", have="none")
    ink = [[_mark_ink(key) for key, _cap in items] for _l, items in groups]
    any_wide = any(a > 1.0 for inks in ink for a, _c in inks)
    tab_h = 0.40 * k
    pad = 0.16 * k
    gut = 0.16 * k
    margin = MARK_MARGIN * k
    body_h = box.h - tab_h
    cap_face = _face(th)
    n_g = len(groups)
    avail = box.w - gut * (n_g - 1) - 2 * pad * n_g
    cap_max = measure.line_in(caption_pt, cap_face) * 2 * measure.SLACK + 0.02

    def cap_paras(cap):
        return [Para((Run(str(cap), "ink", size=caption_pt),), align="center")]

    def cap_fits(cap, col_w):
        try:
            return _need(cap_paras(cap), col_w - 0.08 * k, th) <= cap_max
        except Overflows:                       # a word wider than its column
            return False

    def cap_min(cap):
        """The narrowest column that sets the caption in two lines."""
        lo, hi = 0.10 * k, max(0.5 * k, measure.width_in(str(cap), caption_pt, cap_face) * measure.SLACK + 0.3 * k)
        while not cap_fits(cap, hi):
            hi *= 1.5
        for _ in range(30):
            mid = (lo + hi) / 2.0
            lo, hi = (lo, mid) if cap_fits(cap, mid) else (mid, hi)
        return hi + 0.02 * k                    # clear of the boundary, so rounding to EMU never tips it

    def caps_need(colw):
        hs = []
        for (_l, items), g in zip(groups, colw):
            c = len(g)
            for j, (_key, cap) in enumerate(items):
                if not cap_fits(cap, g[j % c]):
                    return None
                hs.append(_need(cap_paras(cap), g[j % c] - 0.08 * k, th))
        return max(hs)

    def pieces(rows):
        """Per group, per column: the linear needs (u, v) of what it holds, width = u * side / SQUARE_SHARE + v."""
        out = []
        for (_l, items), inks in zip(groups, ink):
            c = int(math.ceil(len(items) / rows))
            col = [[(1.0, 0.0)] for _ in range(c)]
            for j, ((_key, cap), (aspect, _crop)) in enumerate(zip(items, inks)):
                col[j % c].append((0.0, cap_min(cap)))
                if aspect > 1.0:
                    col[j % c].append((min(math.sqrt(aspect), WIDE_CAP) * SQUARE_SHARE, margin))
            out.append(col)
        return out

    def widths_at(pcs, side):
        """Column widths at mark size side, spare width shared in proportion to need; None when they overflow."""
        cw = side / SQUARE_SHARE
        need = [[max(u * cw + v for u, v in col) for col in g] for g in pcs]
        total = sum(sum(g) for g in need)
        if total > avail + 1e-9:
            return None
        return [[w * avail / total for w in g] for g in need]

    def side_for_width(pcs):
        """The largest mark size at which every column fits: bisected, then exact on the active linear pieces."""
        lo, hi = 0.0, 4.0 * k
        if widths_at(pcs, 1e-6) is None:
            return 0.0
        for _ in range(60):
            mid = (lo + hi) / 2.0
            lo, hi = (mid, hi) if widths_at(pcs, mid) is not None else (lo, mid)
        cw = lo / SQUARE_SHARE
        act = [max(col, key=lambda p: p[0] * cw + p[1]) for g in pcs for col in g]
        units, fixed = sum(u for u, _v in act), sum(v for _u, v in act)
        exact = SQUARE_SHARE * ((avail - fixed) / units) if units else lo
        return exact if widths_at(pcs, exact) is not None else lo

    best = None
    max_n = max(len(items) for _l, items in groups)
    for rows in range(1, max_n + 1):
        cols = [int(math.ceil(len(items) / rows)) for _l, items in groups]
        cell_h = (body_h - 2 * pad) / rows
        if not any_wide:                         # the even grid, exactly as before
            cw = avail / sum(cols)
            if cw < 0.70 * k:
                continue
            colw = [[cw] * c for c in cols]
            cap_h = caps_need(colw)
            if cap_h is None:
                continue
            side = min(max_mark * k, SQUARE_SHARE * cw, cell_h - cap_h - 0.20 * k)
        else:
            pcs = pieces(rows)
            side = min(max_mark * k, side_for_width(pcs))
            colw = cap_h = None
            for _ in range(6):                   # captions wrap with the widths: settle the side against the height
                colw = widths_at(pcs, side) if side > 0 else None
                cap_h = caps_need(colw) if colw is not None else None
                if cap_h is None:
                    break
                nxt = min(side, cell_h - cap_h - 0.20 * k)
                if nxt >= side - 1e-12:
                    break
                side = nxt
            if colw is None or cap_h is None or min(w for g in colw for w in g) < 0.70 * k:
                continue
        if side < 0.34 * k:
            continue
        empty = sum(c * rows - len(items) for c, (_l, items) in zip(cols, groups))
        empty_rows = sum(rows - int(math.ceil(len(items) / c)) for c, (_l, items) in zip(cols, groups))
        score = (round(side - 0.08 * k * empty_rows - 0.01 * k * empty, 3), -empty)   # a card with an empty row is air
        if best is None or score > best[0]:
            best = (score, rows, cols, colw, cap_h, side, cell_h)
    if best is None:
        raise Overflows("mark_wall %s" % name, need="marks of 0.34in with two-line captions",
                        have="%.2f x %.2fin" % (box.w, box.h), hint="fewer marks, shorter captions or a larger box")
    _score, rows, cols, colw, cap_h, side, cell_h = best
    widths = [sum(g) + 2 * pad for g in colw]
    x = box.x
    out = []
    for gi, ((label, items), c, w, g, inks) in enumerate(zip(groups, cols, widths, colw, ink), 1):
        gname = "%s.%d" % (name, gi)
        prims.tab(slide, Box(x, box.y, w, tab_h), label, fill=tab_fill, align="c", name="%s.tab" % gname, theme=th)
        body = Box(x, box.y + tab_h, w, body_h)
        prims.card(slide, body, fill=fill, line=None if fill == "card" else "rule-strong", radius_in=0,
                   name="%s.body" % gname, theme=th)
        y0 = body.y + pad            # every group starts its marks on one line (P07); the fullest sets the air
        for j, ((key, cap), (aspect, crop)) in enumerate(zip(items, inks)):
            r, cc = divmod(j, c)
            in_row = min(c, len(items) - r * c)
            row_x = body.x + pad + (sum(g) - sum(g[:in_row])) / 2.0      # a short last row is centred
            cell = Box(row_x + sum(g[:cc]), y0 + r * cell_h, g[cc], cell_h)
            stack_h = side + 0.10 * k + cap_h
            my = cell.y + (cell.h - stack_h) / 2.0
            kind, val = _mark_kind(key)
            sub = "%s.%d" % (gname.replace(".", "-"), j + 1)
            if kind == "vendor":
                icons.vendor(slide, val, Box(cell.cx - side / 2, my, side, side), name=sub, alt=str(cap), theme=th)
            elif kind == "mark":
                mw, mh = _mark_size(aspect, side, cell.w - margin)
                inkb = Box(cell.cx - mw / 2.0, my + (side - mh) / 2.0, mw, mh)
                if crop is None:
                    icons.mark(slide, val, inkb, name=sub, alt=str(cap))
                else:                # place the whole canvas so its ink lands on inkb, then trim the empty margin
                    l_, t_, r_, b_ = crop
                    fw, fh = mw / (1.0 - l_ - r_), mh / (1.0 - t_ - b_)
                    pic = icons.mark(slide, val, Box(inkb.x - l_ * fw, inkb.y - t_ * fh, fw, fh), name=sub,
                                     alt=str(cap))
                    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = l_, t_, r_, b_
                    pic.left, pic.top, pic.width, pic.height = inkb.emu
            else:
                icons.monogram(slide, Box(cell.cx - side / 2, my, side, side), val, "ink-strong",
                               name="%s.%d.mono" % (gname, j + 1), theme=th)
            textbox(slide, Box(cell.x + 0.04 * k, my + side + 0.10 * k, cell.w - 0.08 * k, cap_h), cap_paras(cap),
                    "%s.%d.caption" % (gname, j + 1), th, background=fill)
        out.append(body)
        x += w + gut
    return {"groups": out, "mark": side, "rows": rows}


# ------------------------------------------------------------------ pictures: portrait, screen

def _image_bytes(image) -> bytes:
    if isinstance(image, (bytes, bytearray)):
        return bytes(image)
    p = Path(image)
    if not p.is_file():
        raise UnsupportedPart("image", need="a PNG or JPEG file", have=str(p))
    return p.read_bytes()


def _image_size(blob) -> tuple[int, int]:
    from PIL import Image
    try:
        with Image.open(io.BytesIO(blob)) as im:
            if im.format not in ("PNG", "JPEG"):
                raise UnsupportedPart("image", need="PNG or JPEG", have=im.format)
            return im.size
    except UnsupportedPart:
        raise
    except Exception as e:  # noqa: BLE001 - PIL raises many types for a non-image
        raise UnsupportedPart("image", need="a PNG or JPEG", have=str(e)[:60]) from None


def _set_geom(pic, prst):
    sp = pic._element.find(_P + "spPr")
    pg = sp.find(_A + "prstGeom")
    pg.set("prst", prst)


def _set_alt(pic, name, alt):
    nv = pic._element.find(".//" + _P + "cNvPr")
    nv.set("name", name)
    nv.set("descr", alt)


@ooxml.transactional
def portrait(slide, box: Box, image=None, initials=None, *, fill="ink-deep", alt=None, name: str, theme=None):
    """A person's portrait in a circle: a photograph (PNG or JPEG, centre-cropped to a square, clipped to an ellipse)
    or, when no photograph is given, an initials disc in a solid fill. Never a generic silhouette: one of the two is
    required. The circle is the largest square centred in box. A photograph must carry at least 192 pixels per
    displayed inch on its cropped side (gate P22) and an alt text naming whose portrait it is."""
    th = theme_of(theme)
    side = min(box.w, box.h)
    sq = Box(box.cx - side / 2, box.cy - side / 2, side, side)
    if image is None and not initials:
        raise GeometryError("portrait %s" % name, need="a photograph or initials", have="neither",
                            hint="a generic silhouette reads as a placeholder; give the person's initials")
    if image is not None:
        blob = _image_bytes(image)
        pw, ph = _image_size(blob)
        s = min(pw, ph)
        if s < side * 96 * icons.RASTER_MIN_SCALE - 1e-6:
            raise GeometryError("portrait %s" % name, need="%d px" % round(side * 96 * icons.RASTER_MIN_SCALE),
                                have="%d px" % s, hint="a larger photograph or a smaller portrait")
        label = " ".join(str(alt or "").split())
        if not label:
            raise GeometryError("portrait %s alt" % name, need="alt text naming the person or role", have="none")
        pic = slide.shapes.add_picture(io.BytesIO(blob), *sq.emu)
        if pw > ph:
            c = (pw - ph) / 2.0 / pw
            pic.crop_left, pic.crop_right = c, c
        elif ph > pw:
            c = (ph - pw) / 2.0 / ph
            pic.crop_top, pic.crop_bottom = c * 0.6, c * 1.4          # faces sit high in a frame
        _set_geom(pic, "ellipse")
        _set_alt(pic, name, "Portrait: %s" % label[:100])
        ooxml.finish(pic)
        return pic
    letters = str(initials).strip()
    if not (1 <= len(letters) <= 3) or not letters.replace("&", "").isalnum():
        raise GeometryError("portrait %s initials" % name, need="1 to 3 letters", have=letters)
    if fill not in STRONG:
        raise ThemeInvariant("portrait %s fill" % name, need=" / ".join(STRONG), have=fill)
    pt = math.floor(side * 72 * (0.36 if len(letters) < 3 else 0.28) * 2) / 2.0
    paras = [Para((Run(letters, ON_FILL[fill], bold=True, size=prims.check_pt(pt, "portrait initials")),),
                  align="center")]
    inner = side * math.sqrt(0.5)
    disc = prims.circle(slide, sq, fill=fill, name=name, theme=th)
    textbox(slide, Box(sq.cx - inner / 2, sq.cy - inner / 2, inner, inner), paras, "%s.initials" % name, th,
            background=fill, anchor="middle")
    return disc


@ooxml.transactional
def quote_portrait(slide, box: Box, text: str, attribution: str, *, image=None, initials=None, alt=None,
                   portrait_fill="ink-deep", fill="card", portrait_in=None, name: str, theme=None) -> Box:
    """A quotation beside its speaker: a portrait (photograph or initials) at the left, the quote panel (prims.quote,
    its state.quote bar marking the words as the speaker's own) to its right, both centred on one line. The panel
    is as tall as its words need, never stretched to the box; Overflows when it needs more than the box.
    Returns the box the pair occupies."""
    th = theme_of(theme)
    k = _k(slide)
    d = min(portrait_in * k if portrait_in else 1.10 * k, box.h)
    gap = 0.24 * k
    qx = box.x + d + gap
    qw = box.r - qx
    if qw < 2.0 * k:
        raise Overflows("quote_portrait %s" % name, need=round(d + gap + 2.0 * k, 3), have=round(box.w, 3), unit="in")
    try:
        prims.quote(slide, Box(qx, box.y, qw, 0.20), text, attribution, fill=fill, name="%s.probe" % name, theme=th)
        raise GeometryError("quote_portrait %s probe" % name, need="a refusal at 0.20in", have="a quote drawn")
    except Overflows as e:
        need = math.ceil(float(e.need) * 1000 + 1) / 1000.0        # the refusal rounds its need to 0.001in
    if need > box.h + 1e-6:
        raise Overflows("quote_portrait %s" % name, need=round(need, 3), have=round(box.h, 3),
                        hint="cut the passage or give the pair more height")
    h = max(d, need)
    top = box.y + (box.h - h) / 2.0
    portrait(slide, Box(box.x, top + (h - d) / 2, d, d), image=image, initials=initials, fill=portrait_fill,
             alt=alt, name="%s.portrait" % name, theme=th)
    prims.quote(slide, Box(qx, top + (h - need) / 2, qw, need), text, attribution, fill=fill,
                name="%s.quote" % name, theme=th)
    return Box(box.x, top, box.w, h)


def capture_record(png) -> dict | None:
    """The sidecar scripts/capture.py writes beside a capture (<stem>.capture.json), checked against the PNG's
    sha256; None when there is no sidecar."""
    p = Path(png)
    side = p.with_name(p.stem + CAPTURE_SUFFIX)
    if not side.is_file():
        return None
    rec = json.loads(side.read_text(encoding="utf-8"))
    sha = hashlib.sha256(p.read_bytes()).hexdigest()
    if rec.get("sha256") != sha:
        raise UnsupportedPart("capture %s" % p.name, need="the PNG its sidecar records (sha256 %s)"
                              % str(rec.get("sha256"))[:12], have=sha[:12], hint="re-run scripts/capture.py")
    return rec


@ooxml.transactional
def screen_frame(slide, box: Box, png, chrome="browser", band=None, *, pill=None, alt="Product screen",
                 type_pt=None, crop_top=0.0, whole=False, name: str, theme=None) -> dict:
    """A product capture staged as running software, filling the box's height (the product-screen card measures
    the stage at about 80% of the slide).

    chrome "browser": a window with a bar of three neutral dots and an address field; "laptop": a near-black bezel
    over a light base; None: the bare capture with a hairline edge. The window hugs the capture: its width follows
    the box's width, its height is the capture's at that width up to the box's height, and a capture taller than
    that is cropped at its foot, never shrunk (crop_top trims a share off its top first); pass a narrower box to
    show a capture whole, or whole=True to narrow the window until the capture fits uncropped (its type then
    shrinks with it and the floor decides). band: a solid role (accent, ink-deep, ...) drawn as a full-width
    band (bleed.band) behind the window's middle third, the window narrowed to about two thirds of the box and set
    left so the band shows to its right, where pill (a short label naming the moment) sits.

    The floor: a capture made by scripts/capture.py carries a sidecar with its smallest CSS font and CSS width;
    screen_frame converts it to points at the placed width and raises BelowFloor under 8pt. Without a sidecar pass
    type_pt, the smallest type you measured on the placed capture, or 0 for a capture with no text. The figure is
    written into the picture's alt text and returned, for the manifest.

    Returns {"window": Box, "screen": Box, "type_pt": float}."""
    th = theme_of(theme)
    k = _k(slide)
    cv = canvas_for(slide)
    if chrome not in ("browser", "laptop", None):
        raise GeometryError("screen_frame %s chrome" % name, need="browser, laptop or None", have=chrome)
    if band is not None and band not in STRONG:
        raise ThemeInvariant("screen_frame %s band" % name, need=" / ".join(STRONG), have=band)
    if pill and band is None:
        raise GeometryError("screen_frame %s pill" % name, need="a band to sit on", have="no band")
    blob = _image_bytes(png)
    pw, ph = _image_size(blob)
    rec = capture_record(png) if not isinstance(png, (bytes, bytearray)) else None
    if rec is None and type_pt is None:
        raise UnsupportedPart("screen_frame %s" % name, need="a capture sidecar or type_pt",
                              have="neither", hint="capture with scripts/capture.py, or pass the smallest type in "
                                                   "points as placed (0 for a capture with no text)")
    if not 0.0 <= crop_top < 0.5:
        raise GeometryError("screen_frame %s crop_top" % name, need="0 to 0.5", have=crop_top)
    # chrome insets: (left, top, right, bottom) around the screen inside the window
    ins = {"browser": (0.10 * k, 0.44 * k, 0.10 * k, 0.10 * k), "laptop": (0.16 * k, 0.16 * k, 0.16 * k, 0.16 * k),
           None: (0.0, 0.0, 0.0, 0.0)}[chrome]
    base_h = 0.22 * k if chrome == "laptop" else 0.0
    base_x = 1.10 if chrome == "laptop" else 1.0           # the laptop base overhangs the bezel by 5% a side
    max_w = box.w * (0.66 if band else 1.0) / base_x
    win_h = box.h - base_h
    scr_h = win_h - ins[1] - ins[3]
    vis = 1.0 - crop_top
    scr_w = max_w - ins[0] - ins[2]                       # the full width: type stays as large as it can
    if whole:
        scr_w = min(scr_w, scr_h * pw / (ph * vis))       # ... unless the capture must show whole
    shown_h = min(scr_h, scr_w * ph * vis / pw)           # a taller capture is cropped at its foot, never shrunk
    crop_bottom = max(0.0, 1.0 - crop_top - (shown_h * pw) / (scr_w * ph))
    win_w = scr_w + ins[0] + ins[2]
    win_h = shown_h + ins[1] + ins[3]
    wx = box.x + (base_x - 1.0) * win_w / 2.0 if band else box.cx - win_w / 2.0
    wy = box.y + (box.h - base_h - win_h) / 2.0
    window = Box(wx, wy, win_w, win_h)
    screen = Box(wx + ins[0], wy + ins[1], scr_w, shown_h)
    # resolution (P22) and the type floor, before anything is drawn
    if pw < scr_w * 96 * icons.RASTER_MIN_SCALE - 1e-6:
        raise GeometryError("screen_frame %s" % name, need="%d px wide" % round(scr_w * 96 * icons.RASTER_MIN_SCALE),
                            have="%d px" % pw, hint="capture at device scale 2 or more")
    if rec is not None:
        css_w = float(rec["css_width"])
        type_pt = float(rec["min_font_px"]) * (scr_w * 72.0 / css_w) if rec.get("min_font_px") else 0.0
    type_pt = round(float(type_pt), 2)
    if 0.0 < type_pt < measure.FLOOR - 1e-9:
        raise BelowFloor("screen_frame %s type" % name, need=measure.FLOOR, have=type_pt,
                         hint="place the capture wider, or capture at a narrower CSS width or with larger type")
    suffix = " (capture, smallest type %.1fpt as placed)" % type_pt
    label = " ".join(str(alt).split())[:120 - len(suffix)] or "Product screen"     # P20 holds alt text to 120
    if band:
        bh = win_h / 3.0
        by = wy + win_h / 3.0
        prims.rect(slide, Box(0.0, by, cv.w, bh), fill=band, name="bleed.band", theme=th)
    if chrome == "browser":
        prims.card(slide, window, fill="card", line="hairline", name="%s.window" % name, theme=th)
        dot = 0.11 * k
        for i in range(3):
            prims.circle(slide, Box(wx + 0.18 * k + i * (dot + 0.08 * k), wy + (ins[1] - dot) / 2, dot, dot),
                         fill=None, line="rule-strong", name="%s.dot.%d" % (name, i + 1), theme=th)
        fw, fh = min(4.2 * k, win_w * 0.42), 0.24 * k
        field = Box(window.cx - fw / 2, wy + (ins[1] - fh) / 2, fw, fh)
        prims.card(slide, field, fill="paper", line=None, name="%s.address" % name, theme=th)
        ls = 0.14 * k
        glyph(slide, "lock", Box(field.x + 0.10 * k, field.cy - ls / 2, ls, ls), role="ink-muted",
              shape_name="%s.address.lock" % name, theme=th)
    elif chrome == "laptop":
        prims.block(slide, window, fill="ink-deep", name="%s.bezel" % name, theme=th)
        bw = win_w * base_x
        prims.block(slide, Box(window.cx - bw / 2, window.b, bw, base_h), fill="hairline-soft",
                    name="%s.base" % name, theme=th)
    pic = slide.shapes.add_picture(io.BytesIO(blob), *screen.emu)
    pic.crop_top, pic.crop_bottom = crop_top, crop_bottom
    _set_alt(pic, "%s.screen" % name, label + suffix)
    if chrome is None:
        prims.set_line(pic, "hairline", cv.hairline_pt, theme=th)
    ooxml.finish(pic)
    if band and pill:
        zone_x = window.r + 0.30 * k
        zw = box.r - zone_x
        pt = 12.0
        f = "accent-deep" if band == "accent" else "accent"
        tw = measure.width_in(str(pill), pt, _face(th)) * measure.SLACK + 0.5 * k
        ph_ = 0.46 * k
        if tw > zw:
            raise Overflows("screen_frame %s pill" % name, need=round(tw, 3), have=round(zw, 3),
                            hint="a shorter pill label")
        prims.pill(slide, Box(zone_x + (zw - tw) / 2, by + (bh - ph_) / 2, tw, ph_), str(pill), fill=f, pt=pt,
                   name="%s.pill" % name, theme=th)
    return {"window": window, "screen": screen, "type_pt": type_pt}


# ------------------------------------------------------------------ cell_grid

def _state(legend, cell, where):
    if isinstance(cell, (tuple, list)):
        key, word = cell[0], cell[1]
    else:
        key, word = cell, None
    if key not in legend:
        raise GeometryError("cell_grid %s state" % where, need=" / ".join(legend), have=key)
    g, w, f = legend[key]
    return key, g, (word if word else w), f


@ooxml.transactional
def cell_grid(slide, box: Box, rows, cols, cells, legend, *, key=None, row_w=None, gap=0.06, name: str,
              theme=None) -> dict:
    """A coverage matrix: column headers in solid tabs, row labels on a near-black rail, and one cell per row and
    column carrying its state as a glyph plus a word (identity never by colour alone).

    rows: [label] or [(label, line)]; cols: [label]; cells: rows x cols of a state key, or (key, word) to print
    another word in that state (a date, a wave). legend: {key: (glyph, word, fill)}, fill one of accent,
    accent-deep, ink-deep, ink-strong, navy, card or paper (paper cells carry a rule-strong outline). key: draw a
    one-line key under the grid (glyph and word per state); by default only when some cell overrides its word.
    Raises Overflows when a label or a cell word does not fit its box.

    Returns {"cells": {(r, c): Box}, "header": Box, "rail": Box}."""
    th = theme_of(theme)
    k = _k(slide)
    rows, cols = list(rows), list(cols)
    nr, nc = len(rows), len(cols)
    if not rows or not cols:
        raise Capacity("cell_grid %s" % name, need="rows and columns", have=(nr, nc))
    if len(cells) != nr or any(len(r) != nc for r in cells):
        raise GeometryError("cell_grid %s cells" % name, need="%d rows of %d" % (nr, nc),
                            have=[len(r) for r in cells])
    legend = {str(kk): tuple(v) for kk, v in legend.items()}
    for kk, (g, w, f) in legend.items():
        if f not in STRONG + LIGHT:
            raise ThemeInvariant("cell_grid %s state %s fill" % (name, kk), need=" / ".join(STRONG + LIGHT), have=f)
    states = [[_state(legend, c, "%s r%d c%d" % (name, ri + 1, ci + 1)) for ci, c in enumerate(r)]
              for ri, r in enumerate(cells)]
    if key is None:
        key = any(isinstance(c, (tuple, list)) and c[1] for r in cells for c in r)
    gp = gap * k
    rpad = 0.14 * k
    row_paras = []
    for r in rows:
        head, line = (r, None) if isinstance(r, str) else (tuple(r) + (None,))[:2]
        row_paras.append(_stack(head, line, "paper", "paper", head_pt=11.0, body_pt=10.0, gap=1.0))
    if row_w is None:
        need_w = max(measure.width_in(w, 11.0, _face(th, True)) for r in rows
                     for w in str(r if isinstance(r, str) else r[0]).split()) * measure.SLACK
        row_w = min(max(1.6 * k, need_w + 2 * rpad + 0.6 * k), 2.8 * k)
    key_h = 0.34 * k if key else 0.0
    head_paras = [[Para((Run(str(c), "paper", bold=True, size=11.0),), align="center")] for c in cols]
    grid_x = box.x + row_w + gp
    cw = (box.r - grid_x - gp * (nc - 1)) / nc
    if cw < 0.8 * k:
        raise Overflows("cell_grid %s columns" % name, need=round(0.8 * k * nc + row_w, 3), have=round(box.w, 3),
                        hint="fewer columns or a narrower rail")
    head_h = max(0.44 * k, max(_need(p, cw - 0.16 * k, th) for p in head_paras) + 0.16 * k)
    body_h = box.h - head_h - gp - key_h - (0.10 * k if key else 0.0)
    rh = (body_h - gp * (nr - 1)) / nr
    gs = 0.22 * k
    for ri, p in enumerate(row_paras):
        if _need(p, row_w - 2 * rpad, th) > rh - 0.08 * k:
            raise Overflows("cell_grid %s row %d label" % (name, ri + 1), need=round(_need(p, row_w - 2 * rpad, th), 3),
                            have=round(rh - 0.08 * k, 3), hint="a shorter row label or a taller box")
    word_paras = {}
    for ri in range(nr):
        for ci in range(nc):
            _key, g, word, f = states[ri][ci]
            role = ON_FILL[f] if f in STRONG else ("ink-muted" if f == "paper" else "ink")
            p = [Para((Run(str(word), role, bold=f in STRONG, size=10.0),))]
            ww = measure.width_in(str(word), 10.0, _face(th, f in STRONG)) * measure.SLACK + 0.02
            if gs + 0.06 * k + ww > cw - 0.16 * k:
                raise Overflows("cell_grid %s r%d c%d word" % (name, ri + 1, ci + 1), need=round(gs + ww + 0.22 * k, 3),
                                have=round(cw, 3), hint="a shorter word")
            word_paras[ri, ci] = (p, ww, role)
    if rh < max(gs, measure.line_in(10.0, _face(th))) + 0.12 * k:
        raise Overflows("cell_grid %s rows" % name, need=round(nr * 0.4 * k + head_h, 3), have=round(box.h, 3))
    # draw
    header = Box(grid_x, box.y, box.r - grid_x, head_h)
    for ci, p in enumerate(head_paras):
        hb = Box(grid_x + ci * (cw + gp), box.y, cw, head_h)
        prims.block(slide, hb, fill="ink-strong", name="%s.col.%d" % (name, ci + 1), theme=th)
        h = _need(p, cw - 0.16 * k, th)
        textbox(slide, Box(hb.x + 0.08 * k, hb.y + (hb.h - h) / 2, cw - 0.16 * k, h), p,
                "%s.col.%d.text" % (name, ci + 1), th, background="ink-strong")
    out = {}
    y = box.y + head_h + gp
    for ri, p in enumerate(row_paras):
        rb = Box(box.x, y, row_w, rh)
        prims.block(slide, rb, fill="ink-deep", name="%s.row.%d" % (name, ri + 1), theme=th)
        h = _need(p, row_w - 2 * rpad, th)
        textbox(slide, Box(rb.x + rpad, rb.y + (rh - h) / 2, row_w - 2 * rpad, h), p,
                "%s.row.%d.text" % (name, ri + 1), th, background="ink-deep")
        for ci in range(nc):
            _key, g, word, f = states[ri][ci]
            cb = Box(grid_x + ci * (cw + gp), y, cw, rh)
            nm = "%s.r%d.c%d" % (name, ri + 1, ci + 1)
            if f in STRONG:
                prims.block(slide, cb, fill=f, name=nm, theme=th)
            else:
                prims.card(slide, cb, fill=f, line="rule-strong" if f == "paper" else None, name=nm, theme=th)
            wp, ww, role = word_paras[ri, ci]
            lh = measure.line_in(10.0, _face(th)) * measure.SLACK
            total = gs + 0.06 * k + ww
            gx = cb.cx - total / 2
            glyph(slide, g, Box(gx, cb.cy - gs / 2, gs, gs), role=ON_FILL[f] if f in STRONG else
                  ("ink-muted" if f == "paper" else "accent"), shape_name="%s.icon" % nm, theme=th)
            textbox(slide, Box(gx + gs + 0.06 * k, cb.cy - lh / 2, ww, lh), wp, "%s.text" % nm, th, background=f)
            out[ri, ci] = cb
        y += rh + gp
    if key:
        ky = box.b - key_h
        x = grid_x
        for kk, (g, w, f) in legend.items():
            ww = measure.width_in(str(w), 10.0, _face(th)) * measure.SLACK + 0.02
            sw = 0.30 * k
            sb = Box(x, ky + (key_h - sw) / 2, sw, sw)
            if f in STRONG:
                prims.block(slide, sb, fill=f, name="%s.key.%s" % (name, kk), theme=th)
            else:
                prims.card(slide, sb, fill=f, line="rule-strong", name="%s.key.%s" % (name, kk), theme=th)
            glyph(slide, g, sb.inset(0.05 * k), role=ON_FILL[f] if f in STRONG else "ink-muted",
                  shape_name="%s.key.%s.icon" % (name, kk), theme=th)
            lh = measure.line_in(10.0, _face(th)) * measure.SLACK
            textbox(slide, Box(sb.r + 0.08 * k, ky + (key_h - lh) / 2, ww, lh),
                    [Para((Run(str(w), "ink-muted", size=10.0),))], "%s.key.%s.text" % (name, kk), th)
            x = sb.r + 0.08 * k + ww + 0.36 * k
        if x - 0.36 * k > box.r + 1e-6:
            raise Overflows("cell_grid %s key" % name, need=round(x - 0.36 * k - grid_x, 3),
                            have=round(box.r - grid_x, 3), hint="shorter state words")
    return {"cells": out, "header": header, "rail": Box(box.x, box.y + head_h + gp, row_w, body_h)}


__all__ = ["cycle_ring", "mark_wall", "screen_frame", "portrait", "quote_portrait", "cell_grid", "capture_record"]
